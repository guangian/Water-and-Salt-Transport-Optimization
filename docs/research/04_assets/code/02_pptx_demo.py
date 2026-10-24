# -*- coding: utf-8 -*-
"""python-pptx 1.0.2 演示：生成 out/02_路演示例.pptx 与模板复用 out/02_模板填充.pptx"""
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

import pptx_helpers as P  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "out")
CHART_IMG = os.path.join(OUT, "charts", "09_竞争矩阵.png")
CANVAS_IMG = os.path.join(OUT, "charts", "11_商业模式画布.png")

prs = P.new_presentation()
years = ["2024A", "2025E", "2026E", "2027E", "2028E"]

P.cover_slide(prs, "智守康联：居家养老 AI 健康监测平台", "商业计划路演 | 2026 全国高校商业精英挑战赛", "智守康联团队", "2026 年 10 月")
P.agenda_slide(prs, ["项目概述", "市场分析", "产品与技术", "商业模式", "财务预测", "团队介绍", "融资计划"], page_no=2)
P.section_slide(prs, 1, "项目概述", "为 3,000 万独居老人提供无感健康守护")
P.bullets_slide(prs, "行业痛点与解决方案", [
    "独居老人意外跌倒后平均 4.5 小时才被发现", ("现有穿戴设备依从性低于 40%，夜间摘除率高", 1),
    "毫米波雷达无感监测 + 边缘 AI 3 分钟预警", ("无需佩戴、无摄像头隐私顾虑", 1), ("家属 App + 社区大屏双端联动", 1),
], subtitle="要点页：两级项目符号", page_no=4, notes="演讲备注：强调 4.5 小时这个数字来自民政部 2024 年抽样调查。")
s = P.image_text_slide(prs, "竞争格局", CHART_IMG, ["本项目定位“低价高能”象限", "A 公司价格高、依赖穿戴设备",
                                                 "C 公司仅提供摄像头方案，隐私争议大"], subtitle="图文页：matplotlib 图片 + 要点，图片四边各裁 3%",
                       page_no=5, crop=(0.03, 0.03, 0.03, 0.03))
s, chart = P.chart_slide(prs, "市场规模：五年 CAGR 37.6%", P.add_bar_chart, subtitle="数据图表页：原生堆叠柱状图（可在 PPT 内编辑数据）",
                         page_no=6, takeaway="智慧养老硬件 + 订阅市场 2026 年将达 790 亿元，软件订阅占比从 19% 提升至 38%。",
                         categories=["2021", "2022", "2023", "2024", "2025", "2026E"],
                         series={"硬件": [120, 150, 190, 240, 300, 370], "软件订阅": [30, 55, 90, 140, 210, 300],
                                 "服务": [10, 18, 30, 50, 80, 120]}, stacked=True)
P.chart_slide(prs, "用户增长曲线", P.add_line_chart, subtitle="原生折线图（标记 + 数据标签）", page_no=7,
              categories=years, series={"付费家庭数（户）": [2500, 8000, 30000, 65000, 110000],
                                        "活跃设备数（台）": [3000, 9500, 36000, 78000, 130000]})
P.chart_slide(prs, "收入结构（2026E）", P.add_pie_chart, subtitle="原生环形图：逐点上色 + 百分比标签", page_no=8,
              categories=["设备销售", "年度订阅", "数据增值", "政府采购"], values=[0.42, 0.35, 0.08, 0.15],
              takeaway="订阅收入占比 35%，经常性收入提升估值倍数。")
P.chart_slide(prs, "营收与毛利率", P.add_combo_chart, subtitle="组合图：柱（营收，万元）+ 折线（毛利率，次坐标轴）", page_no=9,
              categories=years, bar_series={"营业收入": [480, 980, 1800, 3000, 4500]},
              line_series={"毛利率": [0.33, 0.38, 0.42, 0.45, 0.47]})
P.table_slide(prs, "五年财务预测摘要", ["指标（万元）", "2024A", "2025E", "2026E", "2027E", "2028E"],
              [["营业收入", "480", "980", "1,800", "3,000", "4,500"], ["营业成本", "320", "610", "1,040", "1,650", "2,400"],
               ["毛利", "160", "370", "760", "1,350", "2,100"], ["费用合计", "280", "310", "500", "830", "1,200"],
               ["净利润", "-120", "60", "260", "520", "900"]], col_widths_in=[2.8, 1.8, 1.8, 1.8, 1.8, 1.8],
              subtitle="表格页：深色表头 + 斑马纹 + 末行高亮", page_no=10, highlight_last_row=True)
P.team_slide(prs, "核心团队", [
    {"name": "李华", "role": "CEO / 创始人", "bio": "清华大学生物医学工程硕士，前某头部医疗器械公司产品总监，8 年智能硬件经验。"},
    {"name": "王敏", "role": "CTO", "bio": "毫米波雷达算法专家，发表 SCI 论文 6 篇，主导 3 代产品研发。"},
    {"name": "张强", "role": "COO", "bio": "原连锁养老机构区域总经理，熟悉民政采购与社区渠道。"},
    {"name": "赵云", "role": "CFO", "bio": "注册会计师，曾任职四大会计师事务所，负责两轮融资。"},
], page_no=11)
P.financial_slide(prs, "财务亮点", [("2026E 营业收入", "1,800 万", "同比 +84%"), ("2026E 净利润", "260 万", "净利率 14%"),
                                 ("累计现金流转正", "2027 年", "无需追加融资"), ("本轮融资", "1,500 万", "出让 15% 股权")],
                  P.add_bar_chart, page_no=12, categories=years,
                  series={"营业收入": [480, 980, 1800, 3000, 4500], "净利润": [-120, 60, 260, 520, 900]})
s = P.image_text_slide(prs, "商业模式画布", CANVAS_IMG, ["九宫格由 matplotlib 生成后插入", "也可用形状 + 文本框原生绘制"],
                       subtitle="图文页（右图）", page_no=13, image_left=False)
s = P.bullets_slide(prs, "实施路径", ["形状与“图标”示例：无图标库时用 形状 + 数字徽标 / V 形箭头流程"], page_no=14)
P.add_process_flow(s, Inches(0.8), Inches(3.0), Inches(11.7), ["需求调研", "产品研发", "试点运营", "市场推广", "融资扩张"])
for i, t in enumerate(["3 个社区试点", "6 项专利", "民政采购入围", "3 万户覆盖"]):
    P.add_number_badge(s, Inches(1.0) + Inches(3.0) * i, Inches(4.6), i + 1)
    P.add_text(s, Inches(1.65) + Inches(3.0) * i, Inches(4.62), Inches(2.3), Inches(0.5), t, size=16)
P.closing_slide(prs)

# 复制一页（文本/形状/图片）再删掉它，验证 duplicate/delete
dup = P.duplicate_slide(prs, 4)
P.delete_slide(prs, len(prs.slides._sldIdLst) - 1)

path = os.path.join(OUT, "02_路演示例.pptx")
prs.save(path)
print("saved", path, "slides:", len(prs.slides))

# ---------------- 模板 .pptx 复用与占位符替换 ----------------
tpl = Presentation()  # 用默认模板里的版式当“公司模板”
tpl.slide_width, tpl.slide_height = P.SLIDE_W, P.SLIDE_H
s1 = tpl.slides.add_slide(tpl.slide_layouts[0])  # Title Slide: idx0 标题, idx1 副标题
s1.shapes.title.text = "{{项目名称}}"
s1.placeholders[1].text = "{{副标题}}"
s2 = tpl.slides.add_slide(tpl.slide_layouts[1])  # Title and Content: idx0 标题, idx1 正文
s2.shapes.title.text = "要点"
s2.placeholders[1].text = "待填充"
tpl_path = os.path.join(OUT, "02_模板.pptx")
tpl.save(tpl_path)

prs2 = Presentation(tpl_path)
for s in prs2.slides:
    print("layout:", s.slide_layout.name, [(ph.placeholder_format.idx, ph.placeholder_format.type, ph.name) for ph in s.placeholders])
P.replace_text_everywhere(prs2, {"{{项目名称}}": "智守康联", "{{副标题}}": "居家养老 AI 健康监测平台"})
P.fill_placeholders(prs2.slides[1], {0: "三大核心优势", 1: "无感监测\n3 分钟预警\n双端联动"})
out2 = os.path.join(OUT, "02_模板填充.pptx")
prs2.save(out2)
print("saved", out2)
