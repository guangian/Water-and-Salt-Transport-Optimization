# -*- coding: utf-8 -*-
"""Markdown → docx 渲染器（markdown-it-py 4.2.0 + python-docx 1.2.0 + docx_helpers）。

面向 AI 代理逐章生成：代理只输出 Markdown（含图片引用、表格、题注约定），渲染器负责排版。
支持：# 标题层级（自动多级编号）、段落（首行缩进）、**粗体**/*斜体*/`代码`/[链接]、无序/有序/嵌套列表、
GFM 表格（前一段以“表：”开头即为题注，否则取表格上方 `Table:` 行）、图片 ![题注](路径 "width=12cm")、
> 引用块、分页标记 <!-- pagebreak -->、``` 代码块、--- 忽略。
"""
from __future__ import annotations

import os
import re
from typing import Iterator

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from markdown_it import MarkdownIt
from markdown_it.token import Token

import docx_helpers as H

PAGEBREAK_RE = re.compile(r"^\s*(<!--\s*pagebreak\s*-->|\\pagebreak|---\s*pagebreak\s*---)\s*$", re.I)
WIDTH_RE = re.compile(r"width\s*=\s*([\d.]+)\s*(cm|in)?", re.I)
TABLE_CAPTION_RE = re.compile(r"^(表|Table)\s*[:：]\s*(.+)$")


class MarkdownDocxRenderer:
    def __init__(self, doc: Document | None = None, base_dir: str = ".", heading_offset: int = 0,
                 table_style: str = "three_line", image_width_cm: float = 14.0):
        self.doc = doc or H.setup_chinese_document(Document())
        if doc is None:
            H.add_heading_numbering(self.doc, levels=3)
        self.base_dir = base_dir
        self.heading_offset = heading_offset
        self.table_style = table_style
        self.image_width_cm = image_width_cm
        self.captions = H.CaptionCounter()
        self.footnotes: H.Footnotes | None = None
        self.md = MarkdownIt("commonmark", {"html": True}).enable(["table", "strikethrough"])
        self._pending_table_caption: str | None = None
        self._list_stack: list[tuple[str, int | None]] = []  # (type, num_id)

    # ------------------------------------------------------------ 入口 ----
    def render(self, md_text: str):
        tokens = self.md.parse(md_text)
        it = iter(range(len(tokens)))
        i = 0
        while i < len(tokens):
            i = self._dispatch(tokens, i)
        return self.doc

    def save(self, path: str):
        self.doc.save(path)
        return path

    # -------------------------------------------------------- 块级分发 ----
    def _dispatch(self, tokens: list[Token], i: int) -> int:
        t = tokens[i]
        typ = t.type
        if typ == "heading_open":
            level = int(t.tag[1]) + self.heading_offset
            inline = tokens[i + 1]
            p = self.doc.add_heading("", level=min(level, 9))
            self._inline(p, inline.children or [], east_asia="黑体")
            return i + 3
        if typ == "paragraph_open":
            inline = tokens[i + 1]
            text = inline.content.strip()
            # 分页标记（单独一段）
            if PAGEBREAK_RE.match(text):
                self.doc.add_page_break()
                return i + 3
            m = TABLE_CAPTION_RE.match(text)
            if m and i + 3 < len(tokens) and tokens[i + 3].type == "table_open":
                self._pending_table_caption = m.group(2)
                return i + 3
            # 纯图片段落 → 图 + 题注
            children = inline.children or []
            if children and all(c.type in ("image", "softbreak") for c in children):
                for c in children:
                    if c.type == "image":
                        self._image(c)
                return i + 3
            in_list = bool(self._list_stack)
            if in_list:
                p = self._list_paragraph()
            else:
                p = self.doc.add_paragraph()
            self._inline(p, children)
            return i + 3
        if typ == "html_block":
            if PAGEBREAK_RE.match(t.content.strip()):
                self.doc.add_page_break()
            return i + 1
        if typ == "bullet_list_open":
            self._list_stack.append(("bullet", None))
            return i + 1
        if typ == "ordered_list_open":
            self._list_stack.append(("ordered", None))
            return i + 1
        if typ in ("bullet_list_close", "ordered_list_close"):
            self._list_stack.pop()
            return i + 1
        if typ in ("list_item_open", "list_item_close"):
            return i + 1
        if typ == "blockquote_open":
            j = i + 1
            texts = []
            while tokens[j].type != "blockquote_close":
                if tokens[j].type == "inline":
                    texts.append(tokens[j].content)
                j += 1
            H.add_quote_block(self.doc, "\n".join(texts))
            return j + 1
        if typ == "table_open":
            return self._table(tokens, i)
        if typ == "fence" or typ == "code_block":
            self._code_block(t.content)
            return i + 1
        if typ == "hr":
            return i + 1
        return i + 1

    # ------------------------------------------------------------ 列表 ----
    def _list_paragraph(self):
        depth = len(self._list_stack)
        kind, num_id = self._list_stack[-1]
        suffix = "" if depth == 1 else f" {min(depth, 3)}"
        style = ("List Bullet" if kind == "bullet" else "List Number") + suffix
        p = self.doc.add_paragraph(style=style)
        if kind == "ordered":
            if num_id is None:  # 每个有序列表从 1 重新编号
                num_id = H.restart_list_numbering(self.doc, p, style)
                self._list_stack[-1] = (kind, num_id)
            else:
                H.continue_list_numbering(p, num_id)
        return p

    # ------------------------------------------------------------ 行内 ----
    def _inline(self, p, children: list[Token], east_asia="宋体"):
        bold = italic = code = strike = False
        link: str | None = None
        for c in children:
            if c.type == "text":
                if link:
                    H.add_hyperlink(p, c.content, url=link, east_asia=east_asia)
                else:
                    r = p.add_run(c.content)
                    H.set_run_font(r, "等线" if code else east_asia, "Consolas" if code else "Times New Roman",
                                   bold=bold or None, italic=italic or None)
                    if strike:
                        r.font.strike = True
            elif c.type == "strong_open":
                bold = True
            elif c.type == "strong_close":
                bold = False
            elif c.type == "em_open":
                italic = True
            elif c.type == "em_close":
                italic = False
            elif c.type == "s_open":
                strike = True
            elif c.type == "s_close":
                strike = False
            elif c.type == "code_inline":
                r = p.add_run(c.content)
                H.set_run_font(r, "等线", "Consolas", size=10.5)
            elif c.type == "link_open":
                link = c.attrs.get("href")
            elif c.type == "link_close":
                link = None
            elif c.type in ("softbreak",):
                p.add_run(" ")
            elif c.type == "hardbreak":
                p.add_run().add_break()
            elif c.type == "image":
                self._image(c)
            elif c.type == "html_inline":
                if "<br" in c.content:
                    p.add_run().add_break()

    # ------------------------------------------------------------ 图片 ----
    def _image(self, tok: Token):
        from urllib.parse import unquote
        src = unquote(tok.attrs.get("src", ""))  # markdown-it 会把中文路径百分号编码，必须解码
        path = src if os.path.isabs(src) else os.path.join(self.base_dir, src)
        alt = "".join(c.content for c in (tok.children or [])) or tok.content
        width = self.image_width_cm
        title = tok.attrs.get("title") or ""
        m = WIDTH_RE.search(title)
        if m:
            width = float(m.group(1)) * (2.54 if (m.group(2) or "cm").lower() == "in" else 1)
        if not os.path.exists(path):
            p = self.doc.add_paragraph()
            H.set_run_font(p.add_run(f"[缺失图片：{src}]"), color="C00000")
            return
        H.add_figure(self.doc, path, alt, self.captions, width_cm=width)

    # ------------------------------------------------------------ 表格 ----
    def _table(self, tokens: list[Token], i: int) -> int:
        rows: list[list[str]] = []
        aligns: list[str | None] = []
        j = i + 1
        cur: list[str] | None = None
        in_head = False
        while tokens[j].type != "table_close":
            tt = tokens[j]
            if tt.type == "thead_open":
                in_head = True
            elif tt.type == "thead_close":
                in_head = False
            elif tt.type == "tr_open":
                cur = []
            elif tt.type == "tr_close":
                rows.append(cur or [])
            elif tt.type in ("th_open", "td_open"):
                if in_head:
                    style = tt.attrs.get("style", "")
                    aligns.append(style.split(":")[-1].strip() if style else None)
            elif tt.type == "inline" and cur is not None:
                cur.append(tt.content)
            j += 1
        if not rows:
            return j + 1
        header, body = rows[0], rows[1:]
        if self._pending_table_caption:
            H.add_caption(self.doc, "表", self._pending_table_caption, self.captions)
            self._pending_table_caption = None
        if self.table_style == "grid":
            H.add_grid_table(self.doc, header, body)
        else:
            H.add_three_line_table(self.doc, header, body, header_fill="D9E2F3")
        self.doc.add_paragraph()  # 表后空行，避免与下文粘连
        return j + 1

    # ---------------------------------------------------------- 代码块 ----
    def _code_block(self, content: str):
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        p._p.get_or_add_pPr().get_or_add_ind().set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}firstLineChars", "0")
        H.insert_ordered(p._p.get_or_add_pPr(), parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="F5F5F5"/>'))
        lines = content.rstrip("\n").split("\n")
        for k, line in enumerate(lines):
            r = p.add_run(line)
            H.set_run_font(r, "等线", "Consolas", size=9)
            if k < len(lines) - 1:
                r.add_break()


# ---------------------------------------------------------- 组装流水线 ----

def build_document(chapters_md: list[str], out_path: str, base_dir: str, title: str, subtitle: str,
                   meta: list[tuple[str, str]], header_text: str, with_toc: bool = True):
    """代理逐章产出的 Markdown 列表 → 完整 docx（封面 + 目录 + 页眉页脚 + 正文）。"""
    doc = H.setup_chinese_document(Document())
    H.add_heading_numbering(doc, levels=3)
    H.add_cover_page(doc, title, subtitle, meta=meta, new_section_after=False)
    if with_toc:
        doc.add_page_break()
        H.add_toc(doc, levels="1-3")
    sec = doc.add_section()
    H.set_header_text(sec, header_text)
    H.add_page_number_footer(sec, "第 {PAGE} 页 / 共 {NUMPAGES} 页")
    H.restart_page_numbering(sec, start=1)
    renderer = MarkdownDocxRenderer(doc, base_dir=base_dir)
    for md in chapters_md:
        renderer.render(md)
    renderer.save(out_path)
    return out_path
