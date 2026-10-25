"""book.json → 专业 Word 文档（契约 §7.2）。

用法：
    python -m bizkit.docx_render 05_正文/book.json [--no-pdf]

book.json 字段（路径相对工作区根）：
    cover{title, subtitle, teamName, company, group, date, extraLines[]}、abstract、chapters[{file, numbered}]、
    appendices[{file}]、references（参考文献.json）、style{bodyFont, bodySize, headingFont, latinFont, lineSpacing,
    firstLineIndent, pageNumbers, toc, tocDepth, header, tableStyle: 'banded'|'three-line', h1Align: 'center'|'left'}、
    anonymize{forbidden[], whitelist[]}、output、pdf。

支持的 Markdown 子集：# ## ### 标题（自动编号 1 / 1.1 / 1.1.1，附录 A / A.1）、段落、**粗体**、有序 / 无序列表（可缩进嵌套）、
GFM 表格（前一行 `表：表题` → “表 N-M 表题”）、图片行 `![图题](../04_图表/x.png)` → 居中图 + “图 N-M 图题”、`> 引用`、
`---` / `<!-- pagebreak -->` 分页、`[^n]` 脚注（简化为章末注）、`{来源：…}` 小号灰色来源行、`[12]` 文献引用（上标）。

目录：写入 TOC 域（Word 打开时 updateFields 自动刷新），域结果里同时放一份静态目录作为占位；
若本机能转 PDF，会用 PDF 实际页码回填静态目录再重渲染一次，保证 LibreOffice 等不更新域的阅读端也显示正确页码。

输出 JSON：output、pdf、chapters、figures、tables、words、pagesEstimate、pages、anonymizeHits、missingImages、warnings。
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import fonts, theme as theme_mod
from ._util import find_workspace_root, force_utf8_stdout, read_json, read_text

CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"
APPENDIX_LETTERS = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
_CJK_RE = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]")
_SOURCE_RE = re.compile(r"\{\s*((?:资料|数据|图片)?来源)\s*[:：]\s*([^{}]*?)\s*\}")
_CITATION_RE = r"\[(\d{1,3}(?:\s*[,，\-–]\s*\d{1,3})*)\](?!\()"
_INLINE_RE = re.compile(
    r"(?P<bold>\*\*(?P<bold_t>.+?)\*\*)"
    r"|(?P<code>`(?P<code_t>[^`]+)`)"
    r"|(?P<fn>\[\^(?P<fn_k>[^\]\s]+)\])"
    r"|(?P<cite>" + _CITATION_RE + r")"
    r"|(?P<link>\[(?P<link_t>[^\]]+)\]\((?P<link_u>[^)\s]+)(?:\s+\"[^\"]*\")?\))"
    r"|(?P<br><br\s*/?>)"
)


# ----------------------------------------------------------------------------------------------
# Markdown 子集解析
# ----------------------------------------------------------------------------------------------

@dataclass
class Heading:
    level: int
    text: str
    line: int


@dataclass
class Para:
    text: str
    sources: List[str] = field(default_factory=list)


@dataclass
class ListItem:
    level: int
    text: str


@dataclass
class ListBlock:
    ordered: bool
    items: List[ListItem]


@dataclass
class Table:
    header: List[str]
    rows: List[List[str]]
    aligns: List[str]
    caption: str = ""
    sources: List[str] = field(default_factory=list)


@dataclass
class Figure:
    caption: str
    path: str
    width_hint: str = ""
    sources: List[str] = field(default_factory=list)


@dataclass
class Quote:
    text: str


@dataclass
class PageBreak:
    pass


@dataclass
class SourceLine:
    text: str


Block = Any


@dataclass
class ParsedDoc:
    blocks: List[Block]
    footnotes: Dict[str, str]


_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_HR_RE = re.compile(r"^(?:-{3,}|\*{3,}|_{3,})\s*$")
_IMG_RE = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<path>[^)\s]+)(?:\s+\"(?P<title>[^\"]*)\")?\)\s*(?P<rest>.*)$")
_TABLE_CAP_RE = re.compile(r"^(?:表|Table)\s*[:：]\s*(.+?)\s*$")
_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)*\|?\s*$")
_UL_RE = re.compile(r"^(?P<ind>\s*)[-*+]\s+(?P<t>.*)$")
_OL_RE = re.compile(r"^(?P<ind>\s*)(?P<n>\d+)[.)]\s+(?P<t>.*)$")
_FN_DEF_RE = re.compile(r"^\[\^(?P<k>[^\]\s]+)\]:\s*(?P<t>.*)$")
_PAGEBREAK_RE = re.compile(r"^<!--\s*pagebreak\s*-->\s*$", re.I)
_COMMENT_RE = re.compile(r"<!--.*?-->", re.S)
# 只认 1–2 位的节号（1 / 2.3 / 4.1.2），避免把“2026 年目标”里的年份当编号剥掉
_NUM_PREFIX_RE = re.compile(
    r"^(?:第\s*[一二三四五六七八九十百\d]+\s*(?:部分|[章节部篇])\s*[:：、.．]?\s*"
    r"|\d{1,2}(?:\.\d{1,2}){0,3}\.?\s+"
    r"|[A-Z](?:\.\d{1,2}){1,3}\.?\s+"
    r"|附录\s*[A-Z](?:\s*\d{1,2}(?:\.\d{1,2})*)?\s*[:：、.．]?\s*)"
)
# 作者在题注里手写的“图 2-1 / 表1： / Figure 3.”编号（单位词开头的“图 12 个月…”不算）
_CAP_PREFIX_RE = re.compile(
    r"^(?:图|表|Figure|Fig\.?|Table)\s*(?:[A-Z]|\d{1,2})(?:\s*[-–—.]\s*\d{1,2})?\s*[:：.．、]?\s*(?![个年月日家万亿元%])"
)
# 紧跟表格 / 图片之后的“来源：…”“注：…”整段（作者没写花括号语法）视作来源行
_PLAIN_SOURCE_RE = re.compile(r"^(?:(?:资料|数据|图片)?来源|注|备注)\s*[:：]")


def _strip_numbering(title: str) -> str:
    """作者若已手写“1 引言 / 第一章 引言 / 附录 A 问卷”，去掉以免与自动编号重复。"""
    return _NUM_PREFIX_RE.sub("", title).strip() or title


def _strip_caption_prefix(caption: str) -> str:
    """题注里已带“图 2-1 / 表 3：”编号时去掉，由渲染器统一编号，避免“图 2-1 图 2-1 …”。"""
    return _CAP_PREFIX_RE.sub("", caption or "", count=1).strip()


def _attach_plain_sources(blocks: List["Block"]) -> List["Block"]:
    """把表格 / 图片后面紧跟的“来源：…”普通段落并入该表 / 图的来源行（排成小号灰字，不再按正文缩进排）。"""
    out: List[Block] = []
    for b in blocks:
        prev = out[-1] if out else None
        if (isinstance(b, Para) and isinstance(prev, (Table, Figure)) and _PLAIN_SOURCE_RE.match(b.text)
                and len(b.text) <= 160):
            prev.sources.append(b.text)
            prev.sources.extend(b.sources)
            continue
        out.append(b)
    return out


def _split_cells(line: str) -> List[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|") and not s.endswith("\\|"):
        s = s[:-1]
    parts = re.split(r"(?<!\\)\|", s)
    return [p.replace("\\|", "|").strip() for p in parts]


def _join_lines(lines: List[str]) -> str:
    """中文行间不加空格，西文行间加空格。"""
    out = ""
    for ln in lines:
        ln = ln.strip()
        if not out:
            out = ln
            continue
        if out and ln and _CJK_RE.match(out[-1]) and _CJK_RE.match(ln[0]):
            out += ln
        else:
            out += " " + ln
    return out


def _extract_sources(text: str) -> Tuple[str, List[str]]:
    sources: List[str] = []

    def _grab(m: "re.Match[str]") -> str:
        sources.append(f"{m.group(1)}：{m.group(2).strip()}")
        return ""

    cleaned = _SOURCE_RE.sub(_grab, text).strip()
    return cleaned, sources


def parse_markdown(text: str) -> ParsedDoc:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # 多行 HTML 注释：分页标记保留，其余删除
    text = _COMMENT_RE.sub(lambda m: "<!-- pagebreak -->" if re.search(r"pagebreak", m.group(0), re.I) else "", text)
    lines = text.split("\n")
    blocks: List[Block] = []
    footnotes: Dict[str, str] = {}
    pending_caption = ""
    para_buf: List[str] = []
    n = len(lines)
    i = 0

    def flush_para() -> None:
        nonlocal para_buf
        if para_buf:
            joined = _join_lines(para_buf)
            cleaned, srcs = _extract_sources(joined)
            if cleaned:
                blocks.append(Para(cleaned, srcs))
            else:
                blocks.extend(SourceLine(s) for s in srcs)
            para_buf = []

    def _is_block_start(ln: str) -> bool:
        s = ln.strip()
        return bool(_HEADING_RE.match(s) or _HR_RE.match(s) or _IMG_RE.match(s) or _UL_RE.match(ln) or _OL_RE.match(ln)
                    or s.startswith(">") or _PAGEBREAK_RE.match(s) or _FN_DEF_RE.match(s) or _TABLE_CAP_RE.match(s)
                    or (s.startswith("|") and i + 1 < n and _TABLE_SEP_RE.match(lines[i + 1] if i + 1 < n else "")))

    while i < n:
        line = lines[i]
        s = line.strip()
        if not s:
            flush_para()
            i += 1
            continue
        if _PAGEBREAK_RE.match(s):
            flush_para()
            blocks.append(PageBreak())
            i += 1
            continue
        m = _HEADING_RE.match(s)
        if m:
            flush_para()
            blocks.append(Heading(len(m.group(1)), m.group(2).strip(), i + 1))
            i += 1
            continue
        if _HR_RE.match(s):
            flush_para()
            blocks.append(PageBreak())
            i += 1
            continue
        m = _IMG_RE.match(s)
        if m:
            flush_para()
            alt = m.group("alt").strip()
            width_hint = ""
            if "|" in alt:
                alt, width_hint = [x.strip() for x in alt.rsplit("|", 1)]
            rest, srcs = _extract_sources(m.group("rest") or "")
            # 图片下一行若是 {来源：…} 也归到该图
            if i + 1 < n and _SOURCE_RE.fullmatch(lines[i + 1].strip() or "x"):
                _, more = _extract_sources(lines[i + 1].strip())
                srcs.extend(more)
                i += 1
            blocks.append(Figure(alt or (m.group("title") or ""), m.group("path"), width_hint, srcs))
            i += 1
            continue
        m = _TABLE_CAP_RE.match(s)
        if m:
            # 题注与表格之间允许空行（作者习惯不一），只要下一个非空行是表格就归为题注
            j = i + 1
            while j < n and not lines[j].strip():
                j += 1
            if j < n and lines[j].strip().startswith("|"):
                flush_para()
                pending_caption = m.group(1)
                i = j
                continue
        if s.startswith("|") and i + 1 < n and _TABLE_SEP_RE.match(lines[i + 1]):
            flush_para()
            header = _split_cells(s)
            seps = _split_cells(lines[i + 1])
            aligns = []
            for sp in seps:
                sp = sp.strip()
                if sp.startswith(":") and sp.endswith(":"):
                    aligns.append("center")
                elif sp.endswith(":"):
                    aligns.append("right")
                elif sp.startswith(":"):
                    aligns.append("left")
                else:
                    aligns.append("")
            rows: List[List[str]] = []
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                cells = _split_cells(lines[i])
                cells = (cells + [""] * len(header))[: len(header)]
                rows.append(cells)
                i += 1
            srcs: List[str] = []
            if i < n and _SOURCE_RE.fullmatch(lines[i].strip() or "x"):
                _, srcs = _extract_sources(lines[i].strip())
                i += 1
            blocks.append(Table(header, rows, (aligns + [""] * len(header))[: len(header)], pending_caption, srcs))
            pending_caption = ""
            continue
        if s.startswith(">"):
            flush_para()
            q: List[str] = []
            while i < n and lines[i].strip().startswith(">"):
                q.append(lines[i].strip()[1:].strip())
                i += 1
            blocks.append(Quote(_join_lines([x for x in q if x])))
            continue
        m = _FN_DEF_RE.match(s)
        if m:
            flush_para()
            key, body = m.group("k"), m.group("t")
            i += 1
            while i < n and lines[i].startswith((" ", "\t")) and lines[i].strip():
                body += lines[i].strip()
                i += 1
            footnotes[key] = body.strip()
            continue
        mu, mo = _UL_RE.match(line), _OL_RE.match(line)
        if mu or mo:
            flush_para()
            ordered = bool(mo)
            items: List[ListItem] = []
            while i < n:
                ln = lines[i]
                mm = (_OL_RE.match(ln) if ordered else _UL_RE.match(ln))
                other = (_UL_RE.match(ln) if ordered else _OL_RE.match(ln))
                if mm:
                    ind = mm.group("ind").replace("\t", "    ")
                    items.append(ListItem(min(len(ind) // 2, 3), mm.group("t").strip()))
                    i += 1
                elif other and other.group("ind"):
                    # 嵌套的另一种列表类型：作为子项并入（简化处理）
                    ind = other.group("ind").replace("\t", "    ")
                    items.append(ListItem(min(max(len(ind) // 2, 1), 3), other.group("t").strip()))
                    i += 1
                elif ln.strip() and ln.startswith((" ", "\t")) and items and not _is_block_start(ln):
                    items[-1].text += ("" if _CJK_RE.match(items[-1].text[-1:] or "a") else " ") + ln.strip()
                    i += 1
                else:
                    break
            blocks.append(ListBlock(ordered, items))
            continue
        if _SOURCE_RE.fullmatch(s):
            flush_para()
            _, srcs = _extract_sources(s)
            blocks.extend(SourceLine(x) for x in srcs)
            i += 1
            continue
        para_buf.append(line)
        i += 1
    flush_para()
    return ParsedDoc(blocks, footnotes)


# ----------------------------------------------------------------------------------------------
# 匿名清洗 / 字数
# ----------------------------------------------------------------------------------------------

def anonymize_text(text: str, forbidden: Sequence[str], whitelist: Sequence[str], where: str) -> Tuple[str, List[Dict[str, Any]]]:
    """forbidden 词替换为「XX」，whitelist 中的词先用占位符保护。返回 (新文本, 命中列表[{file, line, term, count}])。"""
    terms = [t for t in forbidden if t]
    if not terms:
        return text, []
    hits: List[Dict[str, Any]] = []
    protected: Dict[str, str] = {}
    for k, w in enumerate(w for w in whitelist if w):
        token = f"\ue000{k}\ue001"
        protected[token] = w
        text = text.replace(w, token)
    out_lines = []
    for ln_no, ln in enumerate(text.split("\n"), 1):
        for t in sorted(terms, key=len, reverse=True):
            c = ln.count(t)
            if c:
                hits.append({"file": where, "line": ln_no, "term": t, "count": c})
                ln = ln.replace(t, "XX")
        out_lines.append(ln)
    text = "\n".join(out_lines)
    for token, w in protected.items():
        text = text.replace(token, w)
    return text, hits


def count_words(text: str) -> int:
    """中文按字计、西文按词计、数字串算一个词；忽略 Markdown 符号。"""
    plain = re.sub(r"[#*`>|\-\[\]()!{}]", " ", text)
    cjk = len(_CJK_RE.findall(plain))
    latin = len(re.findall(r"[A-Za-z0-9][A-Za-z0-9.%,]*", plain))
    return cjk + latin


# ----------------------------------------------------------------------------------------------
# 参考文献（GB/T 7714-2015 顺序编码制）
# ----------------------------------------------------------------------------------------------

_TYPE_CODE = {
    "report": "R", "annual-report": "R", "web": "EB", "paper": "J", "book": "M", "gov": "A", "news": "N",
    "standard": "S", "thesis": "D", "database": "DB", "patent": "P", "conference": "C",
}


def _clean_end(s: str) -> str:
    return s.rstrip(" .。，,;；")


def format_reference(ref: Dict[str, Any]) -> str:
    """按 type 生成著录文本（不含 [n] 序号）。缺字段时尽量降级而不是报错。"""
    rtype = str(ref.get("type") or "web").lower()
    code = _TYPE_CODE.get(rtype, "Z")
    url = str(ref.get("url") or "").strip()
    online = bool(url)
    derived_author = not ref.get("author")   # 机构出版物常无作者字段，用出版者顶替后不再在出版项重复
    author = _clean_end(str(ref.get("author") or ref.get("publisher") or "佚名"))
    title = _clean_end(str(ref.get("title") or "无题"))
    publisher = _clean_end(str(ref.get("publisher") or ""))
    place = _clean_end(str(ref.get("place") or ""))
    year = str(ref.get("year") or "").strip()
    pub_date = str(ref.get("date") or "").strip()
    accessed = str(ref.get("accessed") or "").strip()
    marker = f"[{code}/OL]" if online else f"[{code}]"
    if rtype == "paper":
        journal = _clean_end(str(ref.get("journal") or publisher))
        vol = str(ref.get("volume") or "")
        issue = str(ref.get("issue") or "")
        pages = str(ref.get("pages") or "")
        tail = journal
        if year:
            tail += f", {year}"
        if vol:
            tail += f", {vol}"
        if issue:
            tail += f"({issue})"
        if pages:
            tail += f": {pages}"
        text = f"{author}. {title}{marker}. {tail}."
    elif rtype in ("book", "standard", "thesis"):
        if rtype == "standard" and ref.get("number"):
            title = f"{title}: {ref['number']}"
        loc = ": ".join(x for x in (place, publisher) if x and not (derived_author and x == author))
        tail = ", ".join(x for x in (loc, year) if x)
        text = f"{author}. {title}{marker}. {tail}."
    elif rtype == "news":
        paper = _clean_end(str(ref.get("newspaper") or publisher))
        edition = str(ref.get("edition") or "")
        d = pub_date or year
        text = f"{author}. {title}{marker}. {paper}, {d}" + (f"({edition})" if edition else "") + "."
    else:
        if rtype == "gov" and ref.get("number"):
            title = f"{title}: {ref['number']}"
        parts = [f"{author}. {title}{marker}."]
        loc = ": ".join(x for x in (place, publisher) if x and x != author)
        d = f"({pub_date})" if pub_date else (f"({year})" if year and online else year)
        if loc:
            parts.append(f"{loc}" + (f", {d}" if d and not online else ""))
        if online:
            dates = (f"({pub_date})" if pub_date else (f"({year})" if year else "")) + (f"[{accessed}]" if accessed else "")
            parts.append(dates + ".")
        elif not loc and year:
            parts.append(f"{year}.")
        text = " ".join(p for p in parts if p).replace("..", ".")
        if loc and not online and not text.endswith("."):
            text += "."
    if online:
        text = text.rstrip(".") + ". " + url + "."
    note = str(ref.get("note") or "").strip()
    if note:
        text += f" {note}"
    return re.sub(r"\s+", " ", text).replace(" .", ".").strip()


# ----------------------------------------------------------------------------------------------
# docx 底层帮助函数
# ----------------------------------------------------------------------------------------------

def _set_rfonts(rpr_parent, cjk: str, latin: str) -> None:
    """同时写 ascii/hAnsi（西文）与 eastAsia（中文），并清掉主题字体属性，否则主题字体会覆盖显式字体。"""
    rpr = rpr_parent.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), cjk)
    rfonts.set(qn("w:cs"), latin)


def _style_font(style, cjk: str, latin: str, size_pt: float, bold: Optional[bool] = None, color: Optional[str] = None) -> None:
    _set_rfonts(style.element, cjk, latin)
    style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    for tag in ("w:color",):
        el = rpr.find(qn(tag))
        if el is not None and el.get(qn("w:themeColor")) is not None:
            del el.attrib[qn("w:themeColor")]
            for extra in ("w:themeShade", "w:themeTint"):
                if el.get(qn(extra)) is not None:
                    del el.attrib[qn(extra)]


def _ppr_child(paragraph, tag: str) -> Any:
    ppr = paragraph._p.get_or_add_pPr()
    el = ppr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        ppr.append(el)
    return el


def _set_first_line_chars(paragraph_or_style, chars: int) -> None:
    """首行缩进按“字符”计（w:firstLineChars），Word 随字号自适应；同时写 firstLine 兜底给不认识 chars 的阅读端。"""
    if hasattr(paragraph_or_style, "_p"):
        ppr = paragraph_or_style._p.get_or_add_pPr()
    else:
        ppr = paragraph_or_style.element.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:firstLineChars"), str(chars * 100))
    ind.set(qn("w:firstLine"), str(int(chars * 12 * 20)))


def _set_char_spacing(run, pt: float) -> None:
    """字符间距（w:spacing，单位 1/20 pt）：封面文种标签用拉开的字距做弱化处理。"""
    rpr = run._r.get_or_add_rPr()
    for old in rpr.findall(qn("w:spacing")):
        rpr.remove(old)
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:val"), str(int(pt * 20)))
    rpr.append(sp)


def _add_field(paragraph, instr: str, cached: str = "", dirty: bool = False) -> None:
    """复杂域：begin / instrText / separate / 缓存结果 / end。"""
    r1 = paragraph.add_run()
    fc1 = OxmlElement("w:fldChar")
    fc1.set(qn("w:fldCharType"), "begin")
    if dirty:
        fc1.set(qn("w:dirty"), "true")
    r1._r.append(fc1)
    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r2._r.append(it)
    r3 = paragraph.add_run()
    fc3 = OxmlElement("w:fldChar")
    fc3.set(qn("w:fldCharType"), "separate")
    r3._r.append(fc3)
    if cached:
        paragraph.add_run(cached)
    r5 = paragraph.add_run()
    fc5 = OxmlElement("w:fldChar")
    fc5.set(qn("w:fldCharType"), "end")
    r5._r.append(fc5)


def _fld_char(paragraph, kind: str, dirty: bool = False) -> None:
    r = paragraph.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), kind)
    if dirty:
        fc.set(qn("w:dirty"), "true")
    r._r.append(fc)


def _instr_text(paragraph, instr: str) -> None:
    r = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)


def _set_cell_shading(cell, fill_hex: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)


def _set_cell_borders(cell, **edges: Tuple[str, int, str]) -> None:
    """edges: top=(val, sz_eighths, color) …"""
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        if edge in edges:
            val, sz, color = edges[edge]
            el = borders.find(qn(f"w:{edge}"))
            if el is None:
                el = OxmlElement(f"w:{edge}")
                borders.append(el)
            el.set(qn("w:val"), val)
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)


def _set_table_borders(table, spec: Dict[str, Tuple[str, int, str]]) -> None:
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        val, sz, color = spec.get(edge, ("nil", 0, "auto"))
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def _set_table_cell_margins(table, top: int = 30, bottom: int = 30, left: int = 80, right: int = 80) -> None:
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)


def _row_header_repeat(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    trpr.append(el)


def _row_cant_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trpr.append(el)


def _mark_section_break_paragraph(doc) -> None:
    """add_section 产生的空段落尽量不占空间，避免上一节正好满页时多出一张空白页。"""
    p = doc.paragraphs[-1]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)


def _set_pg_num_type(section, fmt: str, start: Optional[int]) -> None:
    sectpr = section._sectPr
    el = sectpr.find(qn("w:pgNumType"))
    if el is None:
        el = OxmlElement("w:pgNumType")
        sectpr.insert_element_before(el, "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg", "w:textDirection",
                                     "w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings", "w:sectPrChange")
    el.set(qn("w:fmt"), fmt)
    if start is not None:
        el.set(qn("w:start"), str(start))


def _enable_update_fields(doc) -> None:
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is not None:
        return
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    compat = settings.find(qn("w:compat"))
    if compat is not None:
        compat.addprevious(el)
    else:
        settings.append(el)


def _set_doc_defaults(doc, cjk: str, latin: str, size_pt: float) -> None:
    styles_el = doc.styles.element
    rpr_default = styles_el.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    rfonts = rpr_default.find(qn("w:rFonts"))
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), cjk)
    rfonts.set(qn("w:cs"), latin)
    for tag in ("w:sz", "w:szCs"):
        el = rpr_default.find(qn(tag))
        if el is not None:
            el.set(qn("w:val"), str(int(size_pt * 2)))
    lang = rpr_default.find(qn("w:lang"))
    if lang is not None:
        lang.set(qn("w:eastAsia"), "zh-CN")
    ppr_default = styles_el.find(qn("w:docDefaults")).find(qn("w:pPrDefault")).find(qn("w:pPr"))
    sp = ppr_default.find(qn("w:spacing"))
    if sp is not None:
        sp.set(qn("w:after"), "0")
        sp.set(qn("w:line"), "240")
        sp.set(qn("w:lineRule"), "auto")


# ----------------------------------------------------------------------------------------------
# 渲染器
# ----------------------------------------------------------------------------------------------

@dataclass
class Numbering:
    label: str = ""          # 章号 "1" / 附录 "A" / 无编号 ""
    numbered: bool = True
    h2: int = 0
    h3: int = 0
    fig: int = 0
    tab: int = 0

    def fig_label(self) -> str:
        self.fig += 1
        return f"图 {self.label}-{self.fig}" if self.label else f"图 {self.fig}"

    def tab_label(self) -> str:
        self.tab += 1
        return f"表 {self.label}-{self.tab}" if self.label else f"表 {self.tab}"


@dataclass
class TocEntry:
    level: int
    text: str
    page: str = ""
    section: str = "body"   # front | body


@dataclass
class RenderStats:
    chapters: int = 0
    figures: int = 0
    tables: int = 0
    table_rows: int = 0
    words: int = 0
    headings: int = 0
    missing_images: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    anonymize_hits: List[Dict[str, Any]] = field(default_factory=list)


class BookRenderer:
    def __init__(self, book: Dict[str, Any], book_path: Path, toc_pages: Optional[Dict[str, str]] = None):
        self.book = book
        self.book_path = book_path
        probe = None
        for key in ("abstract",):
            if book.get(key):
                probe = book[key]
        if not probe and book.get("chapters"):
            probe = (book["chapters"][0] or {}).get("file")
        self.root = find_workspace_root(book_path.parent, probe=probe)
        self.style_cfg = dict(book.get("style") or {})
        self.th = theme_mod.get_theme("bizgold")
        df = fonts.doc_fonts(self.style_cfg.get("bodyFont"), self.style_cfg.get("headingFont"), self.style_cfg.get("latinFont"))
        self.cjk_body, self.cjk_head, self.latin = df["body"], df["heading"], df["latin"]
        self.body_size = float(self.style_cfg.get("bodySize") or 12)
        self.line_spacing = float(self.style_cfg.get("lineSpacing") or 1.5)
        self.first_indent = int(self.style_cfg.get("firstLineIndent", 2))
        self.table_style = str(self.style_cfg.get("tableStyle") or "banded")
        self.toc_depth = int(self.style_cfg.get("tocDepth") or 2)
        self.toc_pages = toc_pages or {}
        anon = book.get("anonymize") or {}
        self.forbidden = list(anon.get("forbidden") or [])
        self.whitelist = list(anon.get("whitelist") or [])
        self.stats = RenderStats()
        self.toc: List[TocEntry] = []
        self.doc = Document()
        self.text_width_cm = 0.0
        self.chapter_no = 0
        self.est_pos = 0.0          # 正文已消耗的估算页数（静态目录页码用）
        self._fn_order: List[str] = []
        self._footnotes: Dict[str, str] = {}
        self._toc_anchor = None

    # ---- 文本处理 ----
    def _anon(self, text: str, where: str) -> str:
        new, hits = anonymize_text(text, self.forbidden, self.whitelist, where)
        self.stats.anonymize_hits.extend(hits)
        return new

    def _rel(self, p: str) -> Path:
        return (self.root / p).resolve()

    # ---- 页面与样式 ----
    def setup(self) -> None:
        doc = self.doc
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        lr = float(self.style_cfg.get("marginLR") or 3.17)
        tb = float(self.style_cfg.get("marginTB") or 2.54)
        sec.left_margin = sec.right_margin = Cm(lr)
        sec.top_margin = sec.bottom_margin = Cm(tb)
        sec.header_distance = Cm(1.5)
        sec.footer_distance = Cm(1.5)
        self.text_width_cm = 21.0 - 2 * lr
        _set_doc_defaults(doc, self.cjk_body, self.latin, self.body_size)
        styles = doc.styles

        normal = styles["Normal"]
        _style_font(normal, self.cjk_body, self.latin, self.body_size, color=self.th.ink)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.line_spacing = self.line_spacing

        body = styles.add_style("BizBody", WD_STYLE_TYPE.PARAGRAPH)
        body.base_style = normal
        body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        body.paragraph_format.space_after = Pt(0)
        if self.first_indent > 0:
            _set_first_line_chars(body, self.first_indent)
        self.s_body = body

        h1_align = WD_ALIGN_PARAGRAPH.CENTER if str(self.style_cfg.get("h1Align") or "center") == "center" else WD_ALIGN_PARAGRAPH.LEFT
        for name, size, bold, before, after, align in (
            ("Heading 1", self.th.typo.doc_h1, True, 0, 18, h1_align),
            ("Heading 2", self.th.typo.doc_h2, True, 14, 8, WD_ALIGN_PARAGRAPH.LEFT),
            ("Heading 3", self.th.typo.doc_h3, True, 10, 6, WD_ALIGN_PARAGRAPH.LEFT),
        ):
            st = styles[name]
            _style_font(st, self.cjk_head, self.latin, size, bold=bold, color=self.th.ink)
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(after)
            st.paragraph_format.line_spacing = 1.5
            st.paragraph_format.keep_with_next = True
            st.paragraph_format.alignment = align
            st.paragraph_format.first_line_indent = Cm(0)

        cap = styles["Caption"]
        _style_font(cap, self.cjk_body, self.latin, self.th.typo.doc_small, bold=True, color=self.th.ink)
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(4)
        cap.paragraph_format.space_after = Pt(8)
        cap.paragraph_format.line_spacing = 1.25
        cap.font.italic = False
        self.s_caption = cap

        src = styles.add_style("BizSource", WD_STYLE_TYPE.PARAGRAPH)
        src.base_style = normal
        _style_font(src, self.cjk_body, self.latin, 9, color=self.th.muted)
        src.paragraph_format.space_after = Pt(6)
        src.paragraph_format.line_spacing = 1.15
        self.s_source = src

        tbl = styles.add_style("BizTable", WD_STYLE_TYPE.PARAGRAPH)
        tbl.base_style = normal
        _style_font(tbl, self.cjk_body, self.latin, self.th.typo.doc_small, color=self.th.ink)
        tbl.paragraph_format.line_spacing = 1.15
        tbl.paragraph_format.space_after = Pt(0)
        self.s_table = tbl

        lst = styles.add_style("BizList", WD_STYLE_TYPE.PARAGRAPH)
        lst.base_style = normal
        lst.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        lst.paragraph_format.space_after = Pt(2)
        self.s_list = lst

        quote = styles.add_style("BizQuote", WD_STYLE_TYPE.PARAGRAPH)
        quote.base_style = normal
        _style_font(quote, self.cjk_body, self.latin, self.body_size - 1, color=self.th.secondary)
        quote.paragraph_format.left_indent = Cm(0.9)
        quote.paragraph_format.right_indent = Cm(0.5)
        quote.paragraph_format.space_before = Pt(6)
        quote.paragraph_format.space_after = Pt(6)
        quote.paragraph_format.line_spacing = 1.3
        ppr = quote.element.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), self.th.accent)
        pbdr.append(left)
        ppr.append(pbdr)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), self.th.paper_alt)
        ppr.append(shd)
        self.s_quote = quote

        ref = styles.add_style("BizReference", WD_STYLE_TYPE.PARAGRAPH)
        ref.base_style = normal
        _style_font(ref, self.cjk_body, self.latin, self.th.typo.doc_small, color=self.th.ink)
        ref.paragraph_format.left_indent = Cm(0.9)
        ref.paragraph_format.first_line_indent = Cm(-0.9)
        ref.paragraph_format.space_after = Pt(3)
        ref.paragraph_format.line_spacing = 1.25
        self.s_ref = ref

        note = styles.add_style("BizNote", WD_STYLE_TYPE.PARAGRAPH)
        note.base_style = normal
        _style_font(note, self.cjk_body, self.latin, 9.5, color=self.th.muted)
        note.paragraph_format.space_after = Pt(2)
        note.paragraph_format.line_spacing = 1.15
        self.s_note = note

        self.s_toc = []
        for lvl in (1, 2, 3):
            st = styles.add_style(f"toc {lvl}", WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = normal
            _style_font(st, self.cjk_body if lvl > 1 else self.cjk_head, self.latin, self.body_size if lvl == 1 else self.body_size - 1, bold=False, color=self.th.ink)
            st.paragraph_format.left_indent = Cm(0.75 * (lvl - 1))
            st.paragraph_format.space_after = Pt(3 if lvl == 1 else 1)
            st.paragraph_format.space_before = Pt(6 if lvl == 1 else 0)
            st.paragraph_format.line_spacing = 1.4
            st.paragraph_format.tab_stops.add_tab_stop(Cm(self.text_width_cm), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            self.s_toc.append(st)

        hdr = styles["Header"]
        _style_font(hdr, self.cjk_body, self.latin, 9, color=self.th.muted)
        ftr = styles["Footer"]
        _style_font(ftr, self.cjk_body, self.latin, 9, color=self.th.ink)

        cp = doc.core_properties
        cover = self.book.get("cover") or {}
        cp.title = str(cover.get("title") or self.book.get("docType") or "")
        cp.subject = str(cover.get("subtitle") or "")
        cp.author = str(cover.get("teamName") or "")
        cp.comments = ""
        cp.last_modified_by = str(cover.get("teamName") or "")
        _enable_update_fields(doc)

    # ---- 行内 ----
    def _add_runs(self, paragraph, text: str, size: Optional[float] = None, bold: bool = False, color: Optional[str] = None,
                  cjk: Optional[str] = None) -> None:
        pos = 0
        for m in _INLINE_RE.finditer(text):
            if m.start() > pos:
                self._plain_run(paragraph, text[pos:m.start()], size, bold, color, cjk)
            if m.group("bold"):
                self._add_runs(paragraph, m.group("bold_t"), size, True, color, cjk)
            elif m.group("code"):
                self._plain_run(paragraph, m.group("code_t"), size, bold, color, cjk)
            elif m.group("fn"):
                key = m.group("fn_k")
                idx = self._footnote_index(key)
                mark = CIRCLED[idx - 1] if 0 < idx <= len(CIRCLED) else f"[注{idx}]"
                r = self._plain_run(paragraph, mark, size, False, self.th.primary, cjk)
                r.font.superscript = True
            elif m.group("cite"):
                label = re.sub(r"\s+", "", m.group("cite")[1:-1])
                r = self._plain_run(paragraph, f"[{label}]", size, False, color, cjk)
                if str(self.style_cfg.get("citationStyle") or "superscript") == "superscript":
                    r.font.superscript = True
            elif m.group("link"):
                self._plain_run(paragraph, m.group("link_t"), size, bold, color, cjk)
            elif m.group("br"):
                paragraph.add_run().add_break(WD_BREAK.LINE)
            pos = m.end()
        if pos < len(text):
            self._plain_run(paragraph, text[pos:], size, bold, color, cjk)

    def _plain_run(self, paragraph, text: str, size: Optional[float], bold: bool, color: Optional[str], cjk: Optional[str]):
        text = text.replace("\\*", "*").replace("\\_", "_").replace("\\#", "#").replace("\\|", "|")
        run = paragraph.add_run(text)
        if bold:
            run.bold = True
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if cjk:
            _set_rfonts(run._r, cjk, self.latin)
        return run

    def _footnote_index(self, key: str) -> int:
        if key not in self._fn_order:
            self._fn_order.append(key)
        return self._fn_order.index(key) + 1

    # ---- 封面 ----
    def add_cover(self) -> None:
        doc = self.doc
        cover = dict(self.book.get("cover") or {})
        for k, v in list(cover.items()):
            if isinstance(v, str):
                cover[k] = self._anon(v, "book.json#cover")
        doc_type = self._anon(str(self.book.get("docType") or ""), "book.json#docType")

        def line(text: str, size: float, bold: bool = False, color: str = self.th.ink, before: float = 0, after: float = 0,
                 cjk: Optional[str] = None, align=WD_ALIGN_PARAGRAPH.CENTER):
            p = doc.add_paragraph()
            p.alignment = align
            p.paragraph_format.space_before = Pt(before)
            p.paragraph_format.space_after = Pt(after)
            p.paragraph_format.line_spacing = 1.2
            if text:
                self._plain_run(p, text, size, bold, color, cjk or self.cjk_head)
            return p

        # 层级：弱化的文种标签（小字、拉开字距）→ 短金线 → 主标题 32pt（超长降到 30pt）→ 副标题 17pt → 细线 → 两列对齐信息块
        line("", 12, before=72)
        if doc_type:
            p = line(doc_type, 11, color=self.th.secondary, after=0)
            for r in p.runs:
                _set_char_spacing(r, 6.0)
        line("", 12, before=54)
        p = line("", 1)
        self._gold_rule(p, half_width_cm=1.6, size_eighths=18)
        title = str(cover.get("title") or "")
        title_pt = 32.0 if self._text_width_cm(title, 32.0) <= self.text_width_cm - 1.0 else 30.0
        line(title, title_pt, bold=True, color=self.th.primary, before=16, after=8)
        if cover.get("subtitle"):
            line(str(cover["subtitle"]), 17, color=self.th.primary_dark, after=10)
        p = line("", 1)
        self._gold_rule(p, half_width_cm=5.0, size_eighths=6)
        line("", 12, before=96)
        rows = [("团队名称", cover.get("teamName")), ("指导企业", cover.get("company")), ("参赛组别", cover.get("group")), ("完成日期", cover.get("date"))]
        # 两列对齐：标签右对齐到同一制表位（冒号对齐），内容左对齐到其右侧制表位；制表位在 Word / LibreOffice 下表现一致
        tab_r = Cm(self.text_width_cm * 0.46)
        tab_l = Cm(self.text_width_cm * 0.46 + 0.45)
        for label, val in rows:
            if not val:
                continue
            val = str(val)
            if val.startswith(label):
                val = val[len(label):].lstrip("：: ")
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(7)
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.tab_stops.add_tab_stop(tab_r, WD_TAB_ALIGNMENT.RIGHT)
            p.paragraph_format.tab_stops.add_tab_stop(tab_l, WD_TAB_ALIGNMENT.LEFT)
            self._plain_run(p, f"\t{label}：", 12, False, self.th.secondary, self.cjk_head)
            self._plain_run(p, f"\t{val}", 13, False, self.th.ink, self.cjk_body)
        if cover.get("extraLines"):
            line("", 6, before=10)
        for extra in cover.get("extraLines") or []:
            line(str(extra), 11, color=self.th.secondary, after=4, cjk=self.cjk_body)

    def _text_width_cm(self, text: str, size_pt: float) -> float:
        """粗估一行文字宽度：中文按 1 em、西文按 0.55 em。"""
        em = sum(1.0 if _CJK_RE.match(ch) else 0.55 for ch in text)
        return em * size_pt * 0.03528

    def _gold_rule(self, paragraph, half_width_cm: float = 4.5, size_eighths: int = 12) -> None:
        """金色装饰线：用段落下边框实现，比插图形更稳。half_width_cm 为线的半宽（居中），size_eighths 为线粗（1/8 pt）。"""
        pbdr = _ppr_child(paragraph, "w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), str(size_eighths))
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), self.th.accent)
        pbdr.append(bottom)
        indent = max(0.5, self.text_width_cm / 2 - half_width_cm)
        paragraph.paragraph_format.left_indent = Cm(indent)
        paragraph.paragraph_format.right_indent = Cm(indent)

    # ---- 页眉页脚 ----
    def _header_footer(self, section, header_text: str, page_fmt: str, link_header: bool) -> None:
        if not link_header:
            section.header.is_linked_to_previous = False
            hp = section.header.paragraphs[0]
            hp.text = ""
            hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            hp.style = self.doc.styles["Header"]
            if header_text:
                self._plain_run(hp, header_text, 9, False, self.th.muted, self.cjk_body)
                pbdr = _ppr_child(hp, "w:pBdr")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "6")
                bottom.set(qn("w:space"), "1")
                bottom.set(qn("w:color"), self.th.primary)
                pbdr.append(bottom)
        if self.style_cfg.get("pageNumbers", True):
            section.footer.is_linked_to_previous = False
            fp = section.footer.paragraphs[0]
            fp.text = ""
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fp.style = self.doc.styles["Footer"]
            _add_field(fp, "PAGE", "1")
            _set_pg_num_type(section, page_fmt, 1)

    # ---- 目录 ----
    def insert_toc(self, anchor) -> None:
        """在锚点段落前插入“目录”标题与 TOC 域；域结果就是静态目录条目（Word 打开时会刷新为真实页码）。"""
        h = anchor.insert_paragraph_before()
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.paragraph_format.space_after = Pt(16)
        self._plain_run(h, "目  录", self.th.typo.doc_h1, True, self.th.ink, self.cjk_head)
        entries = [e for e in self.toc if e.level <= self.toc_depth]
        toc_pages = max(1, math.ceil(len(entries) / 32))
        for e in entries:
            if e.section == "front":
                e.page = _roman(toc_pages + 1)
        if not entries:
            entries = [TocEntry(1, "（目录将在 Word 中打开后自动更新）", "")]
        for k, e in enumerate(entries):
            p = anchor.insert_paragraph_before(style=self.s_toc[min(e.level, 3) - 1])
            if k == 0:
                _fld_char(p, "begin", dirty=True)
                _instr_text(p, f'TOC \\o "1-{self.toc_depth}" \\h \\z \\u')
                _fld_char(p, "separate")
            self._plain_run(p, e.text, None, e.level == 1, None, None)
            page = self.toc_pages.get(e.text, e.page)
            p.add_run("\t" + str(page))
            if k == len(entries) - 1:
                _fld_char(p, "end")

    # ---- 章节 ----
    def _chapter_title_text(self, num: Numbering, text: str, level: int) -> str:
        text = _strip_numbering(text)
        if not num.numbered or not num.label:
            return text
        if level == 1:
            return (f"附录 {num.label}  {text}" if num.label.isalpha() else f"{num.label}  {text}")
        if level == 2:
            return f"{num.label}.{num.h2}  {text}"
        return f"{num.label}.{num.h2}.{num.h3}  {text}"

    def add_markdown_file(self, rel_path: str, num: Numbering, first_in_section: bool, title_fallback: str = "",
                          toc_section: str = "body", toc_max_level: int = 3) -> None:
        path = self._rel(rel_path)
        if not path.exists():
            self.stats.warnings.append(f"章节文件不存在：{rel_path}")
            p = self.doc.add_paragraph(style=self.s_body)
            self._plain_run(p, f"【缺章节文件：{rel_path}】", None, True, self.th.negative, None)
            return
        raw = read_text(path)
        raw = self._anon(raw, rel_path)
        parsed = parse_markdown(raw)
        parsed.blocks = _attach_plain_sources(parsed.blocks)
        self._fn_order = []
        self._footnotes = parsed.footnotes
        base_dir = path.parent
        has_h1 = any(isinstance(b, Heading) and b.level == 1 for b in parsed.blocks)
        if not has_h1:
            # 没有一级标题的文件：把开头的标题提升为章标题，否则用兜底标题
            first = parsed.blocks[0] if parsed.blocks else None
            if isinstance(first, Heading):
                first.level = 1
            else:
                parsed.blocks.insert(0, Heading(1, title_fallback or path.stem, 0))
        h1_seen = 0
        for block in parsed.blocks:
            if isinstance(block, Heading):
                if block.level == 1:
                    if h1_seen and num.numbered and num.label.isdigit():
                        # 同一文件里的第二个一级标题视为新的一章
                        self.chapter_no += 1
                        num.label = str(self.chapter_no)
                        num.fig = num.tab = 0
                    self._render_heading(block, num, page_break=not (first_in_section and h1_seen == 0), toc_section=toc_section, toc_max_level=toc_max_level)
                    h1_seen += 1
                else:
                    self._render_heading(block, num, page_break=False, toc_section=toc_section, toc_max_level=toc_max_level)
            elif isinstance(block, Para):
                self._render_para(block)
            elif isinstance(block, ListBlock):
                self._render_list(block)
            elif isinstance(block, Table):
                self._render_table(block, num)
            elif isinstance(block, Figure):
                self._render_figure(block, num, base_dir)
            elif isinstance(block, Quote):
                p = self.doc.add_paragraph(style=self.s_quote)
                self._add_runs(p, block.text)
                self._count(block.text)
            elif isinstance(block, PageBreak):
                self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                self.est_pos = math.ceil(self.est_pos)
            elif isinstance(block, SourceLine):
                self._render_source(block.text)
        if self._fn_order:
            self._render_footnotes()

    def _count(self, text: str) -> None:
        n = count_words(text)
        self.stats.words += n
        self.est_pos += n / 650

    def _render_heading(self, block: Heading, num: Numbering, page_break: bool, toc_section: str = "body", toc_max_level: int = 3) -> None:
        level = min(block.level, 3)
        if level == 1:
            num.h2 = num.h3 = 0
            if page_break:
                self.est_pos = math.ceil(self.est_pos)
        elif level == 2:
            num.h2 += 1
            num.h3 = 0
        else:
            num.h3 += 1
        text = self._chapter_title_text(num, block.text, level)
        p = self.doc.add_paragraph(style=self.doc.styles[f"Heading {level}"])
        if page_break:
            p.paragraph_format.page_break_before = True
        self._plain_run(p, text, None, None, None, None)
        self.stats.headings += 1
        self.stats.words += count_words(text)
        if level <= min(self.toc_depth, toc_max_level):
            self.toc.append(TocEntry(level, text, str(int(self.est_pos) + 1), toc_section))
        self.est_pos += 0.08

    def _render_para(self, block: Para) -> None:
        p = self.doc.add_paragraph(style=self.s_body)
        self._add_runs(p, block.text)
        self._count(block.text)
        for s in block.sources:
            self._render_source(s)

    def _render_source(self, text: str, center: bool = False) -> None:
        p = self.doc.add_paragraph(style=self.s_source)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        self._add_runs(p, text, 9, False, self.th.muted)

    def _render_list(self, block: ListBlock) -> None:
        counters = [0, 0, 0, 0]
        bullets = ["•", "–", "·", "·"]
        indent_unit_cm = self.body_size * 0.03528 * 2  # 2 字符
        for item in block.items:
            lvl = item.level
            counters[lvl] += 1
            for deeper in range(lvl + 1, 4):
                counters[deeper] = 0
            p = self.doc.add_paragraph(style=self.s_list)
            hanging = 0.75 if block.ordered else 0.5
            # 一级列表项的项目符号与正文首行缩进（2 字符）对齐，层级越深越向右
            p.paragraph_format.left_indent = Cm(indent_unit_cm * (lvl + 1) + hanging)
            p.paragraph_format.first_line_indent = Cm(-hanging)
            prefix = f"{counters[lvl]}. " if block.ordered else f"{bullets[lvl]} "
            if block.ordered and lvl > 0:
                prefix = f"({counters[lvl]}) "
            self._plain_run(p, prefix, None, False, self.th.primary if not block.ordered else None, None)
            text, srcs = _extract_sources(item.text)
            self._add_runs(p, text)
            self._count(text)
            for s in srcs:
                self._render_source(s)

    def _render_table(self, block: Table, num: Numbering) -> None:
        doc = self.doc
        ncols = len(block.header)
        if ncols == 0:
            return
        label = num.tab_label()
        self.stats.tables += 1
        self.stats.table_rows += len(block.rows) + 1
        self.est_pos += (len(block.rows) + 1) / 40
        self.stats.words += sum(count_words(c) for c in block.header) + sum(count_words(c) for r in block.rows for c in r)
        cap = doc.add_paragraph(style=self.s_caption)
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_before = Pt(8)
        cap.paragraph_format.space_after = Pt(4)
        caption = _strip_caption_prefix(block.caption)
        self._add_runs(cap, f"{label}  {caption}" if caption else label, self.th.typo.doc_small, True)
        # 分页控制：≤6 行的短表整体不跨页；长表允许跨页（表头重复），但首尾各至少 3 行在一起，
        # 不出现只剩 1 行孤悬到下一页；最后一行再拉住来源行
        n_rows = len(block.rows)
        keep_idx = set(range(n_rows)) if n_rows <= 6 else set(range(0, min(2, n_rows))) | set(range(max(0, n_rows - 3), n_rows - 1))

        table = doc.add_table(rows=1 + len(block.rows), cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        # 列宽：按列内容最大长度加权，保证窄列不被挤扁
        weights = []
        for c in range(ncols):
            lens = [len(block.header[c])] + [len(r[c]) for r in block.rows]
            weights.append(min(max(max(lens), 3), 28) ** 0.85)
        total_w = sum(weights) or 1
        widths = [self.text_width_cm * w / total_w for w in weights]
        is_numeric = [all(re.fullmatch(r"[\d,.%+\-–—/\s¥$万亿元千百]*", (r[c] or "—")) for r in block.rows) and len(block.rows) > 0 for c in range(ncols)]
        long_text = [max([len(block.header[c])] + [len(r[c]) for r in block.rows]) > 14 for c in range(ncols)]

        tblpr = table._tbl.tblPr
        tblw = tblpr.find(qn("w:tblW"))
        if tblw is None:
            tblw = OxmlElement("w:tblW")
            tblpr.append(tblw)
        tblw.set(qn("w:type"), "dxa")
        tblw.set(qn("w:w"), str(int(self.text_width_cm * 567)))
        _set_table_cell_margins(table)
        three_line = self.table_style == "three-line"
        if three_line:
            _set_table_borders(table, {"top": ("single", 12, "000000"), "bottom": ("single", 12, "000000")})
        else:
            _set_table_borders(table, {"top": ("single", 4, self.th.line), "bottom": ("single", 4, self.th.line),
                                       "left": ("single", 4, self.th.line), "right": ("single", 4, self.th.line),
                                       "insideH": ("single", 4, self.th.line), "insideV": ("single", 4, self.th.line)})

        for c in range(ncols):
            table.columns[c].width = Cm(widths[c])
        for r_idx, row in enumerate(table.rows):
            is_header = r_idx == 0
            cells_text = block.header if is_header else block.rows[r_idx - 1]
            if is_header:
                _row_header_repeat(row)
            _row_cant_split(row)
            last_row = r_idx == len(table.rows) - 1
            keep_next = is_header or (not last_row and (r_idx - 1) in keep_idx) or (last_row and bool(block.sources))
            for c, cell in enumerate(row.cells):
                cell.width = Cm(widths[c])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cell.paragraphs[0]
                p.style = self.s_table
                if keep_next:
                    p.paragraph_format.keep_with_next = True
                align = block.aligns[c]
                if is_header:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif align == "right" or (not align and is_numeric[c]):
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif align == "left" or (not align and long_text[c]):
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                text = cells_text[c] if c < len(cells_text) else ""
                if is_header:
                    self._add_runs(p, text, self.th.typo.doc_small, True, "FFFFFF" if not three_line else self.th.ink)
                    if three_line:
                        _set_cell_borders(cell, bottom=("single", 6, "000000"))
                    else:
                        _set_cell_shading(cell, self.th.primary)
                else:
                    self._add_runs(p, text, self.th.typo.doc_small)
                    if not three_line and r_idx % 2 == 0:
                        _set_cell_shading(cell, self.th.paper_alt)
        for s in block.sources:
            self._render_source(s)
        if not block.sources:
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(0)
            spacer.paragraph_format.line_spacing = Pt(6)

    def _render_figure(self, block: Figure, num: Numbering, base_dir: Path) -> None:
        doc = self.doc
        img_path = Path(block.path)
        if not img_path.is_absolute():
            img_path = (base_dir / img_path).resolve()
        label = num.fig_label()   # 缺图也占用编号，补图后其余图号不变
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        if img_path.exists():
            width_cm = self._figure_width_cm(img_path, block.width_hint)
            try:
                p.add_run().add_picture(str(img_path), width=Cm(width_cm))
                self.stats.figures += 1
                self.est_pos += 0.45
            except Exception as e:
                self.stats.warnings.append(f"图片无法插入 {block.path}: {e}")
                self._plain_run(p, f"【缺图：{block.path}】", None, True, self.th.negative, None)
                self.stats.missing_images.append(block.path)
        else:
            self.stats.missing_images.append(block.path)
            self.stats.warnings.append(f"找不到图片：{block.path}（相对 {base_dir}）")
            self._plain_run(p, f"【缺图：{block.path}】", None, True, self.th.negative, None)
        cap = doc.add_paragraph(style=self.s_caption)
        if block.sources:
            cap.paragraph_format.keep_with_next = True
        caption = _strip_caption_prefix(block.caption)
        self._add_runs(cap, f"{label}  {caption}" if caption else label, self.th.typo.doc_small, True)
        for s in block.sources:
            self._render_source(s, center=True)

    def _figure_width_cm(self, img_path: Path, hint: str) -> float:
        max_w = min(self.text_width_cm, 15.5)
        width = max_w
        hint = (hint or "").strip().lower()
        m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*(%|mm|cm)?", hint)
        if m:
            v, unit = float(m.group(1)), m.group(2) or "%"
            width = max_w * v / 100 if unit == "%" else (v / 10 if unit == "mm" else v)
            width = min(width, max_w)
        try:
            from PIL import Image
            with Image.open(img_path) as im:
                w_px, h_px = im.size
            ratio = h_px / w_px if w_px else 1
            max_h = 18.0
            if width * ratio > max_h:
                width = max_h / ratio
            # 窄高图（如漏斗、定位图）不必撑满版心
            if not m and ratio > 1.05:
                width = min(width, max_w * 0.72)
        except Exception:
            pass
        return width

    def _render_footnotes(self) -> None:
        doc = self.doc
        p = doc.add_paragraph(style=self.s_note)
        p.paragraph_format.space_before = Pt(10)
        pbdr = _ppr_child(p, "w:pBdr")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), "4")
        top.set(qn("w:space"), "4")
        top.set(qn("w:color"), self.th.line)
        pbdr.append(top)
        self._plain_run(p, "本章注释", 9.5, True, self.th.muted, self.cjk_head)
        for i, key in enumerate(self._fn_order, 1):
            mark = CIRCLED[i - 1] if i <= len(CIRCLED) else f"[注{i}]"
            body = self._footnotes.get(key)
            if body is None:
                self.stats.warnings.append(f"脚注 [^{key}] 没有定义")
                body = "（缺脚注内容）"
            q = doc.add_paragraph(style=self.s_note)
            q.paragraph_format.left_indent = Cm(0.6)
            q.paragraph_format.first_line_indent = Cm(-0.6)
            self._add_runs(q, f"{mark} {body}", 9.5, False, self.th.muted)

    # ---- 参考文献 ----
    def add_references(self, rel_path: str, first_in_section: bool) -> None:
        path = self._rel(rel_path)
        if not path.exists():
            self.stats.warnings.append(f"参考文献文件不存在：{rel_path}")
            return
        try:
            refs = read_json(path)
        except Exception as e:
            self.stats.warnings.append(f"参考文献 JSON 无法解析：{e}")
            return
        if isinstance(refs, dict):
            refs = refs.get("references") or refs.get("items") or []
        if not isinstance(refs, list) or not refs:
            self.stats.warnings.append("参考文献为空")
            return
        refs = sorted((r for r in refs if isinstance(r, dict)), key=lambda r: (int(r.get("id") or 0) if str(r.get("id") or "").isdigit() else 10 ** 6))
        p = self.doc.add_paragraph(style=self.doc.styles["Heading 1"])
        if not first_in_section:
            p.paragraph_format.page_break_before = True
            self.est_pos = math.ceil(self.est_pos)
        self._plain_run(p, "参考文献", None, None, None, None)
        self.toc.append(TocEntry(1, "参考文献", str(int(self.est_pos) + 1), "body"))
        for k, ref in enumerate(refs, 1):
            rid = ref.get("id") or k
            text = self._anon(format_reference(ref), f"{rel_path}#id={rid}")
            q = self.doc.add_paragraph(style=self.s_ref)
            self._plain_run(q, f"[{rid}] ", self.th.typo.doc_small, False, None, None)
            self._plain_run(q, text, self.th.typo.doc_small, False, None, None)
            self.est_pos += 2.2 / 40

    # ---- 整本书 ----
    def build(self) -> None:
        """三节结构：封面（无页眉页脚）→ 目录 + 概要（罗马页码）→ 正文 + 附录 + 参考文献（阿拉伯页码从 1 起）。"""
        self.setup()
        doc = self.doc
        header_text = self._anon(str(self.style_cfg.get("header") or (self.book.get("cover") or {}).get("subtitle") or ""), "book.json#style.header")
        want_toc = bool(self.style_cfg.get("toc", True))
        self.add_cover()

        front = doc.add_section(WD_SECTION.NEW_PAGE)
        _mark_section_break_paragraph(doc)
        self._header_footer(front, header_text, "lowerRoman", link_header=False)
        abstract = self.book.get("abstract")
        if want_toc:
            # 目录依赖后文标题，先放一个锚点段落，正文渲染完再把目录插到它前面；锚点兼作目录与概要之间的分页
            self._toc_anchor = doc.add_paragraph()
            if abstract:
                self._toc_anchor.add_run().add_break(WD_BREAK.PAGE)
        if abstract:
            self.add_markdown_file(str(abstract), Numbering(label="", numbered=False), first_in_section=True,
                                   title_fallback="概要", toc_section="front", toc_max_level=1)

        body = doc.add_section(WD_SECTION.NEW_PAGE)
        _mark_section_break_paragraph(doc)
        self._header_footer(body, header_text, "decimal", link_header=True)
        self.est_pos = 0.0
        for idx, ch in enumerate(self.book.get("chapters") or []):
            if isinstance(ch, str):
                ch = {"file": ch}
            numbered = bool(ch.get("numbered", True))
            if numbered:
                self.chapter_no += 1
            num = Numbering(label=str(self.chapter_no) if numbered else "", numbered=numbered)
            self.add_markdown_file(str(ch["file"]), num, first_in_section=(idx == 0), title_fallback=Path(str(ch["file"])).stem)
            self.stats.chapters += 1
        for idx, ap in enumerate(self.book.get("appendices") or []):
            if isinstance(ap, str):
                ap = {"file": ap}
            letter = APPENDIX_LETTERS[idx % 26]
            self.add_markdown_file(str(ap["file"]), Numbering(label=letter, numbered=True), first_in_section=False, title_fallback=f"附录 {letter}")
        refs = self.book.get("references")
        if refs:
            self.add_references(str(refs), first_in_section=False)
        if want_toc and self._toc_anchor is not None:
            self.insert_toc(self._toc_anchor)

    def save(self, out_path: Path) -> None:
        out_path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(out_path))


# ----------------------------------------------------------------------------------------------
# PDF 回填目录页码
# ----------------------------------------------------------------------------------------------

def _pdf_page_texts(pdf_path: Path) -> List[str]:
    texts: List[str] = []
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(pdf_path))
        for page in reader.pages:
            texts.append(page.extract_text() or "")
    except Exception:
        texts = []
    if not any(t.strip() for t in texts):
        try:
            import pdfplumber
            with pdfplumber.open(str(pdf_path)) as pdf:
                texts = [(pg.extract_text() or "") for pg in pdf.pages]
        except Exception:
            texts = []
    return texts


def _norm(s: str) -> str:
    return re.sub(r"\s+", "", s)


def toc_pages_from_pdf(pdf_path: Path, toc: List[TocEntry]) -> Tuple[Dict[str, str], Dict[str, Any]]:
    """用 PDF 文本定位每个标题所在页，换算为正文页码（正文第 1 章所在页 = 1）；前置部分用罗马数字。"""
    texts = _pdf_page_texts(pdf_path)
    if not texts:
        return {}, {"reason": "PDF 无法提取文本"}
    page_lines = [[_norm(ln) for ln in t.split("\n") if ln.strip()] for t in texts]
    page_blobs = [_norm(t) for t in texts]

    def find(text: str, start: int) -> Optional[int]:
        key = _norm(text)
        for i in range(start, len(texts)):
            if key in page_lines[i]:
                return i
        for i in range(start, len(texts)):  # 标题换行时退化为包含匹配（跳过目录页本身）
            if key in page_blobs[i] and not any(ln.startswith(key) and len(ln) > len(key) + 1 and ln[-1].isdigit() for ln in page_lines[i]):
                return i
        return None

    toc_page = next((i for i, lines in enumerate(page_lines) if "目录" in lines), 1)
    found: Dict[str, int] = {}
    cursor = toc_page + 1
    for e in toc:
        idx = find(e.text, cursor)
        if idx is None:
            continue
        found[e.text] = idx
        cursor = idx
    body_entries = [e for e in toc if e.section == "body" and e.level == 1 and e.text in found]
    if not body_entries:
        return {}, {"reason": "未在 PDF 中定位到正文标题", "pages": len(texts)}
    body_start = found[body_entries[0].text]
    result: Dict[str, str] = {}
    for e in toc:
        if e.text not in found:
            continue
        if e.section == "front":
            result[e.text] = _roman(found[e.text] - toc_page + 1)
        else:
            result[e.text] = str(found[e.text] - body_start + 1)
    return result, {"pages": len(texts), "bodyStartPage": body_start + 1, "located": len(found), "total": len(toc)}


def _roman(n: int) -> str:
    vals = [(10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
    out = ""
    for v, s in vals:
        while n >= v:
            out += s
            n -= v
    return out or "i"


# ----------------------------------------------------------------------------------------------
# 入口
# ----------------------------------------------------------------------------------------------

def render_book(book_path: Path, make_pdf: Optional[bool] = None) -> Dict[str, Any]:
    book = read_json(book_path)
    renderer = BookRenderer(book, book_path)
    renderer.build()
    out_rel = str(book.get("output") or "产出/文档.docx")
    out_path = renderer._rel(out_rel)
    renderer.save(out_path)
    stats = renderer.stats
    result: Dict[str, Any] = {
        "ok": True, "output": str(out_path), "chapters": stats.chapters, "figures": stats.figures, "tables": stats.tables,
        "tableRows": stats.table_rows, "words": stats.words,
        "pagesEstimate": round(stats.words / 650 + stats.figures * 0.45 + stats.table_rows / 40, 1),
        "anonymizeHits": stats.anonymize_hits, "missingImages": stats.missing_images, "warnings": list(stats.warnings),
        "tocPageNumbers": "estimated", "pdf": None, "pages": None,
    }
    want_pdf = bool(book.get("pdf")) if make_pdf is None else make_pdf
    if want_pdf:
        from . import pdf as pdf_mod
        conv = pdf_mod.convert(out_path)
        if not conv.get("ok"):
            result["warnings"].append(f"PDF 转换失败：{conv.get('reason')}")
        else:
            pdf_path = Path(conv["pdf"])
            pages_map, info = toc_pages_from_pdf(pdf_path, renderer.toc)
            if pages_map and book.get("style", {}).get("toc", True):
                # 用真实页码重渲染一次，让不更新域的阅读端（LibreOffice / 预览）也显示正确目录
                renderer2 = BookRenderer(book, book_path, toc_pages=pages_map)
                renderer2.build()
                renderer2.save(out_path)
                conv = pdf_mod.convert(out_path)
                if conv.get("ok"):
                    result["tocPageNumbers"] = "pdf"
                else:
                    result["warnings"].append(f"第二次 PDF 转换失败：{conv.get('reason')}")
            elif not pages_map:
                result["warnings"].append(f"目录页码未能回填：{info.get('reason', '')}")
            if conv.get("ok"):
                result["pdf"] = conv["pdf"]
                result["pages"] = conv.get("pages")
                result["converter"] = conv.get("converter")
    return result


def main(argv: Optional[Sequence[str]] = None) -> int:
    force_utf8_stdout()
    ap = argparse.ArgumentParser(description="book.json → docx")
    ap.add_argument("book_json")
    ap.add_argument("--no-pdf", action="store_true", help="即使 book.json.pdf 为 true 也不转 PDF")
    args = ap.parse_args(argv)
    path = Path(args.book_json).resolve()
    try:
        result = render_book(path, make_pdf=False if args.no_pdf else None)
    except FileNotFoundError as e:
        print(json.dumps({"ok": False, "error": f"文件不存在：{e}"}, ensure_ascii=False))
        return 1
    except Exception as e:
        import traceback
        print(json.dumps({"ok": False, "error": f"{type(e).__name__}: {e}", "trace": traceback.format_exc(limit=5)}, ensure_ascii=False))
        return 1
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
