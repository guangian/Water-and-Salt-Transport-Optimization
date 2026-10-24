"""提交包：按赛道命名规则收集产出、生成声明模板 / 数据真实性清单 / README，并打 zip。

用法：
    python -m bizkit.package --track brand --workspace . [--out 08_提交包] [--no-zip]

读取：tracks/<track>/naming.json（工作区内；找不到则用 agent/tracks 或内置默认）与 project.json（工作区根，契约 §2）。
收集：产出/*.docx|pdf|pptx、03_财务/财务模型.xlsx、06_路演/讲稿.md·答辩题库.md·视频分镜.md、
      01_调研/调研工具包/、07_评审/最新一轮 → 复制到 08_提交包/ 并按 fileNameRule 重命名。
生成：声明与授权/（作品原创性声明、作品版权声明、企业授权书，docx 通用模板，注明以组委会当年附件为准）、
      数据真实性清单.md（numbers.json 中 synthetic:true 的条目 + 正文里的「【示例数据」标记）、README.md。
输出 JSON：{ok, dir, zip, baseName, files[{src, dst, size}], missing[], placeholders[], warnings[]}。
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from . import fonts
from ._util import force_utf8_stdout, read_json, write_text

TRACK_NAMES = {"brand": "品牌策划赛道", "startup": "创业计划赛道", "digital": "数字营销赛道", "case": "案例分析赛道", "mice": "商务会奖旅游策划赛道"}
DOC_NAMES = {"brand": "品牌策划书", "startup": "商业计划书", "digital": "数字营销方案", "case": "案例分析报告", "mice": "会奖策划书"}
DEFAULT_NAMING: Dict[str, Any] = {
    "fileNameRule": "{组别}-{学校}-{团队名}-{作品名}",
    "coverFields": ["作品名称", "团队名称", "指导企业名称"],
    "coverForbidden": ["院校", "姓名", "联系方式"],
    "declarations": ["作品原创性声明", "作品版权声明", "企业授权书"],
    "pageLimit": 80,
    "videoMinutes": 10,
    "qaMinutes": 5,
    "videoPeople": "3–5 人",
}
# 这些占位符没有可靠来源，只能先占位并在 README 里提醒用户替换
MANUAL_PLACEHOLDERS = {"学校": "学校名称", "抽签编号": "抽签编号", "抽签编号或学校": "抽签编号或学校", "院校": "学校名称", "编号": "编号"}
SKIP_DIRS = {"toolkit", ".claude", ".claude-home", "node_modules", ".cache", "08_提交包", "inputs"}


def _safe_filename(name: str) -> str:
    name = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", name)
    name = re.sub(r"\s+", " ", name).strip(" ._")
    return name[:120] or "提交包"


def load_naming(workspace: Path, track: str) -> Tuple[Dict[str, Any], str]:
    candidates = [workspace / "tracks" / track / "naming.json", Path(__file__).resolve().parents[2] / "tracks" / track / "naming.json"]
    for c in candidates:
        if c.exists():
            try:
                data = dict(DEFAULT_NAMING)
                data.update(read_json(c))
                return data, str(c)
            except Exception:
                continue
    return dict(DEFAULT_NAMING), "builtin-default"


def build_base_name(rule: str, project: Dict[str, Any], track: str) -> Tuple[str, List[str]]:
    subject = project.get("subject") or {}
    values = {
        "组别": project.get("group") or "组别", "团队名": project.get("teamName") or "团队名", "作品名": project.get("name") or "作品名",
        "赛道": TRACK_NAMES.get(track, track), "企业": subject.get("company") or subject.get("name") or "", "指导企业": subject.get("company") or "",
        "品牌": subject.get("name") or "", "材料": DOC_NAMES.get(track, "参赛作品"),
    }
    placeholders: List[str] = []

    def _sub(m: "re.Match[str]") -> str:
        key = m.group(1).strip()
        if key in values and values[key]:
            return str(values[key])
        if key in MANUAL_PLACEHOLDERS:
            placeholders.append(MANUAL_PLACEHOLDERS[key])
            return MANUAL_PLACEHOLDERS[key]
        placeholders.append(key)
        return key

    base = re.sub(r"\{([^{}]+)\}", _sub, rule)
    return _safe_filename(base), placeholders


def latest_review_dir(workspace: Path) -> Optional[Path]:
    root = workspace / "07_评审"
    if not root.is_dir():
        return None
    rounds = []
    for d in root.iterdir():
        m = re.match(r"第(\d+)轮", d.name)
        if d.is_dir() and m:
            rounds.append((int(m.group(1)), d))
    if not rounds:
        return None
    return max(rounds)[1]


# ---- 声明模板 ----

def _declaration_doc(title: str, intro: str, items: List[str], sign_lines: List[str], out: Path, project: Dict[str, Any]) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    df = fonts.doc_fonts()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(3.0)
    sec.top_margin = sec.bottom_margin = Cm(2.54)

    def run_font(run, size: float, bold: bool = False, cjk: str = df["body"]) -> None:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = df["latin"]
        run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), cjk)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run_font(p.add_run(title), 18, True, df["heading"])
    for text in [intro] + items:
        q = doc.add_paragraph()
        q.paragraph_format.first_line_indent = Cm(0.85)
        q.paragraph_format.line_spacing = 1.5
        q.paragraph_format.space_after = Pt(6)
        run_font(q.add_run(text), 12)
    doc.add_paragraph()
    for line in sign_lines:
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        q.paragraph_format.line_spacing = 1.8
        run_font(q.add_run(line), 12)
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(24)
    run_font(note.add_run("说明：本文件为通用模板，正式提交请以组委会当年发布的附件格式为准；签字、盖章与指印须为手写 / 实体，扫描为 PDF 后随作品一并提交。"), 9)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))


def make_declarations(out_dir: Path, names: List[str], project: Dict[str, Any]) -> List[Path]:
    work = project.get("name") or "（作品名称）"
    team = project.get("teamName") or "（团队名称）"
    company = (project.get("subject") or {}).get("company") or "（企业名称）"
    year = datetime.now().year
    templates = {
        "作品原创性声明": (
            f"本团队郑重声明：所提交的参赛作品《{work}》系本团队在指导教师指导下独立完成的原创成果。",
            [
                "作品中引用的他人数据、观点与图片均已注明出处；除注明部分外，不包含任何他人已发表或撰写的作品内容。",
                "作品中标注为“示例数据”的内容已在提交前全部替换为团队真实调研或企业提供的数据，团队对数据真实性负责。",
                "若有抄袭、剽窃、伪造数据等学术不端行为，本团队愿承担由此产生的全部责任并接受组委会处理。",
            ],
            [f"参赛团队：{team}", "全体成员签名（手写并按指印）：", "", f"{year} 年    月    日"],
        ),
        "作品版权声明": (
            f"本团队同意：参赛作品《{work}》的著作权归团队全体成员共同所有，并作出如下授权。",
            [
                "授权大赛组委会在非营利目的范围内对作品进行展示、汇编、出版与宣传，无需另行支付报酬。",
                "授权大赛组委会将作品提供给命题 / 指导企业参考使用；企业如需商业化使用，须另行与团队协商。",
                "团队保证作品不侵犯任何第三方的著作权、商标权或其他合法权益。",
            ],
            [f"参赛团队：{team}", "全体成员签名（手写并按指印）：", "", f"{year} 年    月    日"],
        ),
        "企业授权书": (
            f"{company}（以下简称“本企业”）同意作为参赛团队“{team}”的指导 / 命题企业，就作品《{work}》作出如下授权与确认。",
            [
                "本企业授权团队在作品中使用本企业的品牌名称、标识、公开经营信息及本企业提供的资料。",
                "本企业确认团队所使用的企业相关数据来源合法；作品中涉及本企业的商业数据如需保密，已由双方另行约定脱敏方式。",
                "本授权仅用于全国高校商业精英挑战赛参赛及评审，不构成任何商业合作承诺。",
            ],
            [f"授权企业（盖章）：{company}", "法定代表人 / 授权代表签字：", "联系电话（仅供组委会核验，不出现在作品中）：", "", f"{year} 年    月    日"],
        ),
    }
    out: List[Path] = []
    for name in names:
        intro, items, signs = templates.get(name, (f"本文件为《{name}》模板，请按组委会当年附件格式填写。", [], [f"参赛团队：{team}", f"{year} 年    月    日"]))
        path = out_dir / f"{_safe_filename(name)}.docx"
        _declaration_doc(name, intro, items, signs, path, project)
        out.append(path)
    return out


# ---- 数据真实性清单 ----

def scan_synthetic(workspace: Path) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    entries: List[Dict[str, Any]] = []
    numbers = workspace / "03_财务" / "numbers.json"
    if numbers.exists():
        try:
            data = read_json(numbers)
            for e in data.get("entries") or []:
                if isinstance(e, dict) and e.get("synthetic"):
                    entries.append(e)
        except Exception:
            pass
    marks: List[Dict[str, Any]] = []
    for md in sorted(workspace.rglob("*.md")):
        rel = md.relative_to(workspace)
        if any(part in SKIP_DIRS for part in rel.parts):
            continue
        try:
            text = md.read_text(encoding="utf-8-sig")
        except Exception:
            continue
        for ln_no, line in enumerate(text.split("\n"), 1):
            if "【示例数据" in line or "【示例" in line and "数据" in line:
                idx = line.find("【示例")
                ctx = line[max(0, idx - 40): idx + 60].strip()
                marks.append({"file": str(rel).replace("\\", "/"), "line": ln_no, "context": ctx})
    return entries, marks


def write_truth_list(path: Path, entries: List[Dict[str, Any]], marks: List[Dict[str, Any]]) -> None:
    lines = ["# 数据真实性清单", "",
             "> 本清单列出材料中标注为示例 / 模拟的数据。**提交前必须用真实调研或企业数据替换**，并同步更新策划书正文、路演 PPT、讲稿与答辩题库，保持数字一致。", "",
             f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}", "",
             "## 一、数字账本（03_财务/numbers.json）中标记为示例的条目", ""]
    if entries:
        lines += ["| key | 名称 | 当前值 | 单位 | 来源 | 被引用位置 |", "|---|---|---:|---|---|---|"]
        for e in entries:
            used = "、".join(str(u) for u in (e.get("usedIn") or [])) or "—"
            lines.append(f"| {e.get('key', '')} | {e.get('label', '')} | {e.get('value', '')} | {e.get('unit', '')} | {e.get('source', '')} | {used} |")
    else:
        lines.append("（无）")
    lines += ["", "## 二、正文 / 路演中的「【示例数据】」标记", ""]
    if marks:
        lines += ["| 文件 | 行 | 上下文 |", "|---|---:|---|"]
        for m in marks:
            lines.append(f"| {m['file']} | {m['line']} | {m['context'].replace('|', '/')} |")
    else:
        lines.append("（无）")
    lines += ["", "## 三、替换步骤", "",
              "1. 完成真实问卷 / 访谈后，把结果写入 01_调研/数据/ 与 03_财务/assumptions.json；",
              "2. 重新运行财务模型与图表（finmodel → charts），再渲染正文与 PPT（docx_render → pptx_render）；",
              "3. 删除所有「【示例数据，须替换为真实调研结果】」标记；",
              "4. 运行 `python -m bizkit.checks` 确认跨材料一致性与匿名检查全部通过；",
              "5. 重新打包（`python -m bizkit.package`）。", ""]
    write_text(path, "\n".join(lines))


# ---- 主流程 ----

def package(workspace: Path, track: str, out_dir: Optional[Path] = None, make_zip: bool = True) -> Dict[str, Any]:
    workspace = Path(workspace).resolve()
    warnings: List[str] = []
    missing: List[str] = []
    files: List[Dict[str, Any]] = []
    project: Dict[str, Any] = {}
    pj = workspace / "project.json"
    if pj.exists():
        try:
            project = read_json(pj)
        except Exception as e:
            warnings.append(f"project.json 无法解析：{e}")
    else:
        warnings.append("未找到 project.json，命名使用占位符")
    track = track or str(project.get("track") or "brand")
    naming, naming_src = load_naming(workspace, track)
    base, placeholders = build_base_name(str(naming.get("fileNameRule") or DEFAULT_NAMING["fileNameRule"]), project, track)

    out_dir = Path(out_dir).resolve() if out_dir else workspace / "08_提交包"
    if out_dir.exists():
        shutil.rmtree(out_dir)   # 提交包是纯生成物，每次重建以免残留旧文件
    out_dir.mkdir(parents=True)

    def copy(src: Path, dst_rel: str, note: str) -> None:
        dst = out_dir / dst_rel
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)
        files.append({"src": str(src.relative_to(workspace)), "dst": dst_rel, "size": dst.stat().st_size, "note": note})

    # 1) 主文档与路演
    produced = workspace / "产出"
    book_rel = str((project.get("artifacts") or {}).get("book") or "")
    book_path = (workspace / book_rel) if book_rel else None
    if produced.is_dir():
        docs = sorted(produced.glob("*.docx")) + sorted(produced.glob("*.pdf")) + sorted(produced.glob("*.pptx"))
        main_stem = book_path.stem if book_path and book_path.exists() else None
        if main_stem is None:
            docx_files = sorted(produced.glob("*.docx"))
            main_stem = docx_files[0].stem if docx_files else None
        for f in docs:
            if f.name.startswith("~$"):
                continue
            ext = f.suffix.lower()
            if ext == ".pptx":
                dst = f"{base}-路演PPT{ext}" if len(list(produced.glob('*.pptx'))) == 1 else f"{base}-{f.stem}{ext}"
                note = "路演 PPT"
            elif f.stem == main_stem:
                dst = f"{base}{ext}"
                note = DOC_NAMES.get(track, "书面材料") + ("（PDF）" if ext == ".pdf" else "（Word 源文件）")
            else:
                dst = f"{base}-{f.stem}{ext}"
                note = "其他产出"
            copy(f, dst, note)
        if not any(f["dst"].endswith(".pdf") for f in files):
            missing.append("产出/*.pdf（书面材料 PDF 版）")
        if not any(f["dst"].endswith(".docx") for f in files):
            missing.append("产出/*.docx（书面材料）")
    else:
        missing.append("产出/（书面材料与路演 PPT）")

    # 2) 财务模型、路演文稿
    for rel, dst, note in (
        ("03_财务/财务模型.xlsx", f"{base}-财务模型.xlsx", "财务模型（含公式）"),
        ("06_路演/讲稿.md", f"{base}-讲稿.md", "10 分钟讲稿"),
        ("06_路演/答辩题库.md", f"{base}-答辩题库.md", "答辩题库"),
        ("06_路演/视频分镜.md", f"{base}-视频分镜.md", "视频分镜脚本"),
    ):
        src = workspace / rel
        if src.exists():
            copy(src, dst, note)
        else:
            missing.append(rel)

    # 3) 调研工具包与最新评审
    def copy_tree(src: Path, dst_rel: str, note: str) -> None:
        for f in sorted(src.rglob("*")):
            if f.is_file():
                copy(f, f"{dst_rel}/{f.relative_to(src).as_posix()}", note)

    kit = workspace / "01_调研" / "调研工具包"
    if kit.is_dir():
        copy_tree(kit, "调研工具包", "问卷 / 访谈提纲 / 执行计划")
    else:
        missing.append("01_调研/调研工具包/")
    review = latest_review_dir(workspace)
    if review:
        copy_tree(review, "评审报告", f"评审材料（{review.name}）")
    else:
        missing.append("07_评审/第N轮/")

    # 4) 生成件
    decl_names = list(naming.get("declarations") or DEFAULT_NAMING["declarations"])
    for p in make_declarations(out_dir / "声明与授权", decl_names, project):
        files.append({"src": "(generated)", "dst": str(p.relative_to(out_dir)).replace("\\", "/"), "size": p.stat().st_size, "note": "声明 / 授权模板（需签字盖章）"})
    entries, marks = scan_synthetic(workspace)
    truth = out_dir / "数据真实性清单.md"
    write_truth_list(truth, entries, marks)
    files.append({"src": "(generated)", "dst": truth.name, "size": truth.stat().st_size, "note": f"示例数据 {len(entries) + len(marks)} 处"})

    pdf_pages = None
    pdf_files = [f for f in files if f["dst"].endswith(".pdf") and f["src"] != "(generated)"]
    if pdf_files:
        try:
            from .pdf import page_count
            pdf_pages = page_count(out_dir / pdf_files[0]["dst"])
        except Exception:
            pdf_pages = None
    readme = out_dir / "README.md"
    write_text(readme, _readme(project, track, naming, naming_src, base, files, missing, placeholders, pdf_pages, len(entries) + len(marks)))
    files.append({"src": "(generated)", "dst": readme.name, "size": readme.stat().st_size, "note": "文件清单与提交前检查"})

    zip_path = None
    if make_zip:
        zip_path = out_dir / f"{base}.zip"
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
            for f in sorted(out_dir.rglob("*")):
                if f.is_file() and f != zip_path:
                    zf.write(f, f.relative_to(out_dir).as_posix())
    return {
        "ok": True, "dir": str(out_dir), "zip": str(zip_path) if zip_path else None, "baseName": base, "naming": naming_src,
        "files": files, "missing": missing, "placeholders": sorted(set(placeholders)), "pdfPages": pdf_pages,
        "syntheticCount": len(entries) + len(marks), "warnings": warnings,
    }


def _readme(project: Dict[str, Any], track: str, naming: Dict[str, Any], naming_src: str, base: str, files: List[Dict[str, Any]],
            missing: List[str], placeholders: List[str], pdf_pages: Optional[int], synthetic_count: int) -> str:
    limit = naming.get("pageLimit") or 80
    lines = [f"# 提交包说明 —— {project.get('name') or base}", "",
             f"- 赛道 / 组别：{TRACK_NAMES.get(track, track)} / {project.get('group') or '—'}",
             f"- 团队：{project.get('teamName') or '—'}　　指导 / 命题企业：{(project.get('subject') or {}).get('company') or '—'}",
             f"- 命名规则：`{naming.get('fileNameRule')}`（来源：{naming_src}）→ `{base}`",
             f"- 生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}", "", "## 文件清单", "", "| 文件 | 大小 | 说明 |", "|---|---:|---|"]
    for f in files:
        size = f["size"]
        human = f"{size / 1024 / 1024:.1f} MB" if size > 1024 * 1024 else f"{max(1, round(size / 1024))} KB"
        lines.append(f"| {f['dst']} | {human} | {f['note']} |")
    if missing:
        lines += ["", "## 缺失项（未找到，提交前请补齐）", ""] + [f"- {m}" for m in missing]
    lines += ["", "## 需要人工完成的事项", ""]
    if placeholders:
        for ph in sorted(set(placeholders)):
            lines.append(f"- 文件名中的「{ph}」是占位符：匿名赛道按组委会分配的编号填写，否则替换为真实名称（**正文与封面仍不得出现院校信息**）。")
    lines += ["- 打印《声明与授权》中的三份文件，手写签名并按指印 / 加盖企业公章，扫描为 PDF 后放入提交包。",
              f"- 《数据真实性清单.md》列出的 {synthetic_count} 处示例数据须替换为真实数据后重新生成全部材料。",
              "", "## 提交前检查清单", "",
              "- [ ] 匿名：正文、封面、页眉页脚、图片、PPT 母版、视频水印均无院校名称 / 校徽 / 姓名 / 联系方式",
              f"- [ ] 页数：书面材料除封面目录外 ≤ {limit} 页" + (f"（当前 PDF 共 {pdf_pages} 页）" if pdf_pages else "（请用 `python -m bizkit.pdf --pages` 核对）"),
              "- [ ] 格式：PDF、A4 竖版、字体嵌入；文件大小符合通知要求（通常 ≤ 30 MB）",
              f"- [ ] 命名：按「{naming.get('fileNameRule')}」命名，与报名信息一致",
              "- [ ] 签名指印：原创性声明、版权声明已全员签字按指印；企业授权书已盖章",
              f"- [ ] 视频：{naming.get('videoPeople', '3–5 人')}真人出镜，时长 ≤ {naming.get('videoMinutes', 10)} 分钟，无院校信息",
              "- [ ] 一致性：正文、PPT、讲稿、财务模型中的关键数字一致（运行 `python -m bizkit.checks`）",
              "- [ ] 示例数据：数据真实性清单全部清零", ""]
    return "\n".join(lines)


def main(argv: Optional[Sequence[str]] = None) -> int:
    force_utf8_stdout()
    ap = argparse.ArgumentParser(description="生成提交包")
    ap.add_argument("--track", default="", help="赛道 id（默认读 project.json）")
    ap.add_argument("--workspace", default=".", help="工作区根目录")
    ap.add_argument("--out", default=None, help="输出目录（默认 <workspace>/08_提交包）")
    ap.add_argument("--no-zip", action="store_true")
    args = ap.parse_args(argv)
    try:
        result = package(Path(args.workspace), args.track, Path(args.out) if args.out else None, make_zip=not args.no_zip)
    except Exception as e:
        import traceback
        print(json.dumps({"ok": False, "error": f"{type(e).__name__}: {e}", "trace": traceback.format_exc(limit=5)}, ensure_ascii=False))
        return 1
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
