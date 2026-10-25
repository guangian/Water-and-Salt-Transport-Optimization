"""assumptions.json → 财务模型.xlsx + numbers.json + 财务摘要.md + 财务图表.charts.json（契约 §7.4 / §7.5，技能 fin-model）。

用法：
    python -m bizkit.finmodel 03_财务/assumptions.json [--out 03_财务]

两种模式
--------
* ``mode = "startup"``：三表模型。第一年按月（12 列）、之后按年；收入由驱动因子推出
  （volume_price / clients_arpu / subscribers_arpu（含流失）/ funnel），成本分 COGS、变动费用、固定、
  营销（费率 × 收入 与 minAnnual 取大）、人力（岗位 × 人数 × 月薪 × (1+社保)）、折旧（CAPEX 直线）、其他运营费用；
  营运资本按周转天数；所得税按累计口径（亏损结转、不退税）；股权融资按期注入。
  另算盈亏平衡、单位经济、三情景、单因子敏感性、DCF / PS / VC 三法估值与融资一致性。
* ``mode = "budget"``：策划案预算表（阶段 × 项目：单价 × 数量 = 金额、小计、机动金、总计与 totalCap 对比）、排期、KPI、可选 ROI。

「公式 + 数值」方案
------------------
全部数字先在 Python 里算出（本模块是唯一的计算口径，numbers.json / 摘要 / 图表都来自它）；
写 Excel 时：假设 / 驱动 / 收入 / 成本 / 人力 页写**数值**（输入蓝字浅黄底），
利润表 / 现金流量表 / 资产负债表 / 检查 页写 **Excel 公式**（同表引用黑字、跨表引用绿字），
每张报表右侧另有「数值校验」列写 Python 的年度值，检查页用公式比较两者并汇总为 Master Check（全部 0 = 通过）。
这样 Excel 打开后既能追溯公式，又能用 Python 值核对；openpyxl 本身不计算公式，所以账本与摘要一律取 Python 值。

金额口径：assumptions 用元；Excel 报表与账本按 ``unit``（默认 10000 = 万元）换算；驱动页的单价 / ARPU 仍用元。
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from statistics import median
from typing import Any, Dict, List, Optional, Sequence, Tuple

from . import theme as theme_mod
from ._util import find_workspace_root, force_utf8_stdout, log, read_json, write_json, write_text

__all__ = ["compute_startup", "run", "main", "StartupResult", "FinModelError"]

Number = float
EPS = 1e-9

DRIVER_NAMES: Dict[str, str] = {
    "volume_price": "销量 × 单价",
    "clients_arpu": "客户数 × ARPU",
    "subscribers_arpu": "订阅数 × ARPU（含流失）",
    "funnel": "流量 × 转化率 × 客单价",
}
# 敏感性驱动因子：键 → （中文名，说明）
SENS_DRIVERS: Dict[str, Tuple[str, str]] = {
    "price": ("单价 / 客单价", "全部收入流的单价、ARPU、客单价同比例变动"),
    "volume": ("销量 / 客户数", "需求侧数量驱动（销量、客户数、流量、新增订阅）同比例变动"),
    "cogsRatio": ("成本率（COGS）", "各收入流 cogsRatio 同比例变动"),
    "cac": ("获客成本（CAC）", "获客成本变动 ⇒ 获取同样客户的营销支出同比例变动；同时影响 LTV:CAC 与回收期"),
    "fixedCost": ("固定成本", "固定成本项同比例变动"),
    "conversion": ("转化率", "funnel 驱动的转化率同比例变动"),
    "marketing": ("营销费用", "营销费用同比例变动"),
    "salary": ("人力成本", "人力成本同比例变动"),
    "churn": ("流失率", "订阅流失率同比例变动"),
    "receivableDays": ("应收账期", "应收周转天数同比例变动（只影响现金）"),
}


class FinModelError(ValueError):
    """assumptions.json 不符合约定；消息面向模型 / 用户，中文可读。"""


# ----------------------------------------------------------------------------------------------
# 基础工具
# ----------------------------------------------------------------------------------------------

def _num(v: Any, default: float = 0.0, name: str = "") -> float:
    if v is None or v == "":
        return float(default)
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        raise FinModelError(f"「{name or '数值'}」必须是数字，收到 {v!r}")
    return float(v)


def _list(v: Any, n: int, default: float = 0.0, name: str = "") -> List[float]:
    """标量或数组 → 长度 n 的数组（不足补最后一个值，没有值补 default）。"""
    if v is None:
        return [float(default)] * n
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return [float(v)] * n
    if not isinstance(v, (list, tuple)):
        raise FinModelError(f"「{name}」应为数字或数组")
    vals = [_num(x, default, name) for x in v]
    if not vals:
        return [float(default)] * n
    if len(vals) < n:
        vals = vals + [vals[-1]] * (n - len(vals))
    return vals[:n]


def _cum_factors(growth: Any, years: int, name: str) -> List[float]:
    """年度变化率（数组：第 1 年忽略；标量：每年同一增速）→ 相对第 1 年的累计倍数。"""
    g = _list(growth, years, 0.0, name)
    if isinstance(growth, (int, float)) and not isinstance(growth, bool):
        g = [0.0] + [float(growth)] * (years - 1)
    out = [1.0]
    for y in range(1, years):
        out.append(out[-1] * (1.0 + g[y]))
    return out


def _ramp_weights(n: int = 12) -> List[float]:
    """默认首年 S 型爬坡权重（和为 1），用于只给了年度总量的驱动。"""
    raw = [1.0 / (1.0 + math.exp(-(i + 1 - (n + 1) / 2.0) / (n / 6.0))) for i in range(n)]
    s = sum(raw)
    return [r / s for r in raw]


def _cumsum(xs: Sequence[float]) -> List[float]:
    out: List[float] = []
    acc = 0.0
    for x in xs:
        acc += x
        out.append(acc)
    return out


def _add(*series: Sequence[float]) -> List[float]:
    n = len(series[0])
    return [sum(s[i] for s in series) for i in range(n)]


def _scale(xs: Sequence[float], k: float) -> List[float]:
    return [x * k for x in xs]


def _mul(a: Sequence[float], b: Sequence[float]) -> List[float]:
    return [x * y for x, y in zip(a, b)]


def _sub(a: Sequence[float], b: Sequence[float]) -> List[float]:
    return [x - y for x, y in zip(a, b)]


def _safe_div(a: float, b: float, default: float = 0.0) -> float:
    return a / b if abs(b) > EPS else default


# ----------------------------------------------------------------------------------------------
# 时间轴
# ----------------------------------------------------------------------------------------------

@dataclass(frozen=True)
class Period:
    idx: int                 # 0 起
    year: int                # 模型年（1 起）
    month: Optional[int]     # 月度期 1–12；年度期 None
    label: str               # "2026-01" / "2027"
    months: int              # 1 或 12
    days: int                # 30 或 365


def build_timeline(start_year: int, years: int, monthly_first: bool) -> List[Period]:
    periods: List[Period] = []
    if monthly_first:
        for m in range(1, 13):
            periods.append(Period(len(periods), 1, m, f"{start_year}-{m:02d}", 1, 30))
        first_annual = 2
    else:
        first_annual = 1
    for y in range(first_annual, years + 1):
        periods.append(Period(len(periods), y, None, f"{start_year + y - 1}", 12, 365))
    return periods


def _year_index(periods: Sequence[Period], years: int) -> List[List[int]]:
    return [[p.idx for p in periods if p.year == y + 1] for y in range(years)]


def _annual_flow(xs: Sequence[float], yidx: Sequence[Sequence[int]]) -> List[float]:
    return [sum(xs[i] for i in ids) for ids in yidx]


def _annual_stock(xs: Sequence[float], yidx: Sequence[Sequence[int]]) -> List[float]:
    return [xs[ids[-1]] for ids in yidx]


# ----------------------------------------------------------------------------------------------
# 收入流
# ----------------------------------------------------------------------------------------------

@dataclass
class StreamResult:
    name: str
    driver: str
    driver_name: str
    source: str
    unit_name: str                      # 数量单位（件 / 家 / 订阅·月 / 单）
    volume_label: str
    price_label: str
    volume: List[float]                 # 每期数量驱动
    price: List[float]                  # 每期单价（元）
    revenue: List[float]                # 元
    cogs: List[float]                   # 元
    cogs_ratio: float
    base_price: float                   # 首年单价（元），单位经济用
    extra: List[Tuple[str, List[float], str]] = field(default_factory=list)   # 驱动页附加行：(名称, 序列, 格式)


def _first_year_profile(obj: Dict[str, Any], key_total: str, periods: Sequence[Period], years: int, name: str) -> Tuple[List[float], float]:
    """首年数量：优先 y1Monthly（12 个月），否则 y1 总量按 S 型权重分到各月；返回（首年各月数量, 首年总量）。"""
    monthly = obj.get("y1Monthly")
    if monthly is not None:
        if not isinstance(monthly, list) or len(monthly) != 12:
            raise FinModelError(f"「{name}.y1Monthly」必须是 12 个数字")
        m = [_num(x, 0, f"{name}.y1Monthly") for x in monthly]
        return m, sum(m)
    total = _num(obj.get(key_total, obj.get("y1")), 0, f"{name}.{key_total}")
    w = _ramp_weights(12)
    return [total * wi for wi in w], total


def _quantity_series(obj: Dict[str, Any], periods: Sequence[Period], years: int, name: str, mult: float) -> List[float]:
    """数量驱动：首年（月度 / 年度）+ yoyGrowth 推后续年度；mult 为需求乘数（情景 / 敏感性）。"""
    if not isinstance(obj, dict):
        raise FinModelError(f"收入流缺少「{name}」对象")
    m12, total = _first_year_profile(obj, "y1", periods, years, name)
    fac = _cum_factors(obj.get("yoyGrowth"), years, f"{name}.yoyGrowth")
    out: List[float] = []
    for p in periods:
        if p.month is not None:
            out.append(m12[p.month - 1] * mult)
        else:
            out.append(total * fac[p.year - 1] * mult)
    return out


def _price_series(obj: Any, periods: Sequence[Period], years: int, name: str, mult: float) -> Tuple[List[float], float]:
    if isinstance(obj, (int, float)) and not isinstance(obj, bool):
        obj = {"base": obj}
    if not isinstance(obj, dict):
        raise FinModelError(f"收入流缺少「{name}」对象（{{base, yoyChange}}）")
    base = _num(obj.get("base"), 0, f"{name}.base")
    fac = _cum_factors(obj.get("yoyChange"), years, f"{name}.yoyChange")
    return [base * fac[p.year - 1] * mult for p in periods], base * mult


def compute_stream(s: Dict[str, Any], periods: Sequence[Period], years: int, demand: float, price_mult: float,
                   cogs_mult: float, conv_mult: float, churn_mult: float) -> StreamResult:
    if not isinstance(s, dict):
        raise FinModelError("revenueStreams 的每一项必须是对象")
    name = str(s.get("name") or "收入流")
    driver = str(s.get("driver") or "volume_price")
    if driver not in DRIVER_NAMES:
        raise FinModelError(f"收入流「{name}」的 driver「{driver}」不支持，可选：{', '.join(DRIVER_NAMES)}")
    cogs_ratio = min(0.999, max(0.0, _num(s.get("cogsRatio"), 0, f"{name}.cogsRatio") * cogs_mult))
    source = str(s.get("source") or "")
    extra: List[Tuple[str, List[float], str]] = []

    if driver == "volume_price":
        volume = _quantity_series(s.get("volume"), periods, years, f"{name}.volume", demand)
        price, base_price = _price_series(s.get("price"), periods, years, f"{name}.price", price_mult)
        unit_name = str(s.get("unitName") or "件")
        vol_label, price_label = f"销量（{unit_name}）", "单价（元）"
    elif driver == "clients_arpu":
        volume = _quantity_series(s.get("clients"), periods, years, f"{name}.clients", demand)
        price, base_price = _price_series(s.get("arpu"), periods, years, f"{name}.arpu", price_mult)
        unit_name = str(s.get("unitName") or "家")
        vol_label, price_label = f"签约客户数（{unit_name}，当期确认）", "ARPU（元/客户/年）"
    elif driver == "funnel":
        traffic = _quantity_series(s.get("traffic"), periods, years, f"{name}.traffic", demand)
        conv, _ = _price_series(s.get("conversion"), periods, years, f"{name}.conversion", conv_mult)
        conv = [min(1.0, c) for c in conv]
        price, base_price = _price_series(s.get("aov"), periods, years, f"{name}.aov", price_mult)
        volume = _mul(traffic, conv)
        unit_name = str(s.get("unitName") or "单")
        vol_label, price_label = f"成交订单（{unit_name}）", "客单价（元）"
        extra = [("流量（人次）", traffic, "int"), ("转化率", conv, "pct")]
    else:  # subscribers_arpu：按月模拟活跃订阅
        sub = s.get("subscribers")
        if not isinstance(sub, dict):
            raise FinModelError(f"收入流「{name}」缺少 subscribers 对象（y1Monthly | y1, yoyGrowth, churnMonthly, opening）")
        churn = min(1.0, max(0.0, _num(sub.get("churnMonthly"), 0.0, f"{name}.churnMonthly") * churn_mult))
        m12, total = _first_year_profile(sub, "y1", periods, years, f"{name}.subscribers")
        fac = _cum_factors(sub.get("yoyGrowth"), years, f"{name}.subscribers.yoyGrowth")
        active = _num(sub.get("opening"), 0.0, f"{name}.subscribers.opening")
        price, base_price = _price_series(s.get("arpu"), periods, years, f"{name}.arpu", price_mult)
        volume, ending, new_series = [], [], []
        for p in periods:
            if p.month is not None:
                news = [m12[p.month - 1] * demand]
            else:
                news = [total * fac[p.year - 1] * demand / 12.0] * 12
            sub_months = 0.0
            for nw in news:
                active = active * (1.0 - churn) + nw
                sub_months += active
            volume.append(sub_months)
            ending.append(active)
            new_series.append(sum(news))
        unit_name = str(s.get("unitName") or "订阅·月")
        vol_label, price_label = "订阅·月（活跃订阅数之和）", "ARPU（元/订阅/月）"
        extra = [("新增订阅", new_series, "int"), ("期末活跃订阅", ending, "int"), ("月流失率", [churn] * len(periods), "pct")]

    revenue = _mul(volume, price)
    cogs = _scale(revenue, cogs_ratio)
    return StreamResult(name=name, driver=driver, driver_name=DRIVER_NAMES[driver], source=source, unit_name=unit_name,
                        volume_label=vol_label, price_label=price_label, volume=volume, price=price, revenue=revenue,
                        cogs=cogs, cogs_ratio=cogs_ratio, base_price=base_price, extra=extra)


# ----------------------------------------------------------------------------------------------
# 三表模型
# ----------------------------------------------------------------------------------------------

@dataclass
class HeadcountResult:
    role: str
    heads: List[float]          # 每期人数
    salary: List[float]         # 每期月薪（元）
    social_rate: float
    cost: List[float]           # 每期人力成本（元）


@dataclass
class CapexResult:
    name: str
    year: int
    month: int
    amount: float
    dep_years: int
    outflow: List[float]
    dep: List[float]


@dataclass
class Breakeven:
    month: Optional[int]                    # 累计经营利润首次转正的月份（1 起）；None = 预测期内未达到
    first_positive_period: Optional[int]    # 单期经营利润首次为正的期序号
    cum_ebit: List[float]
    year_table: List[Dict[str, float]]      # 每年：收入 / 变动成本 / 边际贡献率 / 固定性成本 / 保本收入 / 安全边际
    primary_stream: str
    chart_year: int                         # 盈亏平衡图采用的年份（1 起）
    fixed_cost: float                       # 图：主力产品线需覆盖的固定性成本（元）
    unit_price: float
    unit_var_cost: float
    be_units: float
    actual_units: float
    unit_name: str


@dataclass
class StartupResult:
    """startup 模式的全部计算结果（金额单位：元）。"""
    a: Dict[str, Any]
    periods: List[Period]
    years: int
    start_year: int
    unit: float
    demand: float
    yidx: List[List[int]]
    streams: List[StreamResult]
    fixed_items: List[Tuple[str, List[float]]]
    variable_items: List[Tuple[str, float, List[float]]]
    headcount: List[HeadcountResult]
    capex: List[CapexResult]
    s: Dict[str, List[float]]            # 每期序列
    annual: Dict[str, List[float]]       # 年度序列（流量求和、存量取年末）
    checks: Dict[str, float]
    breakeven: Breakeven
    warnings: List[str] = field(default_factory=list)

    # 便捷取值 ------------------------------------------------------------------------------
    def year_labels(self) -> List[str]:
        return [str(self.start_year + y) for y in range(self.years)]

    @property
    def min_cash(self) -> float:
        """现金最低点（不含融资，含期初现金）——用于与融资额比较。"""
        return self.s["min_cash_prefin"][0]

    @property
    def min_cash_post(self) -> float:
        return min([self.a_opening()] + self.s["cash"])

    def a_opening(self) -> float:
        return _num(self.a.get("openingCash"), 0.0, "openingCash")


def _marketing_series(costs: Dict[str, Any], annual_rev: Sequence[float], periods: Sequence[Period], years: int, mult: float) -> Tuple[List[float], List[float]]:
    mk = costs.get("marketing") or {}
    if not isinstance(mk, dict):
        raise FinModelError("costs.marketing 应为对象 {pctOfRevenue, minAnnual}")
    pct = _list(mk.get("pctOfRevenue"), years, 0.0, "marketing.pctOfRevenue")
    min_annual = _num(mk.get("minAnnual"), 0.0, "marketing.minAnnual")
    annual = [max(pct[y] * annual_rev[y], min_annual) * mult for y in range(years)]
    per_period = []
    for p in periods:
        per_period.append(annual[p.year - 1] / 12.0 if p.month is not None else annual[p.year - 1])
    return per_period, pct


def compute_startup(a: Dict[str, Any], demand: float = 1.0, tweak: Optional[Dict[str, float]] = None) -> StartupResult:
    """按 assumptions 计算 startup 三表模型。

    demand：需求乘数（三情景）；tweak：{驱动键: 变动比例}（敏感性），键见 SENS_DRIVERS。
    """
    tw = dict(tweak or {})
    fac = lambda k: 1.0 + tw.get(k, 0.0)  # noqa: E731
    warnings: List[str] = []

    years = int(_num(a.get("years"), 5, "years"))
    if years < 1 or years > 10:
        raise FinModelError("years 应在 1–10 之间")
    start_year = int(_num(a.get("startYear"), date.today().year, "startYear"))
    monthly_first = bool(a.get("monthlyFirstYear", True))
    unit = _num(a.get("unit"), 10000.0, "unit") or 10000.0
    periods = build_timeline(start_year, years, monthly_first)
    n = len(periods)
    yidx = _year_index(periods, years)

    raw_streams = a.get("revenueStreams")
    if not isinstance(raw_streams, list) or not raw_streams:
        raise FinModelError("revenueStreams 不能为空")
    demand_mult = demand * fac("volume")
    streams = [compute_stream(s, periods, years, demand_mult, fac("price"), fac("cogsRatio"), fac("conversion"), fac("churn")) for s in raw_streams]
    revenue = _add(*[st.revenue for st in streams])
    cogs = _add(*[st.cogs for st in streams])
    annual_rev = _annual_flow(revenue, yidx)

    costs = a.get("costs") or {}
    if not isinstance(costs, dict):
        raise FinModelError("costs 应为对象")
    # 固定成本
    fixed_items: List[Tuple[str, List[float]]] = []
    for it in costs.get("fixed") or []:
        nm = str(it.get("name") or "固定成本")
        monthly = _num(it.get("monthly"), 0.0, f"fixed.{nm}.monthly")
        cf = _cum_factors(it.get("yoyChange"), years, f"fixed.{nm}.yoyChange")
        fixed_items.append((nm, [monthly * p.months * cf[p.year - 1] * fac("fixedCost") for p in periods]))
    fixed_total = _add(*[v for _, v in fixed_items]) if fixed_items else [0.0] * n
    # 变动费用（占收入比）
    variable_items: List[Tuple[str, float, List[float]]] = []
    for it in costs.get("variablePctOfRevenue") or []:
        nm = str(it.get("name") or "变动费用")
        pct = _num(it.get("pct"), 0.0, f"variable.{nm}.pct")
        variable_items.append((nm, pct, _scale(revenue, pct)))
    variable_total = _add(*[v for _, _, v in variable_items]) if variable_items else [0.0] * n
    var_pct_total = sum(p for _, p, _ in variable_items)
    # 营销
    marketing, marketing_pct = _marketing_series(costs, annual_rev, periods, years, fac("marketing") * fac("cac"))
    # 人力
    headcount: List[HeadcountResult] = []
    for it in costs.get("headcount") or []:
        role = str(it.get("role") or "岗位")
        heads_y = _list(it.get("y"), years, 0.0, f"headcount.{role}.y")
        sal_fac = _cum_factors(it.get("salaryYoyChange"), years, f"headcount.{role}.salaryYoyChange")
        base_sal = _num(it.get("monthlySalary"), 0.0, f"headcount.{role}.monthlySalary")
        social = _num(it.get("socialRate"), 0.0, f"headcount.{role}.socialRate")
        m12 = it.get("y1Monthly")
        heads, sal, cost = [], [], []
        for p in periods:
            h = heads_y[p.year - 1]
            if p.month is not None and isinstance(m12, list) and len(m12) == 12:
                h = _num(m12[p.month - 1], 0.0, f"headcount.{role}.y1Monthly")
            s_ = base_sal * sal_fac[p.year - 1]
            heads.append(h)
            sal.append(s_)
            cost.append(h * s_ * (1.0 + social) * p.months * fac("salary"))
        headcount.append(HeadcountResult(role, heads, sal, social, cost))
    hc_total = _add(*[h.cost for h in headcount]) if headcount else [0.0] * n
    # CAPEX 与折旧（按全局月份排程后汇总到期）
    total_months = years * 12
    capex_items: List[CapexResult] = []
    for it in costs.get("capex") or []:
        nm = str(it.get("name") or "资本开支")
        cy = int(_num(it.get("year"), 1, f"capex.{nm}.year"))
        cm = int(_num(it.get("month"), 1, f"capex.{nm}.month"))
        amt = _num(it.get("amount"), 0.0, f"capex.{nm}.amount")
        dy = max(1, int(_num(it.get("depreciationYears"), 5, f"capex.{nm}.depreciationYears")))
        if not (1 <= cy <= years) or not (1 <= cm <= 12):
            warnings.append(f"CAPEX「{nm}」的 year/month 超出预测期，已忽略")
            continue
        gm0 = (cy - 1) * 12 + (cm - 1)
        dep_m = [0.0] * total_months
        per = amt / (dy * 12)
        for g in range(gm0, min(total_months, gm0 + dy * 12)):
            dep_m[g] = per
        out_p = [0.0] * n
        dep_p = [0.0] * n
        for p in periods:
            if p.month is not None:
                g = (p.year - 1) * 12 + (p.month - 1)
                dep_p[p.idx] = dep_m[g]
                if g == gm0:
                    out_p[p.idx] = amt
            else:
                gs = range((p.year - 1) * 12, p.year * 12)
                dep_p[p.idx] = sum(dep_m[g] for g in gs)
                if gm0 in gs:
                    out_p[p.idx] = amt
        capex_items.append(CapexResult(nm, cy, cm, amt, dy, out_p, dep_p))
    capex_total = _add(*[c.outflow for c in capex_items]) if capex_items else [0.0] * n
    dep_total = _add(*[c.dep for c in capex_items]) if capex_items else [0.0] * n
    # 其他运营费用
    other_pct = _num(costs.get("otherOpexPctOfRevenue"), 0.0, "otherOpexPctOfRevenue")
    other_opex = _scale(revenue, other_pct)

    gross_profit = _sub(revenue, cogs)
    opex_total = _add(variable_total, marketing, hc_total, fixed_total, other_opex)
    ebitda = _sub(gross_profit, opex_total)
    ebit = _sub(ebitda, dep_total)
    interest = [0.0] * n
    ebt = _sub(ebit, interest)

    # 所得税：累计口径（亏损结转、不退税）——tax_t = rate × (base_t − base_{t−1})，base_t = max(base_{t−1}, 累计利润总额_t)
    tax_cfg = a.get("tax") or {}
    rate = _num(tax_cfg.get("incomeTaxRate"), 0.25, "tax.incomeTaxRate")
    cum_ebt = _cumsum(ebt)
    tax_base: List[float] = []
    tax: List[float] = []
    prev_base = 0.0
    for c in cum_ebt:
        b = max(prev_base, c)
        tax.append(rate * (b - prev_base))
        tax_base.append(b)
        prev_base = b
    net_profit = _sub(ebt, tax)

    # 营运资本（期末余额 = 当期流量 / 期间天数 × 周转天数）
    wc = a.get("workingCapital") or {}
    dso = _num(wc.get("receivableDays"), 0.0, "receivableDays") * fac("receivableDays")
    dio = _num(wc.get("inventoryDays"), 0.0, "inventoryDays")
    dpo = _num(wc.get("payableDays"), 0.0, "payableDays")
    ar = [revenue[p.idx] / p.days * dso for p in periods]
    inv = [cogs[p.idx] / p.days * dio for p in periods]
    ap = [cogs[p.idx] / p.days * dpo for p in periods]

    def _delta(xs: Sequence[float]) -> List[float]:
        return [xs[i] - (xs[i - 1] if i > 0 else 0.0) for i in range(len(xs))]

    d_ar, d_inv, d_ap = _delta(ar), _delta(inv), _delta(ap)
    cfo = [net_profit[i] + dep_total[i] - d_ar[i] - d_inv[i] + d_ap[i] for i in range(n)]
    cfi = _scale(capex_total, -1.0)

    # 融资注入
    financing = a.get("financing") or []
    fin_in = [0.0] * n
    for rd in financing:
        fy = int(_num(rd.get("year"), 1, "financing.year"))
        fm = int(_num(rd.get("month"), 1, "financing.month"))
        amt = _num(rd.get("amount"), 0.0, "financing.amount")
        hit = [p for p in periods if p.year == fy and (p.month is None or p.month == fm)]
        if not hit:
            warnings.append(f"融资「{rd.get('round', '')}」的时间（第 {fy} 年第 {fm} 月）不在预测期内，已忽略")
            continue
        fin_in[hit[0].idx] += amt
    cff = list(fin_in)

    opening_cash = _num(a.get("openingCash"), 0.0, "openingCash")
    net_change = _add(cfo, cfi, cff)
    cash: List[float] = []
    bal = opening_cash
    for x in net_change:
        bal += x
        cash.append(bal)
    cum_fin = _cumsum(fin_in)
    cash_prefin = _sub(cash, cum_fin)
    min_cash_prefin = min([opening_cash] + cash_prefin)

    # 资产负债表
    fa_gross = _cumsum(capex_total)
    acc_dep = _cumsum(dep_total)
    fa_net = _sub(fa_gross, acc_dep)
    assets = _add(cash, ar, inv, fa_net)
    liabilities = list(ap)
    paid_in = [opening_cash + c for c in cum_fin]
    retained = _cumsum(net_profit)
    equity = _add(paid_in, retained)
    liab_equity = _add(liabilities, equity)

    # 勾稽检查（独立两条路径：现金流量表 vs 资产负债表倒算）
    bs_cash_indep = [liab_equity[i] - ar[i] - inv[i] - fa_net[i] for i in range(n)]
    checks = {
        "balanceSheet": max(abs(assets[i] - liab_equity[i]) for i in range(n)),
        "cashflow": max(abs(cash[i] - bs_cash_indep[i]) for i in range(n)),
        "retained": abs(retained[-1] - sum(net_profit)),
    }

    s: Dict[str, List[float]] = {
        "revenue": revenue, "cogs": cogs, "gross_profit": gross_profit,
        "variable": variable_total, "marketing": marketing, "headcount": hc_total, "fixed": fixed_total,
        "other_opex": other_opex, "opex_total": opex_total, "ebitda": ebitda, "dep": dep_total, "ebit": ebit,
        "interest": interest, "ebt": ebt, "cum_ebt": cum_ebt, "tax_base": tax_base, "tax": tax, "net_profit": net_profit,
        "ar": ar, "inv": inv, "ap": ap, "d_ar": d_ar, "d_inv": d_inv, "d_ap": d_ap,
        "cfo": cfo, "capex": capex_total, "cfi": cfi, "fin_in": fin_in, "cff": cff, "net_change": net_change,
        "cash": cash, "cum_fin": cum_fin, "cash_prefin": cash_prefin, "min_cash_prefin": [min_cash_prefin],
        "fa_gross": fa_gross, "acc_dep": acc_dep, "fa_net": fa_net, "assets": assets, "liabilities": liabilities,
        "paid_in": paid_in, "retained": retained, "equity": equity, "liab_equity": liab_equity,
        "marketing_pct": [marketing_pct[p.year - 1] for p in periods],
        "fcf": _add(cfo, cfi),
    }
    stock_keys = {"ar", "inv", "ap", "cash", "cum_fin", "cash_prefin", "fa_gross", "acc_dep", "fa_net", "assets", "liabilities",
                  "paid_in", "retained", "equity", "liab_equity", "cum_ebt", "tax_base"}
    annual: Dict[str, List[float]] = {}
    for k, v in s.items():
        if len(v) != n:
            continue
        annual[k] = _annual_stock(v, yidx) if k in stock_keys else _annual_flow(v, yidx)
    annual["gross_margin"] = [_safe_div(annual["gross_profit"][y], annual["revenue"][y]) for y in range(years)]
    annual["net_margin"] = [_safe_div(annual["net_profit"][y], annual["revenue"][y]) for y in range(years)]
    annual["marketing_pct"] = list(marketing_pct)

    res = StartupResult(a=a, periods=periods, years=years, start_year=start_year, unit=unit, demand=demand, yidx=yidx,
                        streams=streams, fixed_items=fixed_items, variable_items=variable_items, headcount=headcount,
                        capex=capex_items, s=s, annual=annual, checks=checks, breakeven=None,  # type: ignore[arg-type]
                        warnings=warnings)
    res.breakeven = compute_breakeven(res, var_pct_total, other_pct)
    return res


# ----------------------------------------------------------------------------------------------
# 盈亏平衡
# ----------------------------------------------------------------------------------------------

def _primary_stream(res: StartupResult) -> StreamResult:
    """保本销量 / 单位经济用的主力产品线：优先有「单价」语义的流中收入最大者。"""
    pref = [st for st in res.streams if st.driver in ("volume_price", "funnel")] or list(res.streams)
    return max(pref, key=lambda st: sum(st.revenue))


def compute_breakeven(res: StartupResult, var_pct_total: float, other_pct: float) -> Breakeven:
    s, years = res.s, res.years
    cum = _cumsum(s["ebit"])
    month: Optional[int] = None
    for p in res.periods:
        if cum[p.idx] > EPS:
            if p.month is not None:
                month = p.month
            else:
                prev = cum[p.idx - 1] if p.idx > 0 else 0.0
                need = -prev if prev < 0 else 0.0
                frac = need / s["ebit"][p.idx] if s["ebit"][p.idx] > EPS else 0.0
                month = (p.year - 1) * 12 + max(1, min(12, math.ceil(frac * 12)))
            break
    first_pos = next((p.idx for p in res.periods if s["ebit"][p.idx] > EPS), None)

    an = res.annual
    table: List[Dict[str, float]] = []
    for y in range(years):
        rev = an["revenue"][y]
        var_cost = an["cogs"][y] + an["variable"][y] + an["other_opex"][y]
        fixed_like = an["fixed"][y] + an["marketing"][y] + an["headcount"][y] + an["dep"][y]
        cm = rev - var_cost
        cmr = _safe_div(cm, rev)
        be_rev = _safe_div(fixed_like, cmr, float("inf")) if cmr > EPS else float("inf")
        margin_safety = _safe_div(rev - be_rev, rev) if math.isfinite(be_rev) else float("-inf")
        table.append({"revenue": rev, "variable_cost": var_cost, "contribution": cm, "cm_ratio": cmr,
                      "fixed_like": fixed_like, "be_revenue": be_rev, "margin_of_safety": margin_safety,
                      "cash_be_revenue": _safe_div(fixed_like - an["dep"][y], cmr, float("inf")) if cmr > EPS else float("inf")})

    primary = _primary_stream(res)
    if month is not None:
        chart_year = min(years, max(1, math.ceil(month / 12)))
    else:
        chart_year = years
    y = chart_year - 1
    ids = res.yidx[y]
    rev_p = sum(primary.revenue[i] for i in ids)
    vol_p = sum(primary.volume[i] for i in ids)
    price = _safe_div(rev_p, vol_p, primary.base_price)
    unit_var = price * (primary.cogs_ratio + var_pct_total + other_pct)
    other_contrib = 0.0
    for st in res.streams:
        if st is primary:
            continue
        other_contrib += sum(st.revenue[i] for i in ids) * (1.0 - st.cogs_ratio - var_pct_total - other_pct)
    fixed_like = table[y]["fixed_like"]
    to_cover = fixed_like - other_contrib
    be_units = max(0.0, to_cover) / (price - unit_var) if price - unit_var > EPS else float("inf")
    return Breakeven(month=month, first_positive_period=first_pos, cum_ebit=cum, year_table=table, primary_stream=primary.name,
                     chart_year=chart_year, fixed_cost=to_cover if to_cover > 0 else fixed_like, unit_price=price,
                     unit_var_cost=unit_var, be_units=be_units, actual_units=vol_p, unit_name=primary.unit_name)


# ----------------------------------------------------------------------------------------------
# 单位经济 / 情景 / 敏感性 / 估值 / 融资
# ----------------------------------------------------------------------------------------------

@dataclass
class UnitEconomics:
    cac: float
    price: float
    margin: float
    lifetime_orders: float
    lifetime_months: float
    ltv: float
    ltv_cac: float
    payback_months: float
    first_order_gp: float
    cac_source: str
    notes: List[str]


def compute_unit_economics(res: StartupResult, tweak: Optional[Dict[str, float]] = None) -> UnitEconomics:
    tw = tweak or {}
    ue = res.a.get("unitEconomics") or {}
    primary = _primary_stream(res)
    notes: List[str] = []
    y1 = res.yidx[0]
    price = _num(ue.get("price"), 0.0, "unitEconomics.price")
    if price <= 0:
        rev1, vol1 = sum(primary.revenue[i] for i in y1), sum(primary.volume[i] for i in y1)
        price = _safe_div(rev1, vol1, primary.base_price)
        if primary.driver == "subscribers_arpu":
            price = primary.base_price  # 月 ARPU
    margin = _num(ue.get("grossMarginRef"), 0.0, "unitEconomics.grossMarginRef") or (1.0 - primary.cogs_ratio)
    churn = None
    if primary.driver == "subscribers_arpu":
        sub = (res.a["revenueStreams"][res.streams.index(primary)].get("subscribers") or {})
        churn = _num(sub.get("churnMonthly"), 0.0, "churnMonthly") * (1.0 + tw.get("churn", 0.0))
    lifetime_orders = _num(ue.get("lifetimeOrders"), 0.0, "unitEconomics.lifetimeOrders")
    if lifetime_orders <= 0:
        repeat = _num(ue.get("repeatRate"), 0.0, "unitEconomics.repeatRate")
        if churn:
            lifetime_orders = 1.0 / churn
            notes.append("lifetimeOrders 取 1/月流失率（订阅模式）")
        elif 0 < repeat < 1:
            lifetime_orders = 1.0 / (1.0 - repeat)
            notes.append("lifetimeOrders = 1/(1−复购率)")
        else:
            lifetime_orders = 1.0
    lifetime_months = _num(ue.get("lifetimeMonths"), 0.0, "unitEconomics.lifetimeMonths")
    if lifetime_months <= 0:
        lifetime_months = (1.0 / churn) if churn else 24.0
        if not churn:
            notes.append("客户生命周期默认 24 个月（可用 unitEconomics.lifetimeMonths 覆盖）")
    cac = _num(ue.get("cac"), 0.0, "unitEconomics.cac")
    cac_source = "假设 unitEconomics.cac"
    if cac <= 0:
        mk1 = sum(res.s["marketing"][i] for i in y1)
        vol1 = sum(primary.volume[i] for i in y1)
        cac = _safe_div(mk1, vol1 / max(lifetime_orders, 1.0), 0.0)
        cac_source = "首年营销费用 ÷ 首年新客数（销量 ÷ lifetimeOrders）"
        notes.append("未给 unitEconomics.cac，按首年营销费用 ÷ 新客数估算")
    cac *= 1.0 + tw.get("cac", 0.0)
    price *= 1.0 + tw.get("price", 0.0)
    ltv = price * margin * lifetime_orders
    ltv_cac = _safe_div(ltv, cac, float("inf"))
    arpu_m = _num(ue.get("arpuMonthly"), 0.0, "unitEconomics.arpuMonthly")
    monthly_gp = arpu_m * margin if arpu_m > 0 else _safe_div(ltv, lifetime_months)
    payback = _safe_div(cac, monthly_gp, float("inf"))
    return UnitEconomics(cac=cac, price=price, margin=margin, lifetime_orders=lifetime_orders, lifetime_months=lifetime_months,
                         ltv=ltv, ltv_cac=ltv_cac, payback_months=payback, first_order_gp=price * margin, cac_source=cac_source, notes=notes)


@dataclass
class ScenarioRow:
    key: str
    name: str
    multiplier: float
    revenue: List[float]
    net_profit: List[float]
    min_cash: float
    breakeven_month: Optional[int]
    funding_gap: float          # 现金最低点为负时的缺口（正数）
    runway_months: Optional[int]   # 融资后现金首次为负的月份；None = 预测期内不缺钱


SCENARIO_NAMES = {"base": "中性", "optimistic": "乐观", "pessimistic": "悲观"}


def compute_scenarios(a: Dict[str, Any]) -> List[ScenarioRow]:
    sc = a.get("scenarios") or {"base": 1.0, "optimistic": 1.2, "pessimistic": 0.8}
    rows: List[ScenarioRow] = []
    order = ["pessimistic", "base", "optimistic"] + [k for k in sc if k not in SCENARIO_NAMES]
    for key in order:
        if key not in sc:
            continue
        mult = _num(sc[key], 1.0, f"scenarios.{key}")
        r = compute_startup(a, demand=mult)
        gap = max(0.0, -r.min_cash)
        runway = None
        for p in r.periods:
            if r.s["cash"][p.idx] < -EPS:
                runway = (p.year - 1) * 12 + (p.month or 12)
                break
        rows.append(ScenarioRow(key, SCENARIO_NAMES.get(key, key), mult, r.annual["revenue"], r.annual["net_profit"],
                                r.min_cash, r.breakeven.month, gap, runway))
    return rows


@dataclass
class SensitivityCell:
    driver: str
    delta: float
    net_profit_y: float
    min_cash: float
    ltv_cac: Optional[float] = None
    payback: Optional[float] = None


@dataclass
class Sensitivity:
    target_year: int                       # 1 起
    base_net_profit: float
    base_min_cash: float
    deltas: List[float]
    drivers: List[str]
    cells: Dict[Tuple[str, float], SensitivityCell]
    ranking: List[Tuple[str, float]]       # (driver, 净利润摆幅) 由大到小
    skipped: List[str]


def compute_sensitivity(a: Dict[str, Any], base: StartupResult) -> Sensitivity:
    cfg = a.get("sensitivity") or {}
    drivers = [str(d) for d in (cfg.get("drivers") or ["price", "volume", "cogsRatio", "fixedCost"])]
    deltas = sorted(set(_num(x, 0.0, "sensitivity.range") for x in (cfg.get("range") or [-0.3, -0.15, 0.15, 0.3])))
    ty = min(3, base.years)
    cells: Dict[Tuple[str, float], SensitivityCell] = {}
    skipped: List[str] = []
    kept: List[str] = []
    for d in drivers:
        if d not in SENS_DRIVERS:
            skipped.append(d)
            continue
        kept.append(d)
        for dl in deltas:
            r = compute_startup(a, demand=1.0, tweak={d: dl})
            ue = compute_unit_economics(r, tweak={d: dl}) if d in ("cac", "price", "churn") else None
            cells[(d, dl)] = SensitivityCell(d, dl, r.annual["net_profit"][ty - 1], r.min_cash,
                                             ue.ltv_cac if ue else None, ue.payback_months if ue else None)
    ranking = []
    for d in kept:
        vals = [cells[(d, dl)].net_profit_y for dl in deltas]
        ranking.append((d, max(vals) - min(vals)))
    ranking.sort(key=lambda t: -t[1])
    return Sensitivity(ty, base.annual["net_profit"][ty - 1], base.min_cash, deltas, kept, cells, ranking, skipped)


@dataclass
class Valuation:
    methods: List[Dict[str, Any]]          # {key, name, value(投前, 元), detail: {...}}
    low: float
    high: float
    mid: float
    chosen_pre: float
    chosen_source: str
    dcf: Dict[str, Any]
    ps: Dict[str, Any]
    vc: Dict[str, Any]
    notes: List[str]


def compute_valuation(res: StartupResult, financing_amount: float, round_year: int, round_month: int, pre_money_assumed: Optional[float]) -> Valuation:
    v = res.a.get("valuation") or {}
    methods_req = [str(m) for m in (v.get("methods") or ["dcf", "ps_comparables", "vc_method"])]
    an, years = res.annual, res.years
    notes: List[str] = []
    out: List[Dict[str, Any]] = []

    # DCF：FCF = 经营现金流 + 投资现金流（无债务 ⇒ 等于股权自由现金流），期末折现，Gordon 终值
    wacc = _num(v.get("wacc"), 0.18, "valuation.wacc")
    g = _num(v.get("terminalGrowth"), 0.03, "valuation.terminalGrowth")
    fcf = an["fcf"]
    dfs = [1.0 / (1.0 + wacc) ** (y + 1) for y in range(years)]
    pv = [fcf[y] * dfs[y] for y in range(years)]
    tv = fcf[-1] * (1.0 + g) / (wacc - g) if wacc > g and fcf[-1] > 0 else 0.0
    if tv == 0.0:
        notes.append("DCF 终值为 0（末年 FCF ≤ 0 或 WACC ≤ 永续增长率），估值偏保守")
    pv_tv = tv * dfs[-1]
    opening = res.a_opening()
    dcf_value = sum(pv) + pv_tv + opening
    dcf = {"wacc": wacc, "g": g, "fcf": fcf, "df": dfs, "pv": pv, "tv": tv, "pv_tv": pv_tv, "opening_cash": opening,
           "value": dcf_value, "tv_share": _safe_div(pv_tv, dcf_value)}
    if "dcf" in methods_req:
        out.append({"key": "dcf", "name": "DCF（现金流折现）", "value": dcf_value,
                    "detail": f"WACC {wacc:.0%}，永续增长 {g:.0%}，终值占比 {dcf['tv_share']:.0%}"})

    # PS：可比市销率 × 融资后 12 个月收入（NTM），可选一级市场折价
    ps_mult = _num(v.get("psMultiple"), 3.0, "valuation.psMultiple")
    ps_disc = _num(v.get("psDiscount"), 0.0, "valuation.psDiscount")
    ps_year = v.get("psYear")
    if ps_year is not None:
        yi = max(1, min(years, int(_num(ps_year, 1, "valuation.psYear"))))
        ntm = an["revenue"][yi - 1]
        ntm_desc = f"第 {yi} 年收入"
    else:
        y0 = max(1, min(years, round_year))
        w_next = (round_month - 1) / 12.0
        nxt = an["revenue"][y0] if y0 < years else an["revenue"][y0 - 1]
        ntm = an["revenue"][y0 - 1] * (1 - w_next) + nxt * w_next
        ntm_desc = f"融资后 12 个月收入（第 {y0} 年起）"
    ps_value = ps_mult * ntm * (1.0 - ps_disc)
    ps = {"multiple": ps_mult, "discount": ps_disc, "ntm": ntm, "ntm_desc": ntm_desc, "value": ps_value}
    if "ps_comparables" in methods_req or "ps" in methods_req:
        out.append({"key": "ps_comparables", "name": "可比公司 PS（市销率）", "value": ps_value,
                    "detail": f"PS {ps_mult:.1f}× × {ntm_desc} {ntm / res.unit:,.0f} 万元" + (f"，折价 {ps_disc:.0%}" if ps_disc else "")})

    # VC 法：退出价值 = 退出年净利 × 退出 PE（净利 ≤ 0 时退化为 收入 × PS），按目标 IRR 折现为投后，减融资额为投前
    exit_year = max(1, min(years, int(_num(v.get("exitYear"), years, "valuation.exitYear"))))
    exit_pe = _num(v.get("exitPE"), 20.0, "valuation.exitPE")
    irr = _num(v.get("targetIrr"), 0.40, "valuation.targetIrr")
    retention = _num(v.get("retention"), 1.0, "valuation.retention") or 1.0
    np_exit = an["net_profit"][exit_year - 1]
    if np_exit > 0:
        exit_value = np_exit * exit_pe
        exit_basis = f"第 {exit_year} 年净利润 {np_exit / res.unit:,.0f} 万元 × PE {exit_pe:.0f}×"
    else:
        exit_value = an["revenue"][exit_year - 1] * ps_mult
        exit_basis = f"第 {exit_year} 年净利润 ≤ 0，改用收入 {an['revenue'][exit_year - 1] / res.unit:,.0f} 万元 × PS {ps_mult:.1f}×"
        notes.append("VC 法退出年净利润 ≤ 0，退出价值改按 PS 计算")
    hold = exit_year - round_year + 1
    post = exit_value * retention / (1.0 + irr) ** hold
    vc_pre = post - financing_amount
    vc = {"exit_year": exit_year, "exit_pe": exit_pe, "irr": irr, "hold": hold, "exit_value": exit_value, "exit_basis": exit_basis,
          "post": post, "pre": vc_pre, "retention": retention}
    if "vc_method" in methods_req or "vc" in methods_req:
        out.append({"key": "vc_method", "name": "VC 法（退出倒推）", "value": vc_pre,
                    "detail": f"{exit_basis}，目标 IRR {irr:.0%}，持有 {hold} 年 ⇒ 投后 {post / res.unit:,.0f} 万元 − 融资额"})
    if len(out) < 2:
        notes.append("估值方法少于 2 种，建议至少两法三角验证")
    vals = [m["value"] for m in out] or [0.0]
    low, high, mid = min(vals), max(vals), float(median(vals))
    if pre_money_assumed is not None and pre_money_assumed > 0:
        chosen, src = pre_money_assumed, "融资条款 preMoney"
        if not (low * 0.5 <= pre_money_assumed <= high * 1.5):
            notes.append(f"条款投前估值 {pre_money_assumed / res.unit:,.0f} 万元 偏离三法区间 [{low / res.unit:,.0f}, {high / res.unit:,.0f}] 万元 较远，请复核")
    else:
        chosen, src = mid, "三法中值"
    return Valuation(out, low, high, mid, chosen, src, dcf, ps, vc, notes)


@dataclass
class FinancingSummary:
    rounds: List[Dict[str, Any]]      # 每轮：round, year, month, amount, pre, post, equity, use[]
    amount: float                     # 本轮（第一轮）融资额
    pre_money: float
    post_money: float
    equity_offered: float
    use: List[Tuple[str, float, float]]   # (用途, pct, 金额)
    use_total_pct: float
    gap: float                        # 现金缺口 = max(0, −现金最低点（融资前）)
    coverage: float                   # 融资额 ÷ 缺口
    notes: List[str]


def compute_financing(res: StartupResult, val: Optional[Valuation]) -> FinancingSummary:
    notes: List[str] = []
    rounds_raw = res.a.get("financing") or []
    rounds: List[Dict[str, Any]] = []
    for rd in rounds_raw:
        amt = _num(rd.get("amount"), 0.0, "financing.amount")
        pre = rd.get("preMoney")
        pre_v = _num(pre, 0.0, "financing.preMoney") if pre is not None else (val.chosen_pre if val else 0.0)
        post = pre_v + amt
        use = []
        for u in rd.get("use") or []:
            pct = _num(u.get("pct"), 0.0, "financing.use.pct")
            use.append((str(u.get("item") or "用途"), pct, amt * pct))
        rounds.append({"round": str(rd.get("round") or "本轮"), "year": int(_num(rd.get("year"), 1)), "month": int(_num(rd.get("month"), 1)),
                       "amount": amt, "pre": pre_v, "post": post, "equity": _safe_div(amt, post), "use": use,
                       "use_total_pct": sum(p for _, p, _ in use)})
    if not rounds:
        notes.append("未配置 financing，融资三件套为 0")
        first = {"amount": 0.0, "pre": val.chosen_pre if val else 0.0, "post": val.chosen_pre if val else 0.0, "equity": 0.0, "use": [], "use_total_pct": 0.0}
    else:
        first = rounds[0]
        if len(rounds) > 1:
            notes.append(f"共 {len(rounds)} 轮融资，账本 financing_amount / pre_money / post_money / equity_offered 取第 1 轮（{first['round']}）")
    if first["use"] and abs(first["use_total_pct"] - 1.0) > 0.005:
        notes.append(f"资金用途合计 {first['use_total_pct']:.1%} ≠ 100%")
    gap = max(0.0, -res.min_cash)
    cov = _safe_div(first["amount"], gap, float("inf")) if gap > EPS else float("inf")
    if res.min_cash + first["amount"] < -EPS:
        notes.append(f"现金最低点（融资前）{res.min_cash / res.unit:,.1f} 万元 + 融资额 {first['amount'] / res.unit:,.1f} 万元 < 0，存在融资缺口")
    if res.min_cash_post < -EPS:
        notes.append(f"融资后现金仍在某期为负（最低 {res.min_cash_post / res.unit:,.1f} 万元），请提高融资额或调整融资时点")
    return FinancingSummary(rounds, first["amount"], first["pre"], first["post"], first["equity"], first["use"], first["use_total_pct"], gap, cov, notes)


# ==============================================================================================
# Excel 输出
# ==============================================================================================
from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import BarChart, LineChart, Reference  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

COLOR_INPUT = "0000FF"      # 硬输入：蓝字
FILL_INPUT = "FFF2CC"       # 硬输入：浅黄底
COLOR_FORMULA = "000000"    # 同表公式：黑字
COLOR_XREF = "008000"       # 跨表引用：绿字
COLOR_CHECK = "6B7280"      # 数值校验列：灰字
FMT_MONEY = "#,##0.00;[Red]-#,##0.00"
FMT_INT = "#,##0;[Red]-#,##0"
FMT_DEC1 = "#,##0.0;[Red]-#,##0.0"
FMT_PCT = "0.0%;[Red]-0.0%"
FMT_PRICE = "#,##0.00"
FMT_X = '0.0"×"'

STARTUP_SHEETS = ["说明", "假设", "驱动", "收入", "成本", "人力", "利润表", "现金流量表", "资产负债表", "盈亏平衡", "单位经济", "敏感性", "估值", "融资", "检查"]
BUDGET_SHEETS = ["说明", "预算总表", "排期", "KPI"]


def unit_label(unit: float) -> str:
    return {1: "元", 1000: "千元", 10000: "万元", 1000000: "百万元", 100000000: "亿元"}.get(int(unit), f"元/{unit:g}")


def _font(kind: str, bold: bool = False, italic: bool = False) -> Font:
    color = {"input": COLOR_INPUT, "xref": COLOR_XREF, "check": COLOR_CHECK}.get(kind, COLOR_FORMULA)
    return Font(name="Arial", size=10, color=color, bold=bold, italic=italic)


def _style(cell: Any, kind: str, fmt: Optional[str] = None, bold: bool = False) -> None:
    cell.font = _font(kind, bold=bold, italic=(kind == "check"))
    if kind == "input":
        cell.fill = PatternFill("solid", fgColor=FILL_INPUT)
    if fmt:
        cell.number_format = fmt


def _header_cell(cell: Any, text: str, fill: str = theme_mod.NAVY, color: str = "FFFFFF") -> None:
    cell.value = text
    cell.font = Font(name="Arial", size=10, bold=True, color=color)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _set_widths(ws: Any, widths: Dict[int, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def _print_setup(ws: Any) -> None:
    """横向、按一页宽打印（导出 PDF 作 BP 附录时不被截成多页）。"""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.gridLines = False


class Layout:
    """所有按期页共用的列布局：A 科目 | B 期初 | 月度列… | 首年合计 | 年度列… | 空列 | 数值校验（每年一列）。"""

    def __init__(self, res: StartupResult):
        self.res = res
        self.periods = res.periods
        self.years = res.years
        self.col_label, self.col_open = 1, 2
        self.pcol: Dict[int, int] = {}
        self.ycol: List[int] = [0] * res.years
        self.first_month_col: Optional[int] = None
        self.last_month_col: Optional[int] = None
        col = 3
        for p in self.periods:
            self.pcol[p.idx] = col
            if p.month is not None:
                if p.month == 1:
                    self.first_month_col = col
                if p.month == 12:
                    self.last_month_col = col
                    col += 1
                    self.ycol[0] = col
            else:
                self.ycol[p.year - 1] = col
            col += 1
        self.last_col = col - 1
        self.total_col: Optional[int] = self.ycol[0] if self.first_month_col else None
        self.check_cols = [self.last_col + 2 + y for y in range(res.years)]
        self.year_labels = res.year_labels()

    @staticmethod
    def L(col: int) -> str:
        return get_column_letter(col)

    def period_cols(self) -> List[int]:
        return [self.pcol[p.idx] for p in self.periods]

    def value_cols(self) -> List[int]:
        """全部数据列（期列 + 合计列），从左到右。"""
        return list(range(3, self.last_col + 1))

    def period_range(self, row: int) -> str:
        """期列区域（跳过合计列），形如 C5:N5,P5:S5。"""
        parts: List[str] = []
        cols = self.period_cols()
        start = cols[0]
        prev = cols[0]
        for c in cols[1:] + [None]:  # type: ignore[list-item]
            if c is None or c != prev + 1:
                parts.append(f"{self.L(start)}{row}:{self.L(prev)}{row}")
                if c is not None:
                    start = c
            if c is not None:
                prev = c
        return ",".join(parts)

    def headers(self, unit_text: str) -> List[Tuple[int, str]]:
        hs = [(self.col_label, f"科目（{unit_text}）"), (self.col_open, "期初")]
        for p in self.periods:
            hs.append((self.pcol[p.idx], p.label))
        if self.total_col:
            hs.append((self.total_col, f"{self.year_labels[0]} 全年"))
        for y, c in enumerate(self.check_cols):
            hs.append((c, f"数值校验 {self.year_labels[y]}"))
        return hs


class PeriodSheet:
    """按期页写入器：一行一个科目，支持数值行 / 公式行 / 跨表引用行，并记录行号供其他页引用。"""

    def __init__(self, wb: Workbook, name: str, lay: Layout, unit_text: str):
        self.ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
        self.name = name
        self.lay = lay
        self.rows: Dict[str, int] = {}
        self.row = 2
        ws = self.ws
        for col, text in lay.headers(unit_text):
            fill = theme_mod.SLATE if col in lay.check_cols else theme_mod.NAVY
            _header_cell(ws.cell(row=1, column=col), text, fill=fill)
        ws.row_dimensions[1].height = 30
        widths = {lay.col_label: 34, lay.col_open: 12}
        for c in lay.value_cols():
            widths[c] = 12.5
        for c in lay.check_cols:
            widths[c] = 14
        _set_widths(ws, widths)
        ws.freeze_panes = "B2"
        ws.sheet_view.zoomScale = 90
        _print_setup(ws)

    def section(self, text: str) -> None:
        c = self.ws.cell(row=self.row, column=1, value=text)
        c.font = Font(name="Arial", size=10, bold=True, color=theme_mod.NAVY)
        c.fill = PatternFill("solid", fgColor=theme_mod.PAPER_ALT)
        for col in range(2, self.lay.last_col + 1):
            self.ws.cell(row=self.row, column=col).fill = PatternFill("solid", fgColor=theme_mod.PAPER_ALT)
        self.row += 1

    def blank(self) -> None:
        self.row += 1

    def add(self, key: str, label: str, *, values: Optional[Sequence[float]] = None, formula: Any = None, kind: str = "value",
            fmt: str = FMT_MONEY, total: str = "sum", scale: Optional[float] = None, open_value: Any = None,
            open_kind: str = "input", check: Optional[Sequence[float]] = None, check_scale: Optional[float] = None,
            bold: bool = False, indent: int = 0) -> int:
        """写一行。

        values：每期数值（元；scale 默认 1/unit 换算为展示单位）。formula：f(col_letter, prev_col_letter, period|None) → 公式。
        total：首年合计列的写法 —— sum（流量 SUM）/ last（存量取 12 月）/ same（套用同一公式）/ first（取 1 月）/ none。
        check：年度 Python 值写到右侧「数值校验」列。
        """
        lay, ws, r = self.lay, self.ws, self.row
        self.rows[key] = r
        lc = ws.cell(row=r, column=lay.col_label, value=("　" * indent) + label)
        lc.font = Font(name="Arial", size=10, bold=bold)
        sc = (1.0 / lay.res.unit) if scale is None else scale
        if open_value is not None:
            oc = ws.cell(row=r, column=lay.col_open)
            oc.value = open_value
            _style(oc, open_kind if not (isinstance(open_value, str) and open_value.startswith("=")) else "formula", fmt, bold)
        for p in lay.periods:
            col = lay.pcol[p.idx]
            cell = ws.cell(row=r, column=col)
            if formula is not None:
                cell.value = formula(Layout.L(col), Layout.L(col - 1), p)
                _style(cell, kind if kind in ("xref", "input") else "formula", fmt, bold)
            elif values is not None:
                cell.value = values[p.idx] * sc
                _style(cell, kind if kind in ("input", "xref") else "formula", fmt, bold)
        if lay.total_col:
            tc = ws.cell(row=r, column=lay.total_col)
            fm, lm = Layout.L(lay.first_month_col or 3), Layout.L(lay.last_month_col or 3)
            tcol = Layout.L(lay.total_col)
            if total == "sum":
                tc.value = f"=SUM({fm}{r}:{lm}{r})"
            elif total == "last":
                tc.value = f"={lm}{r}"
            elif total == "first":
                tc.value = f"={fm}{r}"
            elif total == "same" and formula is not None:
                tc.value = formula(tcol, Layout.L(lay.total_col - 1), None)
            if tc.value is not None:
                _style(tc, "formula", fmt, bold)
        if check is not None:
            csc = (1.0 / lay.res.unit) if check_scale is None else check_scale
            for y, c in enumerate(lay.check_cols):
                if y < len(check):
                    cc = ws.cell(row=r, column=c, value=check[y] * csc)
                    _style(cc, "check", fmt)
        self.row += 1
        return r


class TableSheet:
    """普通表格页（键值 / 多列表），带表头样式与地址记录。"""

    def __init__(self, wb: Workbook, name: str, widths: Sequence[float]):
        self.ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
        self.name = name
        self.row = 1
        self.cells: Dict[str, str] = {}
        _set_widths(self.ws, {i + 1: w for i, w in enumerate(widths)})
        self.ws.sheet_view.zoomScale = 90
        _print_setup(self.ws)

    def title(self, text: str, note: str = "") -> None:
        c = self.ws.cell(row=self.row, column=1, value=text)
        c.font = Font(name="Arial", size=12, bold=True, color=theme_mod.NAVY)
        self.row += 1
        if note:
            n = self.ws.cell(row=self.row, column=1, value=note)
            n.font = Font(name="Arial", size=9, italic=True, color=theme_mod.MUTED)
            self.row += 1

    def header(self, labels: Sequence[str], fill: str = theme_mod.NAVY) -> int:
        for i, t in enumerate(labels):
            _header_cell(self.ws.cell(row=self.row, column=i + 1), t, fill=fill)
        self.row += 1
        return self.row - 1

    def line(self, cells: Sequence[Any], kinds: Optional[Sequence[str]] = None, fmts: Optional[Sequence[Optional[str]]] = None,
             key: Optional[str] = None, key_col: int = 3, bold: bool = False) -> int:
        r = self.row
        for i, v in enumerate(cells):
            if v is None:
                continue
            cell = self.ws.cell(row=r, column=i + 1, value=v)
            kind = (kinds[i] if kinds and i < len(kinds) and kinds[i] else ("formula" if isinstance(v, str) and v.startswith("=") else "text"))
            fmt = fmts[i] if fmts and i < len(fmts) else None
            if kind == "text":
                cell.font = Font(name="Arial", size=10, bold=bold)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                _style(cell, kind, fmt, bold)
        if key:
            self.cells[key] = f"{self.name}!{get_column_letter(key_col)}{r}"
        self.row += 1
        return r

    def addr(self, row: int, col: int, absolute: bool = True) -> str:
        return f"{self.name}!${get_column_letter(col)}${row}" if absolute else f"{self.name}!{get_column_letter(col)}{row}"

    def blank(self) -> None:
        self.row += 1


def _fmt_list(v: Any, pct: bool = False) -> str:
    if isinstance(v, list):
        return "、".join((f"{x:.0%}" if pct else f"{x:g}") if isinstance(x, (int, float)) else str(x) for x in v)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return f"{v:.0%}" if pct else f"{v:g}"
    return "" if v is None else str(v)


# ----------------------------------------------------------------------------------------------
# startup：各页
# ----------------------------------------------------------------------------------------------

@dataclass
class StartupBundle:
    res: StartupResult
    ue: UnitEconomics
    scenarios: List[ScenarioRow]
    sens: Sensitivity
    val: Valuation
    fin: FinancingSummary
    warnings: List[str]


def _write_assumptions(wb: Workbook, b: StartupBundle) -> TableSheet:
    res, a = b.res, b.res.a
    t = TableSheet(wb, "假设", [14, 34, 16, 12, 60])
    u = unit_label(res.unit)
    t.title("假设页（全部硬输入集中于此；蓝字浅黄底 = 输入）", f"金额单位：元（报表按 {u} 展示）；每条假设须有来源，无来源的假设不应进入模型。")
    t.header(["类别", "参数", "值", "单位", "来源 / 说明"])

    def inp(cat: str, name: str, value: Any, unit_: str, src: str = "", key: Optional[str] = None, fmt: Optional[str] = None) -> int:
        is_num = isinstance(value, (int, float)) and not isinstance(value, bool)
        kinds = ["text", "text", "input" if is_num else "text", "text", "text"]
        return t.line([cat, name, value, unit_, src], kinds=kinds, fmts=[None, None, fmt, None, None], key=key)

    inp("基本", "公司 / 项目", str(a.get("company") or ""), "", "")
    inp("基本", "起始年", res.start_year, "年", "")
    inp("基本", "预测年数", res.years, "年", "")
    inp("基本", "展示单位", res.unit, "元", f"报表金额 = 元 ÷ {res.unit:g}")
    inp("基本", "首年按月", "是" if any(p.month for p in res.periods) else "否", "", "")
    inp("基本", "期初现金", res.a_opening(), "元", "", key="openingCash", fmt=FMT_INT)
    sc = a.get("scenarios") or {}
    for k, nm in SCENARIO_NAMES.items():
        if k in sc:
            inp("情景", f"{nm}情景需求乘数", _num(sc[k], 1.0), "倍", "作用于销量 / 客户数 / 流量 / 新增订阅", key=f"scenario_{k}", fmt="0.00")
    for i, (st, raw) in enumerate(zip(res.streams, a["revenueStreams"])):
        cat = f"收入流 {i + 1}"
        inp(cat, "名称 / 驱动", f"{st.name}（{st.driver_name}）", "", st.source)
        qkey = {"volume_price": "volume", "clients_arpu": "clients", "funnel": "traffic", "subscribers_arpu": "subscribers"}[st.driver]
        q = raw.get(qkey) or {}
        if isinstance(q, dict):
            if q.get("y1Monthly") is not None:
                inp(cat, f"首年各月{ {'volume': '销量', 'clients': '签约客户', 'traffic': '流量', 'subscribers': '新增订阅'}[qkey] }", _fmt_list(q.get("y1Monthly")), st.unit_name, "12 个月，S 型爬坡")
            else:
                inp(cat, "首年总量", _num(q.get("y1"), 0), st.unit_name, "按默认 S 型权重分月", fmt=FMT_INT)
            inp(cat, "逐年增长率", _fmt_list(q.get("yoyGrowth"), pct=True), "", "第 1 年忽略")
            if st.driver == "subscribers_arpu":
                inp(cat, "月流失率", _num(q.get("churnMonthly"), 0), "", "", fmt=FMT_PCT)
        pkey = {"volume_price": "price", "clients_arpu": "arpu", "funnel": "aov", "subscribers_arpu": "arpu"}[st.driver]
        pr = raw.get(pkey)
        pr = pr if isinstance(pr, dict) else {"base": pr}
        inp(cat, st.price_label, _num(pr.get("base"), 0), "元", "", fmt=FMT_PRICE)
        inp(cat, "单价逐年变化", _fmt_list(pr.get("yoyChange"), pct=True), "", "")
        if st.driver == "funnel":
            cv = raw.get("conversion")
            cv = cv if isinstance(cv, dict) else {"base": cv}
            inp(cat, "转化率", _num(cv.get("base"), 0), "", "", fmt=FMT_PCT)
        inp(cat, "成本率（COGS / 收入）", st.cogs_ratio, "", "行业对标", fmt=FMT_PCT)
    costs = a.get("costs") or {}
    for it in costs.get("fixed") or []:
        inp("固定成本", str(it.get("name")), _num(it.get("monthly"), 0), "元/月", f"年变化 {_fmt_list(it.get('yoyChange'), pct=True) or '0%'}；{it.get('source', '')}", fmt=FMT_INT)
    for it in costs.get("variablePctOfRevenue") or []:
        inp("变动费用", str(it.get("name")), _num(it.get("pct"), 0), "占收入", str(it.get("source", "")), fmt=FMT_PCT)
    mk = costs.get("marketing") or {}
    inp("营销", "营销费率（逐年）", _fmt_list(mk.get("pctOfRevenue"), pct=True), "占收入", "与 minAnnual 取大")
    inp("营销", "年度最低营销投入", _num(mk.get("minAnnual"), 0), "元/年", str(mk.get("source", "")), fmt=FMT_INT)
    for it in costs.get("headcount") or []:
        inp("人力", f"{it.get('role')} 人数（逐年）", _fmt_list(it.get("y")), "人", str(it.get("source", "")))
        inp("人力", f"{it.get('role')} 月薪", _num(it.get("monthlySalary"), 0), "元/月", f"社保公积金系数 {_num(it.get('socialRate'), 0):.0%}", fmt=FMT_INT)
    for it in costs.get("capex") or []:
        inp("资本开支", str(it.get("name")), _num(it.get("amount"), 0), "元", f"第 {int(_num(it.get('year'), 1))} 年第 {int(_num(it.get('month'), 1))} 月投入，直线折旧 {int(_num(it.get('depreciationYears'), 5))} 年；{it.get('source', '')}", fmt=FMT_INT)
    inp("其他", "其他运营费用占收入比", _num(costs.get("otherOpexPctOfRevenue"), 0), "占收入", "", fmt=FMT_PCT)
    wc = a.get("workingCapital") or {}
    inp("营运资本", "应收账款周转天数（DSO）", _num(wc.get("receivableDays"), 0), "天", "期末应收 = 当期收入 ÷ 期间天数 × DSO", key="dso", fmt=FMT_INT)
    inp("营运资本", "存货周转天数（DIO）", _num(wc.get("inventoryDays"), 0), "天", "期末存货 = 当期营业成本 ÷ 期间天数 × DIO", key="dio", fmt=FMT_INT)
    inp("营运资本", "应付账款周转天数（DPO）", _num(wc.get("payableDays"), 0), "天", "期末应付 = 当期营业成本 ÷ 期间天数 × DPO", key="dpo", fmt=FMT_INT)
    inp("税", "企业所得税率", _num((a.get("tax") or {}).get("incomeTaxRate"), 0.25), "", "累计口径计税：亏损结转、不退税", key="taxRate", fmt=FMT_PCT)
    for rd in a.get("financing") or []:
        inp("融资", f"{rd.get('round', '本轮')} 融资额", _num(rd.get("amount"), 0), "元", f"第 {int(_num(rd.get('year'), 1))} 年第 {int(_num(rd.get('month'), 1))} 月到账", fmt=FMT_INT)
        if rd.get("preMoney") is not None:
            inp("融资", f"{rd.get('round', '本轮')} 投前估值", _num(rd.get("preMoney"), 0), "元", "", fmt=FMT_INT)
        for uitem in rd.get("use") or []:
            inp("融资", f"用途：{uitem.get('item')}", _num(uitem.get("pct"), 0), "占融资额", "", fmt=FMT_PCT)
    v = a.get("valuation") or {}
    for k, nm, fmt in (("wacc", "折现率（WACC）", FMT_PCT), ("terminalGrowth", "永续增长率", FMT_PCT), ("psMultiple", "可比 PS 倍数", FMT_X),
                       ("exitPE", "退出 PE 倍数", FMT_X), ("exitYear", "退出年份（第 N 年）", FMT_INT), ("targetIrr", "目标 IRR（VC 法）", FMT_PCT)):
        if k in v:
            inp("估值", nm, _num(v.get(k), 0), "", "", key=f"val_{k}", fmt=fmt)
    ue = a.get("unitEconomics") or {}
    for k, nm, unit_, fmt in (("cac", "获客成本 CAC", "元/人", FMT_PRICE), ("repeatRate", "复购率", "", FMT_PCT), ("grossMarginRef", "毛利率参考", "", FMT_PCT),
                              ("lifetimeOrders", "生命周期购买次数", "次", FMT_DEC1), ("lifetimeMonths", "客户生命周期", "月", FMT_INT), ("arpuMonthly", "月 ARPU", "元/月", FMT_PRICE)):
        if k in ue:
            inp("单位经济", nm, _num(ue.get(k), 0), unit_, "", key=f"ue_{k}", fmt=fmt)
    sens = a.get("sensitivity") or {}
    inp("敏感性", "驱动因子", "、".join(SENS_DRIVERS.get(str(d), (str(d),))[0] for d in (sens.get("drivers") or [])), "", "")
    inp("敏感性", "变动幅度", _fmt_list(sens.get("range"), pct=True), "", "")
    t.ws.freeze_panes = "A4"
    return t


def _write_drivers(wb: Workbook, lay: Layout, b: StartupBundle) -> PeriodSheet:
    res = b.res
    sh = PeriodSheet(wb, "驱动", lay, "数量：自然单位；单价：元")
    for i, st in enumerate(res.streams):
        sh.section(f"收入流 {i + 1}：{st.name}（{st.driver_name}）")
        for nm, series, f in st.extra:
            fmt = {"int": FMT_INT, "pct": FMT_PCT}.get(f, FMT_DEC1)
            sh.add(f"s{i}_extra_{nm}", nm, values=series, scale=1.0, fmt=fmt, total=("sum" if f == "int" and "期末" not in nm and "流失" not in nm else "last"), indent=1)
        is_int = all(abs(v - round(v)) < 1e-6 for v in st.volume)
        sh.add(f"s{i}_volume", st.volume_label, values=st.volume, scale=1.0, fmt=FMT_INT if is_int else FMT_DEC1, total="sum", indent=1)
        sh.add(f"s{i}_price", st.price_label, values=st.price, scale=1.0, fmt=FMT_PRICE, total="first", indent=1)
    sh.section("资金驱动")
    sh.add("fin_in", f"股权融资注入（{unit_label(res.unit)}）", values=res.s["fin_in"], kind="input", total="sum", check=res.annual["fin_in"])
    sh.add("demand", "需求乘数（当前情景）", values=[res.demand] * len(res.periods), scale=1.0, fmt="0.00", total="first")
    return sh


def _write_revenue(wb: Workbook, lay: Layout, b: StartupBundle) -> PeriodSheet:
    res = b.res
    u = unit_label(res.unit)
    sh = PeriodSheet(wb, "收入", lay, u)
    sh.section("分产品线收入（数值 = 驱动页数量 × 单价）")
    rev_rows, cogs_rows = [], []
    for i, st in enumerate(res.streams):
        rev_rows.append(sh.add(f"rev_{i}", f"{st.name} 收入", values=st.revenue, indent=1, check=_annual_flow(st.revenue, res.yidx)))
    sh.section("分产品线营业成本（COGS = 收入 × 成本率）")
    for i, st in enumerate(res.streams):
        cogs_rows.append(sh.add(f"cogs_{i}", f"{st.name} 营业成本（成本率 {st.cogs_ratio:.0%}）", values=st.cogs, indent=1))
    sh.section("合计")
    sh.add("revenue", "营业收入合计", formula=lambda c, pc, p: "=" + "+".join(f"{c}{r}" for r in rev_rows), total="same", bold=True, check=res.annual["revenue"])
    sh.add("cogs", "营业成本合计", formula=lambda c, pc, p: "=" + "+".join(f"{c}{r}" for r in cogs_rows), total="same", bold=True, check=res.annual["cogs"])
    sh.add("gross_profit", "毛利", formula=lambda c, pc, p: f"={c}{sh.rows['revenue']}-{c}{sh.rows['cogs']}", total="same", bold=True, check=res.annual["gross_profit"])
    sh.add("gross_margin", "毛利率", formula=lambda c, pc, p: f"=IF({c}{sh.rows['revenue']}=0,0,{c}{sh.rows['gross_profit']}/{c}{sh.rows['revenue']})",
           fmt=FMT_PCT, total="same", check=res.annual["gross_margin"], check_scale=1.0)
    return sh


def _write_headcount(wb: Workbook, lay: Layout, b: StartupBundle) -> PeriodSheet:
    res = b.res
    u = unit_label(res.unit)
    sh = PeriodSheet(wb, "人力", lay, f"人数：人；月薪：元；成本：{u}")
    cost_rows, head_rows = [], []
    for i, h in enumerate(res.headcount):
        sh.section(f"{h.role}（社保公积金系数 {h.social_rate:.0%}）")
        head_rows.append(sh.add(f"hc{i}_heads", "人数", values=h.heads, scale=1.0, fmt=FMT_INT, total="last", indent=1))
        sh.add(f"hc{i}_salary", "月薪（元）", values=h.salary, scale=1.0, fmt=FMT_INT, total="first", indent=1)
        cost_rows.append(sh.add(f"hc{i}_cost", f"人力成本（人数 × 月薪 × (1+{h.social_rate:.0%}) × 月数）", values=h.cost, indent=1))
    sh.section("合计")
    if head_rows:
        sh.add("heads_total", "总人数", formula=lambda c, pc, p: "=" + "+".join(f"{c}{r}" for r in head_rows), scale=1.0, fmt=FMT_INT, total="same", bold=True)
        sh.add("headcount", "人力成本合计", formula=lambda c, pc, p: "=" + "+".join(f"{c}{r}" for r in cost_rows), total="same", bold=True, check=res.annual["headcount"])
    else:
        sh.add("heads_total", "总人数", values=[0.0] * len(res.periods), scale=1.0, fmt=FMT_INT, total="last", bold=True)
        sh.add("headcount", "人力成本合计", values=[0.0] * len(res.periods), total="sum", bold=True, check=res.annual["headcount"])
    return sh


def _write_costs(wb: Workbook, lay: Layout, b: StartupBundle, hc: PeriodSheet) -> PeriodSheet:
    res = b.res
    u = unit_label(res.unit)
    sh = PeriodSheet(wb, "成本", lay, u)
    sh.section("固定成本（月额 × 月数 × 年变化）")
    fixed_rows = [sh.add(f"fixed_{i}", nm, values=v, indent=1) for i, (nm, v) in enumerate(res.fixed_items)]
    if fixed_rows:
        sh.add("fixed", "固定成本合计", formula=lambda c, pc, p: "=" + "+".join(f"{c}{r}" for r in fixed_rows), total="same", bold=True, check=res.annual["fixed"])
    else:
        sh.add("fixed", "固定成本合计", values=[0.0] * len(res.periods), bold=True, check=res.annual["fixed"])
    sh.section("变动费用（占收入比 × 收入）")
    var_rows = [sh.add(f"var_{i}", f"{nm}（{pct:.1%}）", values=v, indent=1) for i, (nm, pct, v) in enumerate(res.variable_items)]
    if var_rows:
        sh.add("variable", "变动费用合计", formula=lambda c, pc, p: "=" + "+".join(f"{c}{r}" for r in var_rows), total="same", bold=True, check=res.annual["variable"])
    else:
        sh.add("variable", "变动费用合计", values=[0.0] * len(res.periods), bold=True, check=res.annual["variable"])
    sh.section("营销与其他")
    sh.add("marketing_pct", "营销费率（当年）", values=res.s["marketing_pct"], scale=1.0, fmt=FMT_PCT, total="first", indent=1)
    sh.add("marketing", "营销费用（max(费率 × 年收入, 年度最低投入)，首年均摊到月）", values=res.s["marketing"], bold=True, check=res.annual["marketing"])
    sh.add("other_opex", "其他运营费用（占收入比）", values=res.s["other_opex"], bold=True, check=res.annual["other_opex"])
    sh.add("headcount", "人力成本合计（引用人力页）", formula=lambda c, pc, p: f"={hc.name}!{c}{hc.rows['headcount']}",
           kind="xref", total="same", bold=True, check=res.annual["headcount"])
    sh.section("资本开支与折旧（直线法）")
    capex_rows, dep_rows = [], []
    for i, cx in enumerate(res.capex):
        capex_rows.append(sh.add(f"capex_{i}", f"{cx.name} 投入（第 {cx.year} 年第 {cx.month} 月）", values=cx.outflow, indent=1))
        dep_rows.append(sh.add(f"dep_{i}", f"{cx.name} 折旧（{cx.dep_years} 年）", values=cx.dep, indent=1))
    if capex_rows:
        sh.add("capex", "资本开支合计", formula=lambda c, pc, p: "=" + "+".join(f"{c}{r}" for r in capex_rows), total="same", bold=True, check=res.annual["capex"])
        sh.add("dep", "折旧摊销合计", formula=lambda c, pc, p: "=" + "+".join(f"{c}{r}" for r in dep_rows), total="same", bold=True, check=res.annual["dep"])
    else:
        sh.add("capex", "资本开支合计", values=[0.0] * len(res.periods), bold=True, check=res.annual["capex"])
        sh.add("dep", "折旧摊销合计", values=[0.0] * len(res.periods), bold=True, check=res.annual["dep"])
    sh.section("合计")
    sh.add("opex_total", "营业费用合计（变动 + 营销 + 人力 + 固定 + 其他，不含折旧）",
           formula=lambda c, pc, p: f"={c}{sh.rows['variable']}+{c}{sh.rows['marketing']}+{c}{sh.rows['headcount']}+{c}{sh.rows['fixed']}+{c}{sh.rows['other_opex']}",
           total="same", bold=True, check=res.annual["opex_total"])
    return sh


def _write_income_statement(wb: Workbook, lay: Layout, b: StartupBundle, rev: PeriodSheet, cost: PeriodSheet, hc: PeriodSheet, asm: TableSheet) -> PeriodSheet:
    res, an = b.res, b.res.annual
    sh = PeriodSheet(wb, "利润表", lay, unit_label(res.unit))
    X = lambda sheet, key: (lambda c, pc, p: f"={sheet.name}!{c}{sheet.rows[key]}")  # noqa: E731
    sh.add("revenue", "营业收入", formula=X(rev, "revenue"), kind="xref", bold=True, check=an["revenue"])
    sh.add("cogs", "减：营业成本", formula=X(rev, "cogs"), kind="xref", check=an["cogs"])
    sh.add("gross_profit", "毛利", formula=lambda c, pc, p: f"={c}{sh.rows['revenue']}-{c}{sh.rows['cogs']}", total="same", bold=True, check=an["gross_profit"])
    sh.add("gross_margin", "毛利率", formula=lambda c, pc, p: f"=IF({c}{sh.rows['revenue']}=0,0,{c}{sh.rows['gross_profit']}/{c}{sh.rows['revenue']})",
           fmt=FMT_PCT, total="same", check=an["gross_margin"], check_scale=1.0)
    sh.add("variable", "减：变动费用（佣金 / 物流等）", formula=X(cost, "variable"), kind="xref", check=an["variable"])
    sh.add("marketing", "减：营销费用", formula=X(cost, "marketing"), kind="xref", check=an["marketing"])
    sh.add("headcount", "减：人力成本", formula=X(hc, "headcount"), kind="xref", check=an["headcount"])
    sh.add("fixed", "减：固定成本", formula=X(cost, "fixed"), kind="xref", check=an["fixed"])
    sh.add("other_opex", "减：其他运营费用", formula=X(cost, "other_opex"), kind="xref", check=an["other_opex"])
    sh.add("ebitda", "EBITDA",
           formula=lambda c, pc, p: f"={c}{sh.rows['gross_profit']}-{c}{sh.rows['variable']}-{c}{sh.rows['marketing']}-{c}{sh.rows['headcount']}-{c}{sh.rows['fixed']}-{c}{sh.rows['other_opex']}",
           total="same", bold=True, check=an["ebitda"])
    sh.add("dep", "减：折旧摊销", formula=X(cost, "dep"), kind="xref", check=an["dep"])
    sh.add("ebit", "营业利润（EBIT）", formula=lambda c, pc, p: f"={c}{sh.rows['ebitda']}-{c}{sh.rows['dep']}", total="same", bold=True, check=an["ebit"])
    sh.add("interest", "减：财务费用（无债务融资，为 0）", values=res.s["interest"], kind="input", check=an["interest"])
    sh.add("ebt", "利润总额", formula=lambda c, pc, p: f"={c}{sh.rows['ebit']}-{c}{sh.rows['interest']}", total="same", bold=True, check=an["ebt"])
    sh.add("cum_ebt", "　累计利润总额（计税用）", formula=lambda c, pc, p: f"={pc}{sh.rows['cum_ebt']}+{c}{sh.rows['ebt']}", total="last", open_value=0, check=an["cum_ebt"])
    sh.add("tax_base", "　累计计税基数（历史最大累计利润）", formula=lambda c, pc, p: f"=MAX({pc}{sh.rows['tax_base']},{c}{sh.rows['cum_ebt']})", total="last", open_value=0, check=an["tax_base"])
    rate_ref = _abs_ref(asm.cells["taxRate"])
    sh.add("tax", "减：所得税（税率 × 计税基数增量；亏损结转、不退税）",
           formula=lambda c, pc, p: f"={rate_ref}*({c}{sh.rows['tax_base']}-{pc}{sh.rows['tax_base']})", kind="xref", total="sum", check=an["tax"])
    sh.add("net_profit", "净利润", formula=lambda c, pc, p: f"={c}{sh.rows['ebt']}-{c}{sh.rows['tax']}", total="same", bold=True, check=an["net_profit"])
    sh.add("net_margin", "净利率", formula=lambda c, pc, p: f"=IF({c}{sh.rows['revenue']}=0,0,{c}{sh.rows['net_profit']}/{c}{sh.rows['revenue']})",
           fmt=FMT_PCT, total="same", check=an["net_margin"], check_scale=1.0)
    sh.blank()
    note = sh.ws.cell(row=sh.row, column=1, value="说明：月度 / 年度列为公式（绿字 = 引用收入 / 成本 / 人力 / 假设页）；右侧灰色斜体为 Python 计算的年度值，检查页逐项比较两者。")
    note.font = Font(name="Arial", size=9, italic=True, color=theme_mod.MUTED)
    sh.row += 2
    _chart_income(sh, lay)
    return sh


def _abs_ref(ref: str) -> str:
    """假设!C12 → 假设!$C$12。"""
    sheet, cell = ref.split("!")
    m = re.fullmatch(r"\$?([A-Z]+)\$?(\d+)", cell)
    return f"{sheet}!${m.group(1)}${m.group(2)}" if m else ref


def _chart_income(sh: PeriodSheet, lay: Layout) -> None:
    ws = sh.ws
    c0, c1 = lay.ycol[0], lay.ycol[-1]
    bar = BarChart()
    bar.type = "col"
    bar.title = "年度营业收入与净利润"
    bar.y_axis.title = "金额"
    bar.x_axis.tickLblPos = "low"
    bar.height, bar.width = 9.0, 22.0
    cats = Reference(ws, min_col=c0, max_col=c1, min_row=1, max_row=1)
    from openpyxl.chart.series import SeriesLabel  # noqa: E402

    for key, nm in (("revenue", "营业收入"), ("net_profit", "净利润")):
        r = sh.rows[key]
        bar.add_data(Reference(ws, min_col=c0, max_col=c1, min_row=r, max_row=r), from_rows=True, titles_from_data=False)
        bar.series[-1].tx = SeriesLabel(v=nm)
    bar.set_categories(cats)
    line = LineChart()
    for key, nm in (("gross_margin", "毛利率"), ("net_margin", "净利率")):
        r = sh.rows[key]
        line.add_data(Reference(ws, min_col=c0, max_col=c1, min_row=r, max_row=r), from_rows=True, titles_from_data=False)
        line.series[-1].tx = SeriesLabel(v=nm)
    line.y_axis.axId = 200
    line.y_axis.title = "比率"
    line.y_axis.number_format = "0%"
    line.y_axis.crosses = "max"
    bar += line
    ws.add_chart(bar, f"B{sh.row}")
    sh.row += 20


def _chart_cashflow(sh: PeriodSheet, lay: Layout) -> None:
    from openpyxl.chart.series import SeriesLabel  # noqa: E402

    ws = sh.ws
    c0, c1 = lay.ycol[0], lay.ycol[-1]
    bar = BarChart()
    bar.type = "col"
    bar.grouping = "clustered"
    bar.title = "年度现金流与期末现金"
    bar.y_axis.title = "金额"
    bar.x_axis.tickLblPos = "low"
    bar.height, bar.width = 9.0, 22.0
    for key, nm in (("cfo", "经营活动现金流"), ("cfi", "投资活动现金流"), ("cff", "筹资活动现金流")):
        r = sh.rows[key]
        bar.add_data(Reference(ws, min_col=c0, max_col=c1, min_row=r, max_row=r), from_rows=True, titles_from_data=False)
        bar.series[-1].tx = SeriesLabel(v=nm)
    bar.set_categories(Reference(ws, min_col=c0, max_col=c1, min_row=1, max_row=1))
    line = LineChart()
    r = sh.rows["cash_end"]
    line.add_data(Reference(ws, min_col=c0, max_col=c1, min_row=r, max_row=r), from_rows=True, titles_from_data=False)
    line.series[-1].tx = SeriesLabel(v="期末现金")
    line.y_axis.axId = 200
    line.y_axis.title = "期末现金"
    line.y_axis.crosses = "max"
    bar += line
    ws.add_chart(bar, f"B{sh.row}")
    sh.row += 20


def _write_cashflow(wb: Workbook, lay: Layout, b: StartupBundle, is_: PeriodSheet, cost: PeriodSheet, drv: PeriodSheet, bs_rows: Dict[str, int]) -> PeriodSheet:
    """现金流量表：间接法；资产负债表的应收 / 存货 / 应付行号由调用方预先约定（两表互相引用）。"""
    res, an = b.res, b.res.annual
    u = unit_label(res.unit)
    sh = PeriodSheet(wb, "现金流量表", lay, u)
    BS = "资产负债表"
    sh.section("经营活动")
    sh.add("net_profit", "净利润", formula=lambda c, pc, p: f"=利润表!{c}{is_.rows['net_profit']}", kind="xref", check=an["net_profit"], indent=1)
    sh.add("dep", "加：折旧摊销", formula=lambda c, pc, p: f"=成本!{c}{cost.rows['dep']}", kind="xref", check=an["dep"], indent=1)
    sh.add("d_ar", "减：应收账款增加", formula=lambda c, pc, p: f"=-({BS}!{c}{bs_rows['ar']}-{BS}!{pc}{bs_rows['ar']})", kind="xref", check=_scale(an["d_ar"], -1.0), indent=1)
    sh.add("d_inv", "减：存货增加", formula=lambda c, pc, p: f"=-({BS}!{c}{bs_rows['inv']}-{BS}!{pc}{bs_rows['inv']})", kind="xref", check=_scale(an["d_inv"], -1.0), indent=1)
    sh.add("d_ap", "加：应付账款增加", formula=lambda c, pc, p: f"={BS}!{c}{bs_rows['ap']}-{BS}!{pc}{bs_rows['ap']}", kind="xref", check=an["d_ap"], indent=1)
    sh.add("cfo", "经营活动现金流净额",
           formula=lambda c, pc, p: f"={c}{sh.rows['net_profit']}+{c}{sh.rows['dep']}+{c}{sh.rows['d_ar']}+{c}{sh.rows['d_inv']}+{c}{sh.rows['d_ap']}",
           total="same", bold=True, check=an["cfo"])
    sh.section("投资活动")
    sh.add("capex", "减：资本开支", formula=lambda c, pc, p: f"=-成本!{c}{cost.rows['capex']}", kind="xref", check=_scale(an["capex"], -1.0), indent=1)
    sh.add("cfi", "投资活动现金流净额", formula=lambda c, pc, p: f"={c}{sh.rows['capex']}", total="same", bold=True, check=an["cfi"])
    sh.section("筹资活动")
    sh.add("fin_in", "加：股权融资到账", formula=lambda c, pc, p: f"=驱动!{c}{drv.rows['fin_in']}", kind="xref", check=an["fin_in"], indent=1)
    sh.add("cff", "筹资活动现金流净额", formula=lambda c, pc, p: f"={c}{sh.rows['fin_in']}", total="same", bold=True, check=an["cff"])
    sh.section("现金汇总")
    sh.add("net_change", "现金净增加额", formula=lambda c, pc, p: f"={c}{sh.rows['cfo']}+{c}{sh.rows['cfi']}+{c}{sh.rows['cff']}", total="same", bold=True, check=an["net_change"])
    r_open = sh.row
    r_end = r_open + 1
    sh.add("cash_open", "期初现金", formula=lambda c, pc, p: f"={pc}{r_end}", total="first", check=[res.a_opening()] + an["cash"][:-1])
    sh.add("cash_end", "期末现金", formula=lambda c, pc, p: f"={c}{r_open}+{c}{sh.rows['net_change']}", total="same", open_value=res.a_opening() / res.unit, bold=True, check=an["cash"])
    sh.add("cum_fin", "　累计股权融资", formula=lambda c, pc, p: f"={pc}{sh.rows['cum_fin']}+{c}{sh.rows['fin_in']}", total="last", open_value=0, check=an["cum_fin"])
    sh.add("cash_prefin", "融资前现金余额（期末现金 − 累计融资）", formula=lambda c, pc, p: f"={c}{sh.rows['cash_end']}-{c}{sh.rows['cum_fin']}", total="same", check=an["cash_prefin"])
    sh.blank()
    r = sh.row
    ws = sh.ws
    ws.cell(row=r, column=1, value="现金最低点（融资前，含期初）").font = Font(name="Arial", size=10, bold=True)
    c = ws.cell(row=r, column=2, value=f"=MIN({Layout.L(lay.col_open)}{sh.rows['cash_end']},{lay.period_range(sh.rows['cash_prefin'])})")
    _style(c, "formula", FMT_MONEY, True)
    cc = ws.cell(row=r, column=3, value=res.min_cash / res.unit)
    _style(cc, "check", FMT_MONEY)
    ws.cell(row=r, column=4, value="← 右侧灰字为 Python 值").font = Font(name="Arial", size=9, italic=True, color=theme_mod.MUTED)
    sh.rows["min_cash"] = r
    r += 1
    ws.cell(row=r, column=1, value="现金最低点（融资后）").font = Font(name="Arial", size=10, bold=True)
    c = ws.cell(row=r, column=2, value=f"=MIN({Layout.L(lay.col_open)}{sh.rows['cash_end']},{lay.period_range(sh.rows['cash_end'])})")
    _style(c, "formula", FMT_MONEY, True)
    _style(ws.cell(row=r, column=3, value=res.min_cash_post / res.unit), "check", FMT_MONEY)
    sh.rows["min_cash_post"] = r
    sh.row = r + 2
    _chart_cashflow(sh, lay)
    return sh


def _write_balance_sheet(wb: Workbook, lay: Layout, b: StartupBundle, is_: PeriodSheet, cf: PeriodSheet, rev: PeriodSheet, cost: PeriodSheet,
                         drv: PeriodSheet, asm: TableSheet, bs_rows: Dict[str, int]) -> PeriodSheet:
    res, an = b.res, b.res.annual
    sh = PeriodSheet(wb, "资产负债表", lay, unit_label(res.unit))
    ws = sh.ws
    # 第 2 行：期间天数（输入），应收 / 存货 / 应付公式引用
    r_days = sh.row
    sh.rows["days"] = r_days
    ws.cell(row=r_days, column=1, value="期间天数（月 30 / 年 365）").font = Font(name="Arial", size=10)
    for p in res.periods:
        _style(ws.cell(row=r_days, column=lay.pcol[p.idx], value=p.days), "input", FMT_INT)
    if lay.total_col:
        _style(ws.cell(row=r_days, column=lay.total_col, value=365), "input", FMT_INT)
    sh.row += 1
    dso, dio, dpo = (_abs_ref(asm.cells[k]) for k in ("dso", "dio", "dpo"))
    sh.section("资产")
    assert sh.row == bs_rows["cash"], "资产负债表行号与现金流量表约定不一致"
    sh.add("cash", "货币资金", formula=lambda c, pc, p: f"=现金流量表!{c}{cf.rows['cash_end']}", kind="xref", total="last", open_value=res.a_opening() / res.unit, check=an["cash"], indent=1)
    sh.add("ar", "应收账款（收入 ÷ 天数 × DSO）", formula=lambda c, pc, p: f"=收入!{c}{rev.rows['revenue']}/{c}${r_days}*{dso}", kind="xref", total="last", open_value=0, check=an["ar"], indent=1)
    sh.add("inv", "存货（营业成本 ÷ 天数 × DIO）", formula=lambda c, pc, p: f"=收入!{c}{rev.rows['cogs']}/{c}${r_days}*{dio}", kind="xref", total="last", open_value=0, check=an["inv"], indent=1)
    sh.add("current_assets", "流动资产合计", formula=lambda c, pc, p: f"={c}{sh.rows['cash']}+{c}{sh.rows['ar']}+{c}{sh.rows['inv']}", total="same", bold=True)
    sh.add("fa_gross", "固定资产原值", formula=lambda c, pc, p: f"={pc}{sh.rows['fa_gross']}+成本!{c}{cost.rows['capex']}", kind="xref", total="last", open_value=0, check=an["fa_gross"], indent=1)
    sh.add("acc_dep", "减：累计折旧", formula=lambda c, pc, p: f"={pc}{sh.rows['acc_dep']}+成本!{c}{cost.rows['dep']}", kind="xref", total="last", open_value=0, check=an["acc_dep"], indent=1)
    sh.add("fa_net", "固定资产净值", formula=lambda c, pc, p: f"={c}{sh.rows['fa_gross']}-{c}{sh.rows['acc_dep']}", total="same", bold=True, check=an["fa_net"])
    sh.add("assets", "资产总计", formula=lambda c, pc, p: f"={c}{sh.rows['current_assets']}+{c}{sh.rows['fa_net']}", total="same", bold=True, check=an["assets"])
    sh.section("负债")
    sh.add("ap", "应付账款（营业成本 ÷ 天数 × DPO）", formula=lambda c, pc, p: f"=收入!{c}{rev.rows['cogs']}/{c}${r_days}*{dpo}", kind="xref", total="last", open_value=0, check=an["ap"], indent=1)
    sh.add("liabilities", "负债合计", formula=lambda c, pc, p: f"={c}{sh.rows['ap']}", total="same", bold=True, check=an["liabilities"])
    sh.section("所有者权益")
    sh.add("paid_in", "实收资本及资本公积（期初现金 + 累计融资）", formula=lambda c, pc, p: f"={pc}{sh.rows['paid_in']}+驱动!{c}{drv.rows['fin_in']}", kind="xref", total="last",
           open_value=f"=B{sh.rows['cash']}", check=an["paid_in"], indent=1)
    sh.add("retained", "留存收益（累计净利润）", formula=lambda c, pc, p: f"={pc}{sh.rows['retained']}+利润表!{c}{is_.rows['net_profit']}", kind="xref", total="last", open_value=0, check=an["retained"], indent=1)
    sh.add("equity", "所有者权益合计", formula=lambda c, pc, p: f"={c}{sh.rows['paid_in']}+{c}{sh.rows['retained']}", total="same", bold=True, check=an["equity"])
    sh.add("liab_equity", "负债及所有者权益总计", formula=lambda c, pc, p: f"={c}{sh.rows['liabilities']}+{c}{sh.rows['equity']}", total="same", bold=True, check=an["liab_equity"])
    sh.blank()
    sh.add("diff", "检查：资产 − 负债及权益（应为 0）", formula=lambda c, pc, p: f"={c}{sh.rows['assets']}-{c}{sh.rows['liab_equity']}", total="same",
           open_value=f"=B{sh.rows['assets']}-B{sh.rows['liab_equity']}", check=[0.0] * res.years)
    return sh


def _bs_row_plan(res: StartupResult) -> Dict[str, int]:
    """资产负债表关键行号（现金流量表先写、需要引用它们）：第 2 行天数、第 3 行分区标题、之后依次 cash / ar / inv。"""
    return {"cash": 4, "ar": 5, "inv": 6, "ap": 13}


def _write_breakeven(wb: Workbook, lay: Layout, b: StartupBundle, is_: PeriodSheet) -> Tuple[PeriodSheet, Dict[str, str]]:
    res, be = b.res, b.res.breakeven
    u = unit_label(res.unit)
    sh = PeriodSheet(wb, "盈亏平衡", lay, u)
    cells: Dict[str, str] = {}
    sh.section("累计经营利润（公式引用利润表）")
    sh.add("ebit", "营业利润（EBIT）", formula=lambda c, pc, p: f"=利润表!{c}{is_.rows['ebit']}", kind="xref", check=res.annual["ebit"])
    sh.add("cum_ebit", "累计营业利润", formula=lambda c, pc, p: f"={pc}{sh.rows['cum_ebit']}+{c}{sh.rows['ebit']}", total="last", open_value=0,
           check=_annual_stock(be.cum_ebit, res.yidx))
    sh.add("positive", "累计转正标记（1 = 已转正）", formula=lambda c, pc, p: f"=IF({c}{sh.rows['cum_ebit']}>0,1,0)", fmt="0", total="last")
    sh.blank()
    ws = sh.ws
    r = sh.row
    ws.cell(row=r, column=1, value="盈亏平衡月份（累计经营利润首次转正；年度期内按线性推算）").font = Font(name="Arial", size=10, bold=True)
    c = ws.cell(row=r, column=2, value=be.month if be.month is not None else "未达到")
    _style(c, "formula", FMT_INT, True)
    cells["breakeven_month"] = f"{sh.name}!B{r}"
    r += 1
    fp = be.first_positive_period
    ws.cell(row=r, column=1, value="单期经营利润首次为正").font = Font(name="Arial", size=10)
    ws.cell(row=r, column=2, value=res.periods[fp].label if fp is not None else "未达到")
    r += 2
    # 年度表：收入 / 变动成本 / 边际贡献 / 固定性成本 / 保本收入 / 安全边际
    heads = ["年度盈亏平衡", *lay.year_labels]
    for i, t in enumerate(heads):
        _header_cell(ws.cell(row=r, column=1 + i), t)
    r += 1
    rows_spec = [("revenue", "营业收入", FMT_MONEY, 1 / res.unit), ("variable_cost", "变动成本（COGS + 变动费用 + 其他运营）", FMT_MONEY, 1 / res.unit),
                 ("contribution", "边际贡献", FMT_MONEY, 1 / res.unit), ("cm_ratio", "边际贡献率", FMT_PCT, 1.0),
                 ("fixed_like", "固定性成本（固定 + 营销 + 人力 + 折旧）", FMT_MONEY, 1 / res.unit), ("be_revenue", "保本收入 = 固定性成本 ÷ 边际贡献率", FMT_MONEY, 1 / res.unit),
                 ("cash_be_revenue", "现金保本收入（不含折旧）", FMT_MONEY, 1 / res.unit), ("margin_of_safety", "安全边际 = (收入 − 保本收入) ÷ 收入", FMT_PCT, 1.0)]
    for key, label, fmt, sc in rows_spec:
        ws.cell(row=r, column=1, value=label).font = Font(name="Arial", size=10)
        for y in range(res.years):
            v = be.year_table[y][key]
            cell = ws.cell(row=r, column=2 + y, value=(v * sc if math.isfinite(v) else "—"))
            _style(cell, "formula", fmt)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value=f"保本销量（主力产品线：{be.primary_stream}，第 {be.chart_year} 年口径）").font = Font(name="Arial", size=10, bold=True, color=theme_mod.NAVY)
    r += 1
    spec2 = [("需覆盖的固定性成本（扣除其他产品线边际贡献）", be.fixed_cost / res.unit, FMT_MONEY, "fixed"),
             (f"单价（元/{be.unit_name}）", be.unit_price, FMT_PRICE, "price"),
             (f"单位变动成本（元/{be.unit_name}，单价 × (成本率 + 变动费率 + 其他运营费率)）", be.unit_var_cost, FMT_PRICE, "uvc"),
             (f"保本销量（{be.unit_name}）= 固定性成本 ÷ (单价 − 单位变动成本)", None, FMT_INT, "be_units"),
             (f"预计销量（{be.unit_name}）", be.actual_units, FMT_INT, "actual"),
             ("安全边际（销量口径）", None, FMT_PCT, "mos")]
    addr: Dict[str, int] = {}
    for label, v, fmt, k in spec2:
        ws.cell(row=r, column=1, value=label).font = Font(name="Arial", size=10)
        addr[k] = r
        if k == "be_units":
            v = f"=B{addr['fixed']}*{res.unit:g}/(B{addr['price']}-B{addr['uvc']})"
        elif k == "mos":
            v = f"=IF(B{addr['actual']}=0,0,(B{addr['actual']}-B{addr['be_units']})/B{addr['actual']})"
        cell = ws.cell(row=r, column=2, value=v)
        _style(cell, "formula", fmt)
        if k == "be_units":
            _style(ws.cell(row=r, column=3, value=be.be_units if math.isfinite(be.be_units) else "—"), "check", fmt)
        r += 1
    sh.row = r
    return sh, cells


def _write_unit_economics(wb: Workbook, b: StartupBundle) -> TableSheet:
    res, ue = b.res, b.ue
    t = TableSheet(wb, "单位经济", [44, 16, 16, 60])
    t.title("单位经济（CAC / LTV / 回收期）", "LTV = 客单价 × 毛利率 × 生命周期购买次数；回收期 = CAC ÷ 月均毛利贡献（LTV ÷ 生命周期月数）。")
    t.header(["指标", "值（公式）", "Python 值", "说明"])
    r_cac = t.line(["获客成本 CAC（元/人）", ue.cac, ue.cac, ue.cac_source], kinds=["text", "input", "check", "text"], fmts=[None, FMT_PRICE, FMT_PRICE, None], key="cac", key_col=2)
    r_price = t.line(["客单价（元）", ue.price, ue.price, f"主力产品线「{res.breakeven.primary_stream}」首年单价"], kinds=["text", "input", "check", "text"], fmts=[None, FMT_PRICE, FMT_PRICE, None])
    r_margin = t.line(["毛利率", ue.margin, ue.margin, "unitEconomics.grossMarginRef 或 1 − 成本率"], kinds=["text", "input", "check", "text"], fmts=[None, FMT_PCT, FMT_PCT, None])
    r_orders = t.line(["生命周期购买次数", ue.lifetime_orders, ue.lifetime_orders, "；".join(ue.notes) or "unitEconomics.lifetimeOrders"], kinds=["text", "input", "check", "text"], fmts=[None, FMT_DEC1, FMT_DEC1, None])
    r_months = t.line(["客户生命周期（月）", ue.lifetime_months, ue.lifetime_months, "unitEconomics.lifetimeMonths（默认 24）"], kinds=["text", "input", "check", "text"], fmts=[None, FMT_INT, FMT_INT, None])
    t.line(["首单毛利（元）", f"=B{r_price}*B{r_margin}", ue.first_order_gp, "客单价 × 毛利率；≥ CAC 则首单回本"], kinds=["text", "formula", "check", "text"], fmts=[None, FMT_PRICE, FMT_PRICE, None])
    r_ltv = t.line(["LTV（元）", f"=B{r_price}*B{r_margin}*B{r_orders}", ue.ltv, "客单价 × 毛利率 × 购买次数"], kinds=["text", "formula", "check", "text"], fmts=[None, FMT_PRICE, FMT_PRICE, None], key="ltv", key_col=2)
    t.line(["LTV : CAC（倍）", f"=IF(B{r_cac}=0,0,B{r_ltv}/B{r_cac})", ue.ltv_cac if math.isfinite(ue.ltv_cac) else 0, "健康 ≥ 3；> 5 通常意味着获客投入不足"], kinds=["text", "formula", "check", "text"], fmts=[None, FMT_X, FMT_X, None], key="ltv_cac", key_col=2)
    t.line(["CAC 回收期（月）", f"=IF(B{r_ltv}=0,0,B{r_cac}/(B{r_ltv}/B{r_months}))", ue.payback_months if math.isfinite(ue.payback_months) else 0, "消费品 ≤ 6–9 个月为佳；> 24 个月为红线"],
           kinds=["text", "formula", "check", "text"], fmts=[None, FMT_DEC1, FMT_DEC1, None], key="payback_months", key_col=2)
    return t


def _write_sensitivity(wb: Workbook, b: StartupBundle) -> TableSheet:
    res, se, u = b.res, b.sens, unit_label(b.res.unit)
    t = TableSheet(wb, "敏感性", [26] + [14] * 8)
    t.title(f"单因子敏感性（第 {se.target_year} 年净利润 / 现金最低点（融资前），{u}）", "每个因子单独变动、其他假设不变；重跑整套模型得到结果，不是手填。")
    heads = ["驱动因子"] + [f"{d:+.0%}" for d in se.deltas] + ["基准", "摆幅（最大 − 最小）"]
    t.header(heads)
    for d, swing in se.ranking:
        vals = [se.cells[(d, dl)].net_profit_y / res.unit for dl in se.deltas]
        t.line([f"{SENS_DRIVERS[d][0]} → 净利润"] + vals + [se.base_net_profit / res.unit, swing / res.unit], kinds=["text"] + ["formula"] * (len(vals) + 2), fmts=[None] + [FMT_MONEY] * (len(vals) + 2))
    for d, _ in se.ranking:
        vals = [se.cells[(d, dl)].min_cash / res.unit for dl in se.deltas]
        t.line([f"{SENS_DRIVERS[d][0]} → 现金最低点"] + vals + [se.base_min_cash / res.unit, (max(vals) - min(vals))], kinds=["text"] + ["formula"] * (len(vals) + 2), fmts=[None] + [FMT_MONEY] * (len(vals) + 2))
    for d, _ in se.ranking:
        c0 = se.cells[(d, se.deltas[0])]
        if c0.ltv_cac is not None:
            vals = [se.cells[(d, dl)].ltv_cac for dl in se.deltas]
            t.line([f"{SENS_DRIVERS[d][0]} → LTV:CAC"] + vals + [b.ue.ltv_cac, max(vals) - min(vals)], kinds=["text"] + ["formula"] * (len(vals) + 2), fmts=[None] + [FMT_X] * (len(vals) + 2))
            vals = [se.cells[(d, dl)].payback for dl in se.deltas]
            t.line([f"{SENS_DRIVERS[d][0]} → 回收期（月）"] + vals + [b.ue.payback_months, max(vals) - min(vals)], kinds=["text"] + ["formula"] * (len(vals) + 2), fmts=[None] + [FMT_DEC1] * (len(vals) + 2))
    if se.skipped:
        t.line([f"未识别的驱动因子已跳过：{', '.join(se.skipped)}"])
    t.blank()
    t.line(["因子口径："] + [""])
    for d in se.drivers:
        t.line([SENS_DRIVERS[d][0], SENS_DRIVERS[d][1]])
    t.blank()
    t.title(f"三情景（需求乘数作用于销量 / 客户数 / 流量 / 新增订阅；{u}）")
    yl = res.year_labels()
    t.header(["情景", "需求乘数"] + [f"收入 {y}" for y in yl] + ["净利润 " + yl[0], "净利润 " + yl[-1], "现金最低点（融资前）", "盈亏平衡月", "资金缺口", "融资后首次缺钱（月）"])
    for sc in b.scenarios:
        t.line([sc.name, sc.multiplier] + [v / res.unit for v in sc.revenue] + [sc.net_profit[0] / res.unit, sc.net_profit[-1] / res.unit, sc.min_cash / res.unit,
               sc.breakeven_month if sc.breakeven_month is not None else "未达到", sc.funding_gap / res.unit, sc.runway_months if sc.runway_months is not None else "不缺钱"],
               kinds=["text", "input"] + ["formula"] * (len(yl) + 6), fmts=[None, "0.00"] + [FMT_MONEY] * (len(yl) + 3) + [FMT_INT, FMT_MONEY, FMT_INT])
    return t


def _write_valuation(wb: Workbook, b: StartupBundle) -> TableSheet:
    res, val, u = b.res, b.val, unit_label(b.res.unit)
    t = TableSheet(wb, "估值", [40, 16, 16, 16, 16, 16, 16, 40])
    t.title(f"估值（三法三角验证，{u}；全部基于同一套预测）")
    yl = res.year_labels()
    d = val.dcf
    t.header(["DCF（现金流折现）"] + yl + ["", "说明"])
    t.line(["自由现金流 FCF（经营 + 投资现金流）"] + [v / res.unit for v in d["fcf"]], kinds=["text"] + ["formula"] * res.years, fmts=[None] + [FMT_MONEY] * res.years)
    t.line(["折现因子 1/(1+WACC)^t"] + d["df"], kinds=["text"] + ["formula"] * res.years, fmts=[None] + ["0.000"] * res.years)
    t.line(["FCF 现值"] + [v / res.unit for v in d["pv"]], kinds=["text"] + ["formula"] * res.years, fmts=[None] + [FMT_MONEY] * res.years)
    t.line(["终值 TV = FCF_N × (1+g) ÷ (WACC − g)", d["tv"] / res.unit, None, None, None, None, None, f"WACC {d['wacc']:.0%}，g {d['g']:.0%}"], kinds=["text", "formula"], fmts=[None, FMT_MONEY])
    t.line(["终值现值", d["pv_tv"] / res.unit, None, None, None, None, None, f"终值占比 {d['tv_share']:.0%}"], kinds=["text", "formula"], fmts=[None, FMT_MONEY])
    t.line(["期初现金", d["opening_cash"] / res.unit], kinds=["text", "formula"], fmts=[None, FMT_MONEY])
    t.line(["DCF 估值（投前口径）", d["value"] / res.unit, None, None, None, None, None, "Σ FCF 现值 + 终值现值 + 期初现金"], kinds=["text", "formula"], fmts=[None, FMT_MONEY], bold=True, key="dcf", key_col=2)
    t.blank()
    p = val.ps
    t.header(["可比公司 PS（市销率）", "值", "", "", "", "", "", "说明"])
    t.line([p["ntm_desc"], p["ntm"] / res.unit], kinds=["text", "formula"], fmts=[None, FMT_MONEY])
    t.line(["PS 倍数", p["multiple"]], kinds=["text", "input"], fmts=[None, FMT_X])
    t.line(["一级市场折价", p["discount"]], kinds=["text", "input"], fmts=[None, FMT_PCT])
    t.line(["PS 估值（投前口径）", p["value"] / res.unit, None, None, None, None, None, "NTM 收入 × PS × (1 − 折价)"], kinds=["text", "formula"], fmts=[None, FMT_MONEY], bold=True, key="ps", key_col=2)
    t.blank()
    v = val.vc
    t.header(["VC 法（退出倒推）", "值", "", "", "", "", "", "说明"])
    t.line(["退出价值", v["exit_value"] / res.unit, None, None, None, None, None, v["exit_basis"]], kinds=["text", "formula"], fmts=[None, FMT_MONEY])
    t.line(["目标 IRR", v["irr"]], kinds=["text", "input"], fmts=[None, FMT_PCT])
    t.line(["持有年数", v["hold"]], kinds=["text", "formula"], fmts=[None, FMT_INT])
    t.line(["投后估值 = 退出价值 ÷ (1+IRR)^n", v["post"] / res.unit], kinds=["text", "formula"], fmts=[None, FMT_MONEY])
    t.line(["VC 法投前估值 = 投后 − 融资额", v["pre"] / res.unit], kinds=["text", "formula"], fmts=[None, FMT_MONEY], bold=True, key="vc", key_col=2)
    t.blank()
    t.header(["估值区间汇总", f"投前估值（{u}）", "", "", "", "", "", "说明"])
    for m in val.methods:
        t.line([m["name"], m["value"] / res.unit, None, None, None, None, None, m["detail"]], kinds=["text", "formula"], fmts=[None, FMT_MONEY])
    t.line(["区间下限 / 中值 / 上限", val.low / res.unit, val.mid / res.unit, val.high / res.unit], kinds=["text", "formula", "formula", "formula"], fmts=[None, FMT_MONEY, FMT_MONEY, FMT_MONEY], bold=True)
    t.line(["本轮采用投前估值", val.chosen_pre / res.unit, None, None, None, None, None, val.chosen_source], kinds=["text", "formula"], fmts=[None, FMT_MONEY], bold=True, key="chosen_pre", key_col=2)
    for n in val.notes:
        t.line([f"注：{n}"])
    return t


def _write_financing(wb: Workbook, b: StartupBundle) -> TableSheet:
    res, fin, u = b.res, b.fin, unit_label(b.res.unit)
    t = TableSheet(wb, "融资", [22, 14, 14, 16, 16, 16, 14, 40])
    t.title(f"融资计划（{u}）", "出让比例 = 融资额 ÷ 投后估值；投后 = 投前 + 融资额。三者在正文 / 路演 / 讲稿中必须一致。")
    t.header(["轮次", "时间", "融资额", "投前估值", "投后估值（公式）", "出让比例（公式）", "出让比例（Python）", "说明"])
    first_row = None
    for rd in fin.rounds:
        r = t.row
        t.line([rd["round"], f"第 {rd['year']} 年第 {rd['month']} 月", rd["amount"] / res.unit, rd["pre"] / res.unit, f"=C{r}+D{r}", f"=IF(E{r}=0,0,C{r}/E{r})", rd["equity"], ""],
               kinds=["text", "text", "input", "input", "formula", "formula", "check", "text"], fmts=[None, None, FMT_MONEY, FMT_MONEY, FMT_MONEY, FMT_PCT, FMT_PCT, None])
        if first_row is None:
            first_row = r
    if first_row is None:
        first_row = t.line(["本轮", "—", 0, fin.pre_money / res.unit, "=C{0}+D{0}".format(t.row), "=IF(E{0}=0,0,C{0}/E{0})".format(t.row), 0, "未配置 financing"],
                           kinds=["text", "text", "input", "input", "formula", "formula", "check", "text"], fmts=[None, None, FMT_MONEY, FMT_MONEY, FMT_MONEY, FMT_PCT, FMT_PCT, None])
    t.cells["financing_amount"] = f"融资!C{first_row}"
    t.cells["pre_money"] = f"融资!D{first_row}"
    t.cells["post_money"] = f"融资!E{first_row}"
    t.cells["equity_offered"] = f"融资!F{first_row}"
    t.blank()
    t.header(["资金用途（本轮）", "占比", "金额（公式）", "", "", "", "", "对应里程碑 / 成本科目"])
    use_rows = []
    for item, pct, amt in fin.use:
        r = t.row
        use_rows.append(r)
        t.line([item, pct, f"=B{r}*融资!$C${first_row}", None, None, None, None, ""], kinds=["text", "input", "formula"], fmts=[None, FMT_PCT, FMT_MONEY])
    if use_rows:
        t.line(["合计", f"=SUM(B{use_rows[0]}:B{use_rows[-1]})", f"=SUM(C{use_rows[0]}:C{use_rows[-1]})"], kinds=["text", "formula", "formula"], fmts=[None, FMT_PCT, FMT_MONEY], bold=True)
    t.blank()
    t.header(["融资额与现金缺口", f"值（{u}）", "", "", "", "", "", "说明"])
    t.line(["现金最低点（融资前，含期初现金）", res.min_cash / res.unit, None, None, None, None, None, "见现金流量表"], kinds=["text", "formula"], fmts=[None, FMT_MONEY])
    t.line(["资金缺口 = max(0, −现金最低点)", fin.gap / res.unit], kinds=["text", "formula"], fmts=[None, FMT_MONEY])
    t.line(["融资额 ÷ 缺口（倍）", fin.coverage if math.isfinite(fin.coverage) else "无缺口", None, None, None, None, None, "建议 1.1–1.2（缺口 + 10–20% 安全垫）"], kinds=["text", "formula"], fmts=[None, FMT_X])
    t.line(["融资后现金最低点", res.min_cash_post / res.unit, None, None, None, None, None, "应 ≥ 0"], kinds=["text", "formula"], fmts=[None, FMT_MONEY])
    for n in fin.notes:
        t.line([f"注：{n}"])
    return t


def _write_checks(wb: Workbook, lay: Layout, b: StartupBundle, is_: PeriodSheet, cf: PeriodSheet, bs: PeriodSheet, rev: PeriodSheet) -> PeriodSheet:
    res = b.res
    sh = PeriodSheet(wb, "检查", lay, "差异，应全为 0")
    BS, CF, IS = bs.name, cf.name, is_.name
    stream_terms = "+".join(f"收入!{{c}}{rev.rows[f'rev_{i}']}" for i in range(len(res.streams)))
    sh.add("c_bs", "资产负债表平衡：资产 − 负债及权益", formula=lambda c, pc, p: f"={BS}!{c}{bs.rows['assets']}-{BS}!{c}{bs.rows['liab_equity']}", kind="xref", total="same")
    sh.add("c_cash", "现金勾稽：现金流量表期末现金 − 资产负债表货币资金", formula=lambda c, pc, p: f"={CF}!{c}{cf.rows['cash_end']}-{BS}!{c}{bs.rows['cash']}", kind="xref", total="same")
    sh.add("c_re", "留存收益滚动：Δ留存收益 − 净利润", formula=lambda c, pc, p: f"={BS}!{c}{bs.rows['retained']}-{BS}!{pc}{bs.rows['retained']}-{IS}!{c}{is_.rows['net_profit']}", kind="xref", total="sum")
    sh.add("c_rev", "收入分项加总：利润表收入 − Σ 产品线收入", formula=lambda c, pc, p: f"={IS}!{c}{is_.rows['revenue']}-(" + stream_terms.format(c=c) + ")", kind="xref", total="same")
    sh.add("c_cf", "现金流闭合：期初 + 净增加 − 期末", formula=lambda c, pc, p: f"={CF}!{c}{cf.rows['cash_open']}+{CF}!{c}{cf.rows['net_change']}-{CF}!{c}{cf.rows['cash_end']}", kind="xref", total="same")
    # 公式 vs Python（年度列）
    ws = sh.ws
    for key, label, sheet, rk in (("c_py_np", "公式 vs Python：净利润（年度）", is_, "net_profit"), ("c_py_cash", "公式 vs Python：期末现金（年度）", cf, "cash_end"),
                                  ("c_py_assets", "公式 vs Python：资产总计（年度）", bs, "assets"), ("c_py_rev", "公式 vs Python：营业收入（年度）", is_, "revenue")):
        r = sh.row
        sh.rows[key] = r
        ws.cell(row=r, column=1, value=label).font = Font(name="Arial", size=10)
        for y in range(res.years):
            yc, cc = lay.ycol[y], lay.check_cols[y]
            cell = ws.cell(row=r, column=yc, value=f"={sheet.name}!{Layout.L(yc)}{sheet.rows[rk]}-{sheet.name}!{Layout.L(cc)}{sheet.rows[rk]}")
            _style(cell, "xref", FMT_MONEY)
        sh.row += 1
    last = sh.row - 1
    sh.blank()
    r = sh.row
    rng = f"{Layout.L(3)}2:{Layout.L(lay.last_col)}{last}"
    ws.cell(row=r, column=1, value="Master Check：Σ|差异|（公式）").font = Font(name="Arial", size=10, bold=True)
    _style(ws.cell(row=r, column=2, value=f"=SUMPRODUCT(ABS({rng}))"), "formula", "0.000000", True)
    sh.rows["master"] = r
    r += 1
    ws.cell(row=r, column=1, value="Master Check 状态").font = Font(name="Arial", size=10, bold=True)
    st = ws.cell(row=r, column=2, value=f'=IF(B{r - 1}<0.0001,"通过","未通过")')
    st.font = Font(name="Arial", size=10, bold=True, color=theme_mod.POS)
    sh.rows["master_status"] = r
    r += 2
    ws.cell(row=r, column=1, value="Python 侧勾稽（元，两条独立路径的最大绝对差）").font = Font(name="Arial", size=10, bold=True, color=theme_mod.NAVY)
    r += 1
    for k, label in (("balanceSheet", "资产 = 负债 + 权益"), ("cashflow", "现金流量表期末现金 = 资产负债表倒算现金"), ("retained", "留存收益 = Σ 净利润")):
        ws.cell(row=r, column=1, value=label).font = Font(name="Arial", size=10)
        _style(ws.cell(row=r, column=2, value=res.checks[k]), "check", "0.000000")
        ws.cell(row=r, column=3, value="通过" if res.checks[k] < 1.0 else "未通过").font = Font(name="Arial", size=10, bold=True, color=theme_mod.POS if res.checks[k] < 1.0 else theme_mod.NEG)
        r += 1
    sh.row = r
    return sh


def _write_cover(wb: Workbook, b: StartupBundle, checks: PeriodSheet, warnings: List[str]) -> None:
    res, u = b.res, unit_label(b.res.unit)
    ws = wb["说明"]
    _set_widths(ws, {1: 30, 2: 22, 3: 70})
    _print_setup(ws)
    r = 1
    c = ws.cell(row=r, column=1, value=f"{res.a.get('company', '')} 财务模型")
    c.font = Font(name="Arial", size=16, bold=True, color=theme_mod.NAVY)
    r += 1
    ws.cell(row=r, column=1, value=f"版本 v1 · 生成日期 {date.today().isoformat()} · 由 bizkit.finmodel 按 assumptions.json 生成 · 金额单位 {u}").font = Font(name="Arial", size=9, color=theme_mod.MUTED)
    r += 2
    ws.cell(row=r, column=1, value="Master Check（检查页）").font = Font(name="Arial", size=11, bold=True)
    st = ws.cell(row=r, column=2, value=f"=检查!B{checks.rows['master_status']}")
    st.font = Font(name="Arial", size=11, bold=True, color=theme_mod.POS)
    ws.cell(row=r, column=3, value="全部勾稽差异为 0 时显示「通过」；资产 = 负债 + 权益、现金流期末 = 货币资金、Δ留存收益 = 净利润、收入分项加总、公式 vs Python").font = Font(name="Arial", size=9, color=theme_mod.MUTED)
    r += 2
    ws.cell(row=r, column=1, value="颜色约定").font = Font(name="Arial", size=11, bold=True)
    r += 1
    for kind, text in (("input", "蓝字浅黄底 = 硬输入（假设页 / 期初 / 天数 / 融资注入）"), ("formula", "黑字 = 同表公式"), ("xref", "绿字 = 跨表引用"), ("check", "灰色斜体 = Python 数值校验（与公式结果应一致）")):
        cell = ws.cell(row=r, column=2, value="示例 1,234.56")
        _style(cell, kind)
        ws.cell(row=r, column=3, value=text).font = Font(name="Arial", size=10)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="页签目录").font = Font(name="Arial", size=11, bold=True)
    r += 1
    desc = {"说明": "版本、口径、目录、颜色约定、关键结果", "假设": "全部硬输入（值 / 单位 / 来源）", "驱动": "量价 / 客户 / 漏斗 / 订阅驱动与融资注入",
            "收入": "分产品线收入与营业成本、毛利", "成本": "固定 / 变动 / 营销 / 其他 / 资本开支与折旧", "人力": "岗位 × 人数 × 月薪 × (1+社保)",
            "利润表": "公式：收入 → 毛利 → EBITDA → EBIT → 税 → 净利", "现金流量表": "公式：间接法三段 + 期初期末 + 融资前余额",
            "资产负债表": "公式：货币资金 / 应收 / 存货 / 固定资产 / 应付 / 实收资本 / 留存收益", "盈亏平衡": "累计经营利润、保本收入、保本销量、安全边际",
            "单位经济": "CAC / LTV / LTV:CAC / 回收期", "敏感性": "单因子 ±15% / ±30% 与三情景", "估值": "DCF / PS / VC 三法与区间", "融资": "融资额 / 投前 / 投后 / 出让比例 / 资金用途 / 缺口",
            "检查": "勾稽差异与 Master Check"}
    for i, name in enumerate(STARTUP_SHEETS):
        ws.cell(row=r, column=1, value=f"{i + 1}. {name}").font = Font(name="Arial", size=10)
        ws.cell(row=r, column=3, value=desc.get(name, "")).font = Font(name="Arial", size=10, color=theme_mod.INK)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value=f"关键结果（{u}，Python 值）").font = Font(name="Arial", size=11, bold=True)
    r += 1
    an = res.annual
    yl = res.year_labels()
    for i, y in enumerate(yl):
        ws.cell(row=r, column=1, value=f"{y} 收入 / 净利润 / 毛利率").font = Font(name="Arial", size=10)
        _style(ws.cell(row=r, column=2, value=an["revenue"][i] / res.unit), "check", FMT_MONEY)
        ws.cell(row=r, column=3, value=f"净利润 {an['net_profit'][i] / res.unit:,.1f}；毛利率 {an['gross_margin'][i]:.1%}；净利率 {an['net_margin'][i]:.1%}").font = Font(name="Arial", size=10)
        r += 1
    be = res.breakeven
    kv = [("盈亏平衡月份", be.month if be.month is not None else "未达到", "累计经营利润首次转正"),
          ("现金最低点（融资前）", f"{res.min_cash / res.unit:,.1f}", f"融资额 {b.fin.amount / res.unit:,.1f} → 融资后最低 {res.min_cash_post / res.unit:,.1f}"),
          ("融资三件套", f"{b.fin.amount / res.unit:,.0f} / {b.fin.pre_money / res.unit:,.0f} / {b.fin.post_money / res.unit:,.0f}", f"融资额 / 投前 / 投后；出让 {b.fin.equity_offered:.1%}"),
          ("估值区间", f"{b.val.low / res.unit:,.0f} – {b.val.high / res.unit:,.0f}", f"{len(b.val.methods)} 种方法，中值 {b.val.mid / res.unit:,.0f}"),
          ("单位经济", f"LTV:CAC {b.ue.ltv_cac:.1f}×", f"CAC {b.ue.cac:,.0f} 元，LTV {b.ue.ltv:,.0f} 元，回收期 {b.ue.payback_months:.1f} 月")]
    for k, v, note in kv:
        ws.cell(row=r, column=1, value=k).font = Font(name="Arial", size=10)
        ws.cell(row=r, column=2, value=v).font = Font(name="Arial", size=10, bold=True)
        ws.cell(row=r, column=3, value=note).font = Font(name="Arial", size=10, color=theme_mod.MUTED)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="口径说明").font = Font(name="Arial", size=11, bold=True)
    r += 1
    for line in ("第一年按月（12 列）+ 首年合计列，之后按年；期初列为 t0 余额。",
                 "营销费用 = max(费率 × 年收入, 年度最低投入)，首年均摊到各月；人力 = 人数 × 月薪 × (1+社保) × 月数；折旧直线法按月排程。",
                 "营运资本：期末应收 / 存货 / 应付 = 当期收入或营业成本 ÷ 期间天数 × 周转天数。所得税按累计口径：亏损结转、不退税。",
                 "现金最低点（融资前）= min(期初现金, 期末现金 − 累计融资)，用于判断融资额是否覆盖缺口。",
                 "openpyxl 不计算公式：Excel / WPS 打开时自动重算；账本 numbers.json 与摘要取 Python 值，检查页比较两者。"):
        ws.cell(row=r, column=1, value="•").font = Font(name="Arial", size=10)
        ws.cell(row=r, column=2, value=line).font = Font(name="Arial", size=10)
        r += 1
    if warnings:
        r += 1
        ws.cell(row=r, column=1, value="警告").font = Font(name="Arial", size=11, bold=True, color=theme_mod.NEG)
        r += 1
        for w in warnings:
            ws.cell(row=r, column=2, value=w).font = Font(name="Arial", size=10, color=theme_mod.NEG)
            r += 1


def write_startup_xlsx(b: StartupBundle, path: Path) -> Dict[str, str]:
    """写 startup 模式的 财务模型.xlsx；返回账本 key → 单元格地址（不含文件名）。"""
    res = b.res
    wb = Workbook()
    wb.remove(wb.active)
    for name in STARTUP_SHEETS:
        wb.create_sheet(name)
    lay = Layout(res)
    asm = _write_assumptions(wb, b)
    drv = _write_drivers(wb, lay, b)
    rev = _write_revenue(wb, lay, b)
    hc = _write_headcount(wb, lay, b)
    cost = _write_costs(wb, lay, b, hc)
    is_ = _write_income_statement(wb, lay, b, rev, cost, hc, asm)
    bs_rows = _bs_row_plan(res)
    cf = _write_cashflow(wb, lay, b, is_, cost, drv, bs_rows)
    bs = _write_balance_sheet(wb, lay, b, is_, cf, rev, cost, drv, asm, bs_rows)
    for k in ("cash", "ar", "inv", "ap"):
        assert bs.rows[k] == bs_rows[k], f"资产负债表 {k} 行号 {bs.rows[k]} ≠ 约定 {bs_rows[k]}"
    be_sheet, be_cells = _write_breakeven(wb, lay, b, is_)
    ue_sheet = _write_unit_economics(wb, b)
    _write_sensitivity(wb, b)
    _write_valuation(wb, b)
    fin_sheet = _write_financing(wb, b)
    checks = _write_checks(wb, lay, b, is_, cf, bs, rev)
    _write_cover(wb, b, checks, b.warnings)
    wb.active = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)

    cells: Dict[str, str] = {}
    for y in range(res.years):
        yc = Layout.L(lay.ycol[y])
        cells[f"revenue_y{y + 1}"] = f"利润表!{yc}{is_.rows['revenue']}"
        cells[f"net_profit_y{y + 1}"] = f"利润表!{yc}{is_.rows['net_profit']}"
        cells[f"gross_margin_y{y + 1}"] = f"利润表!{yc}{is_.rows['gross_margin']}"
    cells["breakeven_month"] = be_cells["breakeven_month"]
    cells["min_cash"] = f"现金流量表!B{cf.rows['min_cash']}"
    cells.update(fin_sheet.cells)
    cells.update(ue_sheet.cells)
    return cells


# ==============================================================================================
# startup：账本 / 摘要 / 图表
# ==============================================================================================

def _r(v: float, nd: int = 1) -> float:
    return round(float(v), nd) if math.isfinite(v) else 0.0


def _entry(key: str, label: str, value: Any, unit_: str, source: str) -> Dict[str, Any]:
    return {"key": key, "label": label, "value": value, "unit": unit_, "source": source, "usedIn": []}


def build_startup_numbers(b: StartupBundle, cells: Dict[str, str], xlsx_rel: str) -> Dict[str, Any]:
    res, an, u = b.res, b.res.annual, unit_label(b.res.unit)
    src = lambda k: f"{xlsx_rel}!{cells[k]}" if k in cells else xlsx_rel  # noqa: E731
    entries: List[Dict[str, Any]] = []
    for y in range(res.years):
        n = y + 1
        entries.append(_entry(f"revenue_y{n}", f"第 {n} 年营业收入", _r(an["revenue"][y] / res.unit), u, src(f"revenue_y{n}")))
    for y in range(res.years):
        n = y + 1
        entries.append(_entry(f"net_profit_y{n}", f"第 {n} 年净利润", _r(an["net_profit"][y] / res.unit), u, src(f"net_profit_y{n}")))
    for y in range(res.years):
        n = y + 1
        entries.append(_entry(f"gross_margin_y{n}", f"第 {n} 年毛利率", _r(an["gross_margin"][y], 4), "%", src(f"gross_margin_y{n}")))
    be = res.breakeven
    entries.append(_entry("breakeven_month", "盈亏平衡月份（累计经营利润转正）", be.month if be.month is not None else None, "月", src("breakeven_month")))
    entries.append(_entry("min_cash", "现金最低点（融资前，含期初现金）", _r(res.min_cash / res.unit), u, src("min_cash")))
    entries.append(_entry("financing_amount", "融资额", _r(b.fin.amount / res.unit), u, src("financing_amount")))
    entries.append(_entry("pre_money", "投前估值", _r(b.fin.pre_money / res.unit), u, src("pre_money")))
    entries.append(_entry("post_money", "投后估值", _r(b.fin.post_money / res.unit), u, src("post_money")))
    entries.append(_entry("equity_offered", "出让比例", _r(b.fin.equity_offered, 4), "%", src("equity_offered")))
    entries.append(_entry("cac", "获客成本 CAC", _r(b.ue.cac), "元", src("cac")))
    entries.append(_entry("ltv", "客户生命周期价值 LTV", _r(b.ue.ltv), "元", src("ltv")))
    entries.append(_entry("ltv_cac", "LTV : CAC", _r(b.ue.ltv_cac, 2), "倍", src("ltv_cac")))
    entries.append(_entry("payback_months", "CAC 回收期", _r(b.ue.payback_months), "月", src("payback_months")))
    return {"generatedAt": date.today().isoformat(), "unit": u, "mode": "startup", "company": res.a.get("company", ""), "entries": entries}


def _md_table(headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> str:
    out = ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    for r in rows:
        out.append("| " + " | ".join(str(x) for x in r) + " |")
    return "\n".join(out)


def _m(v: float, unit: float, nd: int = 1) -> str:
    return f"{v / unit:,.{nd}f}" if math.isfinite(v) else "—"


def _fmt_qty(v: float) -> str:
    """单价 / 数量：整数不带小数，否则保留 2 位。"""
    return f"{v:,.0f}" if abs(v - round(v)) < 1e-9 else f"{v:,.2f}"


def _acceptance(b: StartupBundle) -> List[Tuple[str, bool, str]]:
    """技能 fin-model §4 的验收规则 → (规则, 是否通过, 说明)。"""
    res, be, fin, val, ue = b.res, b.res.breakeven, b.fin, b.val, b.ue
    u = res.unit
    rules: List[Tuple[str, bool, str]] = []
    ok3 = all(v < 1.0 for v in res.checks.values())
    rules.append(("三表勾稽（误差 < 1 元）", ok3, "；".join(f"{k} {v:.2e}" for k, v in res.checks.items())))
    rules.append(("盈亏平衡 12–36 个月", be.month is not None and 12 <= be.month <= 36, f"第 {be.month} 月" if be.month is not None else "预测期内未转正"))
    rules.append(("现金最低点（融资前）+ 融资额 ≥ 0", res.min_cash + fin.amount >= -EPS, f"{_m(res.min_cash, u)} + {_m(fin.amount, u)} = {_m(res.min_cash + fin.amount, u)} {unit_label(u)}"))
    implied = _safe_div(fin.amount, fin.post_money)
    rules.append(("出让比例 = 融资额 ÷ 投后（±1 pp）", abs(implied - fin.equity_offered) <= 0.01, f"{fin.equity_offered:.1%} vs {implied:.1%}"))
    rules.append(("估值方法 ≥ 2 种", len(val.methods) >= 2, f"{len(val.methods)} 种，区间 {_m(val.low, u, 0)}–{_m(val.high, u, 0)} {unit_label(u)}"))
    rules.append(("资金用途合计 = 100%", (not fin.use) or abs(fin.use_total_pct - 1.0) <= 0.005, f"{fin.use_total_pct:.1%}" if fin.use else "未配置"))
    y1_np, y1_rev = res.annual["net_profit"][0], res.annual["revenue"][0]
    rules.append(("首年亏损或微利（净利率 ≤ 10%）", y1_np <= 0 or _safe_div(y1_np, y1_rev) <= 0.10, f"首年净利润 {_m(y1_np, u)} {unit_label(u)}"))
    rules.append(("LTV:CAC ≥ 3", ue.ltv_cac >= 3.0, f"{ue.ltv_cac:.1f}×"))
    return rules


def build_startup_summary(b: StartupBundle, files: Dict[str, str]) -> str:
    res, an, u, ul = b.res, b.res.annual, b.res.unit, unit_label(b.res.unit)
    yl = res.year_labels()
    be, ue, fin, val, se = res.breakeven, b.ue, b.fin, b.val, b.sens
    rules = _acceptance(b)
    passed = all(ok for _, ok, _ in rules)
    L: List[str] = [f"# 财务摘要 —— {res.a.get('company', '')}", "",
                    f"生成日期 {date.today().isoformat()} · 模式 startup · 金额单位 {ul} · 预测期 {yl[0]}–{yl[-1]}（首年按月） · 验收 {'全部通过' if passed else '有未通过项'}", "",
                    "## 1. 五年预测（基准情形）", ""]
    L.append(_md_table(["年份", "营业收入", "毛利", "毛利率", "净利润", "净利率", "期末现金", "融资前现金余额"],
                       [[yl[y], _m(an["revenue"][y], u), _m(an["gross_profit"][y], u), f"{an['gross_margin'][y]:.1%}", _m(an["net_profit"][y], u), f"{an['net_margin'][y]:.1%}",
                         _m(an["cash"][y], u), _m(an["cash_prefin"][y], u)] for y in range(res.years)]))
    L += ["", "收入结构：" + "；".join(f"{st.name} {_m(sum(st.revenue[i] for i in res.yidx[0]), u)} → {_m(sum(st.revenue[i] for i in res.yidx[-1]), u)}" for st in res.streams) + f" {ul}（首年 → 末年）", ""]
    L += ["## 2. 盈亏平衡", ""]
    if be.month is not None:
        L.append(f"- 累计经营利润首次转正：第 **{be.month}** 个月（{res.start_year + (be.month - 1) // 12} 年 {(be.month - 1) % 12 + 1} 月）；"
                 + (f"单期经营利润首次为正：{res.periods[be.first_positive_period].label}" if be.first_positive_period is not None else "单期经营利润预测期内未转正"))
    else:
        L.append("- 预测期内累计经营利润未转正（请检查成本结构或增长假设）")
    yt = be.year_table[be.chart_year - 1]
    L.append(f"- 第 {be.chart_year} 年口径：边际贡献率 {yt['cm_ratio']:.1%}，固定性成本 {_m(yt['fixed_like'], u)} {ul}，保本收入 {_m(yt['be_revenue'], u)} {ul}，安全边际 "
             + (f"{yt['margin_of_safety']:.1%}" if math.isfinite(yt["margin_of_safety"]) else "—"))
    L.append(f"- 保本销量（主力产品线「{be.primary_stream}」）：{be.be_units:,.0f} {be.unit_name}（单价 {be.unit_price:,.2f} 元，单位变动成本 {be.unit_var_cost:,.2f} 元），预计销量 {be.actual_units:,.0f} {be.unit_name}"
             if math.isfinite(be.be_units) else "- 保本销量：单价不高于单位变动成本，不存在保本点")
    L += ["", "## 3. 单位经济", ""]
    L.append(_md_table(["指标", "值", "口径"], [["CAC", f"{ue.cac:,.0f} 元", ue.cac_source], ["客单价 × 毛利率", f"{ue.price:,.0f} 元 × {ue.margin:.0%}", f"主力产品线「{be.primary_stream}」"],
                                             ["LTV", f"{ue.ltv:,.0f} 元", f"客单价 × 毛利率 × {ue.lifetime_orders:.1f} 次"], ["LTV : CAC", f"{ue.ltv_cac:.1f}×", "≥ 3 健康；> 5 可能获客投入不足"],
                                             ["回收期", f"{ue.payback_months:.1f} 月", f"CAC ÷ (LTV ÷ {ue.lifetime_months:.0f} 月)"]]))
    L += ["", "## 4. 三情景", ""]
    L.append(_md_table(["情景", "需求乘数", f"收入 {yl[0]}", f"收入 {yl[-1]}", f"净利润 {yl[0]}", f"净利润 {yl[-1]}", "现金最低点（融资前）", "盈亏平衡月", "资金缺口"],
                       [[s.name, f"{s.multiplier:.2f}", _m(s.revenue[0], u), _m(s.revenue[-1], u), _m(s.net_profit[0], u), _m(s.net_profit[-1], u), _m(s.min_cash, u),
                         s.breakeven_month if s.breakeven_month is not None else "未达到", _m(s.funding_gap, u)] for s in b.scenarios]))
    pes = next((s for s in b.scenarios if s.key == "pessimistic"), None)
    if pes:
        L.append("")
        L.append(f"悲观情景：现金最低点 {_m(pes.min_cash, u)} {ul}，" + (f"融资后在第 {pes.runway_months} 个月出现现金为负，需追加融资或降本" if pes.runway_months else f"融资 {_m(fin.amount, u)} {ul} 后预测期内现金不为负")
                 + (f"；盈亏平衡推迟到第 {pes.breakeven_month} 月" if pes.breakeven_month else "；预测期内未达盈亏平衡"))
    L += ["", f"## 5. 敏感性排序（第 {se.target_year} 年净利润，{ul}）", ""]
    L.append(_md_table(["驱动因子"] + [f"{d:+.0%}" for d in se.deltas] + ["摆幅"],
                       [[SENS_DRIVERS[d][0]] + [_m(se.cells[(d, dl)].net_profit_y, u) for dl in se.deltas] + [_m(sw, u)] for d, sw in se.ranking]))
    if se.ranking:
        L.append("")
        L.append(f"最敏感因子：{SENS_DRIVERS[se.ranking[0][0]][0]}；现金最低点对 " + "、".join(SENS_DRIVERS[d][0] for d, _ in se.ranking[:2]) + " 最敏感。")
    L += ["", f"## 6. 估值区间（投前，{ul}）", ""]
    L.append(_md_table(["方法", "结果", "说明"], [[m["name"], _m(m["value"], u, 0), m["detail"]] for m in val.methods]))
    L.append("")
    L.append(f"区间 **{_m(val.low, u, 0)} – {_m(val.high, u, 0)}** {ul}，中值 {_m(val.mid, u, 0)}；本轮采用投前 **{_m(val.chosen_pre, u, 0)}** {ul}（{val.chosen_source}）。")
    for n in val.notes:
        L.append(f"- 注：{n}")
    L += ["", "## 7. 融资三件套", ""]
    L.append(f"- 融资额 **{_m(fin.amount, u, 0)}** {ul} · 投前 **{_m(fin.pre_money, u, 0)}** · 投后 **{_m(fin.post_money, u, 0)}** · 出让 **{fin.equity_offered:.1%}**（= 融资额 ÷ 投后）")
    L.append(f"- 现金缺口（融资前最低点取反）{_m(fin.gap, u)} {ul}；融资额 ÷ 缺口 = " + (f"{fin.coverage:.1f}×" if math.isfinite(fin.coverage) else "无缺口") + f"；融资后现金最低点 {_m(res.min_cash_post, u)} {ul}")
    if fin.use:
        L.append("")
        L.append(_md_table(["资金用途", "占比", f"金额（{ul}）"], [[i, f"{p:.0%}", _m(amt, u)] for i, p, amt in fin.use] + [["合计", f"{fin.use_total_pct:.0%}", _m(sum(x for _, _, x in fin.use), u)]]))
    for n in fin.notes:
        L.append(f"- 注：{n}")
    L += ["", "## 8. 检查结果", ""]
    L.append(_md_table(["规则", "结果", "说明"], [[r_, "✅ 通过" if ok else "❌ 未通过", d] for r_, ok, d in rules]))
    L.append("")
    L.append("Excel 检查页：资产 = 负债 + 权益、现金流期末 = 货币资金、Δ留存收益 = 净利润、收入分项加总、公式 vs Python 均以公式计算，打开 Excel 重算后 Master Check 应显示「通过」。")
    L += ["", "## 9. 逻辑解释", ""]
    growth = [(an["revenue"][y] / an["revenue"][y - 1] - 1) if an["revenue"][y - 1] > EPS else 0.0 for y in range(1, res.years)]
    L.append(f"收入由 {len(res.streams)} 条产品线的驱动因子推出（" + "；".join(f"{st.name}：{st.driver_name}" for st in res.streams) + "），首年按月爬坡，之后逐年增速 "
             + "、".join(f"{g:.0%}" for g in growth) + f"。首年营销费用 {_m(an['marketing'][0], u)} {ul}（费率 {an['marketing_pct'][0]:.0%} 与最低投入取大）与人力成本 {_m(an['headcount'][0], u)} {ul} 构成主要固定性支出，"
             + f"首年净利润 {_m(an['net_profit'][0], u)} {ul}；"
             + (f"累计经营利润在第 {be.month} 个月转正，" if be.month else "预测期内累计经营利润未转正，")
             + f"末年收入 {_m(an['revenue'][-1], u)} {ul}、净利率 {an['net_margin'][-1]:.1%}。营运资本（DSO/DIO/DPO）使现金流滞后于利润，融资前现金最低点 {_m(res.min_cash, u)} {ul}，"
             + f"融资 {_m(fin.amount, u, 0)} {ul} 可覆盖缺口并留出安全垫；估值取三法区间 {_m(val.low, u, 0)}–{_m(val.high, u, 0)} {ul} 的中值附近，出让 {fin.equity_offered:.1%} 与融资额 ÷ 投后一致。")
    if b.warnings:
        L += ["", "## 警告", ""] + [f"- {w}" for w in b.warnings]
    L += ["", "---", f"文件：{files.get('xlsx', '')} · {files.get('numbers', '')} · {files.get('charts', '')}", ""]
    return "\n".join(L)


def build_startup_charts(b: StartupBundle, xlsx_rel: str, warnings: List[str]) -> Dict[str, Any]:
    res, an, u, ul = b.res, b.res.annual, b.res.unit, unit_label(b.res.unit)
    yl = res.year_labels()
    company = str(res.a.get("company") or "")
    src = lambda page: f"资料来源：{xlsx_rel} {page}（基准情形）"  # noqa: E731
    charts: List[Dict[str, Any]] = []
    charts.append({"id": "fig_fin_revenue_mix", "type": "stacked_bar", "title": f"{company}收入结构预测（{yl[0]}—{yl[-1]}）",
                   "data": {"categories": yl, "series": [{"name": st.name, "values": [_r(v / u) for v in _annual_flow(st.revenue, res.yidx)]} for st in res.streams]},
                   "yLabel": "营业收入", "unit": ul, "valueLabels": True, "source": src("收入页")})
    charts.append({"id": "fig_fin_revenue_profit", "type": "line", "title": f"{company}营业收入与净利润预测",
                   "data": {"categories": yl, "series": [{"name": "营业收入", "values": [_r(v / u) for v in an["revenue"]]}, {"name": "净利润", "values": [_r(v / u) for v in an["net_profit"]]}]},
                   "yLabel": "金额", "unit": ul, "source": src("利润表")})
    cats = [(p.label[2:] if p.month is not None else p.label) for p in res.periods]
    fcf = [_r(v / u) for v in res.s["fcf"]]
    cum = [_r(v / u) for v in res.s["cash_prefin"]]
    charts.append({"id": "fig_fin_cashflow", "type": "cashflow_combo", "title": f"{company}现金流预测（融资前，{yl[0]} 按月）",
                   "data": {"categories": cats, "values": fcf, "cumulative": cum}, "yLabel": "现金流", "unit": ul,
                   "note": f"累计线 = 融资前现金余额（含期初现金 {_m(res.a_opening(), u)} {ul}）；融资 {_m(b.fin.amount, u, 0)} {ul} 覆盖最低点 {_m(res.min_cash, u)} {ul}",
                   "source": src("现金流量表"), "width": 170, "height": 90})
    be = res.breakeven
    if math.isfinite(be.be_units) and be.unit_price > be.unit_var_cost and be.fixed_cost > 0:
        charts.append({"id": "fig_fin_breakeven", "type": "breakeven", "title": f"「{be.primary_stream}」盈亏平衡分析（第 {be.chart_year} 年口径）",
                       "data": {"fixedCost": round(be.fixed_cost, 2), "unitPrice": round(be.unit_price, 2), "unitVarCost": round(be.unit_var_cost, 2),
                                "maxUnits": int(max(be.be_units * 1.6, be.actual_units * 1.2, 10)), "unitName": be.unit_name},
                       "unit": "元", "source": src("盈亏平衡页")})
    else:
        warnings.append("盈亏平衡图未生成：单价不高于单位变动成本或固定成本为 0")
    se = b.sens
    if se.ranking:
        items = [{"name": SENS_DRIVERS[d][0], "low": _r(se.cells[(d, se.deltas[0])].net_profit_y / u), "high": _r(se.cells[(d, se.deltas[-1])].net_profit_y / u)} for d, _ in se.ranking]
        charts.append({"id": "fig_fin_tornado", "type": "tornado", "title": f"{yl[se.target_year - 1]} 年净利润敏感性分析（{se.deltas[0]:+.0%} / {se.deltas[-1]:+.0%}）",
                       "data": {"base": _r(se.base_net_profit / u), "items": items}, "xLabel": "净利润", "unit": ul, "source": src("敏感性页")})
    if b.fin.use and sum(amt for _, _, amt in b.fin.use) > 0:
        charts.append({"id": "fig_fin_use_of_funds", "type": "pie_as_bar", "title": f"融资 {_m(b.fin.amount, u, 0)} {ul} 资金用途",
                       "data": {"categories": [i for i, _, _ in b.fin.use], "values": [_r(amt / u) for _, _, amt in b.fin.use]}, "yLabel": "金额", "unit": ul, "source": src("融资页")})
    return {"outputDir": "04_图表", "dpi": 300, "theme": "bizgold", "charts": charts}


def _startup_warnings(b: StartupBundle) -> List[str]:
    res, an = b.res, b.res.annual
    w: List[str] = list(res.warnings)
    for rule, ok, detail in _acceptance(b):
        if not ok:
            w.append(f"验收未通过：{rule}（{detail}）")
    if any(an["net_margin"][y] > 0.30 for y in range(res.years)):
        w.append("某年净利率 > 30%，评委可能质疑盈利能力假设")
    if b.fin.equity_offered > 0.20 + EPS:
        w.append(f"单轮出让 {b.fin.equity_offered:.1%} > 20%")
    if b.ue.ltv_cac > 5.0:
        w.append(f"LTV:CAC = {b.ue.ltv_cac:.1f} > 5，通常意味着获客投入不足或 CAC 偏低")
    growth = [(an["revenue"][y] / an["revenue"][y - 1] - 1) if an["revenue"][y - 1] > EPS else 0.0 for y in range(1, res.years)]
    if len(growth) >= 2 and all(abs(growth[i] - growth[i - 1]) < 0.05 for i in range(1, len(growth))):
        w.append("相邻年份收入增速差异均 < 5 个百分点，接近等比增长，建议调整增长假设")
    if b.sens.skipped:
        w.append(f"敏感性驱动因子未识别已跳过：{', '.join(b.sens.skipped)}（可选：{', '.join(SENS_DRIVERS)}）")
    for n in b.val.notes + b.fin.notes + b.ue.notes:
        if n not in w:
            w.append(n)
    return w


# ==============================================================================================
# budget 模式
# ==============================================================================================

@dataclass
class BudgetItem:
    item: str
    unit_price: float
    quantity: float
    unit: str
    amount: float
    source: str
    category: str


@dataclass
class BudgetPhase:
    name: str
    items: List[BudgetItem]
    subtotal: float
    start: Optional[Tuple[int, int]]
    end: Optional[Tuple[int, int]]


@dataclass
class BudgetResult:
    a: Dict[str, Any]
    unit: float
    phases: List[BudgetPhase]
    subtotal: float
    contingency_pct: float
    contingency: float
    total: float
    total_cap: Optional[float]
    kpis: List[Dict[str, Any]]
    roi: Optional[Dict[str, Any]]
    headcount: Optional[float]
    months: List[Tuple[int, int]]
    schedule: List[List[float]]            # phases × months（元）
    warnings: List[str] = field(default_factory=list)

    @property
    def within_cap(self) -> bool:
        return self.total_cap is None or self.total <= self.total_cap + EPS


_RANGE_RE = re.compile(r"(\d{4})\s*[.\-/年]\s*(\d{1,2})\s*月?\s*[–\-—~至到]\s*(?:(\d{4})\s*[.\-/年]\s*)?(\d{1,2})")
_SINGLE_RE = re.compile(r"(\d{4})\s*[.\-/年]\s*(\d{1,2})")


def _parse_ym(v: Any) -> Optional[Tuple[int, int]]:
    if not v:
        return None
    m = _SINGLE_RE.search(str(v))
    if not m:
        return None
    y, mo = int(m.group(1)), int(m.group(2))
    return (y, mo) if 1 <= mo <= 12 else None


def _phase_window(ph: Dict[str, Any]) -> Tuple[Optional[Tuple[int, int]], Optional[Tuple[int, int]]]:
    s, e = _parse_ym(ph.get("start")), _parse_ym(ph.get("end"))
    if s and e:
        return s, e
    m = _RANGE_RE.search(str(ph.get("name") or ""))
    if m:
        y1, m1 = int(m.group(1)), int(m.group(2))
        y2 = int(m.group(3)) if m.group(3) else y1
        m2 = int(m.group(4))
        if 1 <= m1 <= 12 and 1 <= m2 <= 12:
            return (y1, m1), (y2, m2)
    return s, e


def _month_span(a: Tuple[int, int], b_: Tuple[int, int]) -> List[Tuple[int, int]]:
    out = []
    y, m = a
    while (y, m) <= b_ and len(out) < 120:
        out.append((y, m))
        m += 1
        if m > 12:
            y, m = y + 1, 1
    return out


def compute_budget(a: Dict[str, Any]) -> BudgetResult:
    unit = _num(a.get("unit"), 10000.0, "unit") or 10000.0
    warnings: List[str] = []
    raw_phases = a.get("phases")
    if not isinstance(raw_phases, list) or not raw_phases:
        raise FinModelError("budget 模式需要 phases 数组")
    phases: List[BudgetPhase] = []
    for pi, ph in enumerate(raw_phases):
        if not isinstance(ph, dict):
            raise FinModelError("phases 的每一项必须是对象")
        name = str(ph.get("name") or f"第 {pi + 1} 阶段")
        items: List[BudgetItem] = []
        for it in ph.get("items") or []:
            nm = str(it.get("item") or "项目")
            up = _num(it.get("unitPrice"), 0.0, f"{nm}.unitPrice")
            qty = _num(it.get("quantity"), 0.0, f"{nm}.quantity")
            items.append(BudgetItem(nm, up, qty, str(it.get("unit") or ""), up * qty, str(it.get("source") or ""), str(it.get("category") or nm)))
            if not it.get("source"):
                warnings.append(f"预算项「{nm}」缺少来源（刊例 / 报价）")
        s, e = _phase_window(ph)
        phases.append(BudgetPhase(name, items, sum(i.amount for i in items), s, e))
    subtotal = sum(p.subtotal for p in phases)
    c_pct = _num(a.get("contingencyPct"), 0.0, "contingencyPct")
    contingency = subtotal * c_pct
    total = subtotal + contingency
    cap = a.get("totalCap")
    total_cap = _num(cap, 0.0, "totalCap") if cap is not None else None
    if total_cap is not None and total > total_cap + EPS:
        warnings.append(f"预算总计 {total / unit:,.1f} {unit_label(unit)} 超过上限 {total_cap / unit:,.1f} {unit_label(unit)}（超出 {(total - total_cap) / unit:,.1f}）")
    kpis = [k for k in (a.get("kpis") or []) if isinstance(k, dict)]
    roi_cfg = a.get("roi") if isinstance(a.get("roi"), dict) else None
    roi = None
    if roi_cfg:
        inc = _num(roi_cfg.get("incrementalRevenue"), 0.0, "roi.incrementalRevenue")
        gm = _num(roi_cfg.get("grossMargin"), 0.0, "roi.grossMargin")
        inc_gp = inc * gm
        roi = {"incremental_revenue": inc, "gross_margin": gm, "incremental_gp": inc_gp, "budget": total,
               "roi": _safe_div(inc_gp - total, total), "roas": _safe_div(inc, total), "note": str(roi_cfg.get("brandAssetNote") or "")}
    hc = a.get("headcount")
    headcount = _num(hc, 0.0, "headcount") if hc is not None else None
    if headcount is not None and headcount <= 0:
        headcount = None
    # 排期：有起止月的阶段均摊到各月；没有的阶段按顺序接续（每阶段 3 个月）
    windows: List[List[Tuple[int, int]]] = []
    cursor: Optional[Tuple[int, int]] = None
    for p in phases:
        if p.start and p.end and p.start <= p.end:
            win = _month_span(p.start, p.end)
        else:
            if cursor is None:
                cursor = (int(_num(a.get("startYear"), date.today().year)), 1)
            win = _month_span(cursor, (cursor[0] + (cursor[1] + 1) // 12, (cursor[1] + 2 - 1) % 12 + 1))
            warnings.append(f"阶段「{p.name}」未给出起止月（start/end 或名称中的 2026.4–6），排期按顺序各 3 个月假定")
        p.start, p.end = win[0], win[-1]
        windows.append(win)
        nxt = (win[-1][0], win[-1][1] + 1)
        cursor = (nxt[0] + 1, 1) if nxt[1] > 12 else nxt
    all_months = sorted({m for w in windows for m in w})
    schedule = [[(p.subtotal / len(w) if m in w else 0.0) for m in all_months] for p, w in zip(phases, windows)]
    return BudgetResult(a, unit, phases, subtotal, c_pct, contingency, total, total_cap, kpis, roi, headcount, all_months, schedule, warnings)


def write_budget_xlsx(br: BudgetResult, path: Path) -> Dict[str, str]:
    wb = Workbook()
    wb.remove(wb.active)
    for name in BUDGET_SHEETS:
        wb.create_sheet(name)
    cells: Dict[str, str] = {}
    # ---- 预算总表
    t = TableSheet(wb, "预算总表", [30, 30, 14, 12, 10, 16, 10, 44, 3, 16])
    t.title("预算总表（金额 = 单价 × 数量，公式）", "蓝字浅黄底 = 输入（单价 / 数量 / 机动金比例 / 预算上限）；黑字 = 公式；右侧灰字 = Python 校验值。金额单位：元。")
    t.header(["阶段", "项目", "单价（元）", "数量", "单位", "金额（元）", "占总计", "来源", "", "数值校验（元）"])
    sub_rows: List[int] = []
    item_rows: List[int] = []
    for ph in br.phases:
        for it in ph.items:
            r = t.row
            item_rows.append(r)
            t.line([ph.name, it.item, it.unit_price, it.quantity, it.unit, f"=C{r}*D{r}", None, it.source],
                   kinds=["text", "text", "input", "input", "text", "formula", None, "text"], fmts=[None, None, FMT_PRICE, FMT_INT, None, FMT_INT, FMT_PCT, None])
        r = t.row
        first_item = r - len(ph.items)
        t.line([f"{ph.name} 小计", None, None, None, None, f"=SUM(F{first_item}:F{r - 1})" if ph.items else 0, None, None, None, ph.subtotal],
               kinds=["text", None, None, None, None, "formula", None, None, None, "check"], fmts=[None, None, None, None, None, FMT_INT, FMT_PCT, None, None, FMT_INT], bold=True)
        sub_rows.append(r)
    t.blank()
    r_sub = t.line(["合计（不含机动金）", None, None, None, None, "=" + "+".join(f"F{r}" for r in sub_rows), None, None, None, br.subtotal],
                   kinds=["text", None, None, None, None, "formula", None, None, None, "check"], fmts=[None, None, None, None, None, FMT_INT, None, None, None, FMT_INT], bold=True)
    r_cont = t.line(["机动金（不可预见费）", "合计 × 比例", None, br.contingency_pct, "", f"=F{r_sub}*D{t.row}", None, "运营费用的 5–10%", None, br.contingency],
                    kinds=["text", "text", None, "input", "text", "formula", None, "text", None, "check"], fmts=[None, None, None, FMT_PCT, None, FMT_INT, None, None, None, FMT_INT])
    r_total = t.line(["预算总计", None, None, None, None, f"=F{r_sub}+F{r_cont}", None, None, None, br.total],
                     kinds=["text", None, None, None, None, "formula", None, None, None, "check"], fmts=[None, None, None, None, None, FMT_INT, None, None, None, FMT_INT], bold=True)
    cells["budget_total"] = f"预算总表!F{r_total}"
    cells["budget_contingency"] = f"预算总表!F{r_cont}"
    if br.total_cap is not None:
        r_cap = t.line(["预算上限（命题 / totalCap）", None, None, None, None, br.total_cap, None, "不得超过"], kinds=["text", None, None, None, None, "input", None, "text"],
                       fmts=[None, None, None, None, None, FMT_INT, None, None])
        t.line(["余量 = 上限 − 总计", None, None, None, None, f"=F{r_cap}-F{r_total}", None, f'=IF(F{r_cap}>=F{r_total},"在预算内","超出预算")'],
               kinds=["text", None, None, None, None, "formula", None, "formula"], fmts=[None, None, None, None, None, FMT_INT, None, None], bold=True)
    if br.headcount:
        r_hc = t.line(["团队人数", None, None, br.headcount, "人", None, None, "人均预算 = 总计 ÷ 人数"], kinds=["text", None, None, "input", "text"], fmts=[None, None, None, FMT_INT, None])
        r_pp = t.line(["人均预算（元/人）", None, None, None, None, f"=F{r_total}/D{r_hc}", None, None, None, br.total / br.headcount],
                      kinds=["text", None, None, None, None, "formula", None, None, None, "check"], fmts=[None, None, None, None, None, FMT_INT, None, None, None, FMT_INT])
        cells["budget_per_person"] = f"预算总表!F{r_pp}"
    # 占总计列（公式）
    for r in item_rows + sub_rows + [r_sub, r_cont]:
        c = t.ws.cell(row=r, column=7, value=f"=IF($F${r_total}=0,0,F{r}/$F${r_total})")
        _style(c, "formula", FMT_PCT)
    t.ws.freeze_panes = "A4"
    # ---- 排期
    s = TableSheet(wb, "排期", [30, 10, 10, 16, 10] + [11] * len(br.months))
    s.title("执行排期与月度支出计划（阶段预算在起止月之间均摊）", "金额单位：元。")
    s.header(["阶段", "开始", "结束", "阶段预算（元）", "占比"] + [f"{y}-{m:02d}" for y, m in br.months])
    first = s.row
    for ph, row in zip(br.phases, br.schedule):
        s.line([ph.name, f"{ph.start[0]}-{ph.start[1]:02d}" if ph.start else "", f"{ph.end[0]}-{ph.end[1]:02d}" if ph.end else "", ph.subtotal,
                _safe_div(ph.subtotal, br.subtotal)] + row, kinds=["text", "text", "text", "formula", "formula"] + ["formula"] * len(row),
               fmts=[None, None, None, FMT_INT, FMT_PCT] + [FMT_INT] * len(row))
    last = s.row - 1
    if br.months:
        cols = [get_column_letter(6 + i) for i in range(len(br.months))]
        s.line(["月度合计", None, None, f"=SUM(D{first}:D{last})", None] + [f"=SUM({c}{first}:{c}{last})" for c in cols],
               kinds=["text", None, None, "formula", None] + ["formula"] * len(cols), fmts=[None, None, None, FMT_INT, None] + [FMT_INT] * len(cols), bold=True)
        r_tot = s.row - 1
        s.line(["累计支出", None, None, None, None] + [f"=SUM($F${r_tot}:{c}{r_tot})" for c in cols],
               kinds=["text", None, None, None, None] + ["formula"] * len(cols), fmts=[None, None, None, None, None] + [FMT_INT] * len(cols))
    s.ws.freeze_panes = "B4"
    # ---- KPI
    k = TableSheet(wb, "KPI", [22, 16, 18, 20, 10, 30])
    k.title("效果评估 KPI", "每个指标：基线 → 目标，监测工具与周期；与正文传播目标一致。")
    k.header(["指标", "基线", "目标", "监测工具", "周期", "对应阶段 / 说明"])
    for i, kp in enumerate(br.kpis):
        r = k.line([str(kp.get("metric") or f"KPI {i + 1}"), str(kp.get("baseline") or "—"), str(kp.get("target") or ""), str(kp.get("tool") or ""), str(kp.get("cycle") or ""), str(kp.get("phase") or kp.get("note") or "")])
        cells[f"kpi_{i + 1}"] = f"KPI!C{r}"
    if not br.kpis:
        k.line(["（未配置 kpis）"])
    # ---- ROI（可选）
    if br.roi:
        ro = TableSheet(wb, "ROI", [34, 18, 50])
        ro.title("投入产出（ROI）", "ROI = (增量毛利 − 预算总计) ÷ 预算总计；ROAS = 增量收入 ÷ 预算总计。金额单位：元。")
        ro.header(["指标", "值", "说明"])
        r_inc = ro.line(["增量收入", br.roi["incremental_revenue"], "roi.incrementalRevenue"], kinds=["text", "input", "text"], fmts=[None, FMT_INT, None])
        r_gm = ro.line(["毛利率", br.roi["gross_margin"], "roi.grossMargin"], kinds=["text", "input", "text"], fmts=[None, FMT_PCT, None])
        r_gp = ro.line(["增量毛利", f"=B{r_inc}*B{r_gm}", "增量收入 × 毛利率"], kinds=["text", "formula", "text"], fmts=[None, FMT_INT, None])
        r_b = ro.line(["预算总计", f"=预算总表!F{r_total}", "引用预算总表"], kinds=["text", "xref", "text"], fmts=[None, FMT_INT, None])
        r_roi = ro.line(["ROI", f"=IF(B{r_b}=0,0,(B{r_gp}-B{r_b})/B{r_b})", f"Python 值 {br.roi['roi']:.1%}"], kinds=["text", "formula", "text"], fmts=[None, FMT_PCT, None], bold=True)
        ro.line(["ROAS（倍）", f"=IF(B{r_b}=0,0,B{r_inc}/B{r_b})", f"Python 值 {br.roi['roas']:.2f}×"], kinds=["text", "formula", "text"], fmts=[None, FMT_X, None])
        if br.roi["note"]:
            ro.line(["品牌资产目标", br.roi["note"], "不可直接货币化的效果"])
        cells["roi"] = f"ROI!B{r_roi}"
    # ---- 说明
    ws = wb["说明"]
    _set_widths(ws, {1: 30, 2: 22, 3: 70})
    _print_setup(ws)
    ul = unit_label(br.unit)
    ws.cell(row=1, column=1, value="预算模型（budget 模式）").font = Font(name="Arial", size=16, bold=True, color=theme_mod.NAVY)
    ws.cell(row=2, column=1, value=f"生成日期 {date.today().isoformat()} · 由 bizkit.finmodel 按 assumptions.json 生成 · 表内金额单位：元；账本单位：{ul}").font = Font(name="Arial", size=9, color=theme_mod.MUTED)
    r = 4
    for label, v, note in ((f"预算总计（{ul}）", br.total / br.unit, f"含机动金 {br.contingency / br.unit:,.1f}（{br.contingency_pct:.0%}）"),
                           (f"预算上限（{ul}）", (br.total_cap / br.unit) if br.total_cap is not None else "未设", "在预算内" if br.within_cap else "超出预算！"),
                           ("阶段数 / 项目数", f"{len(br.phases)} / {sum(len(p.items) for p in br.phases)}", ""),
                           ("KPI 数", len(br.kpis), "")):
        ws.cell(row=r, column=1, value=label).font = Font(name="Arial", size=10)
        c = ws.cell(row=r, column=2, value=v)
        c.font = Font(name="Arial", size=10, bold=True)
        if isinstance(v, float):
            c.number_format = FMT_DEC1
        ws.cell(row=r, column=3, value=note).font = Font(name="Arial", size=10, color=theme_mod.NEG if "超出" in note else theme_mod.MUTED)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="颜色约定").font = Font(name="Arial", size=11, bold=True)
    r += 1
    for kind, text in (("input", "蓝字浅黄底 = 硬输入（单价 / 数量 / 比例 / 上限）"), ("formula", "黑字 = 公式"), ("xref", "绿字 = 跨表引用"), ("check", "灰色斜体 = Python 校验值")):
        _style(ws.cell(row=r, column=2, value="示例 1,234"), kind)
        ws.cell(row=r, column=3, value=text).font = Font(name="Arial", size=10)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="页签目录").font = Font(name="Arial", size=11, bold=True)
    r += 1
    for i, (nm, d) in enumerate((("说明", "口径与目录"), ("预算总表", "阶段 × 项目：单价 × 数量 = 金额、小计、机动金、总计与上限对比"), ("排期", "阶段起止与月度支出计划"), ("KPI", "效果评估指标")) + ((("ROI", "投入产出"),) if br.roi else ())):
        ws.cell(row=r, column=1, value=f"{i + 1}. {nm}").font = Font(name="Arial", size=10)
        ws.cell(row=r, column=3, value=d).font = Font(name="Arial", size=10)
        r += 1
    if br.warnings:
        r += 1
        ws.cell(row=r, column=1, value="警告").font = Font(name="Arial", size=11, bold=True, color=theme_mod.NEG)
        r += 1
        for w in br.warnings:
            ws.cell(row=r, column=2, value=w).font = Font(name="Arial", size=10, color=theme_mod.NEG)
            r += 1
    wb.active = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return cells


_NUM_RE = re.compile(r"^\s*([-+]?\d[\d,]*(?:\.\d+)?)\s*(.*)$")


def _parse_target(s: Any) -> Tuple[Any, str]:
    """'5000 万' → (5000, '万')；'+120%' → (1.2, '%')；解析失败返回 (原文, '')。"""
    if isinstance(s, (int, float)) and not isinstance(s, bool):
        return s, ""
    m = _NUM_RE.match(str(s or ""))
    if not m:
        return str(s or ""), ""
    v = float(m.group(1).replace(",", ""))
    rest = m.group(2).strip()
    if rest.startswith("%") or rest.startswith("％"):
        return round(v / 100.0, 4), "%"
    return (int(v) if v.is_integer() else v), rest


def build_budget_numbers(br: BudgetResult, cells: Dict[str, str], xlsx_rel: str) -> Dict[str, Any]:
    ul = unit_label(br.unit)
    src = lambda k: f"{xlsx_rel}!{cells[k]}" if k in cells else xlsx_rel  # noqa: E731
    entries = [_entry("budget_total", "预算总额（含机动金）", _r(br.total / br.unit), ul, src("budget_total")),
               _entry("budget_contingency", f"机动金（{br.contingency_pct:.0%}）", _r(br.contingency / br.unit), ul, src("budget_contingency"))]
    if br.headcount:
        entries.append(_entry("budget_per_person", "人均预算", _r(br.total / br.headcount / br.unit, 2), ul, src("budget_per_person")))
    for i, kp in enumerate(br.kpis):
        v, unit_ = _parse_target(kp.get("target"))
        entries.append(_entry(f"kpi_{i + 1}", str(kp.get("metric") or f"KPI {i + 1}"), v, unit_, src(f"kpi_{i + 1}")))
    if br.roi:
        entries.append(_entry("roi", "投入产出比 ROI", _r(br.roi["roi"], 4), "%", src("roi")))
    return {"generatedAt": date.today().isoformat(), "unit": ul, "mode": "budget", "entries": entries}


def build_budget_summary(br: BudgetResult, files: Dict[str, str]) -> str:
    u, ul = br.unit, unit_label(br.unit)
    L = ["# 财务摘要 —— 预算模型", "", f"生成日期 {date.today().isoformat()} · 模式 budget · 金额单位 {ul}", "", "## 1. 预算总览", ""]
    L.append(_md_table(["项目", f"金额（{ul}）", "说明"],
                       [["合计（不含机动金）", _m(br.subtotal, u), f"{len(br.phases)} 个阶段、{sum(len(p.items) for p in br.phases)} 个项目"],
                        [f"机动金（{br.contingency_pct:.0%}）", _m(br.contingency, u), "不可预见费"],
                        ["**预算总计**", f"**{_m(br.total, u)}**", ("在上限内" if br.within_cap else "**超出上限**") if br.total_cap is not None else "未设上限"],
                        ["预算上限", _m(br.total_cap, u) if br.total_cap is not None else "—", f"余量 {_m(br.total_cap - br.total, u)}" if br.total_cap is not None else ""]]
                       + ([["人均预算", _m(br.total / br.headcount, u, 2), f"团队 {br.headcount:.0f} 人"]] if br.headcount else [])))
    L += ["", "## 2. 分阶段预算", ""]
    L.append(_md_table(["阶段", "起止", f"预算（{ul}）", "占比", "项目数"],
                       [[p.name, f"{p.start[0]}-{p.start[1]:02d} 至 {p.end[0]}-{p.end[1]:02d}" if p.start and p.end else "—", _m(p.subtotal, u), f"{_safe_div(p.subtotal, br.subtotal):.0%}", len(p.items)] for p in br.phases]))
    items = sorted([(it, p.name) for p in br.phases for it in p.items], key=lambda t: -t[0].amount)[:8]
    L += ["", "## 3. 主要支出项（前 8）", ""]
    L.append(_md_table(["项目", "阶段", "单价 × 数量", f"金额（{ul}）", "来源"], [[it.item, ph, f"{_fmt_qty(it.unit_price)} 元 × {_fmt_qty(it.quantity)} {it.unit}", _m(it.amount, u), it.source or "【待补来源】"] for it, ph in items]))
    L += ["", "## 4. KPI", ""]
    if br.kpis:
        L.append(_md_table(["指标", "基线", "目标", "监测工具", "周期"], [[k.get("metric", ""), k.get("baseline", "—"), k.get("target", ""), k.get("tool", ""), k.get("cycle", "")] for k in br.kpis]))
    else:
        L.append("（未配置 kpis）")
    if br.roi:
        L += ["", "## 5. ROI", "", f"增量收入 {_m(br.roi['incremental_revenue'], u)} {ul} × 毛利率 {br.roi['gross_margin']:.0%} = 增量毛利 {_m(br.roi['incremental_gp'], u)} {ul}；"
              f"ROI = (增量毛利 − 预算) ÷ 预算 = **{br.roi['roi']:.1%}**，ROAS = {br.roi['roas']:.2f}×" + (f"；品牌资产：{br.roi['note']}" if br.roi["note"] else "")]
    L += ["", "## 6. 检查结果", ""]
    L.append(_md_table(["规则", "结果", "说明"], [["金额 = 单价 × 数量（公式）", "✅", "预算总表 F 列公式，右侧灰字为 Python 校验"],
                                             ["总计 ≤ 预算上限", ("✅ 通过" if br.within_cap else "❌ 未通过") if br.total_cap is not None else "— 未设上限", f"{_m(br.total, u)} vs {_m(br.total_cap, u) if br.total_cap is not None else '—'} {ul}"],
                                             ["每项有来源", "✅" if not any(not it.source for p in br.phases for it in p.items) else "❌", "缺来源项见警告"]]))
    L += ["", "## 7. 逻辑解释", "", f"预算按阶段 × 项目自下而上汇总：每项金额 = 单价 × 数量（单价来自刊例 / 报价，见来源列），阶段小计相加得 {_m(br.subtotal, u)} {ul}，"
          f"再按 {br.contingency_pct:.0%} 计提机动金 {_m(br.contingency, u)} {ul}，总计 {_m(br.total, u)} {ul}"
          + (f"，{'未超过' if br.within_cap else '超过'}命题上限 {_m(br.total_cap, u)} {ul}" if br.total_cap is not None else "") + "。排期页把各阶段预算在起止月之间均摊，用于现金安排与进度监控；KPI 页给出每个阶段的效果指标与监测工具。"]
    if br.warnings:
        L += ["", "## 警告", ""] + [f"- {w}" for w in br.warnings]
    L += ["", "---", f"文件：{files.get('xlsx', '')} · {files.get('numbers', '')} · {files.get('charts', '')}", ""]
    return "\n".join(L)


def build_budget_charts(br: BudgetResult, xlsx_rel: str) -> Dict[str, Any]:
    u, ul = br.unit, unit_label(br.unit)
    pairs = sorted([(it.item, it.amount) for p in br.phases for it in p.items], key=lambda t: -t[1])
    if len(pairs) > 8:
        head, tail = pairs[:7], pairs[7:]
        pairs = head + [("其他", sum(v for _, v in tail))]
    if br.contingency > 0:
        pairs.append(("机动金", br.contingency))
    pairs = [(n, v) for n, v in pairs if v > 0]
    charts: List[Dict[str, Any]] = []
    if pairs:
        charts.append({"id": "fig_fin_budget_mix", "type": "pie_as_bar", "title": f"预算结构（总计 {_m(br.total, u)} {ul}）",
                       "data": {"categories": [n for n, _ in pairs], "values": [_r(v / u) for _, v in pairs]}, "yLabel": "预算", "unit": ul,
                       "source": f"资料来源：{xlsx_rel} 预算总表"})
    cats_all: Dict[str, float] = {}
    for p in br.phases:
        for it in p.items:
            cats_all[it.category] = cats_all.get(it.category, 0.0) + it.amount
    top = [c for c, _ in sorted(cats_all.items(), key=lambda t: -t[1])[:5]]
    series_names = top + (["其他"] if len(cats_all) > len(top) else [])
    series = []
    for nm in series_names:
        vals = []
        for p in br.phases:
            if nm == "其他":
                vals.append(sum(it.amount for it in p.items if it.category not in top))
            else:
                vals.append(sum(it.amount for it in p.items if it.category == nm))
        series.append({"name": nm, "values": [_r(v / u) for v in vals]})
    if series and any(any(v > 0 for v in s_["values"]) for s_ in series):
        charts.append({"id": "fig_fin_budget_phases", "type": "stacked_bar", "title": "分阶段预算构成",
                       "data": {"categories": [re.sub(r"[（(].*?[)）]", "", p.name).strip() or p.name for p in br.phases], "series": series},
                       "yLabel": "预算", "unit": ul, "source": f"资料来源：{xlsx_rel} 预算总表 / 排期"})
    return {"outputDir": "04_图表", "dpi": 300, "theme": "bizgold", "charts": charts}


# ==============================================================================================
# 入口
# ==============================================================================================

def _resolve_out(assumptions_path: Path, out: Optional[str]) -> Path:
    if not out:
        return assumptions_path.parent
    p = Path(out)
    if p.is_absolute():
        return p
    root = find_workspace_root(assumptions_path.parent, probe=p)
    return (root / p).resolve()


def _xlsx_rel(assumptions_path: Path, out_dir: Path) -> str:
    root = find_workspace_root(assumptions_path.parent, probe=out_dir.name)
    try:
        rel = out_dir.resolve().relative_to(root.resolve())
        return (rel.as_posix() + "/" if str(rel) != "." else "") + "财务模型.xlsx"
    except ValueError:
        return f"{out_dir.name}/财务模型.xlsx"


def run(assumptions_path: Path, out: Optional[str] = None) -> Dict[str, Any]:
    """读取 assumptions.json，生成四个文件，返回结果字典（见模块文档）。"""
    assumptions_path = Path(assumptions_path).resolve()
    a = read_json(assumptions_path)
    if not isinstance(a, dict):
        raise FinModelError("assumptions.json 顶层必须是对象")
    mode = str(a.get("mode") or "startup")
    out_dir = _resolve_out(assumptions_path, out)
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx_rel = _xlsx_rel(assumptions_path, out_dir)
    files = {"xlsx": str(out_dir / "财务模型.xlsx"), "numbers": str(out_dir / "numbers.json"),
             "summary": str(out_dir / "财务摘要.md"), "charts": str(out_dir / "财务图表.charts.json")}
    if mode == "startup":
        res = compute_startup(a)
        ue = compute_unit_economics(res)
        scenarios = compute_scenarios(a)
        sens = compute_sensitivity(a, res)
        rounds = a.get("financing") or []
        first = rounds[0] if rounds else {}
        val = compute_valuation(res, _num(first.get("amount"), 0.0), int(_num(first.get("year"), 1)), int(_num(first.get("month"), 1)),
                                _num(first.get("preMoney"), 0.0) if first.get("preMoney") is not None else None)
        fin = compute_financing(res, val)
        b = StartupBundle(res, ue, scenarios, sens, val, fin, [])
        b.warnings = _startup_warnings(b)
        log(f"startup 模型：{len(res.periods)} 期，收入流 {len(res.streams)} 条，盈亏平衡 {res.breakeven.month} 月，现金最低点 {res.min_cash / res.unit:,.1f} {unit_label(res.unit)}")
        cells = write_startup_xlsx(b, Path(files["xlsx"]))
        write_json(files["numbers"], build_startup_numbers(b, cells, xlsx_rel))
        charts_cfg = build_startup_charts(b, xlsx_rel, b.warnings)
        write_json(files["charts"], charts_cfg)
        write_text(files["summary"], build_startup_summary(b, files))
        checks = {k: round(v, 6) for k, v in res.checks.items()}
        ok = all(v < 1.0 for v in res.checks.values())
        return {"ok": ok, "mode": mode, "files": files, "checks": checks, "warnings": b.warnings,
                "breakevenMonth": res.breakeven.month, "minCash": _r(res.min_cash / res.unit), "charts": len(charts_cfg["charts"])}
    if mode == "budget":
        br = compute_budget(a)
        log(f"budget 模型：{len(br.phases)} 阶段，总计 {br.total / br.unit:,.1f} {unit_label(br.unit)}" + ("" if br.within_cap else "（超出上限！）"))
        cells = write_budget_xlsx(br, Path(files["xlsx"]))
        write_json(files["numbers"], build_budget_numbers(br, cells, xlsx_rel))
        charts_cfg = build_budget_charts(br, xlsx_rel)
        write_json(files["charts"], charts_cfg)
        write_text(files["summary"], build_budget_summary(br, files))
        checks = {"balanceSheet": None, "cashflow": None, "retained": None, "budgetTotal": _r(br.total / br.unit),
                  "totalCap": _r(br.total_cap / br.unit) if br.total_cap is not None else None, "withinCap": br.within_cap}
        return {"ok": True, "mode": mode, "files": files, "checks": checks, "warnings": br.warnings, "charts": len(charts_cfg["charts"])}
    raise FinModelError(f"不支持的 mode「{mode}」，可选 startup | budget")


def main(argv: Optional[Sequence[str]] = None) -> int:
    force_utf8_stdout()
    ap = argparse.ArgumentParser(description="assumptions.json → 财务模型.xlsx + numbers.json + 财务摘要.md + 财务图表.charts.json")
    ap.add_argument("assumptions", help="03_财务/assumptions.json 路径")
    ap.add_argument("--out", default="", help="输出目录（默认与 assumptions.json 同目录；相对路径相对工作区根）")
    args = ap.parse_args(argv)
    try:
        result = run(Path(args.assumptions), args.out or None)
    except FinModelError as e:
        print(json.dumps({"ok": False, "error": str(e)}, ensure_ascii=False))
        return 1
    except FileNotFoundError as e:
        print(json.dumps({"ok": False, "error": f"文件不存在：{e}"}, ensure_ascii=False))
        return 1
    except Exception as e:  # 其他异常也以 JSON 报告，便于上层解析
        print(json.dumps({"ok": False, "error": f"{type(e).__name__}: {e}"}, ensure_ascii=False))
        return 1
    print(json.dumps(result, ensure_ascii=False))
    return 0 if result.get("ok") else 1


if __name__ == "__main__":
    sys.exit(main())
