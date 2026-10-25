"""机器检查器（契约 §7.7；标准 §6.6–§6.7）：按规则扫描工作区材料，写出 `07_评审/第N轮/机器检查.json|md`。

用法：
    python -m bizkit.checks --track brand --workspace . [--round N] [--variant accounting|retail]
    python -m bizkit.checks --track brand --workspace . --chapter 05_正文/02_品牌现状分析.md   # 单章快检，只打印 JSON

设计原则：
- 容错优先：输入文件缺失 → 对应检查项 `status=skipped`；单项内部异常 → 该项 skipped 且 `summary.degraded=true`；进程不崩溃。
- 阈值来自 `tracks/<id>/checks.json`（`track_defaults.load_track_config()` 已把两套字段名补齐），`requiredEvidence{}`
  决定赛道专项检查是否启用；`--variant` 合并 `requiredEvidenceByVariant` / `coverFieldsByVariant`。
- 每项结果 `{id, value, threshold, status(pass|warn|fail|skipped), detail, locations?}`；硬门槛按 `rubric.json.hardGates` 命名输出；
  `crossDoc` 给出账本（numbers.json）与正文 / PPT / 讲稿 / 题库之间的数字不一致清单。
- 启发式一律保守：宁可漏报也不乱报；所有命中带文件 / 行号 / 原句，便于智能体直接修复。
"""
from __future__ import annotations

import argparse
import json
import re
import statistics
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional, Sequence, Tuple

from . import track_defaults as td
from ._util import force_utf8_stdout, read_json, write_json, write_text

PASS, WARN, FAIL, SKIP = "pass", "warn", "fail", "skipped"
MAX_LOCATIONS = 40

CheckFn = Callable[["Workspace"], Dict[str, Any]]

# ---------------------------------------------------------------------------------------------- 正则与文本工具
_CJK_RE = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff]")
_LATIN_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9.%,]*")
_CITE_RE = re.compile(r"\[\d+(?:\s*[-–—,，]\s*\d+)*\]")
_FOOTNOTE_RE = re.compile(r"\[\^[^\]]*\]")
_IMG_RE = re.compile(r"!\[([^\]]*)\]\(([^)\s]+)[^)]*\)")
_TBL_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(?:\|\s*:?-{2,}:?\s*)*\|?\s*$")
_TBL_TITLE_RE = re.compile(r"^\s*(?:表|Table)\s*[:：]\s*(.+?)\s*$")
_HEADING_RE = re.compile(r"^\s{0,3}(#{1,6})\s+(.*?)\s*#*\s*$")
_HEAD_NUM_RE = re.compile(
    r"^(?:第\s*[一二三四五六七八九十百\d]+\s*[章节部分篇]|[一二三四五六七八九十]+\s*[、.．]|\d+(?:\.\d+)*\s*[、.．]?|[(（]\s*[一二三四五六七八九十\d]+\s*[)）])\s*"
)
_SRC_INLINE_RE = re.compile(r"\{[^{}]*来源[^{}]*\}")
_SRC_LINE_RE = re.compile(r"^\s*(?:资料来源|数据来源|来源|Source)\s*[:：]")
_SRC_ANY_RE = re.compile(r"来源|Source|source")
_TIME_MARK_RE = re.compile(r"\[\s*\d{1,2}[:：]\d{2}\s*[–—\-~～]\s*\d{1,2}[:：]\d{2}\s*\]")
_UNIT = (r"亿人次|亿次|亿人|万人次|万人|万次|万份|万件|万单|万粉|万元|万|亿元|亿|元|％|%|个百分点|pp|份|人|月|倍|套|件|家|次|年|天|小时|公里|km|㎡|平方米|吨|kg|克|位|名|页|张|条|款|场"
         r"|章|节|期|轮|批|项|步|部分|层|级|段")
_NUM_RE = re.compile(r"(?<![A-Za-z0-9_.])([-−–]?\d[\d,，]*(?:\.\d+)?)\s*(" + _UNIT + r")?")
# 不是“数量”的数字：序数（第 2 章）、标识符 / 路径 / 时刻（03_财务、16:9、2026-03）
_NOT_QUANTITY_BEFORE = re.compile(r"(第\s*|\d\s*[:：]|[_/\\]|\d[\-–—~～])$")
_NOT_QUANTITY_AFTER = re.compile(r"^\s*[_/\\:：\-–][0-9A-Za-z\u4e00-\u9fff]")
_DIGIT_RE = re.compile(r"\d")
_MARKER_RE = re.compile(r"示例数据")
_SOURCE_HINT_RE = re.compile(r"来源|据《|\[\d+\]|引自")
_SENT_END = "。！？；!?;\n"
_MONEY_UNITS = {"亿元": 1e8, "亿": 1e8, "万元": 1e4, "万": 1e4, "元": 1.0}
_PCT_UNITS = {"%", "％", "个百分点", "pp"}
_CN_NUM = {1: "一", 2: "二", 3: "三", 4: "四", 5: "五", 6: "六", 7: "七", 8: "八", 9: "九", 10: "十"}


def _count_words(text: str) -> int:
    """中文按字、西文 / 数字串按词计（与 docx_render.count_words 口径一致）。"""
    plain = re.sub(r"[#*`>|\-\[\]()!{}]", " ", text)
    return len(_CJK_RE.findall(plain)) + len(_LATIN_RE.findall(plain))


def _clean_body(text: str) -> str:
    """正文清洗：去 HTML 注释、图片路径（保留图题）、文献号、来源行、表格分隔线、时间标记与 Markdown 符号。"""
    t = re.sub(r"<!--.*?-->", " ", text, flags=re.S)
    t = _IMG_RE.sub(lambda m: m.group(1), t)
    t = _CITE_RE.sub(" ", t)
    t = _FOOTNOTE_RE.sub(" ", t)
    t = _SRC_INLINE_RE.sub(" ", t)
    out: List[str] = []
    for line in t.splitlines():
        if _TBL_SEP_RE.match(line) or _SRC_LINE_RE.match(line):
            continue
        hm = _HEADING_RE.match(line)
        if hm:
            line = hm.group(2)
        out.append(_TIME_MARK_RE.sub(" ", line))
    return re.sub(r"[#*`>|_]", " ", "\n".join(out))


def _norm(s: str) -> str:
    """关键词匹配用的归一化：去空白、小写、全角冒号 / 括号转半角。"""
    return re.sub(r"\s+", "", str(s or "")).lower().replace("：", ":").replace("（", "(").replace("）", ")")


def _strip_head_num(s: str) -> str:
    return _HEAD_NUM_RE.sub("", str(s or "").strip()).strip()


def _protect(text: str, terms: Iterable[str]) -> str:
    """把白名单词替换为等长占位符（保留位置），避免 “大学生” 命中 “大学”。"""
    for t in sorted({str(x) for x in terms if x}, key=len, reverse=True):
        if t in text:
            text = text.replace(t, "□" * len(t))
    return text


def _sentence_at(text: str, start: int, end: int, limit: int = 160) -> str:
    s = start
    while s > 0 and text[s - 1] not in _SENT_END:
        s -= 1
    e = end
    while e < len(text) and text[e] not in _SENT_END:
        e += 1
    if e < len(text) and text[e] != "\n":
        e += 1
    sent = text[s:e].strip()
    return sent if len(sent) <= limit else sent[: limit - 1] + "…"


def _split_sentences(text: str) -> List[str]:
    return [s.strip() for s in re.split(r"(?<=[。！？；!?;])|\n", text) if s and s.strip()]


def _parse_num(s: str) -> Optional[float]:
    try:
        return float(s.replace(",", "").replace("，", "").replace("−", "-").replace("–", "-"))
    except (TypeError, ValueError):
        return None


def _fmt_num(v: float) -> str:
    if isinstance(v, bool):
        return str(v)
    if float(v).is_integer():
        return str(int(v))
    return f"{v:.4f}".rstrip("0").rstrip(".")


def _median(xs: Sequence[float]) -> float:
    return float(statistics.median(xs)) if xs else 0.0


def _ratio(a: float, b: float) -> float:
    return round(a / b, 4) if b else 0.0


def _status_min(value: float, pass_min: float, warn_min: float) -> str:
    return PASS if value >= pass_min else WARN if value >= warn_min else FAIL


def _status_max(value: float, pass_max: float, warn_max: float) -> str:
    return PASS if value <= pass_max else WARN if value <= warn_max else FAIL


def _worst(statuses: Iterable[str]) -> str:
    order = {FAIL: 3, WARN: 2, PASS: 1, SKIP: 0}
    best = PASS
    for s in statuses:
        if order.get(s, 0) > order.get(best, 0):
            best = s
    return best


def _mk(cid: str, value: Any, threshold: Any, status: str, detail: str, locations: Optional[List[Dict[str, Any]]] = None, **extra: Any) -> Dict[str, Any]:
    """统一的检查项结构（契约 §7.7）。locations 超过 MAX_LOCATIONS 条时截断并记录总数。"""
    d: Dict[str, Any] = {"id": cid, "value": value, "threshold": threshold, "status": status, "detail": detail}
    if locations is not None:
        d["locations"] = locations[:MAX_LOCATIONS]
        if len(locations) > MAX_LOCATIONS:
            d["locationsTotal"] = len(locations)
    d.update(extra)
    return d


def _skip(cid: str, detail: str, threshold: Any = None) -> Dict[str, Any]:
    return _mk(cid, None, threshold, SKIP, detail)


def _read_json_safe(path: Path) -> Any:
    try:
        return read_json(path) if path.is_file() else None
    except Exception:
        return None


def _read_text_safe(path: Path) -> Optional[str]:
    try:
        if path.is_file():
            with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
                return f.read()
    except Exception:
        return None
    return None


def _iter_strings(obj: Any, prefix: str = "") -> Iterable[Tuple[str, str]]:
    """递归产出 JSON 里的所有字符串 (路径, 值)。"""
    if isinstance(obj, str):
        yield prefix, obj
    elif isinstance(obj, dict):
        for k, v in obj.items():
            yield from _iter_strings(v, f"{prefix}.{k}" if prefix else str(k))
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            yield from _iter_strings(v, f"{prefix}[{i}]")


# ---------------------------------------------------------------------------------------------- Markdown 结构
@dataclass
class Heading:
    level: int
    text: str
    line: int


@dataclass
class Figure:
    file: str
    line: int
    alt: str
    path: str

    @property
    def fig_id(self) -> str:
        return Path(self.path.split("?")[0]).stem


@dataclass
class Table:
    file: str
    line: int
    header: List[str]
    rows: List[List[str]]
    title: str = ""
    heading: str = ""
    end_line: int = 0

    @property
    def nrows(self) -> int:
        return len([r for r in self.rows if any(c.strip() for c in r)])

    @property
    def context(self) -> str:
        return " ".join([self.title, self.heading] + self.header)

    def col(self, pattern: str) -> Optional[int]:
        rx = re.compile(pattern, re.I)
        for i, h in enumerate(self.header):
            if rx.search(h):
                return i
        return None

    def cols(self, pattern: str) -> int:
        rx = re.compile(pattern, re.I)
        return sum(1 for h in self.header if rx.search(h))

    def cell(self, row: Sequence[str], i: Optional[int]) -> str:
        return row[i].strip() if i is not None and i < len(row) else ""

    def nums(self, i: Optional[int]) -> List[float]:
        out: List[float] = []
        for r in self.rows:
            m = _NUM_RE.search(self.cell(r, i))
            v = _parse_num(m.group(1)) if m else None
            if v is not None:
                out.append(v)
        return out


@dataclass
class Doc:
    rel: str
    text: str
    lines: List[str] = field(default_factory=list)
    headings: List[Heading] = field(default_factory=list)
    tables: List[Table] = field(default_factory=list)
    figures: List[Figure] = field(default_factory=list)
    _body: Optional[str] = None

    @property
    def body(self) -> str:
        if self._body is None:
            self._body = _clean_body(self.text)
        return self._body

    @property
    def words(self) -> int:
        return _count_words(self.body)

    @property
    def table_rows(self) -> int:
        return sum(t.nrows for t in self.tables)

    @property
    def stem(self) -> str:
        return Path(self.rel).stem

    def is_table_line(self, idx: int) -> bool:
        return self.lines[idx].strip().startswith("|")


def _split_cells(line: str) -> List[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [re.sub(r"\*\*|__|`", "", c).strip() for c in s.split("|")]


def parse_doc(rel: str, text: str) -> Doc:
    """解析一章 Markdown：标题层级、表格（含上一行 `表：题注` 与所属标题）、图片行。"""
    lines = text.splitlines()
    doc = Doc(rel=rel, text=text, lines=lines)
    last_heading = ""
    i = 0
    while i < len(lines):
        line = lines[i]
        hm = _HEADING_RE.match(line)
        if hm:
            last_heading = hm.group(2).strip()
            doc.headings.append(Heading(len(hm.group(1)), last_heading, i + 1))
            i += 1
            continue
        for m in _IMG_RE.finditer(line):
            doc.figures.append(Figure(rel, i + 1, m.group(1).strip(), m.group(2).strip()))
        if line.strip().startswith("|") and i + 1 < len(lines) and _TBL_SEP_RE.match(lines[i + 1]):
            header = _split_cells(line)
            j = i + 2
            rows: List[List[str]] = []
            while j < len(lines) and lines[j].strip().startswith("|") and not _TBL_SEP_RE.match(lines[j]):
                rows.append(_split_cells(lines[j]))
                j += 1
            k = i - 1
            while k >= 0 and not lines[k].strip():
                k -= 1
            title = ""
            if k >= 0:
                tm = _TBL_TITLE_RE.match(lines[k])
                title = tm.group(1) if tm else ""
            doc.tables.append(Table(rel, i + 1, header, rows, title, last_heading, end_line=j))
            i = j
            continue
        i += 1
    return doc


# ---------------------------------------------------------------------------------------------- 工作区
class Workspace:
    """一次检查运行的全部输入（缺什么就记 None，检查项自行判定 skipped）。"""

    def __init__(self, root: Path, cfg: Dict[str, Any], chapter: Optional[str] = None) -> None:
        self.root = Path(root).resolve()
        self.cfg = cfg
        self.ev: Dict[str, Any] = dict(cfg.get("requiredEvidence") or {})
        self.chapter_mode = chapter is not None
        self.project: Dict[str, Any] = _read_json_safe(self.root / "project.json") or {}
        self.book: Dict[str, Any] = _read_json_safe(self.root / "05_正文" / "book.json") or {}
        self.cover: Dict[str, Any] = self.book.get("cover") if isinstance(self.book.get("cover"), dict) else {}
        self.appendix_rels: List[str] = [str(a.get("file") or a) for a in (self.book.get("appendices") or []) if a]
        self.chapters: List[Doc] = []
        if chapter is not None:
            p = Path(chapter)
            p = p if p.is_absolute() else (self.root / p)
            text = _read_text_safe(p)
            if text is not None:
                self.chapters = [parse_doc(self._rel(p), text)]
        else:
            for p in sorted((self.root / "05_正文").glob("*.md")) if (self.root / "05_正文").is_dir() else []:
                text = _read_text_safe(p)
                if text is not None:
                    self.chapters.append(parse_doc(self._rel(p), text))
        self.body_docs: List[Doc] = [d for d in self.chapters if not self._is_appendix(d)]
        self.appendix_docs: List[Doc] = [d for d in self.chapters if self._is_appendix(d)]
        self.pitch_docs: List[Doc] = []
        self.script: Optional[Doc] = None
        self.qa: Optional[Doc] = None
        self.slides: Optional[Dict[str, Any]] = None
        self.preview: Optional[List[Dict[str, Any]]] = None
        self.render: Optional[Dict[str, Any]] = None
        self.numbers: Optional[Dict[str, Any]] = None
        self.refs: List[Dict[str, Any]] = []
        self.chart_index: Dict[str, Any] = {}
        self.pdfs: List[Path] = []
        self.xlsx: Optional[Path] = None
        self.questionnaire_doc: Optional[Doc] = None
        if not self.chapter_mode:
            pitch_dir = self.root / "06_路演"
            for p in sorted(pitch_dir.glob("*.md")) if pitch_dir.is_dir() else []:
                text = _read_text_safe(p)
                if text is None:
                    continue
                d = parse_doc(self._rel(p), text)
                self.pitch_docs.append(d)
                if "讲稿" in p.stem:
                    self.script = d
                elif "题库" in p.stem or "答辩" in p.stem:
                    self.qa = d
            s = _read_json_safe(pitch_dir / "slides.json")
            self.slides = s if isinstance(s, dict) else None
            pv = _read_json_safe(self.root / "产出" / "路演.pptx.preview.json")
            self.preview = pv if isinstance(pv, list) else None
            rd = _read_json_safe(self.root / "产出" / "路演.render.json")
            self.render = rd if isinstance(rd, dict) else None
            n = _read_json_safe(self.root / "03_财务" / "numbers.json")
            self.numbers = n if isinstance(n, dict) else None
            self.refs = self._load_refs()
            ci = _read_json_safe(self.root / "04_图表" / "index.json")
            self.chart_index = ci if isinstance(ci, dict) else {}
            out_dir = self.root / "产出"
            self.pdfs = sorted(out_dir.glob("*.pdf")) if out_dir.is_dir() else []
            x = self.root / "03_财务" / "财务模型.xlsx"
            self.xlsx = x if x.is_file() else None
            q = self.root / "01_调研" / "调研工具包" / "问卷.md"
            qt = _read_text_safe(q)
            self.questionnaire_doc = parse_doc(self._rel(q), qt) if qt is not None else None
        else:
            ci = _read_json_safe(self.root / "04_图表" / "index.json")
            self.chart_index = ci if isinstance(ci, dict) else {}

    # ---- 基础访问
    def _rel(self, p: Path) -> str:
        try:
            return p.resolve().relative_to(self.root).as_posix()
        except ValueError:
            return p.name

    def _is_appendix(self, d: Doc) -> bool:
        if d.rel in self.appendix_rels:
            return True
        return bool(re.match(r"^[A-Za-z]\d*_", d.stem)) or "附录" in d.stem or "问卷" in d.stem

    def _load_refs(self) -> List[Dict[str, Any]]:
        data = _read_json_safe(self.root / "01_调研" / "参考文献.json")
        if isinstance(data, dict):
            for k in ("entries", "references", "items", "refs"):
                if isinstance(data.get(k), list):
                    data = data[k]
                    break
        return [x for x in data if isinstance(x, dict)] if isinstance(data, list) else []

    @property
    def entries(self) -> List[Dict[str, Any]]:
        if not self.numbers:
            return []
        ents = self.numbers.get("entries")
        return [e for e in ents if isinstance(e, dict) and e.get("key")] if isinstance(ents, list) else []

    def entry(self, key: str) -> Optional[Dict[str, Any]]:
        for e in self.entries:
            if e.get("key") == key:
                return e
        return None

    def entry_value(self, key: str) -> Optional[float]:
        e = self.entry(key)
        v = e.get("value") if e else None
        return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None

    @property
    def all_tables(self) -> List[Table]:
        return [t for d in self.chapters for t in d.tables]

    @property
    def all_figures(self) -> List[Figure]:
        return [f for d in self.chapters for f in d.figures]

    @property
    def chapter_text(self) -> str:
        return "\n".join(d.text for d in self.chapters)

    @property
    def body_text(self) -> str:
        return "\n".join(d.body for d in self.body_docs)

    @property
    def body_words(self) -> int:
        return sum(d.words for d in self.body_docs)

    @property
    def is_assigned_case(self) -> bool:
        subj = self.project.get("subject") if isinstance(self.project.get("subject"), dict) else {}
        return bool(subj.get("isAssignedCase"))

    def param(self, key: str, default: Any = None) -> Any:
        params = self.project.get("params") if isinstance(self.project.get("params"), dict) else {}
        return params.get(key, default)

    def self_names(self) -> List[str]:
        """本品牌 / 本项目的称呼（竞品矩阵列头排除用）。"""
        names: List[str] = ["本品牌", "本项目", "本公司", "本企业", "我们", "本方案", "我方", "自身"]
        subj = self.project.get("subject") if isinstance(self.project.get("subject"), dict) else {}
        for k in ("name", "company"):
            if subj.get(k):
                names.append(str(subj[k]))
        sub = str(self.cover.get("subtitle") or "")
        brand = re.sub(r"(品牌|数字营销|案例|会奖|商业|创业|营销)?(策划书|策划案|计划书|分析报告|方案|报告|策划)$", "", sub).strip()
        if len(brand) >= 2:
            names.append(brand)
        return names

    # ---- 扫描单元：(文件, 行号或页码, 文本)
    def line_units(self, docs: Sequence[Doc]) -> List[Tuple[str, int, str]]:
        out: List[Tuple[str, int, str]] = []
        for d in docs:
            for i, line in enumerate(d.lines):
                if line.strip() and not _TBL_SEP_RE.match(line):
                    out.append((d.rel, i + 1, line))
        return out

    def slide_pages(self) -> List[Dict[str, Any]]:
        """统一的页列表 [{index, layout, title, texts[], note, visual}]，优先 slides.json，其次 pptx 预览。"""
        pages: List[Dict[str, Any]] = []
        if self.slides and isinstance(self.slides.get("slides"), list):
            for i, s in enumerate(self.slides["slides"]):
                if not isinstance(s, dict):
                    continue
                texts = _slide_texts(s)
                visual = s.get("layout") in ("chart", "image", "table", "kpi") or any(s.get(k) for k in ("image", "chart", "columns"))
                pages.append({"index": i + 1, "layout": str(s.get("layout") or ""), "title": str(s.get("title") or ""), "texts": texts,
                              "note": str(s.get("note") or ""), "visual": bool(visual), "fields": _slide_fields(s)})
        elif self.preview:
            for i, p in enumerate(self.preview):
                if not isinstance(p, dict):
                    continue
                texts = [str(t) for t in (p.get("texts") or []) if t]
                visual = p.get("layout") in ("chart", "image", "table", "kpi") or any(p.get(k) for k in ("image", "chart", "table", "kpis"))
                pages.append({"index": int(p.get("index") or i + 1), "layout": str(p.get("layout") or ""), "title": str(p.get("title") or ""),
                              "texts": texts, "note": str(p.get("note") or ""), "visual": bool(visual), "fields": [("texts", t) for t in texts]})
        return pages

    def slide_units(self) -> List[Tuple[str, int, str]]:
        """PPT 的扫描单元：每页标题 + 各字段 + 备注，各作为独立单元（页码当行号）。"""
        file = "06_路演/slides.json" if self.slides else "产出/路演.pptx.preview.json"
        out: List[Tuple[str, int, str]] = []
        for p in self.slide_pages():
            if p["title"]:
                out.append((file, p["index"], p["title"]))
            for _, t in p["fields"]:
                out.append((file, p["index"], t))
            for s in _split_sentences(p["note"]):
                out.append((file, p["index"], s))
        return out

    def cover_units(self) -> List[Tuple[str, int, str]]:
        if not self.cover:
            return []
        parts = [f"{k}：{v}" for k, v in self.cover.items() if isinstance(v, str) and v.strip()]
        parts += [str(x) for x in (self.cover.get("extraLines") or []) if isinstance(x, str)]
        return [("05_正文/book.json", 0, "；".join(parts))]


def _slide_fields(s: Dict[str, Any]) -> List[Tuple[str, str]]:
    """一页 slides.json 的正文字段（不含 title / note / 表格单元格）。"""
    out: List[Tuple[str, str]] = []

    def add(name: str, v: Any) -> None:
        if isinstance(v, str) and v.strip():
            out.append((name, v.strip()))
        elif isinstance(v, list):
            for i, x in enumerate(v):
                if isinstance(x, str) and x.strip():
                    out.append((f"{name}[{i}]", x.strip()))

    for k in ("subtitle", "kicker", "bullets", "items", "caption", "insight", "quote", "by", "contact", "highlight", "highlightLabel"):
        add(k, s.get(k))
    for side in ("left", "right"):
        col = s.get(side)
        if isinstance(col, dict):
            add(f"{side}.heading", col.get("heading"))
            add(f"{side}.bullets", col.get("bullets"))
    for k, a, b in (("kpis", "value", "label"), ("steps", "when", "what"), ("members", "role", "duty")):
        for i, item in enumerate(s.get(k) or []):
            if isinstance(item, dict):
                out.append((f"{k}[{i}]", f"{item.get(a) or ''} {item.get(b) or ''}".strip()))
    return out


def _slide_texts(s: Dict[str, Any]) -> List[str]:
    return [t for _, t in _slide_fields(s)]


# ---------------------------------------------------------------------------------------------- 通用扫描
def _scan(units: Sequence[Tuple[str, int, str]], patterns: Sequence[Tuple[str, "re.Pattern[str]", bool]], whitelist: Sequence[str]) -> List[Dict[str, Any]]:
    """在扫描单元里找词：白名单先遮蔽；needSource 的词同句有来源短语则放过。返回 {file, line, term, text}。"""
    hits: List[Dict[str, Any]] = []
    for file, line, text in units:
        prot = _protect(text, whitelist)
        spans: List[Tuple[int, int]] = []
        for label, rx, need_source in patterns:
            for m in rx.finditer(prot):
                if any(m.start() < e and m.end() > s for s, e in spans):  # 同一处被正则与字面词重复命中只记一次
                    continue
                sent = _sentence_at(text, m.start(), m.end())
                if need_source and _SOURCE_HINT_RE.search(sent):
                    continue
                spans.append((m.start(), m.end()))
                hits.append({"file": file, "line": line, "term": text[m.start():m.end()] or label, "text": sent})
    return hits


def _text_units(ws: Workspace, include_pitch: bool = True, include_slides: bool = True) -> List[Tuple[str, int, str]]:
    units = ws.line_units(ws.chapters)
    if not ws.chapter_mode and include_pitch:
        units += ws.line_units(ws.pitch_docs)
    if not ws.chapter_mode and include_slides:
        units += ws.slide_units()
    return units


def _find_sentences(ws: Workspace, pattern: str, docs: Optional[Sequence[Doc]] = None, need_number: bool = False) -> List[Dict[str, Any]]:
    """在章节（默认全部章节）里找匹配句子；need_number 时要求同句含数字。"""
    rx = re.compile(pattern, re.I)
    out: List[Dict[str, Any]] = []
    for file, line, text in ws.line_units(docs if docs is not None else ws.chapters):
        for m in rx.finditer(text):
            sent = _sentence_at(text, m.start(), m.end())
            if need_number and not _DIGIT_RE.search(sent):
                continue
            out.append({"file": file, "line": line, "term": m.group(0), "text": sent})
            break
    return out


def _has(ws: Workspace, pattern: str) -> bool:
    return re.search(pattern, ws.chapter_text, re.I) is not None


def _table_loc(t: Table, note: str = "") -> Dict[str, Any]:
    return {"file": t.file, "line": t.line, "text": (t.title or t.heading or " | ".join(t.header))[:120], "note": note}


# ---------------------------------------------------------------------------------------------- 核心检查项
def check_struct_coverage(ws: Workspace) -> Dict[str, Any]:
    """必备要素命中率：一级标题 / 文件名 与 requiredElements 的名称、别名互相包含即命中；「目录」恒命中。"""
    thr = {"pass": 1.0, "warn": 0.85}
    elements = [e for e in (ws.cfg.get("requiredElements") or []) if isinstance(e, dict)]
    if not ws.chapters:
        return _skip("struct_coverage", "05_正文/ 下没有章节 Markdown", thr)
    heads = [_strip_head_num(h.text) for d in ws.chapters for h in d.headings if h.level == 1]
    stems = [re.sub(r"^[A-Za-z]?\d+_", "", d.stem) for d in ws.chapters]
    cands = [_norm(x) for x in heads + stems if x]

    def match(name: str) -> bool:
        n = _norm(name)
        return bool(n) and any(n in c or (len(c) >= 2 and c in n) for c in cands)

    hit: List[str] = []
    miss: List[str] = []
    for el in elements:
        names = [str(el.get("name") or "")] + [str(a) for a in (el.get("aliases") or [])]
        eid = str(el.get("id") or names[0])
        ok = eid == "toc" or "目录" in names or any(match(n) for n in names if n)
        if not ok and eid in ("appendix",) and (ws.appendix_rels or ws.refs):
            ok = True
        (hit if ok else miss).append(names[0] or eid)
    total = len(elements)
    value = round(len(hit) / total, 3) if total else 1.0
    detail = f"{len(hit)}/{total} 要素命中" + (f"；缺：{'、'.join(miss)}" if miss else "")
    return _mk("struct_coverage", value, thr, _status_min(value, 1.0, 0.85), detail, missing=miss, headings=heads)


def _pick_book_pdf(ws: Workspace) -> Optional[Path]:
    out = str(ws.book.get("output") or "")
    if out:
        cand = ws.root / Path(out).with_suffix(".pdf")
        if cand.is_file():
            return cand
    cands = [p for p in ws.pdfs if not re.search(r"路演|deck|ppt|slides", p.stem, re.I)]
    return max(cands, key=lambda p: p.stat().st_mtime) if cands else None


def _pdf_pages(path: Path) -> Optional[int]:
    try:
        from pypdf import PdfReader  # 延迟导入：缺依赖时退回估算

        return len(PdfReader(str(path), strict=False).pages)
    except Exception:
        return None


def _estimate_pages(docs: Sequence[Doc]) -> float:
    words = sum(d.words for d in docs)
    figs = sum(len(d.figures) for d in docs)
    rows = sum(d.table_rows for d in docs)
    return round(words / 650 + figs * 0.45 + rows / 40, 1)


def check_page_count(ws: Workspace) -> Dict[str, Any]:
    """页数：有 产出/*.pdf 读真实页数，否则 字数/650 + 图×0.45 + 表行/40 估算；超 pageLimit 为 fail，偏离 pageTarget±tol 为 warn。"""
    limit = int(ws.cfg.get("pageLimit") or 80)
    target = ws.param("pageTarget")
    tol = float(ws.cfg.get("pageTargetTolerance") or 0.1)
    thr: Dict[str, Any] = {"limit": limit, "target": target, "tolerance": tol}
    pdf = _pick_book_pdf(ws)
    pages: Optional[float] = None
    source = ""
    if pdf is not None:
        n = _pdf_pages(pdf)
        if n is not None:
            pages, source = float(n), ws._rel(pdf)
    if pages is None:
        if not ws.chapters:
            return _skip("page_count", "无 产出/*.pdf 且无章节 Markdown，无法统计", thr)
        pages, source = _estimate_pages(ws.chapters), "估算（字数/650 + 图×0.45 + 表行/40）"
    status = PASS
    notes = [f"{_fmt_num(pages)} 页（{source}）"]
    if pages > limit:
        status = FAIL
        notes.append(f"超过上限 {limit}")
    elif isinstance(target, (int, float)) and target > 0 and not (target * (1 - tol) <= pages <= target * (1 + tol)):
        status = WARN
        notes.append(f"偏离目标 {target}±{int(tol * 100)}%")
    return _mk("page_count", pages, thr, status, "；".join(notes), pages=pages, source=source)


def check_word_count(ws: Workspace) -> Dict[str, Any]:
    """正文字数（不含附录）；wordMin=0 时只报数不判定。"""
    wmin = int(ws.cfg.get("wordMin") or 0)
    thr = {"min": wmin}
    if not ws.chapters:
        return _skip("word_count", "无章节 Markdown", thr)
    words = ws.body_words
    if wmin <= 0:
        return _mk("word_count", words, thr, PASS, f"{words} 字（赛道未设下限）")
    status = _status_min(words, wmin, wmin * 0.9)
    return _mk("word_count", words, thr, status, f"{words} 字，下限 {wmin}")


def check_anonymity(ws: Workspace) -> Dict[str, Any]:
    """匿名：赛道正则 + book/slides 的 anonymize.forbidden 字面词；白名单（赛道 + book + slides）先遮蔽。"""
    thr = {"pass": 0}
    units = ws.line_units(ws.chapters)
    if not ws.chapter_mode:
        units += ws.line_units(ws.pitch_docs) + ws.slide_units() + ws.cover_units()
    if not units:
        return _skip("anonymity_hits", "没有可扫描的文本（正文 / 路演 / 封面均缺）", thr)
    try:
        rx = re.compile(str(ws.cfg.get("anonymityRegex") or td.ANONYMITY_REGEX_DEFAULT))
    except re.error:
        rx = re.compile(td.ANONYMITY_REGEX_DEFAULT)
    whitelist: List[str] = list(ws.cfg.get("anonymityWhitelist") or [])
    literals: List[str] = []
    for src in (ws.book.get("anonymize"), (ws.slides or {}).get("anonymize")):
        if isinstance(src, dict):
            whitelist += [str(x) for x in (src.get("whitelist") or [])]
            literals += [str(x) for x in (src.get("forbidden") or []) if x]
    patterns: List[Tuple[str, "re.Pattern[str]", bool]] = [("anonymity", rx, False)]
    patterns += [(t, re.compile(re.escape(t)), False) for t in dict.fromkeys(literals)]
    hits = _scan(units, patterns, whitelist)
    detail = "无匿名违规" if not hits else f"{len(hits)} 处疑似院校 / 人员信息：" + "；".join(f"{h['file']}:{h['line']}「{h['term']}」" for h in hits[:5])
    return _mk("anonymity_hits", len(hits), thr, PASS if not hits else FAIL, detail, hits)


def _density(text: str) -> Tuple[float, int, int]:
    words = _count_words(text)
    tokens = len(_NUM_RE.findall(text))
    return (round(tokens / words * 1000, 1) if words else 0.0), tokens, words


def _chapter_threshold(stem: str, cfg: Dict[str, Any]) -> Tuple[float, bool]:
    base = float(cfg.get("numericDensity") or 15)
    s = _norm(re.sub(r"^[A-Za-z]?\d+_", "", stem))
    for key, v in (cfg.get("numericDensityChapters") or {}).items():
        k = _norm(re.sub(r"^[A-Za-z]?\d+_", "", str(key)))
        if _norm(key) == _norm(stem) or (k and (k in s or s in k)):
            try:
                return float(v), True
            except (TypeError, ValueError):
                break
    return base, False


def check_numeric_density(ws: Workspace) -> Dict[str, Any]:
    """每千字数值 token（含 %、万、亿、元、倍）：全文 + 按章；章阈值取 numericDensityChapters，缺省 numericDensity。"""
    base = float(ws.cfg.get("numericDensity") or 15)
    thr: Dict[str, Any] = {"pass": base, "warn": round(base / 2, 1), "chapters": ws.cfg.get("numericDensityChapters") or {}}
    docs = ws.body_docs or ws.chapters
    if not docs:
        return _skip("numeric_density", "无章节 Markdown", thr)
    chapters: List[Dict[str, Any]] = []
    for d in docs:
        dens, tokens, words = _density(d.body)
        t, specific = _chapter_threshold(d.stem, ws.cfg)
        st = _status_min(dens, t, t / 2) if words >= 50 else PASS
        chapters.append({"file": d.rel, "density": dens, "tokens": tokens, "words": words, "threshold": t, "specific": specific, "status": st})
    dens, tokens, words = _density("\n".join(d.body for d in docs))
    overall = _status_min(dens, base, base / 2)
    low = [c for c in chapters if c["status"] != PASS]
    status = overall if overall != PASS else (WARN if low else PASS)
    detail = f"全文 {dens}/千字（{tokens} 个数值 / {words} 字，阈值 {base}）"
    if low:
        detail += "；偏低章节：" + "、".join(f"{Path(c['file']).stem} {c['density']}/{_fmt_num(c['threshold'])}" for c in low)
    return _mk("numeric_density", dens, thr, status, detail, chapters=chapters)


def _ref_year(r: Dict[str, Any]) -> Optional[int]:
    for k in ("year", "date", "accessed"):
        m = re.search(r"(19|20)\d{2}", str(r.get(k) or ""))
        if m:
            return int(m.group(0))
    return None


def _ref_domain(r: Dict[str, Any]) -> str:
    url = str(r.get("url") or "")
    m = re.match(r"^\s*(?:https?://)?([^/\s:?#]+)", url)
    if m:
        host = m.group(1).lower()
        host = host[4:] if host.startswith("www.") else host
        parts = host.split(".")
        if len(parts) >= 3 and parts[-2] in ("com", "org", "gov", "edu", "net", "ac", "co"):
            return ".".join(parts[-3:])
        return ".".join(parts[-2:]) if len(parts) >= 2 else host
    for k in ("publisher", "journal", "newspaper", "author"):
        if r.get(k):
            return f"机构:{_norm(str(r[k]))}"
    return ""


def _ref_is_official(r: Dict[str, Any]) -> bool:
    if str(r.get("type") or "").lower() in td.OFFICIAL_TYPES:
        return True
    blob = " ".join(str(r.get(k) or "") for k in ("url", "publisher", "author", "title", "note")).lower()
    return any(h.lower() in blob for h in td.OFFICIAL_DOMAIN_HINTS)


def check_sources(ws: Workspace) -> List[Dict[str, Any]]:
    """参考文献四项：条数、独立域名（无 URL 的书 / 论文按出版机构计）、官方占比、近 N 年占比。"""
    smin = int(ws.cfg.get("sourceMin") or 15)
    dmin = int(ws.cfg.get("sourceDomainsMin") or 8)
    omin = float(ws.cfg.get("sourceOfficialRatio") or 0.4)
    fmin = float(ws.cfg.get("sourceFreshRatio") or 0.7)
    years = int(ws.cfg.get("sourceFreshYears") or 3)
    thr_c = {"pass": smin, "warn": round(smin * 2 / 3)}
    thr_d = {"pass": dmin, "warn": round(dmin * 2 / 3)}
    thr_o = {"pass": omin, "warn": round(omin * 0.75, 2)}
    thr_f = {"pass": fmin, "warn": round(fmin * 0.75, 2), "years": years}
    refs = ws.refs
    if not refs:
        why = "01_调研/参考文献.json 缺失或为空"
        return [_skip("source_count", why, thr_c), _skip("source_domains", why, thr_d), _skip("source_official_ratio", why, thr_o), _skip("source_freshness", why, thr_f)]
    n = len(refs)
    domains = sorted({d for d in (_ref_domain(r) for r in refs) if d})
    official = [r for r in refs if _ref_is_official(r)]
    cur_year = datetime.now().year
    dated = [(r, _ref_year(r)) for r in refs]
    known = [y for _, y in dated if y]
    fresh = [y for y in known if y >= cur_year - years]
    o_ratio = _ratio(len(official), n)
    f_ratio = _ratio(len(fresh), len(known)) if known else 0.0
    out = [
        _mk("source_count", n, thr_c, _status_min(n, smin, smin * 2 / 3), f"{n} 条参考文献（下限 {smin}）"),
        _mk("source_domains", len(domains), thr_d, _status_min(len(domains), dmin, dmin * 2 / 3), f"{len(domains)} 个独立来源域（下限 {dmin}）", domains=domains),
        _mk("source_official_ratio", o_ratio, thr_o, _status_min(o_ratio, omin, omin * 0.75), f"官方 / 年报 / 协会来源 {len(official)}/{n} = {o_ratio:.0%}（下限 {omin:.0%}）"),
        _mk("source_freshness", f_ratio, thr_f, _status_min(f_ratio, fmin, fmin * 0.75) if known else WARN,
            f"近 {years} 年来源 {len(fresh)}/{len(known)} = {f_ratio:.0%}（下限 {fmin:.0%}；{n - len(known)} 条无年份）"),
    ]
    return out


def check_source_pairing(ws: Workspace) -> Dict[str, Any]:
    """图表来源配对：图片行 / 表格之后 3 行内有「来源」，或图表 id 在 04_图表/index.json 里有 source。"""
    thr = {"pass": 0.95, "warn": 0.8}
    total = 0
    unpaired: List[Dict[str, Any]] = []
    for d in ws.chapters:
        lines = d.lines
        for f in d.figures:
            total += 1
            window = " ".join(lines[f.line - 1: f.line + 3])
            idx_src = ws.chart_index.get(f.fig_id) if isinstance(ws.chart_index.get(f.fig_id), dict) else None
            ok = bool(_SRC_ANY_RE.search(window)) or bool(idx_src and idx_src.get("source"))
            if not ok:
                unpaired.append({"file": d.rel, "line": f.line, "text": f"图：{f.alt or f.path}"})
        for t in d.tables:
            total += 1
            start = t.line - 2 if t.title else t.line - 1
            window = " ".join(lines[max(start, 0): t.end_line + 3])
            if not _SRC_ANY_RE.search(window):
                unpaired.append({"file": d.rel, "line": t.line, "text": f"表：{t.title or t.heading or ' | '.join(t.header)}"[:120]})
    if total == 0:
        return _skip("source_pairing", "正文中没有图片 / 表格", thr)
    value = _ratio(total - len(unpaired), total)
    detail = f"{total - len(unpaired)}/{total} 图表带来源（{value:.0%}）"
    return _mk("source_pairing", value, thr, _status_min(value, 0.95, 0.8), detail, unpaired)


def check_chart_count(ws: Workspace) -> Dict[str, Any]:
    """图表数 = 正文图片数 + 表格数。"""
    cmin = int(ws.cfg.get("chartMin") or 15)
    thr = {"pass": cmin, "warn": round(cmin * 2 / 3)}
    if not ws.chapters:
        return _skip("chart_count", "无章节 Markdown", thr)
    figs, tabs = len(ws.all_figures), len(ws.all_tables)
    value = figs + tabs
    detail = f"图 {figs} + 表 {tabs} = {value}（下限 {cmin}；04_图表/index.json 有 {len(ws.chart_index)} 张）"
    return _mk("chart_count", value, thr, _status_min(value, cmin, cmin * 2 / 3), detail, figures=figs, tables=tabs)


def _forbidden_patterns(cfg: Dict[str, Any]) -> List[Tuple[str, "re.Pattern[str]", bool]]:
    pats: List[Tuple[str, "re.Pattern[str]", bool]] = []
    for t in cfg.get("forbiddenTerms") or []:
        if isinstance(t, dict):
            term, need = str(t.get("term") or ""), bool(t.get("needSource"))
        else:
            term = str(t)
            need = term in td.FORBIDDEN_NEED_SOURCE
        if term:
            pats.append((term, re.compile(re.escape(term)), need))
    return pats


def check_forbidden_terms(ws: Workspace) -> Dict[str, Any]:
    """禁语：纯词命中即计；needSource 词（据统计 / 填补空白 / 预计将）只在同句无「来源 / 据《 / [n]」时命中。"""
    thr = {"pass": 0, "warn": 2}
    units = _text_units(ws)
    if not units:
        return _skip("forbidden_terms", "没有可扫描的文本", thr)
    hits = _scan(units, _forbidden_patterns(ws.cfg), list(ws.cfg.get("advertisingLawWhitelist") or []))
    detail = "无禁语" if not hits else f"{len(hits)} 处：" + "；".join(f"{h['file']}:{h['line']}「{h['term']}」" for h in hits[:6])
    return _mk("forbidden_terms", len(hits), thr, _status_max(len(hits), 0, 2), detail, hits)


_ADLAW_100_RE = r"100\s*[%％](?=\s*(?:有效|安全|纯|天然|保证|满意|成功|无|正品|真|好评|治愈|有机|放心))"
# 「第一」：排除序数后缀（第一年 / 第一阶段…）与列举用法（“第一，…；第二，…”“第一是”）
_ADLAW_FIRST_RE = r"第一(?!" + td.ADLAW_ORDINAL_SUFFIX + r"|\s*[，,、：:；;是]|\s*$)"
_ADLAW_EXTRA_WHITELIST = ["绝对值", "绝对数", "绝对额", "极致化"]


def _adlaw_patterns(cfg: Dict[str, Any]) -> List[Tuple[str, "re.Pattern[str]", bool]]:
    pats: List[Tuple[str, "re.Pattern[str]", bool]] = []
    for t in cfg.get("advertisingLawTerms") or []:
        t = str(t).strip()
        if not t:
            continue
        if t == "100%":
            pats.append((t, re.compile(_ADLAW_100_RE), False))
        elif t == "第一":
            pats.append((t, re.compile(_ADLAW_FIRST_RE), False))
        else:
            pats.append((t, re.compile(td.adlaw_patterns([t])[0]), False))
    return pats


def check_adlaw(ws: Workspace) -> Dict[str, Any]:
    """广告法绝对化用语（硬门槛）：advertisingLawTerms − advertisingLawWhitelist；「第一」排除序数，「100%」只认宣传搭配。"""
    thr = {"pass": 0}
    units = _text_units(ws)
    if not units:
        return _skip("forbidden_terms_adlaw", "没有可扫描的文本", thr)
    whitelist = list(ws.cfg.get("advertisingLawWhitelist") or []) + _ADLAW_EXTRA_WHITELIST
    hits = _scan(units, _adlaw_patterns(ws.cfg), whitelist)
    detail = "无广告法绝对化用语" if not hits else f"{len(hits)} 处：" + "；".join(f"{h['file']}:{h['line']}「{h['term']}」" for h in hits[:6])
    return _mk("forbidden_terms_adlaw", len(hits), thr, PASS if not hits else FAIL, detail, hits)


def check_weak_terms(ws: Workspace) -> Dict[str, Any]:
    """弱评价词密度（每百字），专业搭配（有效问卷 / 显著性…）不计。"""
    per100 = float(ws.cfg.get("weakTermsPer100") or 0.1)
    thr = {"pass": per100, "warn": round(per100 * 2, 3)}
    docs = ws.body_docs or ws.chapters
    if not docs:
        return _skip("weak_terms", "无章节 Markdown", thr)
    pats = [(t, re.compile(re.escape(str(t))), False) for t in (ws.cfg.get("weakTerms") or []) if t]
    hits = _scan(ws.line_units(docs), pats, td.WEAK_TERM_EXCEPTIONS)
    words = sum(d.words for d in docs)
    value = round(len(hits) / words * 100, 3) if words else 0.0
    detail = f"{len(hits)} 处 / {words} 字 = {value}/百字（上限 {per100}）"
    return _mk("weak_terms", value, thr, _status_max(value, per100, per100 * 2), detail, hits, hits=len(hits))


# ---------------------------------------------------------------------------------------------- 赛道专项（requiredEvidence 开关）
_SAMPLE_PATTERNS = [
    r"N\s*[=＝]\s*(\d{2,6})",
    r"样本(?:量|数|总量|规模)?[^\d\n]{0,8}?(\d{2,6})\s*(?:份|人|个|名|位)?",
    r"(?:回收|有效|发放)[^\d\n]{0,6}?(\d{2,6})\s*份",
    r"(\d{2,6})\s*份(?:有效)?(?:问卷|样本|答卷)",
    r"有效问卷[^\d\n]{0,6}?(\d{2,6})",
]


def check_primary_research(ws: Workspace) -> Dict[str, Any]:
    """一手调研：问卷 / 访谈 / 实地关键词 + 样本量数值 + 方法词（抽样 / 有效率 / 信度…）。"""
    smin = int(ws.cfg.get("sampleSizeMin") or 384)
    swarn = int(ws.cfg.get("sampleSizeWarn") or 200)
    thr = {"sampleMin": smin, "sampleWarn": swarn, "needs": ["关键词", "样本量", "方法词"]}
    if not ws.chapters:
        return _skip("has_primary_research", "无章节 Markdown", thr)
    text = ws.chapter_text
    kw = re.search(r"问卷|访谈|实地|焦点小组|深访|田野|观察法|神秘顾客", text)
    samples: List[int] = []
    valid: List[int] = []  # 明确写成“有效”的样本数优先于回收 / 发放数
    for p in _SAMPLE_PATTERNS:
        for m in re.finditer(p, text):
            v = int(m.group(1))
            if 20 <= v <= 100000:
                samples.append(v)
                if "有效" in text[max(0, m.start() - 6): m.end()]:
                    valid.append(v)
    sample = max(valid) if valid else max(samples) if samples else 0
    method = re.search(r"抽样|配额|随机|分层|滚雪球|问卷星|定量|定性|信度|效度|有效率|Cronbach|α|KMO|交叉分析|样本结构|置信", text)
    dated = re.search(r"20\d\d\s*年\s*\d{1,2}\s*月", text)
    if not kw:
        status, why = FAIL, "未见问卷 / 访谈 / 实地等一手调研关键词"
    elif sample >= smin and method:
        status, why = PASS, f"样本 {sample} ≥ {smin}，含方法词" + ("" if dated else "（缺调研时间）")
    elif sample >= smin:
        status, why = WARN, f"样本 {sample} 达标但缺方法 / 质控说明"
    elif sample >= swarn:
        status, why = WARN, f"样本 {sample} 介于 {swarn}–{smin}"
    else:
        status, why = FAIL, ("未提取到样本量" if not sample else f"样本 {sample} < {swarn}")
    return _mk("has_primary_research", sample, thr, status, why, method=bool(method), dated=bool(dated))


def check_tam_sam_som(ws: Workspace) -> Dict[str, Any]:
    """TAM / SAM / SOM 三词各伴数值；双法（自上而下 + 自下而上）为 pass。"""
    thr = {"terms": 3, "methods": 2}
    if not ws.chapters:
        return _skip("has_tam_sam_som", "无章节 Markdown", thr)
    words = {"TAM": r"\bTAM\b|潜在市场(?:规模|总量)?", "SAM": r"\bSAM\b|可服务市场|可触达市场", "SOM": r"\bSOM\b|可获得市场|可获取市场|可得市场"}
    found = {k: bool(_find_sentences(ws, p, need_number=True)) for k, p in words.items()}
    text = ws.chapter_text
    methods = sum(1 for p in (r"自上而下|top[- ]?down", r"自下而上|bottom[- ]?up") if re.search(p, text, re.I))
    n = sum(found.values())
    if n == 3 and methods >= 2:
        status = PASS
    elif n == 3 or (n >= 2 and methods >= 1):
        status = WARN
    else:
        status = FAIL
    detail = f"命中 {n}/3（{'、'.join(k for k, v in found.items() if v) or '无'}），测算方法词 {methods}/2"
    return _mk("has_tam_sam_som", n, thr, status, detail, methods=methods)


_GENERIC_HEADER_RE = re.compile(r"^(维度|指标|项目|说明|备注|权重|得分|评分|小计|合计|单位|来源|对比项|比较项|参数|要素|类别|分类|序号|项|/|—|-)?$")
# 属性类表头（基线 / 目标 / 金额 / 年份…）不算“竞品实体”
_ATTR_HEADER_RE = re.compile(
    r"基线|目标|工具|周期|频率|现状|负责人|时间|预算|金额|数量|单价|占比|比例|措施|对策|概率|影响|等级|阶段|渠道|内容|KPI|KGI|年份|20\d\d|Q[1-4]|"
    r"\d+\s*月|Y\d|第\s*[一二三四五\d]\s*年|层级|数据来源|口径|状态|结果|结论|差距|评价|优势|劣势|机会|威胁|问题|原因|建议|行动|方案|进度|里程碑|"
    r"职责|分工|背景|成果|姓名|角色|岗位|价格|成本|费用|收入|利润|销量|人数|场地|供应商|报价|日期|天|时段|画面|台词|镜头|字幕|时长|总计|小计|变化|增长",
    re.I,
)


def _entity_cols(t: Table, self_names: Sequence[str]) -> int:
    n = 0
    for h in t.header[1:]:
        h2 = h.strip()
        if not h2 or _GENERIC_HEADER_RE.match(h2) or _ATTR_HEADER_RE.search(h2):
            continue
        if any((s in h2) or (len(h2) >= 2 and h2 in s) for s in self_names if s):
            continue
        n += 1
    return n


def check_competitor_matrix(ws: Workspace) -> Dict[str, Any]:
    """竞品矩阵：列头 ≥ 3 个非本品实体且行 ≥ 5；或题注 / 所属标题含「竞品 / 竞争对手」且行 ≥ 5。"""
    thr = {"entities": 3, "rows": 5}
    if not ws.chapters:
        return _skip("has_competitor_matrix", "无章节 Markdown", thr)
    names = ws.self_names()
    best: Optional[Table] = None
    partial: Optional[Table] = None
    for t in ws.all_tables:
        ents = _entity_cols(t, names)
        titled = re.search(r"竞品|竞争对手|竞争者|对标|竞争", t.title + " " + t.heading)
        # 实体列路径：首列是“维度 / 指标 / 对比项”之类的属性名，或上下文提到竞争 / 对比，避免把 KPI 表当竞品表
        dim_first = bool(t.header) and bool(re.match(r"^(维度|指标|对比项|比较项|比较维度|对比维度|项目|要素|评价维度)$", t.header[0].strip()))
        entity_ok = ents >= 3 and (dim_first or re.search(r"竞|对比|对标|vs", t.context, re.I))
        if t.nrows >= 5 and (entity_ok or titled):
            best = t
            break
        if titled or entity_ok:
            partial = partial or t
    if best is not None:
        return _mk("has_competitor_matrix", best.nrows, thr, PASS, f"{best.file}:{best.line} 「{best.title or best.heading}」{_entity_cols(best, names)} 列实体 × {best.nrows} 行", [_table_loc(best)])
    if partial is not None:
        return _mk("has_competitor_matrix", partial.nrows, thr, WARN, f"疑似竞品表仅 {partial.nrows} 行（需 ≥ 5）或实体列不足", [_table_loc(partial)])
    return _mk("has_competitor_matrix", 0, thr, FAIL, "未发现竞品对比表（需 ≥ 3 家竞品 × ≥ 5 维）")


def check_unit_economics(ws: Workspace) -> Dict[str, Any]:
    """单位经济四指标（CAC / LTV / 毛利率 / 回收期）各伴数值。"""
    thr = {"pass": 4, "warn": 2}
    if not ws.chapters:
        return _skip("has_unit_economics", "无章节 Markdown", thr)
    terms = {"CAC": r"\bCAC\b|获客成本", "LTV": r"\bLTV\b|生命周期价值|终身价值", "毛利率": r"毛利率", "回收期": r"回收期|回本周期|回本期|回收周期"}
    found = {k: bool(_find_sentences(ws, p, need_number=True)) for k, p in terms.items()}
    n = sum(found.values())
    detail = f"{n}/4 指标伴数值（{'、'.join(k for k, v in found.items() if v) or '无'}）"
    return _mk("has_unit_economics", n, thr, _status_min(n, 4, 2), detail)


def check_sensitivity(ws: Workspace) -> Dict[str, Any]:
    """敏感性 / 情景：关键词 + 表格或图（正文），财务模型「敏感性」页也算证据。"""
    thr = {"needs": ["敏感性(单因子)", "情景(三情景)"]}
    if not ws.chapters and ws.xlsx is None:
        return _skip("has_sensitivity", "无章节 Markdown 与财务模型", thr)
    text = ws.chapter_text
    evid = [t for t in ws.all_tables if re.search(r"敏感|情景|悲观|乐观|±", t.context)]
    evid_f = [f for f in ws.all_figures if re.search(r"敏感|情景|龙卷风|tornado|悲观|乐观", f.alt + " " + f.path, re.I)]
    sheet = False
    if ws.xlsx is not None:
        try:
            import openpyxl  # noqa: WPS433

            sheet = "敏感性" in openpyxl.load_workbook(ws.xlsx, read_only=True).sheetnames
        except Exception:
            sheet = False
    has_sens = (bool(re.search(r"敏感性|单因素|单因子", text)) and bool(evid or evid_f or sheet)) or sheet
    has_scen = bool(re.search(r"情景|悲观|乐观|基准情形|中性", text)) and bool(evid or evid_f or sheet)
    n = int(has_sens) + int(has_scen)
    status = PASS if n == 2 else WARN if n == 1 else FAIL
    detail = f"敏感性 {'✓' if has_sens else '✗'}，情景 {'✓' if has_scen else '✗'}（表 {len(evid)} / 图 {len(evid_f)}{' / 模型敏感性页' if sheet else ''}）"
    return _mk("has_sensitivity", n, thr, status, detail, [_table_loc(t) for t in evid[:3]])


def check_three_statement(ws: Workspace) -> Dict[str, Any]:
    """三表勾稽：numbers.json 的 checks{} / check_* 条目，或 财务模型.xlsx「检查」页 Python 侧勾稽行（差异 < 1 元）。"""
    thr = {"maxDiffYuan": 1.0}
    diffs: Dict[str, float] = {}
    verdicts: List[str] = []
    src = ""
    if ws.numbers:
        chk = ws.numbers.get("checks")
        if isinstance(chk, dict):
            diffs = {k: float(v) for k, v in chk.items() if isinstance(v, (int, float))}
            src = "numbers.json.checks"
        for e in ws.entries:
            if str(e.get("key") or "").startswith("check_") and isinstance(e.get("value"), (int, float)):
                diffs[str(e["key"])] = float(e["value"])
                src = src or "numbers.json.entries"
    if not diffs and ws.xlsx is not None:
        try:
            import openpyxl

            wb = openpyxl.load_workbook(ws.xlsx, read_only=True, data_only=True)
            if "检查" in wb.sheetnames:
                sh = wb["检查"]
                for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row or 60, 200), max_col=3, values_only=True):
                    label = str(row[0] or "")
                    if re.search(r"资产\s*=\s*负债|现金流量表期末现金|留存收益\s*=", label):
                        v = row[1]
                        if isinstance(v, (int, float)):
                            diffs[label] = float(v)
                        elif isinstance(v, str) and _parse_num(v) is not None:
                            diffs[label] = float(_parse_num(v) or 0)
                        if isinstance(row[2], str) and row[2].strip():
                            verdicts.append(row[2].strip())
                    if label.startswith("Master Check 状态") and isinstance(row[1], str) and row[1].strip() in ("通过", "未通过"):
                        verdicts.append(row[1].strip())
                src = "财务模型.xlsx!检查"
        except Exception as e:  # 读不了模型 → skipped，不中断
            return _skip("three_statement_check", f"无法读取 财务模型.xlsx：{e}", thr)
    if not diffs and not verdicts:
        return _skip("three_statement_check", "numbers.json 无 checks 且无 财务模型.xlsx 检查页", thr)
    worst = max((abs(v) for v in diffs.values()), default=0.0)
    failed = worst >= 1.0 or any(v == "未通过" for v in verdicts)
    detail = f"{src}：最大差异 {worst:.4f} 元；" + "、".join(f"{k} {abs(v):.4f}" for k, v in list(diffs.items())[:4])
    return _mk("three_statement_check", round(worst, 6), thr, FAIL if failed else PASS, detail, diffs=diffs)


def check_financing_consistency(ws: Workspace) -> Dict[str, Any]:
    """融资三件套：|融资额 ÷ 投后 − 出让比例| ≤ 1 pp；投后 = 投前 + 融资额（±1%）。"""
    thr = {"equityTolerancePp": 1.0, "postMoneyTolerance": 0.01}
    amt, post, pre, eq = (ws.entry_value(k) for k in ("financing_amount", "post_money", "pre_money", "equity_offered"))
    if amt is None or post is None or eq is None:
        missing = [k for k, v in (("financing_amount", amt), ("post_money", post), ("equity_offered", eq)) if v is None]
        return _skip("financing_consistency", f"numbers.json 缺 {', '.join(missing)}", thr)
    eq_frac = eq / 100 if eq > 1.5 else eq
    implied = amt / post if post else 0.0
    gap_pp = abs(implied - eq_frac) * 100
    notes = [f"融资 {_fmt_num(amt)} ÷ 投后 {_fmt_num(post)} = {implied:.1%} vs 出让 {eq_frac:.1%}（差 {gap_pp:.2f} pp）"]
    ok = gap_pp <= 1.0
    if pre is not None:
        gap_post = abs(post - (pre + amt)) / post if post else 0.0
        ok = ok and gap_post <= 0.01
        notes.append(f"投前 {_fmt_num(pre)} + 融资 = {_fmt_num(pre + amt)} vs 投后 {_fmt_num(post)}（差 {gap_post:.2%}）")
    return _mk("financing_consistency", round(gap_pp, 3), thr, PASS if ok else FAIL, "；".join(notes))


def check_revenue_shape(ws: Workspace) -> Dict[str, Any]:
    """收入曲线非机械：相邻年增长率极差 > 5 pp 为 pass；等差 / 等比（极差 < 0.5 pp）为 fail。"""
    thr = {"growthSpreadPp": 5.0}
    rev = [ws.entry_value(f"revenue_y{i}") for i in range(1, 6)]
    rev = [r for r in rev if r is not None]
    if len(rev) < 3:
        return _skip("revenue_shape", "numbers.json 不足 3 年收入（revenue_y1..y5）", thr)
    if any(r <= 0 for r in rev[:-1]):
        return _mk("revenue_shape", None, thr, WARN, "存在非正收入年份，无法计算增长率")
    growth = [rev[i + 1] / rev[i] - 1 for i in range(len(rev) - 1)]
    spread = (max(growth) - min(growth)) * 100
    diffs = [rev[i + 1] - rev[i] for i in range(len(rev) - 1)]
    arith = (max(diffs) - min(diffs)) <= 0.01 * max(abs(x) for x in diffs) if diffs else False
    if spread > 5:
        status = PASS
    elif spread < 0.5 or arith:
        status = FAIL
    else:
        status = WARN
    detail = f"增长率 {'、'.join(f'{g:.0%}' for g in growth)}，极差 {spread:.1f} pp" + ("（等差）" if arith else "")
    return _mk("revenue_shape", round(spread, 2), thr, status, detail, growth=[round(g, 4) for g in growth])


# ---- 预算
_SUM_ROW_RE = re.compile(r"合计|总计|总额|总预算|汇总")
_SUBTOTAL_RE = re.compile(r"小计")
_CONT_RE = re.compile(r"机动|应急|不可预见|备用金|预备费")


def _table_unit_scale(t: Table) -> Optional[float]:
    ctx = t.context
    if re.search(r"万元", ctx):
        return 1e4
    if re.search(r"亿元", ctx):
        return 1e8
    if re.search(r"[（(]元[)）]|元\b|金额\s*\(?元", ctx):
        return 1.0
    return None


def _amount_col(t: Table) -> Optional[int]:
    for pat in (r"金额", r"预算", r"费用|成本|支出", r"小计|合计|总额"):
        i = t.col(pat)
        if i is not None and i > 0:
            return i
    return None


def _budget_tables(ws: Workspace) -> List[Table]:
    out = []
    for t in ws.all_tables:
        if re.search(r"预算|费用|经费|投入", t.context) and _amount_col(t) is not None and t.nrows >= 2:
            out.append(t)
    return out


def _recalc_table(t: Table) -> Dict[str, Any]:
    """逐行复算：items_sum（含机动金）、total_row（最后一行合计）、contingency、err。"""
    ci = _amount_col(t)
    items: List[float] = []
    total_row: Optional[float] = None
    subtotal: Optional[float] = None
    cont: Optional[float] = None
    for r in t.rows:
        label = " ".join(r[: min(2, len(r))])
        m = _NUM_RE.search(t.cell(r, ci))
        v = _parse_num(m.group(1)) if m else None
        if v is None:
            continue
        if _SUM_ROW_RE.search(label) and not _CONT_RE.search(label):
            total_row = v
        elif _SUBTOTAL_RE.search(label):
            subtotal = v
        else:
            items.append(v)
            if _CONT_RE.search(label):
                cont = v
    items_sum = sum(items)
    err = abs(items_sum - total_row) / total_row if total_row else None
    base = (items_sum - cont) if cont is not None else items_sum
    cont_ratio = (cont / base) if (cont is not None and base) else None
    return {"itemsSum": items_sum, "totalRow": total_row, "subtotal": subtotal, "contingency": cont, "contingencyRatio": cont_ratio, "err": err, "rows": len(items)}


def _xlsx_budget_total(ws: Workspace) -> Optional[float]:
    if ws.xlsx is None:
        return None
    try:
        import openpyxl

        wb = openpyxl.load_workbook(ws.xlsx, read_only=True, data_only=True)
        if "预算总表" not in wb.sheetnames:
            return None
        sh = wb["预算总表"]
        best: Optional[float] = None
        for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row or 100, 300), values_only=True):
            label = " ".join(str(c) for c in row[:2] if c)
            if re.search(r"总计|合计", label) and not re.search(r"小计|阶段", label):
                nums = [c for c in row[2:] if isinstance(c, (int, float))]
                if nums:
                    best = float(nums[-1])
        return best
    except Exception:
        return None


def _budget_total_yuan(ws: Workspace) -> Tuple[Optional[float], str]:
    """账本预算总额（元）：numbers.json budget_total（按 unit 换算），否则 财务模型 预算总表，否则正文预算表合计。"""
    e = ws.entry("budget_total")
    if e and isinstance(e.get("value"), (int, float)):
        scale = _MONEY_UNITS.get(str(e.get("unit") or "万元"), 1e4)
        return float(e["value"]) * scale, "numbers.json"
    x = _xlsx_budget_total(ws)
    if x is not None:
        return x, "财务模型.xlsx!预算总表"
    for t in _budget_tables(ws):
        rc = _recalc_table(t)
        total = rc["totalRow"] or rc["itemsSum"]
        if total:
            return total * (_table_unit_scale(t) or 1.0), f"{t.file}:{t.line}"
    return None, ""


def check_budget_recalc(ws: Workspace) -> Dict[str, Any]:
    """预算可复算：正文预算表 Σ行 vs 合计（≤1%）、合计 vs 账本 budget_total（≤1%，自动试 元/万元 换算）、机动金 5–10%。"""
    thr = {"pass": 0.01, "warn": 0.05, "contingency": [0.05, 0.10]}
    ledger, ledger_src = _budget_total_yuan(ws)
    tables = _budget_tables(ws)
    if not tables and ledger is None:
        return _skip("budget_recalc", "无 numbers.json budget_total 也无正文预算表", thr)
    if not tables:
        return _mk("budget_recalc", None, thr, FAIL, f"账本预算 {ledger_src} 存在，但正文无可复算的预算表")
    best: Optional[Dict[str, Any]] = None
    for t in tables:
        rc = _recalc_table(t)
        total = rc["totalRow"] if rc["totalRow"] else rc["itemsSum"]
        err_ledger: Optional[float] = None
        if ledger and total:
            cands = [total * s for s in (1.0, 1e4, 1e-4, 1e8)]
            err_ledger = min(abs(c - ledger) / ledger for c in cands)
        err_sum = rc["err"] if rc["err"] is not None else 0.0
        err = max(err_sum, err_ledger or 0.0)
        cand = {"table": t, "rc": rc, "err": err, "errSum": err_sum, "errLedger": err_ledger, "total": total}
        if best is None or cand["err"] < best["err"]:
            best = cand
    assert best is not None
    rc, t = best["rc"], best["table"]
    err = best["err"]
    status = _status_max(err, 0.01, 0.05)
    notes = [f"{t.file}:{t.line} Σ行 {_fmt_num(rc['itemsSum'])}" + (f" vs 合计行 {_fmt_num(rc['totalRow'])}（差 {best['errSum']:.2%}）" if rc["totalRow"] else "（无合计行）")]
    if best["errLedger"] is not None:
        notes.append(f"vs 账本 {ledger_src} 差 {best['errLedger']:.2%}")
    cr = rc["contingencyRatio"]
    if cr is not None:
        notes.append(f"机动金 {cr:.1%}")
        if not (0.05 <= cr <= 0.10) and status == PASS:
            status = WARN
            notes[-1] += "（应为 5–10%）"
    else:
        notes.append("未见机动金行")
    return _mk("budget_recalc", round(err, 4), thr, status, "；".join(notes), [_table_loc(t)], recalc={k: v for k, v in rc.items()}, ledgerYuan=ledger)


def check_budget_in_range(ws: Workspace) -> Dict[str, Any]:
    """预算总额落在赛道 / 命题区间内（元）。"""
    rng = ws.ev.get("budget_in_range") or ws.cfg.get("budgetRange")
    thr = {"range": rng}
    if not (isinstance(rng, (list, tuple)) and len(rng) == 2):
        return _skip("budget_in_range", "赛道未设预算区间", thr)
    total, src = _budget_total_yuan(ws)
    if total is None:
        return _skip("budget_in_range", "未找到预算总额（numbers.json / 财务模型 / 正文预算表）", thr)
    lo, hi = float(rng[0]), float(rng[1])
    ok = lo <= total <= hi
    return _mk("budget_in_range", total, thr, PASS if ok else FAIL, f"预算 {_fmt_num(total)} 元（{src}）{'在' if ok else '不在'}区间 [{_fmt_num(lo)}, {_fmt_num(hi)}]")


def check_kpi_table(ws: Workspace) -> Dict[str, Any]:
    """KPI 表：表头含 指标 / 基线 / 目标 / 工具 / 周期 中 ≥ 3 列且行 ≥ 5。"""
    thr = {"cols": 3, "rows": 5}
    if not ws.chapters:
        return _skip("kpi_table", "无章节 Markdown", thr)
    groups = [r"指标|KPI|KGI", r"基线|现状|基准|当前", r"目标|目标值", r"工具|测量|监测|数据来源|口径", r"周期|频率|复盘|时间"]
    best: Optional[Tuple[int, Table]] = None
    for t in ws.all_tables:
        n = sum(1 for g in groups if t.cols(g) > 0)
        if n >= 2 and (best is None or (n, t.nrows) > (best[0], best[1].nrows)):
            best = (n, t)
    if best is None:
        return _mk("kpi_table", 0, thr, FAIL, "未发现 KPI 表（需 指标 / 基线 / 目标 / 工具 / 周期 ≥ 3 列、≥ 5 行）")
    n, t = best
    status = PASS if (n >= 3 and t.nrows >= 5) else WARN
    return _mk("kpi_table", n, thr, status, f"{t.file}:{t.line} 命中 {n}/5 类表头 × {t.nrows} 行", [_table_loc(t)], rows=t.nrows)


def check_positioning_statement(ws: Workspace) -> Dict[str, Any]:
    """定位陈述：「对于……是……品牌」句式或「定位陈述 / 定位句」字样。"""
    thr = {"pattern": "对于.*是.*品牌 | 定位陈述"}
    if not ws.chapters:
        return _skip("positioning_statement", "无章节 Markdown", thr)
    hits = _find_sentences(ws, r"对于[^。！？\n]{2,60}?是[^。！？\n]{1,60}?品牌|定位陈述|定位句|定位声明|定位语")
    return _mk("positioning_statement", len(hits), thr, PASS if hits else FAIL, (f"{hits[0]['file']}:{hits[0]['line']} {hits[0]['text'][:60]}" if hits else "未见定位陈述句"), hits[:3])


def check_big_idea(ws: Workspace) -> Dict[str, Any]:
    """Big Idea / 创意主张 + 视觉锤 / 语言钉。"""
    thr = {"needs": ["Big Idea|创意主张", "视觉锤|语言钉"]}
    if not ws.chapters:
        return _skip("big_idea", "无章节 Markdown", thr)
    idea = _has(ws, r"big\s*idea|创意主张|核心创意")
    hammer = _has(ws, r"视觉锤|语言钉")
    status = PASS if (idea and hammer) else WARN if idea else FAIL
    return _mk("big_idea", int(idea) + int(hammer), thr, status, f"创意主张 {'✓' if idea else '✗'}，视觉锤 / 语言钉 {'✓' if hammer else '✗'}")


def check_kol_tiers(ws: Workspace) -> Dict[str, Any]:
    """KOL 分层：头部 / 腰部 / 尾部 / KOC ≥ 3 种。"""
    thr = {"pass": 3, "warn": 2}
    if not ws.chapters:
        return _skip("kol_tiers", "无章节 Markdown", thr)
    tiers = [t for t in ("头部", "腰部", "尾部", "KOC") if _has(ws, re.escape(t))]
    return _mk("kol_tiers", len(tiers), thr, _status_min(len(tiers), 3, 2), f"命中 {'、'.join(tiers) or '无'}")


def check_offline_event(ws: Workspace) -> Dict[str, Any]:
    thr = {"pass": 1}
    if not ws.chapters:
        return _skip("offline_event", "无章节 Markdown", thr)
    hits = _find_sentences(ws, r"线下(?:活动|快闪|展|体验|门店|巡展|市集|发布会|路演|沙龙|联动)|快闪店|地推|展会|体验店|品鉴会|线下场景")
    return _mk("offline_event", len(hits), thr, PASS if hits else FAIL, (f"{len(hits)} 处线下动作" if hits else "未见线下活动"), hits[:3])


def check_risk_matrix(ws: Workspace) -> Dict[str, Any]:
    """风险矩阵：风险表行数 ≥ risk_matrix_min 且含对策列。"""
    rmin = int(ws.ev.get("risk_matrix_min") or ws.cfg.get("riskMatrixMin") or 6)
    thr = {"rows": rmin, "needs": "对策列"}
    if not ws.chapters:
        return _skip("risk_matrix", "无章节 Markdown", thr)
    best: Optional[Table] = None
    for t in ws.all_tables:
        if re.search(r"风险|预案|应急", t.context) and (best is None or t.nrows > best.nrows):
            best = t
    if best is None:
        return _mk("risk_matrix", 0, thr, FAIL, "未发现风险表")
    has_resp = best.col(r"对策|应对|措施|预案|处置") is not None
    status = PASS if (best.nrows >= rmin and has_resp) else WARN if (best.nrows >= rmin / 2) else FAIL
    return _mk("risk_matrix", best.nrows, thr, status, f"{best.file}:{best.line} {best.nrows} 行（需 ≥ {rmin}），对策列 {'✓' if has_resp else '✗'}", [_table_loc(best)])


def check_brief_response(ws: Workspace) -> Dict[str, Any]:
    """命题回应表：仅命题案例（project.subject.isAssignedCase）且 requiredEvidence 为 if_assigned / true 时检查。"""
    thr = {"when": "isAssignedCase && requiredEvidence.brief_response_table in (true, if_assigned)"}
    flag = ws.ev.get("brief_response_table")
    if not (flag is True or flag == "if_assigned") or not ws.is_assigned_case:
        return _skip("brief_response_table", "非命题案例或赛道未要求", thr)
    if not ws.chapters:
        return _skip("brief_response_table", "无章节 Markdown", thr)
    for t in ws.all_tables:
        if re.search(r"命题", t.context) and re.search(r"回应|对应|响应|落点|方案|措施", t.context):
            return _mk("brief_response_table", t.nrows, thr, PASS, f"{t.file}:{t.line} 「{t.title or t.heading}」{t.nrows} 行", [_table_loc(t)])
    return _mk("brief_response_table", 0, thr, FAIL, "未见「命题回应」表（命题要求 × 本案回应）")


def _count_questions(doc: Doc) -> int:
    n = 0
    for line in doc.lines:
        if re.match(r"^\s*(?:\d{1,3}[.、．)）]|Q\s*\d{1,3}[.、．:：]?|第\s*\d{1,3}\s*题)\s*\S", line):
            n += 1
    return n


def check_appendix_questionnaire(ws: Workspace) -> Dict[str, Any]:
    """附录含问卷且题数 ≥ 25；仅在 01_调研/调研工具包/问卷.md 时为 warn（未入附录）。"""
    thr = {"questions": 25}
    if not ws.chapters:
        return _skip("appendix_questionnaire", "无章节 Markdown", thr)
    cands = [d for d in ws.appendix_docs if re.search(r"问卷", d.text)] or [d for d in ws.chapters if re.search(r"问卷", d.stem)]
    if cands:
        d = max(cands, key=_count_questions)
        n = _count_questions(d)
        return _mk("appendix_questionnaire", n, thr, PASS if n >= 25 else WARN, f"{d.rel} 共 {n} 题（需 ≥ 25）")
    if ws.questionnaire_doc is not None:
        n = _count_questions(ws.questionnaire_doc)
        return _mk("appendix_questionnaire", n, thr, WARN, f"问卷在 {ws.questionnaire_doc.rel}（{n} 题）但未进入附录")
    return _mk("appendix_questionnaire", 0, thr, FAIL, "附录未见问卷")


def check_honor_ratio(ws: Workspace) -> Dict[str, Any]:
    """荣誉 / 获奖 / 认证 / 称号句子的字数占比 ≤ honor_ratio_max。"""
    hmax = float(ws.ev.get("honor_ratio_max") or 0.05)
    thr = {"pass": hmax, "warn": round(hmax * 2, 3)}
    docs = ws.body_docs or ws.chapters
    if not docs:
        return _skip("honor_ratio", "无章节 Markdown", thr)
    total = 0
    honor = 0
    examples: List[Dict[str, Any]] = []
    for d in docs:
        for i, line in enumerate(d.lines):
            for s in _split_sentences(_clean_body(line)):
                w = _count_words(s)
                total += w
                if re.search(r"荣誉|获奖|认证|称号|奖项|殊荣|荣获", s):
                    honor += w
                    if len(examples) < 5:
                        examples.append({"file": d.rel, "line": i + 1, "text": s[:120]})
    value = round(honor / total, 4) if total else 0.0
    return _mk("honor_ratio", value, thr, _status_max(value, hmax, hmax * 2), f"荣誉类句子 {honor}/{total} 字 = {value:.1%}（上限 {hmax:.0%}）", examples)


def check_title_pattern(ws: Workspace) -> Dict[str, Any]:
    """标题范式：封面 title + subtitle 含破折号或「策划书 / 计划书 / 策划案 / 方案 / 报告」。"""
    thr = {"pattern": "——|策划书|计划书|策划案|方案|报告"}
    if not ws.cover:
        return _skip("title_pattern", "无 book.json 封面", thr)
    full = f"{ws.cover.get('title') or ''} {ws.cover.get('subtitle') or ''}".strip()
    ok = bool(re.search(r"——|—|–|－|策划书|计划书|策划案|方案|报告", full))
    return _mk("title_pattern", int(ok), thr, PASS if ok else FAIL, f"「{full}」" + ("" if ok else " 不含破折号 / 文体词"))


def _value_patterns(value: float, unit: str) -> List["re.Pattern[str]"]:
    """账本数值在文本中的写法：原值、千分位；百分比类另加 27.5% 写法。单字符数字过于含糊，跳过。"""
    pats: List[str] = []
    bare = _fmt_num(value)
    if len(bare.replace("-", "")) >= 2:
        pats.append(re.escape(bare))
        if abs(value) >= 1000 and float(value).is_integer():
            pats.append(re.escape(f"{int(value):,}"))
    if unit in _PCT_UNITS and abs(value) <= 1.5:
        p = _fmt_num(round(value * 100, 2))
        pats.append(re.escape(p) + r"\s*[%％]")
    return [re.compile(r"(?<![\d.,])" + p + r"(?![\d.])") for p in pats]


def check_synthetic_markers(ws: Workspace) -> Dict[str, Any]:
    """示例数据标记覆盖率：numbers.json 中 synthetic:true 的数值，每次出现处 ±3 行（PPT 为同页）内须有「示例数据」。"""
    target = float(ws.cfg.get("syntheticMarkerCoverage") or 1.0)
    thr = {"pass": target, "warn": round(target * 0.8, 2), "window": 3}
    synth = [e for e in ws.entries if e.get("synthetic") and isinstance(e.get("value"), (int, float))]
    if not synth:
        return _skip("synthetic_marker_coverage", "numbers.json 无 synthetic:true 条目", thr)
    docs = ws.chapters + ws.pitch_docs
    total = 0
    missing: List[Dict[str, Any]] = []
    skipped_keys: List[str] = []
    for e in synth:
        pats = _value_patterns(float(e["value"]), str(e.get("unit") or ""))
        if not pats:
            skipped_keys.append(str(e["key"]))
            continue
        for d in docs:
            spans = [((t.line - 2 if t.title else t.line - 1), t.end_line) for t in d.tables]
            for i, line in enumerate(d.lines):
                if not any(p.search(line) for p in pats):
                    continue
                total += 1
                lo, hi = i - 3, i + 4
                for s0, s1 in spans:  # 表格内的数字：标记可在题注前 3 行到表后 3 行（来源行）之间
                    if s0 <= i < s1:
                        lo, hi = min(lo, s0 - 3), max(hi, s1 + 3)
                window = "\n".join(d.lines[max(0, lo): hi])
                if not _MARKER_RE.search(window):
                    missing.append({"file": d.rel, "line": i + 1, "key": e["key"], "text": line.strip()[:120]})
        for p in ws.slide_pages():
            blob = "\n".join([p["title"], p["note"]] + p["texts"])
            if any(pt.search(blob) for pt in pats):
                total += 1
                if not _MARKER_RE.search(blob):
                    missing.append({"file": "06_路演/slides.json", "line": p["index"], "key": e["key"], "text": p["title"][:120]})
    value = _ratio(total - len(missing), total) if total else 1.0
    detail = f"{total - len(missing)}/{total} 处示例数值带「示例数据」标记（条目 {', '.join(str(e['key']) for e in synth)}）"
    if skipped_keys:
        detail += f"；值过短未检索：{', '.join(skipped_keys)}"
    return _mk("synthetic_marker_coverage", value, thr, _status_min(value, target, target * 0.8), detail, missing, occurrences=total)


# ---- 关键词 / 表格启发式（创业 / 案例 / 数字营销 / 会奖 专项）
def _kw_check(cid: str, ws: Workspace, pattern: str, pass_min: int = 1, warn_min: int = 0, label: str = "") -> Dict[str, Any]:
    thr = {"pass": pass_min, "warn": warn_min}
    if not ws.chapters:
        return _skip(cid, "无章节 Markdown", thr)
    hits = _find_sentences(ws, pattern)
    n = len(hits)
    status = PASS if n >= pass_min else WARN if n > warn_min else FAIL
    return _mk(cid, n, thr, status, f"{label or cid} 命中 {n} 处", hits[:5])


def _table_check(cid: str, ws: Workspace, ctx_pattern: Optional[str], header_groups: Sequence[str], need_groups: int, min_rows: int, label: str) -> Dict[str, Any]:
    """找一张「上下文匹配 + 表头命中 ≥ need_groups 组 + 行数 ≥ min_rows」的表。"""
    thr = {"headerGroups": need_groups, "rows": min_rows}
    if not ws.chapters:
        return _skip(cid, "无章节 Markdown", thr)
    best: Optional[Tuple[int, Table]] = None
    for t in ws.all_tables:
        if ctx_pattern and not re.search(ctx_pattern, t.context, re.I):
            continue
        n = sum(1 for g in header_groups if t.cols(g) > 0)
        if best is None or (n, t.nrows) > (best[0], best[1].nrows):
            best = (n, t)
    if best is None:
        return _mk(cid, 0, thr, FAIL, f"未发现{label}")
    n, t = best
    status = PASS if (n >= need_groups and t.nrows >= min_rows) else WARN if n >= max(1, need_groups - 1) else FAIL
    return _mk(cid, n, thr, status, f"{t.file}:{t.line} 表头命中 {n}/{len(header_groups)} 组 × {t.nrows} 行", [_table_loc(t)], rows=t.nrows)


def check_ip_list(ws: Workspace) -> Dict[str, Any]:
    """知识产权：专利 / 软著 / 商标 + 编号格式（ZL…、第…号、申请号）或知识产权表。"""
    thr = {"needs": ["专利|软著|商标", "编号格式或清单表"]}
    if not ws.chapters:
        return _skip("ip_list", "无章节 Markdown", thr)
    kw = _find_sentences(ws, r"专利|软著|软件著作权|商标注册|知识产权")
    numbered = _find_sentences(ws, r"ZL\s*\d{6,}|第\s*\d{5,}\s*号|申请号|注册号|受理号|专利号")
    table = [t for t in ws.all_tables if re.search(r"专利|知识产权|商标|软著", t.context)]
    status = PASS if kw and (numbered or table) else WARN if kw else FAIL
    return _mk("ip_list", len(kw), thr, status, f"知识产权关键词 {len(kw)} 处，编号 {len(numbered)} 处，清单表 {len(table)} 张", (numbered or kw)[:3])


def check_param_compare(ws: Workspace) -> Dict[str, Any]:
    """参数 / 指标对比表：≥ 2 列竞品（即 ≥ 3 个实体列含本品）或题注含「对比」。"""
    thr = {"entities": 3}
    if not ws.chapters:
        return _skip("param_compare_table", "无章节 Markdown", thr)
    names = ws.self_names()
    for t in ws.all_tables:
        if re.search(r"参数|指标|性能|对比|比较", t.context) and (_entity_cols(t, names) >= 2 or re.search(r"参数对比|指标对比|性能对比|技术对比", t.context)) and t.nrows >= 3:
            return _mk("param_compare_table", t.nrows, thr, PASS, f"{t.file}:{t.line} 「{t.title or t.heading}」{t.nrows} 行", [_table_loc(t)])
    return _mk("param_compare_table", 0, thr, FAIL, "未发现参数 / 指标对比表")


def check_milestone_gantt(ws: Workspace) -> Dict[str, Any]:
    """甘特图 / 里程碑：图题含甘特 / 里程碑 / 排期 / 路线图，或带日期的里程碑表（≥ 3 行）。"""
    thr = {"needs": "甘特图或里程碑表"}
    if not ws.chapters:
        return _skip("milestone_gantt", "无章节 Markdown", thr)
    figs = [f for f in ws.all_figures if re.search(r"甘特|里程碑|排期|时间表|路线图|gantt|roadmap", f.alt + " " + f.path, re.I)]
    tabs = [t for t in ws.all_tables if re.search(r"里程碑|排期|时间表|路线|阶段|计划", t.context) and t.nrows >= 3 and re.search(r"20\d\d|Q[1-4]|\d+\s*月|M\d+", " ".join(" ".join(r) for r in t.rows))]
    ok = bool(figs or tabs)
    locs = [{"file": f.file, "line": f.line, "text": f.alt} for f in figs[:2]] + [_table_loc(t) for t in tabs[:2]]
    return _mk("milestone_gantt", len(figs) + len(tabs), thr, PASS if ok else FAIL, f"甘特 / 路线图 {len(figs)} 张，里程碑表 {len(tabs)} 张", locs)


def check_equity_table(ws: Workspace) -> Dict[str, Any]:
    """股权表：持股比例列合计 ≈ 100%（±1）。"""
    thr = {"sum": 100, "tolerance": 1}
    if not ws.chapters:
        return _skip("equity_table", "无章节 Markdown", thr)
    for t in ws.all_tables:
        if not re.search(r"股权|持股|股份|出资|股比", t.context):
            continue
        ci = t.col(r"比例|占比|股比|持股|%")
        if ci is None:
            continue
        vals = []
        for r in t.rows:
            if _SUM_ROW_RE.search(" ".join(r[:1])):
                continue
            m = _NUM_RE.search(t.cell(r, ci))
            v = _parse_num(m.group(1)) if m else None
            if v is not None:
                vals.append(v)
        s = sum(vals)
        if s <= 1.5:
            s *= 100
        ok = abs(s - 100) <= 1
        return _mk("equity_table", round(s, 2), thr, PASS if ok else WARN, f"{t.file}:{t.line} 持股合计 {s:.1f}%（{len(vals)} 行）", [_table_loc(t)])
    return _mk("equity_table", 0, thr, FAIL, "未发现含比例列的股权表")


def check_team_matrix(ws: Workspace) -> Dict[str, Any]:
    return _table_check("team_matrix", ws, r"团队|成员|分工|组织", [r"姓名|成员|角色|岗位|职位", r"背景|专业|经历|履历", r"分工|职责|负责", r"成果|业绩|经验|资源"], 3, 2, "团队 背景–分工–成果 表")


def check_short_video_script(ws: Workspace) -> Dict[str, Any]:
    """短视频脚本：关键词 + 分镜表（镜头 / 画面 / 台词 / 时长 / 字幕 ≥ 2 列）。"""
    thr = {"needs": ["短视频脚本|分镜", "分镜表 ≥ 2 列"]}
    if not ws.chapters:
        return _skip("short_video_script", "无章节 Markdown", thr)
    kw = _has(ws, r"短视频脚本|分镜|脚本")
    tabs = [t for t in ws.all_tables if sum(1 for g in (r"镜头|镜号|分镜", r"画面|景别", r"台词|文案|旁白|口播", r"时长|秒", r"字幕|BGM|音乐") if t.cols(g) > 0) >= 2]
    status = PASS if (kw and tabs) else WARN if kw else FAIL
    return _mk("short_video_script", len(tabs), thr, status, f"脚本关键词 {'✓' if kw else '✗'}，分镜表 {len(tabs)} 张", [_table_loc(t) for t in tabs[:2]])


_PLATFORM_RE = r"抖音|小红书|微信|视频号|微博|B\s*站|哔哩哔哩|快手|天猫|京东|淘宝|知乎|美团|大众点评|拼多多|得物|QQ|支付宝|百度|腾讯|TikTok|Instagram|YouTube|Facebook|LinkedIn"


def check_platform_matrix(ws: Workspace) -> Dict[str, Any]:
    """平台矩阵：表头含 ≥ 3 个平台名，或「平台」列下 ≥ 3 行平台名。"""
    thr = {"platforms": 3}
    if not ws.chapters:
        return _skip("platform_matrix", "无章节 Markdown", thr)
    best: Optional[Tuple[int, Table]] = None
    for t in ws.all_tables:
        n_head = len({m.group(0) for h in t.header for m in re.finditer(_PLATFORM_RE, h, re.I)})
        pc = t.col(r"平台|渠道|媒介")
        n_rows = len({m.group(0) for r in t.rows for m in re.finditer(_PLATFORM_RE, t.cell(r, pc), re.I)}) if pc is not None else 0
        n = max(n_head, n_rows)
        if best is None or n > best[0]:
            best = (n, t)
    if best is None or best[0] == 0:
        return _mk("platform_matrix", 0, thr, FAIL, "未发现平台 × 内容 / 预算矩阵")
    n, t = best
    return _mk("platform_matrix", n, thr, _status_min(n, 3, 2), f"{t.file}:{t.line} 覆盖 {n} 个平台", [_table_loc(t)])


def check_ratio_table_formula(ws: Workspace) -> Dict[str, Any]:
    return _table_check("ratio_table_with_formula", ws, r"比率|财务指标|指标|偿债|盈利|营运", [r"比率|指标", r"公式|计算|算法|定义", r"20\d\d|本年|上年|数值|行业"], 2, 5, "带公式列的财务比率表")


def check_comprehensive_income(ws: Workspace) -> Dict[str, Any]:
    """综合收益 / 利润预测表：题注或标题含关键词且表有 ≥ 2 个年份列。"""
    thr = {"yearCols": 2}
    if not ws.chapters:
        return _skip("comprehensive_income_forecast", "无章节 Markdown", thr)
    for t in ws.all_tables:
        if re.search(r"综合收益|利润预测|盈利预测|预测利润表|预计利润|利润表", t.context):
            yc = sum(1 for h in t.header if re.search(r"20\d\d|第\s*[一二三四五\d]\s*年|Y\d|E\b", h))
            if yc >= 2:
                return _mk("comprehensive_income_forecast", yc, thr, PASS, f"{t.file}:{t.line} 「{t.title or t.heading}」{yc} 个年份列", [_table_loc(t)])
    return _mk("comprehensive_income_forecast", 0, thr, FAIL, "未发现带年份列的综合收益 / 利润预测表")


def check_annual_report_refs(ws: Workspace) -> Dict[str, Any]:
    return _kw_check("annual_report_page_refs", ws, r"年报[^。\n]{0,30}?第?\s*\d{1,3}\s*页|年度报告[^。\n]{0,30}?第?\s*\d{1,3}\s*页|[Pp]\.\s*\d{1,3}\b|\bp\d{1,3}\b", 3, 0, "年报页码引用")


def check_contingency(ws: Workspace) -> Dict[str, Any]:
    """突发事件预案条数：预案表行数或应急章节的列表项 ≥ contingency_min。"""
    cmin = int(ws.ev.get("contingency_min") or 8)
    thr = {"pass": cmin, "warn": round(cmin / 2)}
    if not ws.chapters:
        return _skip("contingency", "无章节 Markdown", thr)
    rows = max((t.nrows for t in ws.all_tables if re.search(r"预案|应急|突发|风险", t.context)), default=0)
    items = 0
    for d in ws.chapters:
        in_sec = False
        for line in d.lines:
            hm = _HEADING_RE.match(line)
            if hm:
                in_sec = bool(re.search(r"突发|应急|预案", hm.group(2)))
                continue
            if in_sec and re.match(r"^\s*(?:[-*+]|\d+[.、．)])\s+\S", line):
                items += 1
    n = max(rows, items)
    return _mk("contingency", n, thr, _status_min(n, cmin, cmin / 2), f"预案 {n} 条（表 {rows} 行 / 应急章节列表 {items} 项，需 ≥ {cmin}）")


def check_itinerary_timeline(ws: Workspace) -> Dict[str, Any]:
    """行程时间轴：hh:mm 时刻点 ≥ 10。"""
    thr = {"pass": 10, "warn": 5}
    if not ws.chapters:
        return _skip("itinerary_timeline", "无章节 Markdown", thr)
    n = len(re.findall(r"(?<!\d)(?:[01]?\d|2[0-3])[:：][0-5]\d(?!\d)", ws.chapter_text))
    return _mk("itinerary_timeline", n, thr, _status_min(n, 10, 5), f"{n} 个 hh:mm 时刻点（需 ≥ 10）")


def check_itinerary_grid(ws: Workspace) -> Dict[str, Any]:
    """行程表：表头含时间 + 日期 / Day，或 ≥ 3 个 Day 列。"""
    thr = {"needs": "时间 × 日程 表"}
    if not ws.chapters:
        return _skip("itinerary_grid", "无章节 Markdown", thr)
    for t in ws.all_tables:
        days = sum(1 for h in t.header if re.search(r"D\s*\d|Day\s*\d|第\s*[一二三四五六七\d]\s*天|\d{1,2}\s*月\s*\d{1,2}\s*日", h, re.I))
        if (t.col(r"时间|时段|时刻") is not None and (days >= 1 or t.col(r"日期|Day|天") is not None) and t.nrows >= 3) or days >= 3:
            return _mk("itinerary_grid", t.nrows, thr, PASS, f"{t.file}:{t.line} 「{t.title or t.heading}」{t.nrows} 行", [_table_loc(t)])
    return _mk("itinerary_grid", 0, thr, FAIL, "未发现 时间 × 日程 行程表")


def check_venue_capacity(ws: Workspace) -> Dict[str, Any]:
    """场地容量 ≥ 参会人数（numbers.json headcount 或正文人数）。"""
    thr = {"rule": "容量 ≥ 人数"}
    if not ws.chapters:
        return _skip("venue_capacity_check", "无章节 Markdown", thr)
    caps = [int(m.group(1)) for m in re.finditer(r"(?:容纳|可容纳|容量|可接待|接待能力|最大承载)[^\d\n]{0,10}?(\d{2,6})\s*人", ws.chapter_text)]
    head = ws.entry_value("headcount")
    if head is None:
        m = re.search(r"(?:参会|参与|团队|总)人数[^\d\n]{0,8}?(\d{2,6})\s*人", ws.chapter_text)
        head = float(m.group(1)) if m else None
    if not caps:
        return _mk("venue_capacity_check", 0, thr, FAIL, "未见场地容量说明（可容纳 N 人）")
    if head is None:
        return _mk("venue_capacity_check", max(caps), thr, WARN, f"容量 {max(caps)} 人，但未找到参会人数可比对")
    ok = max(caps) >= head
    return _mk("venue_capacity_check", max(caps), thr, PASS if ok else FAIL, f"场地容量 {max(caps)} 人 vs 人数 {_fmt_num(head)}")


def check_vendor_quote(ws: Workspace) -> Dict[str, Any]:
    return _table_check("vendor_quote_table", ws, r"供应商|报价|采购|询价", [r"供应商|服务商|厂商", r"报价|单价|价格", r"来源|报价单|询价|日期|联系方式|备注"], 3, 3, "供应商报价表")


def check_regional_feature(ws: Workspace) -> Dict[str, Any]:
    return _kw_check("regional_feature", ws, r"地域特色|区域特色|本地特色|本土|地方特色|地理标志|原产地|地域文化|地方文化|区域经济", 3, 0, "地域特色")


def check_problem_identification(ws: Workspace) -> Dict[str, Any]:
    return _kw_check("problem_identification", ws, r"问题识别|核心问题|关键问题|问题诊断|问题树", 1, 0, "问题识别")


def check_industry_cluster(ws: Workspace) -> Dict[str, Any]:
    return _kw_check("industry_cluster_analysis", ws, r"产业集群|集群分析|产业链|上下游", 2, 0, "产业集群分析")


def check_analysis_framework(ws: Workspace) -> Dict[str, Any]:
    """分析框架图：图题含框架 / 模型 / 逻辑图。"""
    thr = {"pass": 1}
    if not ws.chapters:
        return _skip("analysis_framework_diagram", "无章节 Markdown", thr)
    figs = [f for f in ws.all_figures if re.search(r"框架|模型|逻辑|路径图|分析图|framework", f.alt + " " + f.path, re.I)]
    return _mk("analysis_framework_diagram", len(figs), thr, PASS if figs else FAIL, f"分析框架图 {len(figs)} 张", [{"file": f.file, "line": f.line, "text": f.alt} for f in figs[:2]])


def check_novel_method(ws: Workspace) -> Dict[str, Any]:
    return _kw_check("novel_method", ws, r"创新方法|新方法|方法创新|新颖|首次引入|改进.{0,6}模型", 1, 0, "方法创新")


# ---------------------------------------------------------------------------------------------- 路演：PPT / 讲稿 / 题库
_STRUCT_LAYOUTS = {"cover", "agenda", "section", "closing"}
_VERB_RE = re.compile(r"\d|是|将|让|使|达|增|降|超|占|比|提升|实现|成为|可|能|需|应|要|从|到|不足|领先|落后|翻|倍|高|低|仅|只|率|靠|先|再|而|但|却|否|非|无|有|并|把|用|做|成|变|破|抢|赢|打|建|造|卖|买|帮|给|定|选|走|进|入|来|去|回|跑|涨|跌|缺|少|多|快|慢|好|差|新|老")


def check_slide_metrics(ws: Workspace) -> Dict[str, Any]:
    """PPT 指标：正文页数区间（不含 appendix）、备用页上限、每页字数 ≤ 120、结论句标题率 ≥ 80%、含数字页 ≥ 80%、图表页 ≥ 8。"""
    pitch = ws.cfg.get("pitch") or {}
    rng = list(pitch.get("deckPages") or ws.cfg.get("deckPages") or [16, 20])
    app_max = int(pitch.get("deckPagesAppendixMax") or ws.cfg.get("deckPagesAppendixMax") or 4)
    words_max = int(ws.cfg.get("slideWordsMax") or 120)
    chart_min = int(ws.cfg.get("deckChartMin") or 8)
    thr = {"pages": rng, "appendixMax": app_max, "slideWordsMax": words_max, "conclusionRatio": 0.8, "numericRatio": 0.8, "chartMin": chart_min}
    pages = ws.slide_pages()
    if not pages:
        return _skip("slide_metrics", "无 06_路演/slides.json 与 产出/路演.pptx.preview.json", thr)
    body = [p for p in pages if p["layout"] != "appendix"]
    appendix = [p for p in pages if p["layout"] == "appendix"]
    content = [p for p in body if p["layout"] not in _STRUCT_LAYOUTS]
    over: List[Dict[str, Any]] = []
    for p in body:
        if p["layout"] == "table":
            continue
        chars = len(re.sub(r"\s", "", "".join(p["texts"])))
        if chars > words_max:
            over.append({"file": "06_路演/slides.json", "line": p["index"], "text": f"{chars} 字：{p['title'][:40]}"})
    concl = [p for p in content if _VERB_RE.search(p["title"])]
    numeric = [p for p in content if _DIGIT_RE.search(p["title"] + " ".join(p["texts"]))]
    charts = [p for p in body if p["visual"]]
    c_ratio = _ratio(len(concl), len(content)) if content else 1.0
    n_ratio = _ratio(len(numeric), len(content)) if content else 1.0
    sub: Dict[str, str] = {
        "pages": PASS if rng[0] <= len(body) <= rng[1] else FAIL,
        "appendix": PASS if len(appendix) <= app_max else WARN,
        "words": PASS if not over else WARN,
        "conclusion": _status_min(c_ratio, 0.8, 0.6),
        "numeric": _status_min(n_ratio, 0.8, 0.6),
        "charts": _status_min(len(charts), chart_min, chart_min / 2),
    }
    metrics = {"bodyPages": len(body), "appendixPages": len(appendix), "overWordsPages": len(over), "conclusionRatio": c_ratio, "numericRatio": n_ratio, "chartPages": len(charts)}
    detail = (f"正文 {len(body)} 页（区间 {rng[0]}–{rng[1]}）+ 备用 {len(appendix)}；超 {words_max} 字 {len(over)} 页；"
              f"结论句标题 {c_ratio:.0%}；含数字页 {n_ratio:.0%}；图表页 {len(charts)}（需 ≥ {chart_min}）")
    return _mk("slide_metrics", len(body), thr, _worst(sub.values()), detail, over, metrics=metrics, sub=sub)


def check_deck_font_min(ws: Workspace) -> Dict[str, Any]:
    """PPT 最小字号：读 pptx_render 结果 产出/路演.render.json 的 perPage[].minFont（≥ 24 pass / ≥ 18 warn）。"""
    thr = {"pass": 24, "warn": 18}
    per = (ws.render or {}).get("perPage") if ws.render else None
    if not isinstance(per, list) or not per:
        return _skip("deck_font_min", "无 产出/路演.render.json（pptx_render 的结果 JSON）", thr)
    fonts = [(p.get("index"), float(p["minFont"])) for p in per if isinstance(p, dict) and isinstance(p.get("minFont"), (int, float))]
    if not fonts:
        return _skip("deck_font_min", "render.json 无 perPage.minFont", thr)
    vmin = min(f for _, f in fonts)
    low = [{"file": "产出/路演.render.json", "line": i, "text": f"minFont {f}"} for i, f in fonts if f < 18]
    return _mk("deck_font_min", vmin, thr, _status_min(vmin, 24, 18), f"最小字号 {vmin} pt；< 18 pt 的页 {len(low)}", low)


def _script_text(doc: Doc) -> str:
    """讲稿正文：去标题、表格、时间标记；「关键数字卡」之后的内容不计。"""
    out: List[str] = []
    for line in doc.lines:
        hm = _HEADING_RE.match(line)
        if hm:
            if re.search(r"关键数字|数字卡|附录", hm.group(2)):
                break
            continue
        if line.strip().startswith("|") or _TBL_SEP_RE.match(line):
            continue
        out.append(_TIME_MARK_RE.sub(" ", line))
    return _clean_body("\n".join(out))


def check_script_timing(ws: Workspace) -> Dict[str, Any]:
    """讲稿时长：字数 ÷ 190 字/分钟；区间 pitch.scriptWords。"""
    pitch = ws.cfg.get("pitch") or {}
    rng = list(pitch.get("scriptWords") or ws.cfg.get("scriptWords") or [1650, 1950])
    cpm = float(ws.cfg.get("scriptCharsPerMinute") or 190)
    thr = {"range": rng, "charsPerMinute": cpm}
    if ws.script is None:
        return _skip("script_timing", "无 06_路演/讲稿.md", thr)
    text = _script_text(ws.script)
    words = _count_words(text)
    minutes = round(words / cpm, 1) if cpm else 0.0
    paras = len([p for p in re.split(r"\n\s*\n", text) if p.strip()])
    if rng[0] <= words <= rng[1]:
        status = PASS
    elif rng[0] * 0.9 <= words <= rng[1] * 1.1:
        status = WARN
    else:
        status = FAIL
    return _mk("script_timing", words, thr, status, f"{words} 字 ≈ {minutes} 分钟（区间 {rng[0]}–{rng[1]} 字），{paras} 段", minutes=minutes, paragraphs=paras)


def check_sentence_length(ws: Workspace) -> Dict[str, Any]:
    """讲稿句长：中位数 ≤ 25 且最长 ≤ 40 为 pass。"""
    thr = {"median": 25, "max": 40}
    if ws.script is None:
        return _skip("sentence_length", "无 06_路演/讲稿.md", thr)
    sents = [s for s in re.split(r"[。！？；!?;\n]", _script_text(ws.script)) if _count_words(s) >= 2]
    if not sents:
        return _skip("sentence_length", "讲稿无可切分句子", thr)
    lens = [_count_words(s) for s in sents]
    med, mx = _median(lens), max(lens)
    longest = max(sents, key=_count_words).strip()
    status = PASS if (med <= 25 and mx <= 40) else WARN if med <= 35 else FAIL
    long_locs = [{"file": ws.script.rel, "line": 0, "text": s.strip()[:120]} for s in sents if _count_words(s) > 40][:10]
    return _mk("sentence_length", med, thr, status, f"中位数 {_fmt_num(med)} 字，最长 {mx} 字（{len(sents)} 句）：「{longest[:40]}…」", long_locs, max=mx, sentences=len(sents))


_QA_HEAD_RE = re.compile(r"^\s*#{2,4}\s*Q\s*(\d+)\s*[【\[]([^】\]]+)[】\]]\s*(.*)$")


def check_qa_bank(ws: Workspace) -> Dict[str, Any]:
    """答辩题库：`### Qn【维度】` 题数 ≥ qaMin、qaDimensions 各 ≥ 4、每题含 主答 + 结论 / 证据 / 行动。"""
    qmin = int(ws.cfg.get("qaMin") or 40)
    qwarn = int(ws.cfg.get("qaWarn") or 20)
    dims = [str(d) for d in (ws.cfg.get("qaDimensions") or [])]
    thr = {"pass": qmin, "warn": qwarn, "perDimension": 4, "dimensions": dims}
    if ws.qa is None:
        return _skip("qa_bank", "无 06_路演/答辩题库.md", thr)
    blocks: List[Tuple[str, str, int]] = []
    cur: Optional[Tuple[str, int]] = None
    buf: List[str] = []
    for i, line in enumerate(ws.qa.lines):
        m = _QA_HEAD_RE.match(line)
        if m:
            if cur is not None:
                blocks.append((cur[0], "\n".join(buf), cur[1]))
            cur, buf = (m.group(2).strip(), i + 1), []
        elif cur is not None:
            buf.append(line)
    if cur is not None:
        blocks.append((cur[0], "\n".join(buf), cur[1]))
    if not blocks:
        n_fallback = _count_questions(ws.qa)
        if n_fallback == 0:
            return _mk("qa_bank", 0, thr, FAIL, "题库未使用 `### Qn【维度】` 格式且无编号题目")
        return _mk("qa_bank", n_fallback, thr, WARN if n_fallback >= qwarn else FAIL, f"{n_fallback} 道编号题目，但未使用 `### Qn【维度】` 格式（无法统计维度 / 三段式）")
    n = len(blocks)
    per_dim: Dict[str, int] = {d: 0 for d in dims}
    other = 0
    for dim, _, _ in blocks:
        matched = next((d for d in dims if _norm(d) == _norm(dim)), None) or next((d for d in dims if len(set(_norm(d)) & set(_norm(dim))) >= 2), None)
        if matched:
            per_dim[matched] += 1
        else:
            other += 1
    weak_dims = [d for d, c in per_dim.items() if c < 4]
    no_owner = [ln for _, body, ln in blocks if not re.search(r"主答\s*[:：]\s*\S", body)]
    no_three = [ln for _, body, ln in blocks if not all(re.search(k + r"\s*[:：（(]", body) for k in ("结论", "证据", "行动"))]
    complete = n >= qmin and not weak_dims and not no_owner and not no_three
    status = PASS if complete else WARN if n >= qwarn else FAIL
    detail = f"{n} 题（需 ≥ {qmin}）；维度 " + "、".join(f"{d} {c}" for d, c in per_dim.items()) + (f"，其他 {other}" if other else "")
    if weak_dims:
        detail += f"；不足 4 题：{'、'.join(weak_dims)}"
    if no_owner:
        detail += f"；缺主答 {len(no_owner)} 题"
    if no_three:
        detail += f"；缺三段 {len(no_three)} 题"
    locs = [{"file": ws.qa.rel, "line": ln, "text": "缺主答"} for ln in no_owner[:5]] + [{"file": ws.qa.rel, "line": ln, "text": "缺 结论/证据/行动"} for ln in no_three[:5]]
    return _mk("qa_bank", n, thr, status, detail, locs, perDimension=per_dim, missingOwner=len(no_owner), missingThreePart=len(no_three))


# ---------------------------------------------------------------------------------------------- 跨材料数字一致性
_KEY_KEYWORDS: Dict[str, List[str]] = {
    "financing_amount": ["融资额", "融资金额", "融资需求", "本轮融资", "拟融资", "融资"],
    "pre_money": ["投前估值", "投前"],
    "post_money": ["投后估值", "投后"],
    "equity_offered": ["出让比例", "出让股权", "出让", "释放股权", "让渡"],
    "breakeven_month": ["盈亏平衡月", "盈亏平衡点", "盈亏平衡", "保本点", "实现盈亏平衡"],
    "min_cash": ["现金最低点", "最低现金", "现金流最低", "资金缺口"],
    "cac": ["cac", "获客成本"],
    "ltv": ["ltv", "客户终身价值", "生命周期价值", "用户终身价值"],
    "ltv_cac": ["ltv/cac", "ltv:cac"],
    "payback_months": ["回收期", "回本周期", "回本期"],
    "tam": ["tam", "潜在市场规模", "潜在市场"],
    "sam": ["sam", "可服务市场"],
    "som": ["som", "可获得市场"],
    "som_y1": ["som", "可获得市场", "首年可获得"],
    "sample_size": ["有效样本量", "有效样本", "有效问卷", "样本量", "份有效", "有效回收", "有效答卷", "有效"],
    "budget_total": ["预算总额", "总预算", "预算合计", "预算总计", "传播预算", "推广预算", "营销预算", "整体预算", "总投入", "预算规模", "预算"],
    "budget_per_person": ["人均预算", "人均费用", "人均"],
    "headcount": ["参会人数", "参与人数", "总人数", "人数规模", "人数"],
    "roi": ["roi", "投入产出比", "投资回报率"],
    "kpi_exposure": ["曝光量", "曝光"],
    "kpi_conversion": ["转化量", "成交订单", "转化"],
    "kpi_interaction": ["互动量", "互动"],
    "kpi_satisfaction": ["满意度"],
    "roe_y3": ["第3年roe", "roe", "净资产收益率"],
}
# 含这些词的句子不参与比对（理论样本量、预算上限、阶段拆分等不是“同一个数字”）
_KEY_EXCLUDE: Dict[str, "re.Pattern[str]"] = {
    "sample_size": re.compile(r"公式|理论|最少|至少|不少于|需要|应达|要求|≥|>=|计划|目标|拟|发放"),
    "budget_total": re.compile(r"阶段|其中|占比|分配|单价|上限|不超过|以内|≤|封顶|明细|第[一二三四五]|机动|比例"),
    "financing_amount": re.compile(r"融资后|融资前|融资方式|融资轮次|融资历史|融资进度|已完成|天使轮已|用途"),
    "headcount": re.compile(r"容纳|容量|接待|上限|每桌|每车|每组"),
    "kpi_exposure": re.compile(r"阶段|其中|单日|日均|每|第[一二三四五]"),
    "kpi_conversion": re.compile(r"阶段|其中|单日|日均|每|第[一二三四五]|率"),
}
_BASE_WORDS = {"revenue": ["营业收入", "营收", "收入"], "net_profit": ["净利润", "净利"], "gross_margin": ["毛利率"]}
# 泛化关键词（预算 / 融资 / 出让…）只在数字紧邻时才比对：关键词后 15 字内或前 8 字内
_GENERIC_KWS = {"预算", "融资", "出让", "转化", "曝光", "互动", "人均", "人数", "投后", "投前", "收入", "营收", "让渡", "roe", "有效"}
_NEAR_AFTER, _NEAR_BEFORE = 15, 8
# 情景 / 敏感性句（“若成本上浮 20%，ROI 降至 2.1 倍”）说的本来就不是账本上的那个数，所有键一律不比对
_SCENARIO_RE = re.compile(r"若|如果|假设|假如|情景|情境|敏感性|悲观|乐观|保守|激进|上浮|下浮|降至|升至|底线|压力测试")


def _keywords_for(key: str, label: str) -> List[str]:
    kws: set[str] = set()
    lab = _norm(re.sub(r"[（(【\[][^）)】\]]*[）)】\]]", "", str(label or "")))
    if len(lab) >= 2:
        kws.add(lab)
    m = re.match(r"第(\d)年(.+)", lab)
    if m:
        n, rest = int(m.group(1)), m.group(2)
        kws |= {f"第{n}年{rest}", f"第{_CN_NUM.get(n, n)}年{rest}", f"y{n}{rest}", f"{rest}y{n}"}
        if n == 1:
            kws.add(f"首年{rest}")
    km = re.match(r"(revenue|net_profit|gross_margin)_y(\d)$", key)
    if km:
        n = int(km.group(2))
        for b in _BASE_WORDS[km.group(1)]:
            kws |= {f"第{n}年{b}", f"第{_CN_NUM.get(n, n)}年{b}"}
            if n == 1:
                kws.add(f"首年{b}")
    kws |= set(_KEY_KEYWORDS.get(key, []))
    return sorted(kws, key=len, reverse=True)


def _entry_class(unit: str) -> str:
    if unit in _MONEY_UNITS:
        return "money"
    if unit in _PCT_UNITS:
        return "pct"
    return "other"


def _expected_base(value: float, unit: str) -> float:
    cls = _entry_class(unit)
    if cls == "money":
        return abs(value) * _MONEY_UNITS[unit]
    if cls == "pct":
        return abs(value) / 100 if abs(value) > 1.5 else abs(value)
    return abs(value)


def _candidates(text: str, unit: str) -> List[Tuple[List[float], str, int]]:
    """句中与账本条目同类的数字：返回 (可能的基准值列表, 原文, 位置)。"""
    cls = _entry_class(unit)
    out: List[Tuple[List[float], str, int]] = []
    for m in _NUM_RE.finditer(text):
        v = _parse_num(m.group(1))
        if v is None:
            continue
        if _NOT_QUANTITY_BEFORE.search(text[: m.start()]) or _NOT_QUANTITY_AFTER.match(text[m.end():]):
            continue
        v = abs(v)
        u = m.group(2) or ""
        if u in ("章", "节", "期", "轮", "批", "项", "步", "部分", "层", "级", "段", "页", "条", "张", "场", "次", "年", "天", "小时", "位", "名") and u != unit:
            continue
        if cls == "money":
            if u in _MONEY_UNITS:
                vals = [v * _MONEY_UNITS[u]]
            elif not u:
                vals = [v * _MONEY_UNITS[unit], v]
            else:
                continue
        elif cls == "pct":
            if u in _PCT_UNITS:
                vals = [v / 100]
            elif not u:
                vals = [v, v / 100]
            else:
                continue
        else:
            if u and u != unit and not (unit in ("月", "份", "人", "倍") and u == unit):
                continue
            vals = [v]
        out.append((vals, m.group(0).strip(), m.start()))
    return out


def _display(base: float, unit: str) -> float:
    cls = _entry_class(unit)
    if cls == "money":
        return round(base / _MONEY_UNITS[unit], 4)
    return round(base, 4)


def run_cross_doc(ws: Workspace) -> Tuple[Dict[str, Any], Dict[str, Any]]:
    """账本 numbers.json（numberKeys 筛选）↔ 正文 / PPT / 讲稿 / 题库：同句出现 label 关键词时比对其中的同类数字（±1%，万 / 亿换算）。"""
    thr = {"mismatches": 0, "tolerance": 0.01}
    keys = [str(k) for k in (ws.cfg.get("numberKeys") or [])]
    sel = [e for e in ws.entries if (not keys or e.get("key") in keys) and isinstance(e.get("value"), (int, float)) and not isinstance(e.get("value"), bool)]
    empty = {"consistent": True, "checked": 0, "keys": [], "mismatches": []}
    if not sel:
        return _skip("cross_doc_consistency", "numbers.json 缺失或无可比对条目（numberKeys）", thr), empty
    groups: Dict[str, List[Tuple[str, int, str]]] = {"book": [], "deck": ws.slide_units(), "script": [], "qa": []}

    def units_of(doc: Doc) -> List[Tuple[str, int, str]]:
        out: List[Tuple[str, int, str]] = []
        for file, line, text in ws.line_units([doc]):
            hm = _HEADING_RE.match(text)
            if hm:  # 标题去掉章节号（“3.4 预算” 里的 3.4 不是金额）
                text = _strip_head_num(hm.group(2))
            if text.strip().startswith("|"):
                out.append((file, line, text))
            else:
                out += [(file, line, s) for s in _split_sentences(text)]
        return out

    for d in ws.chapters:
        groups["book"] += units_of(d)
    for g, doc in (("script", ws.script), ("qa", ws.qa)):
        if doc is not None:
            groups[g] += units_of(doc)
    if not any(groups.values()):
        return _skip("cross_doc_consistency", "没有可比对的文档（正文 / PPT / 讲稿 / 题库均缺）", thr), empty
    mismatches: List[Dict[str, Any]] = []
    checked: List[str] = []
    for e in sel:
        key, unit = str(e["key"]), str(e.get("unit") or "")
        expected = _expected_base(float(e["value"]), unit)
        tol = max(expected * 0.01, 1e-9)
        kws = _keywords_for(key, str(e.get("label") or ""))
        excl = _KEY_EXCLUDE.get(key)
        found: Dict[str, Optional[float]] = {g: None for g in groups}
        bad: Dict[str, List[Dict[str, Any]]] = {g: [] for g in groups}
        for g, units in groups.items():
            for file, line, text in units:
                n = _norm(text)
                kw = next((k for k in kws if k in n), None)
                if kw is None:
                    continue
                # 排除词（阶段 / 其中 / 机动 / 若…）：句子仍可“找到”账本数，但不据此报不一致
                excluded = bool(excl and excl.search(n)) or bool(_SCENARIO_RE.search(n))
                pos = n.find(kw)
                n2 = n.replace(kw, "□" * len(kw))
                cands = _candidates(n2, unit)
                if kw in _GENERIC_KWS:  # 泛化词只看紧邻数字；计数类单位（份 / 人 / 倍…）还要求数字自带单位
                    cands = [c for c in cands if (pos + len(kw) <= c[2] <= pos + len(kw) + _NEAR_AFTER) or (0 <= pos - c[2] - len(c[1]) <= _NEAR_BEFORE)]
                    if _entry_class(unit) == "other":
                        cands = [c for c in cands if c[1].endswith(unit)]
                if not cands:
                    continue
                ok = [c for c in cands if any(abs(v - expected) <= tol for v in c[0])]
                if ok:
                    found[g] = found[g] if found[g] is not None else _display(expected, unit)
                elif len(cands) <= 2 and not excluded:
                    after = [c for c in cands if c[2] >= pos]
                    pick = min(after or cands, key=lambda c: abs(c[2] - pos))
                    bad[g].append({"file": file, "line": line, "text": text.strip()[:120], "found": pick[1], "value": _display(pick[0][0], unit)})
        if any(found.values()) or any(bad.values()):
            checked.append(key)
        if any(bad.values()):
            item: Dict[str, Any] = {"key": key, "label": e.get("label"), "unit": unit, "expected": e["value"]}
            for g in groups:
                item[g] = bad[g][0]["value"] if bad[g] else found[g]
            item["locations"] = [{**b, "doc": g} for g in groups for b in bad[g]][:MAX_LOCATIONS]
            mismatches.append(item)
    cross = {"consistent": not mismatches, "checked": len(checked), "keys": checked, "mismatches": mismatches}
    if mismatches:
        detail = f"{len(mismatches)} 个账本数字在文档中出现不同值：" + "；".join(f"{m['key']} 账本 {_fmt_num(float(m['expected']))} vs " + "/".join(f"{g} {m[g]}" for g in groups if m[g] is not None and m[g] != _display(_expected_base(float(m['expected']), m['unit']), m['unit'])) for m in mismatches[:4])
    else:
        detail = f"比对 {len(checked)}/{len(sel)} 个账本数字，未发现不一致" + ("" if checked else "（文档中未定位到任何账本数字，无可比对）")
    locs = [loc for m in mismatches for loc in m["locations"]]
    return _mk("cross_doc_consistency", len(mismatches), thr, FAIL if mismatches else PASS, detail, locs, checked=len(checked)), cross


# ---------------------------------------------------------------------------------------------- 硬门槛
_COVER_FIELD_KEYS: Dict[str, List[str]] = {
    "作品名称": ["title", "subtitle"], "作品名": ["title", "subtitle"], "团队名称": ["teamName"], "团队名": ["teamName"],
    "指导企业名称": ["company"], "指导企业": ["company"], "案例企业名称": ["company", "subtitle", "title"], "组别": ["group"],
}


def _cover_fields_required(ws: Workspace, naming: Dict[str, Any]) -> List[str]:
    cf = ws.cfg.get("coverFields")
    if isinstance(cf, list) and cf:
        return [str(x) for x in cf]
    cf = naming.get("coverFields") if isinstance(naming, dict) else None
    if isinstance(cf, list) and cf:
        return [str(x) for x in cf]
    return ["作品名称", "团队名称"]


def gate_cover_fields(ws: Workspace, naming: Dict[str, Any]) -> Dict[str, Any]:
    """封面必填字段：按 coverFields（变体优先）→ naming.json；字段名映射到 book.json cover 键，映射不到时在封面文字里找字段名或占位符。"""
    required = _cover_fields_required(ws, naming)
    if not ws.cover:
        return {"id": "cover_fields", "passed": True, "skipped": True, "required": required, "missing": [], "detail": "无 05_正文/book.json 封面，未校验"}
    blob = "；".join(str(v) for v in _iter_strings(ws.cover) for v in [v[1]])
    missing: List[str] = []
    for f in required:
        keys = _COVER_FIELD_KEYS.get(f, [])
        ok = any(str(ws.cover.get(k) or "").strip() for k in keys)
        if not ok:
            core = re.sub(r"[【】\[\]（）()]", "", f).split("：")[-1]
            ok = f in blob or (len(core) >= 2 and core in blob)
        if not ok:
            missing.append(f)
    return {"id": "cover_fields", "passed": not missing, "required": required, "missing": missing, "detail": ("封面字段齐全" if not missing else f"封面缺：{'、'.join(missing)}")}


def _gate_from_check(gate_id: str, chk: Optional[Dict[str, Any]], **extra: Any) -> Dict[str, Any]:
    if chk is None or chk.get("status") == SKIP:
        return {"id": gate_id, "passed": True, "skipped": True, "detail": (chk or {}).get("detail", "未检查"), **extra}
    return {"id": gate_id, "passed": chk.get("status") != FAIL, "detail": chk.get("detail", ""), **extra}


def build_hard_gates(ws: Workspace, checks: Dict[str, Dict[str, Any]], cross: Dict[str, Any], gate_ids: Sequence[str], naming: Dict[str, Any]) -> List[Dict[str, Any]]:
    """按 rubric.hardGates 的名称逐个产出 {id, passed, ...}。"""
    out: List[Dict[str, Any]] = []
    for gid in gate_ids:
        if gid == "anonymity_hits":
            c = checks.get("anonymity_hits")
            out.append(_gate_from_check(gid, c, hits=(c or {}).get("locations", [])))
        elif gid == "page_limit":
            c = checks.get("page_count")
            g = _gate_from_check(gid, c, pages=(c or {}).get("pages"), limit=int(ws.cfg.get("pageLimit") or 80))
            if c and c.get("status") != SKIP:
                g["passed"] = bool((c.get("pages") or 0) <= g["limit"])
            out.append(g)
        elif gid == "word_min":
            c = checks.get("word_count")
            g = _gate_from_check(gid, c, words=(c or {}).get("value"), min=int(ws.cfg.get("wordMin") or 0))
            if c and c.get("status") != SKIP and g["min"] > 0:
                g["passed"] = bool((c.get("value") or 0) >= g["min"])
            out.append(g)
        elif gid == "cover_fields":
            out.append(gate_cover_fields(ws, naming))
        elif gid == "cross_doc_consistency":
            c = checks.get("cross_doc_consistency")
            out.append(_gate_from_check(gid, c, mismatches=[{k: v for k, v in m.items() if k != "locations"} for m in cross.get("mismatches", [])]))
        elif gid == "financing_consistency":
            out.append(_gate_from_check(gid, checks.get("financing_consistency")))
        elif gid == "forbidden_terms_adlaw":
            c = checks.get("forbidden_terms_adlaw")
            out.append(_gate_from_check(gid, c, hits=(c or {}).get("locations", [])))
        elif gid == "budget_in_range":
            out.append(_gate_from_check(gid, checks.get("budget_in_range")))
        else:
            out.append(_gate_from_check(gid, checks.get(gid)))
    return out


# ---------------------------------------------------------------------------------------------- 运行
CORE_CHECKS: List[Tuple[str, Callable[[Workspace], Any]]] = [
    ("struct_coverage", check_struct_coverage), ("page_count", check_page_count), ("word_count", check_word_count),
    ("anonymity_hits", check_anonymity), ("numeric_density", check_numeric_density), ("sources", check_sources),
    ("source_pairing", check_source_pairing), ("chart_count", check_chart_count), ("forbidden_terms", check_forbidden_terms),
    ("forbidden_terms_adlaw", check_adlaw), ("weak_terms", check_weak_terms), ("synthetic_marker_coverage", check_synthetic_markers),
    ("slide_metrics", check_slide_metrics), ("deck_font_min", check_deck_font_min), ("script_timing", check_script_timing),
    ("sentence_length", check_sentence_length), ("qa_bank", check_qa_bank),
]
# (检查 id, requiredEvidence 键, 函数)
EVIDENCE_CHECKS: List[Tuple[str, str, CheckFn]] = [
    ("has_primary_research", "has_primary_research", check_primary_research),
    ("has_tam_sam_som", "has_tam_sam_som", check_tam_sam_som),
    ("has_competitor_matrix", "has_competitor_matrix", check_competitor_matrix),
    ("has_unit_economics", "has_unit_economics", check_unit_economics),
    ("has_sensitivity", "has_sensitivity", check_sensitivity),
    ("three_statement_check", "three_statement_check", check_three_statement),
    ("financing_consistency", "financing_consistency", check_financing_consistency),
    ("revenue_shape", "revenue_shape", check_revenue_shape),
    ("budget_recalc", "budget_recalc", check_budget_recalc),
    ("budget_in_range", "budget_in_range", check_budget_in_range),
    ("kpi_table", "kpi_table", check_kpi_table),
    ("positioning_statement", "positioning_statement", check_positioning_statement),
    ("big_idea", "big_idea", check_big_idea),
    ("kol_tiers", "kol_tiers", check_kol_tiers),
    ("offline_event", "offline_event", check_offline_event),
    ("risk_matrix", "risk_matrix_min", check_risk_matrix),
    ("brief_response_table", "brief_response_table", check_brief_response),
    ("appendix_questionnaire", "appendix_questionnaire", check_appendix_questionnaire),
    ("honor_ratio", "honor_ratio_max", check_honor_ratio),
    ("title_pattern", "title_pattern", check_title_pattern),
    ("ip_list", "ip_list", check_ip_list),
    ("param_compare_table", "param_compare_table", check_param_compare),
    ("milestone_gantt", "milestone_gantt", check_milestone_gantt),
    ("equity_table", "equity_table", check_equity_table),
    ("team_matrix", "team_matrix", check_team_matrix),
    ("short_video_script", "short_video_script", check_short_video_script),
    ("platform_matrix", "platform_matrix", check_platform_matrix),
    ("ratio_table_with_formula", "ratio_table_with_formula", check_ratio_table_formula),
    ("comprehensive_income_forecast", "comprehensive_income_forecast", check_comprehensive_income),
    ("annual_report_page_refs", "annual_report_page_refs", check_annual_report_refs),
    ("contingency", "contingency_min", check_contingency),
    ("itinerary_timeline", "itinerary_timeline", check_itinerary_timeline),
    ("itinerary_grid", "itinerary_grid", check_itinerary_grid),
    ("venue_capacity_check", "venue_capacity_check", check_venue_capacity),
    ("vendor_quote_table", "vendor_quote_table", check_vendor_quote),
    ("regional_feature", "regional_feature", check_regional_feature),
    ("problem_identification", "problem_identification", check_problem_identification),
    ("industry_cluster_analysis", "industry_cluster_analysis", check_industry_cluster),
    ("analysis_framework_diagram", "analysis_framework_diagram", check_analysis_framework),
    ("novel_method", "novel_method", check_novel_method),
]
CHAPTER_CHECK_IDS = ["numeric_density", "source_pairing", "forbidden_terms", "forbidden_terms_adlaw", "anonymity_hits", "weak_terms"]
DEFAULT_HARD_GATES = ["anonymity_hits", "page_limit", "cover_fields", "cross_doc_consistency", "forbidden_terms_adlaw"]


def _safe_run(cid: str, fn: Callable[[Workspace], Any], ws: Workspace, state: Dict[str, Any]) -> List[Dict[str, Any]]:
    try:
        res = fn(ws)
        return res if isinstance(res, list) else [res]
    except Exception as e:  # 单项崩溃不影响全局
        state["degraded"] = True
        state["errors"].append(f"{cid}: {e!r}")
        return [_mk(cid, None, None, SKIP, f"内部错误（已降级）：{e!r}")]


def _enabled(ws: Workspace, ev_key: str, cid: str, gate_ids: Sequence[str]) -> bool:
    v = ws.ev.get(ev_key)
    if cid == "brief_response_table":
        return v is True or v == "if_assigned"
    if isinstance(v, str):
        return v.strip().lower() in ("true", "yes", "required")
    return bool(v) or cid in gate_ids


def _summary(checks: Sequence[Dict[str, Any]], degraded: bool) -> Dict[str, Any]:
    s = {PASS: 0, WARN: 0, FAIL: 0, SKIP: 0}
    for c in checks:
        s[c.get("status", SKIP)] = s.get(c.get("status", SKIP), 0) + 1
    return {"pass": s[PASS], "warn": s[WARN], "fail": s[FAIL], "skipped": s[SKIP], "degraded": degraded}


def _load_cfg(track: str, root: Path, variant: Optional[str]) -> Dict[str, Any]:
    cfg = td.load_track_config(track, root)
    if cfg.get("_source") == "builtin" and not cfg.get("_warning"):
        cfg["_warning"] = f"未找到赛道配置 tracks/{track}/checks.json，使用内置默认阈值"
    return td.apply_variant(cfg, variant)


def _resolve_track(track: Optional[str], root: Path) -> str:
    if track:
        return track
    proj = _read_json_safe(root / "project.json") or {}
    return str(proj.get("track") or "brand")


def _resolve_variant(variant: Optional[str], root: Path) -> Optional[str]:
    if variant:
        return variant
    proj = _read_json_safe(root / "project.json") or {}
    subj = proj.get("subject") if isinstance(proj.get("subject"), dict) else {}
    params = proj.get("params") if isinstance(proj.get("params"), dict) else {}
    v = subj.get("variant") or params.get("variant")
    return str(v) if v else None


def run_chapter(workspace: Path, track: Optional[str], chapter: str, variant: Optional[str] = None) -> Dict[str, Any]:
    """单章快检：只做 数值密度 / 来源配对 / 禁语 / 广告法 / 匿名 / 弱词，不写文件。"""
    root = Path(workspace).resolve()
    track = _resolve_track(track, root)
    cfg = _load_cfg(track, root, _resolve_variant(variant, root))
    ws = Workspace(root, cfg, chapter=chapter)
    state: Dict[str, Any] = {"degraded": False, "errors": []}
    checks: List[Dict[str, Any]] = []
    warning: Optional[str] = None
    if not ws.chapters:  # 读不到章节：六项全部 skipped，而不是报错退出
        warning = f"读不到章节文件 {chapter}"
        checks = [_skip(cid, warning) for cid in CHAPTER_CHECK_IDS]
    else:
        fns = {cid: fn for cid, fn in CORE_CHECKS}
        for cid in CHAPTER_CHECK_IDS:
            checks += _safe_run(cid, fns[cid], ws, state)
    return {"ok": True, "mode": "chapter", "track": track, "chapter": ws.chapters[0].rel if ws.chapters else chapter, "warning": warning,
            "generatedAt": datetime.now().isoformat(timespec="seconds"), "checks": checks, "summary": _summary(checks, state["degraded"]),
            "errors": state["errors"], "files": None}


def run(workspace: Path, track: Optional[str] = None, round_: Optional[int] = None, variant: Optional[str] = None, write: bool = True) -> Dict[str, Any]:
    """全量检查；write=True 时写 `07_评审/第N轮/机器检查.json|md`（无轮次写 `07_评审/即时/`）。"""
    root = Path(workspace).resolve()
    track = _resolve_track(track, root)
    variant = _resolve_variant(variant, root)
    cfg = _load_cfg(track, root, variant)
    ws = Workspace(root, cfg)
    state: Dict[str, Any] = {"degraded": False, "errors": []}
    rubric = _read_json_safe(td.find_track_file(track, root, "rubric.json") or Path("/nonexistent")) or {}
    naming = _read_json_safe(td.find_track_file(track, root, "naming.json") or Path("/nonexistent")) or {}
    gate_ids = [str(g) for g in (rubric.get("hardGates") or [])] or list(DEFAULT_HARD_GATES)

    checks: List[Dict[str, Any]] = []
    for cid, fn in CORE_CHECKS:
        checks += _safe_run(cid, fn, ws, state)
    for cid, ev_key, fn in EVIDENCE_CHECKS:
        if cid == "brief_response_table":
            if ev_key in ws.ev:
                checks += _safe_run(cid, fn, ws, state)
            continue
        if _enabled(ws, ev_key, cid, gate_ids):
            checks += _safe_run(cid, fn, ws, state)
    try:
        cross_check, cross = run_cross_doc(ws)
    except Exception as e:
        state["degraded"] = True
        state["errors"].append(f"cross_doc_consistency: {e!r}")
        cross_check, cross = _mk("cross_doc_consistency", None, None, SKIP, f"内部错误（已降级）：{e!r}"), {"consistent": True, "checked": 0, "keys": [], "mismatches": []}
    checks.append(cross_check)
    by_id = {c["id"]: c for c in checks}
    try:
        gates = build_hard_gates(ws, by_id, cross, gate_ids, naming)
    except Exception as e:
        state["degraded"] = True
        state["errors"].append(f"hardGates: {e!r}")
        gates = [{"id": g, "passed": True, "skipped": True, "detail": f"内部错误（已降级）：{e!r}"} for g in gate_ids]
    summary = _summary(checks, state["degraded"])
    summary["hardGatePassed"] = all(g.get("passed", True) for g in gates)
    result: Dict[str, Any] = {
        "track": track, "variant": variant, "round": round_, "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "workspace": str(root), "config": {"source": cfg.get("_source"), "warning": cfg.get("_warning")},
        "checks": checks, "hardGates": gates, "crossDoc": cross, "summary": summary, "verdict": _verdict(summary, gates, checks), "errors": state["errors"],
        "files": None,
    }
    if write:
        out_dir = root / "07_评审" / (f"第{round_}轮" if round_ is not None else "即时")
        jp, mp = out_dir / "机器检查.json", out_dir / "机器检查.md"
        try:
            write_json(jp, result)
            write_text(mp, render_markdown(result))
            result["files"] = {"json": str(jp), "md": str(mp)}
        except Exception as e:
            result["files"] = {"json": None, "md": None, "error": repr(e)}
    result["ok"] = True
    return result


def _verdict(summary: Dict[str, Any], gates: Sequence[Dict[str, Any]], checks: Sequence[Dict[str, Any]]) -> str:
    """一句总评。"""
    failed_gates = [g["id"] for g in gates if not g.get("passed", True)]
    fails = [c["id"] for c in checks if c.get("status") == FAIL]
    if failed_gates:
        return f"硬门槛未过（{'、'.join(failed_gates)}），不可提交；另有 {len(fails)} 项 fail、{summary['warn']} 项 warn。"
    if fails:
        return f"硬门槛全部通过；{len(fails)} 项必修（{'、'.join(fails[:6])}{'…' if len(fails) > 6 else ''}）、{summary['warn']} 项建议修。"
    if summary["warn"]:
        return f"硬门槛与必修项全部通过；{summary['warn']} 项建议修，{summary['skipped']} 项因材料缺失跳过。"
    return f"机器检查全部通过（{summary['pass']} 项），{summary['skipped']} 项因材料缺失跳过。"


# ---------------------------------------------------------------------------------------------- Markdown 报告
_STATUS_ICON = {PASS: "✅ pass", WARN: "⚠️ warn", FAIL: "❌ fail", SKIP: "⏭ skipped"}


def _fmt_thr(thr: Any) -> str:
    if thr is None:
        return "—"
    if isinstance(thr, dict):
        return "、".join(f"{k} {_fmt_val(v)}" for k, v in thr.items() if k not in ("chapters", "dimensions"))
    return _fmt_val(thr)


def _fmt_val(v: Any) -> str:
    if isinstance(v, bool):
        return "是" if v else "否"
    if isinstance(v, float):
        return _fmt_num(v)
    if isinstance(v, (list, tuple)):
        return "[" + ", ".join(_fmt_val(x) for x in v) + "]"
    if v is None:
        return "—"
    return str(v)


def _md_cell(s: Any) -> str:
    return str(s).replace("|", "／").replace("\n", " ")


def render_markdown(result: Dict[str, Any]) -> str:
    """机器检查.md：总评 + 硬门槛表 + 检查项表 + 跨材料不一致清单 + 问题定位。"""
    L: List[str] = []
    rnd = f"第 {result['round']} 轮" if result.get("round") is not None else "即时"
    variant = f" / {result['variant']}" if result.get("variant") else ""
    L.append(f"# 机器检查报告 · {result.get('track')}{variant} · {rnd}")
    L.append("")
    L.append(f"生成时间 {result.get('generatedAt')} · 工作区 `{result.get('workspace')}` · 阈值来源 `{(result.get('config') or {}).get('source')}`")
    s = result.get("summary") or {}
    L.append("")
    L.append(f"**总评**：{result.get('verdict', '')}  ")
    L.append(f"pass {s.get('pass', 0)} · warn {s.get('warn', 0)} · fail {s.get('fail', 0)} · skipped {s.get('skipped', 0)}" + ("（有检查项内部出错，已降级）" if s.get("degraded") else ""))
    L.append("")
    L.append("## 硬门槛")
    L.append("")
    L.append("| 门槛 | 结果 | 说明 |")
    L.append("|---|---|---|")
    for g in result.get("hardGates") or []:
        res = "⏭ 未检查" if g.get("skipped") else ("✅ 通过" if g.get("passed") else "❌ 未过")
        L.append(f"| {g['id']} | {res} | {_md_cell(g.get('detail', ''))} |")
    L.append("")
    L.append("## 检查项")
    L.append("")
    L.append("| 检查项 | 值 | 阈值 | 状态 | 说明 |")
    L.append("|---|---|---|---|---|")
    for c in result.get("checks") or []:
        L.append(f"| {c['id']} | {_md_cell(_fmt_val(c.get('value')))} | {_md_cell(_fmt_thr(c.get('threshold')))} | {_STATUS_ICON.get(c.get('status'), c.get('status'))} | {_md_cell(c.get('detail', ''))} |")
    L.append("")
    L.append("## 跨材料数字一致性")
    L.append("")
    cross = result.get("crossDoc") or {}
    if cross.get("mismatches"):
        L.append("| key | 账本 | 正文 | PPT | 讲稿 | 题库 | 位置 |")
        L.append("|---|---|---|---|---|---|---|")
        for m in cross["mismatches"]:
            where = "；".join(f"{x['file']}:{x['line']}「{x.get('found', '')}」" for x in (m.get("locations") or [])[:3])
            L.append(f"| {m['key']}（{_md_cell(m.get('label', ''))}） | {_fmt_val(m.get('expected'))} {m.get('unit', '')} | {_fmt_val(m.get('book'))} | {_fmt_val(m.get('deck'))} | {_fmt_val(m.get('script'))} | {_fmt_val(m.get('qa'))} | {_md_cell(where)} |")
    else:
        L.append(f"一致（比对 {cross.get('checked', 0)} 个账本数字，无不一致）。" if cross.get("checked") else "无可比对数字（numbers.json 缺失或文档中未定位到账本数字）。")
    L.append("")
    probs = [c for c in result.get("checks") or [] if c.get("status") in (FAIL, WARN) and c.get("locations")]
    if probs:
        L.append("## 问题定位")
        L.append("")
        for c in probs:
            L.append(f"- **{c['id']}**（{c['status']}）")
            for loc in c["locations"][:15]:
                term = f"「{loc['term']}」" if loc.get("term") else ""
                L.append(f"  - {loc.get('file', '')}:{loc.get('line', '')} {term}{_md_cell(loc.get('text', ''))}")
            if len(c["locations"]) > 15:
                L.append(f"  - …另有 {len(c['locations']) - 15} 处")
        L.append("")
    return "\n".join(L)


# ---------------------------------------------------------------------------------------------- CLI
def main(argv: Optional[Sequence[str]] = None) -> int:
    force_utf8_stdout()
    ap = argparse.ArgumentParser(prog="python -m bizkit.checks", description="金策 BizGold 机器检查器（契约 §7.7）")
    ap.add_argument("--track", help="赛道 id（brand / startup / digital / case / mice）；缺省读 project.json")
    ap.add_argument("--workspace", default=".", help="工作区根目录")
    ap.add_argument("--round", type=int, default=None, help="评审轮次 N → 07_评审/第N轮/；缺省写 07_评审/即时/")
    ap.add_argument("--chapter", default=None, help="单章快检（相对工作区路径），只打印 JSON 不写文件")
    ap.add_argument("--variant", default=None, help="案例赛道变体 accounting | retail")
    ap.add_argument("--no-write", action="store_true", help="不写出文件，只打印结果")
    args = ap.parse_args(argv)
    try:
        if args.chapter:
            result = run_chapter(Path(args.workspace), args.track, args.chapter, args.variant)
        else:
            result = run(Path(args.workspace), args.track, args.round, args.variant, write=not args.no_write)
            if (result.get("files") or {}).get("json"):
                print(f"[bizkit] 机器检查 → {result['files']['json']}", file=sys.stderr, flush=True)
    except Exception as e:  # 兜底：任何意外都以 JSON 报告而不是堆栈
        result = {"ok": False, "error": repr(e), "track": args.track, "workspace": args.workspace}
    print(json.dumps(result, ensure_ascii=False, default=str))
    return 0 if result.get("ok") else 1


if __name__ == "__main__":
    sys.exit(main())
