"""bizkit.finmodel（assumptions.json → 财务模型.xlsx / numbers.json / 财务摘要.md / 财务图表.charts.json）测试。

运行：cd agent/toolkit && python -m pytest -q tests/test_finmodel.py

策略：
1. 一个可手算的极简案例（单收入流、常量成本、亏损结转），逐项核对收入 / 税 / 净利 / 现金 / 盈亏平衡月份；
2. 示例工作区 examples/ws_demo2/03_财务 跑完整 startup 流程：三表勾稽 < 1 元、出让比例一致、盈亏平衡合理、
   xlsx 重新打开页签齐全且报表为公式、numbers.json key 齐全且 source 指向存在的单元格、charts.json 能被 bizkit.charts 渲染；
3. budget 模式：合计可按 Σ 单价 × 数量 × (1+机动金比例) 复算、KPI / ROI 账本条目、超 totalCap 警告；
4. 本机有 LibreOffice 时，把 xlsx 重算一遍，核对 Excel 公式结果 = Python 数值校验列、Master Check = 通过。
"""
from __future__ import annotations

import json
import math
import re
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Any, Dict

import pytest

TOOLKIT = Path(__file__).resolve().parents[1]
if str(TOOLKIT) not in sys.path:
    sys.path.insert(0, str(TOOLKIT))

from bizkit import charts, finmodel as fm  # noqa: E402
from bizkit._util import read_json, write_json  # noqa: E402

DEMO = TOOLKIT / "examples" / "ws_demo2"
FIN = DEMO / "03_财务"
PY = sys.executable
FIXED_KEYS = [f"revenue_y{i}" for i in range(1, 6)] + [f"net_profit_y{i}" for i in range(1, 6)] + [f"gross_margin_y{i}" for i in range(1, 6)] + [
    "breakeven_month", "min_cash", "financing_amount", "pre_money", "post_money", "equity_offered", "cac", "ltv", "ltv_cac", "payback_months"]


# --------------------------------------------------------------------------- fixtures
@pytest.fixture(scope="session")
def ws(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """示例工作区的干净副本（只带 project.json 与 03_财务 的两个 assumptions）。"""
    dst = tmp_path_factory.mktemp("ws_fin")
    shutil.copy(DEMO / "project.json", dst / "project.json")
    (dst / "03_财务").mkdir()
    for name in ("assumptions.json", "assumptions_budget.json"):
        shutil.copy(FIN / name, dst / "03_财务" / name)
    return dst


@pytest.fixture(scope="session")
def startup(ws: Path) -> Dict[str, Any]:
    return fm.run(ws / "03_财务" / "assumptions.json", "03_财务")


@pytest.fixture(scope="session")
def numbers(startup: Dict[str, Any]) -> Dict[str, Any]:
    return read_json(startup["files"]["numbers"])


@pytest.fixture(scope="session")
def entries(numbers: Dict[str, Any]) -> Dict[str, Dict[str, Any]]:
    return {e["key"]: e for e in numbers["entries"]}


@pytest.fixture(scope="session")
def res(ws: Path) -> fm.StartupResult:
    return fm.compute_startup(read_json(ws / "03_财务" / "assumptions.json"))


@pytest.fixture(scope="session")
def budget(ws: Path, tmp_path_factory: pytest.TempPathFactory) -> Dict[str, Any]:
    out = tmp_path_factory.mktemp("budget_out")
    return fm.run(ws / "03_财务" / "assumptions_budget.json", str(out))


def _simple(**over: Any) -> Dict[str, Any]:
    """可手算案例：每月 100 件 × 10 元，成本率 40%，固定 100 元/月，无其他成本，税 25%，期初现金 10000。"""
    a: Dict[str, Any] = {
        "mode": "startup", "company": "T", "unit": 1, "startYear": 2026, "years": 2, "monthlyFirstYear": True,
        "revenueStreams": [{"name": "A", "driver": "volume_price", "volume": {"y1Monthly": [100] * 12, "yoyGrowth": [0, 0]},
                            "price": {"base": 10}, "cogsRatio": 0.4, "source": "t"}],
        "costs": {"fixed": [{"name": "F", "monthly": 100}], "variablePctOfRevenue": [], "marketing": {"pctOfRevenue": [0, 0], "minAnnual": 0},
                  "headcount": [], "capex": [], "otherOpexPctOfRevenue": 0},
        "workingCapital": {"receivableDays": 0, "inventoryDays": 0, "payableDays": 0},
        "tax": {"incomeTaxRate": 0.25}, "financing": [], "openingCash": 10000,
        "valuation": {"wacc": 0.2, "terminalGrowth": 0.02, "psMultiple": 2, "exitPE": 10, "exitYear": 2, "targetIrr": 0.4},
        "unitEconomics": {"cac": 5, "lifetimeOrders": 2}, "sensitivity": {"drivers": ["price"], "range": [-0.1, 0.1]},
    }
    for k, v in over.items():
        a[k] = v
    return a


# --------------------------------------------------------------------------- 手算案例
def test_hand_calc_profit_tax_cash() -> None:
    r = fm.compute_startup(_simple())
    an = r.annual
    # 月收入 1000、COGS 400、固定 100 → EBIT 500/月；年 6000；税 1500；净利 4500
    assert an["revenue"] == pytest.approx([12000, 12000])
    assert an["ebit"] == pytest.approx([6000, 6000])
    assert an["tax"] == pytest.approx([1500, 1500])
    assert an["net_profit"] == pytest.approx([4500, 4500])
    assert an["cash"] == pytest.approx([14500, 19000])
    assert r.breakeven.month == 1
    assert all(v < 1e-6 for v in r.checks.values())


def test_hand_calc_working_capital() -> None:
    a = _simple(workingCapital={"receivableDays": 30, "inventoryDays": 0, "payableDays": 0})
    r = fm.compute_startup(a)
    # 月末应收 = 1000/30×30 = 1000；首年末现金 = 10000 + 4500 − 1000；第 2 年应收 = 12000/365×30
    assert r.s["ar"][11] == pytest.approx(1000)
    assert r.annual["cash"][0] == pytest.approx(13500)
    ar2 = 12000 / 365 * 30
    assert r.annual["ar"][1] == pytest.approx(ar2)
    assert r.annual["cash"][1] == pytest.approx(13500 + 4500 + (1000 - ar2))


def test_hand_calc_loss_carryforward_and_breakeven() -> None:
    a = _simple()
    a["costs"]["fixed"] = [{"name": "F", "monthly": 700}]           # 月 EBIT = 600 − 700 = −100 → 首年 −1200
    a["revenueStreams"][0]["volume"]["yoyGrowth"] = [0, 1.0]        # 第 2 年销量翻倍：收入 24000，毛利 14400，固定 8400 → EBIT 6000
    r = fm.compute_startup(a)
    assert r.annual["ebit"] == pytest.approx([-1200, 6000])
    assert r.annual["tax"] == pytest.approx([0, 0.25 * 4800])        # 亏损结转：只对累计 4800 计税
    assert r.annual["net_profit"][1] == pytest.approx(6000 - 1200)
    # 累计 −1200 → 第 2 年内按线性推算 ceil(1200/6000×12) = 3 → 第 15 个月
    assert r.breakeven.month == 15
    assert all(v < 1e-6 for v in r.checks.values())


def test_capex_depreciation_schedule() -> None:
    a = _simple()
    a["costs"]["capex"] = [{"name": "设备", "year": 1, "month": 1, "amount": 2400, "depreciationYears": 2}]
    r = fm.compute_startup(a)
    assert r.annual["capex"] == pytest.approx([2400, 0])
    assert r.annual["dep"] == pytest.approx([1200, 1200])
    assert r.annual["fa_net"] == pytest.approx([1200, 0])
    assert all(v < 1e-6 for v in r.checks.values())


def test_other_drivers_funnel_and_subscribers() -> None:
    periods = fm.build_timeline(2026, 2, True)
    funnel = fm.compute_stream({"name": "F", "driver": "funnel", "traffic": {"y1Monthly": [1000] * 12, "yoyGrowth": [0, 0]},
                                "conversion": {"base": 0.1}, "aov": {"base": 50}, "cogsRatio": 0.3}, periods, 2, 1.0, 1.0, 1.0, 1.0, 1.0)
    assert funnel.revenue[0] == pytest.approx(5000)
    assert sum(funnel.revenue[:12]) == pytest.approx(60000)
    assert funnel.revenue[12] == pytest.approx(60000)
    churn = 0.1
    sub = fm.compute_stream({"name": "S", "driver": "subscribers_arpu", "subscribers": {"y1Monthly": [10] * 12, "yoyGrowth": [0, 0], "churnMonthly": churn},
                             "arpu": {"base": 10}, "cogsRatio": 0.2}, periods, 2, 1.0, 1.0, 1.0, 1.0, 1.0)
    active, expect = 0.0, []
    for _ in range(12):
        active = active * (1 - churn) + 10
        expect.append(active * 10)
    assert sub.revenue[:12] == pytest.approx(expect)
    assert sub.extra[1][1][11] == pytest.approx(active)   # 期末活跃订阅


def test_invalid_driver_raises() -> None:
    a = _simple()
    a["revenueStreams"][0]["driver"] = "magic"
    with pytest.raises(fm.FinModelError):
        fm.compute_startup(a)


# --------------------------------------------------------------------------- 示例工作区：startup
def test_run_ok_and_files(startup: Dict[str, Any]) -> None:
    assert startup["ok"] is True and startup["mode"] == "startup"
    for k in ("xlsx", "numbers", "summary", "charts"):
        assert Path(startup["files"][k]).exists(), k
    assert set(startup["checks"]) == {"balanceSheet", "cashflow", "retained"}


def test_three_statements_tie_out(res: fm.StartupResult) -> None:
    assert res.checks["balanceSheet"] < 1.0 and res.checks["cashflow"] < 1.0 and res.checks["retained"] < 1.0
    s = res.s
    for i in range(len(res.periods)):
        assert abs(s["assets"][i] - s["liab_equity"][i]) < 1e-6
        assert abs(s["retained"][i] - sum(s["net_profit"][: i + 1])) < 1e-6
    assert len(res.periods) == 16 and res.periods[11].label == "2026-12" and res.periods[12].label == "2027"


def test_equity_offered_consistent(entries: Dict[str, Dict[str, Any]]) -> None:
    amt, pre, post, eq = (entries[k]["value"] for k in ("financing_amount", "pre_money", "post_money", "equity_offered"))
    assert post == pytest.approx(pre + amt, abs=0.1)
    assert abs(eq - amt / post) <= 0.01
    assert eq == pytest.approx(0.2, abs=0.005)


def test_breakeven_reasonable(res: fm.StartupResult, entries: Dict[str, Dict[str, Any]]) -> None:
    m = entries["breakeven_month"]["value"]
    assert isinstance(m, int) and 12 <= m <= 36
    assert m == res.breakeven.month
    cum = res.breakeven.cum_ebit
    # 第 12 月末累计仍为负，第 2 年末累计为正；首年应亏损
    assert cum[11] < 0 < cum[12]
    assert res.annual["net_profit"][0] < 0
    assert math.isfinite(res.breakeven.be_units) and 0 < res.breakeven.be_units < res.breakeven.actual_units


def test_numbers_keys_complete_and_sources(startup: Dict[str, Any], numbers: Dict[str, Any], entries: Dict[str, Dict[str, Any]]) -> None:
    from openpyxl import load_workbook

    assert numbers["unit"] == "万元"
    missing = [k for k in FIXED_KEYS if k not in entries]
    assert not missing, missing
    wb = load_workbook(startup["files"]["xlsx"])
    for k in FIXED_KEYS:
        e = entries[k]
        assert isinstance(e["value"], (int, float)), k
        assert set(e) >= {"key", "label", "value", "unit", "source"}
        m = re.fullmatch(r"03_财务/财务模型\.xlsx!(.+?)!([A-Z]+\d+)", e["source"])
        assert m, e["source"]
        assert wb[m.group(1)][m.group(2)].value is not None, k
    assert entries["revenue_y1"]["value"] == pytest.approx(304.6, abs=0.1)
    assert entries["min_cash"]["value"] < 0 and entries["min_cash"]["value"] + entries["financing_amount"]["value"] >= 0
    assert entries["ltv_cac"]["value"] == pytest.approx(entries["ltv"]["value"] / entries["cac"]["value"], rel=0.01)
    assert 0 < entries["gross_margin_y3"]["value"] < 1


def test_xlsx_sheets_formulas_styles(startup: Dict[str, Any]) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(startup["files"]["xlsx"])
    assert wb.sheetnames == fm.STARTUP_SHEETS
    is_ = wb["利润表"]
    assert str(is_["C2"].value).startswith("=收入!")                 # 跨表引用
    assert is_["O2"].value == "=SUM(C2:N2)"                          # 首年合计 = 年度 SUM
    assert is_["C2"].font.color.rgb.endswith(fm.COLOR_XREF)          # 绿字
    assert is_["C4"].value == "=C2-C3" and is_["C4"].font.color.rgb.endswith(fm.COLOR_FORMULA)
    assert is_["U2"].value == pytest.approx(304.6, abs=0.1)         # 数值校验列
    assert is_.freeze_panes == "B2"
    for name in ("现金流量表", "资产负债表", "检查"):
        ws = wb[name]
        formulas = [c for row in ws.iter_rows(min_row=2, max_col=19) for c in row if isinstance(c.value, str) and c.value.startswith("=")]
        assert len(formulas) > 30, name
    bs = wb["资产负债表"]
    assert bs["B4"].value == pytest.approx(50.0)                    # 期初现金 50 万元
    assert bs["B4"].fill.fgColor.rgb.endswith(fm.FILL_INPUT) and bs["B4"].font.color.rgb.endswith(fm.COLOR_INPUT)
    asm = wb["假设"]
    inputs = [c for row in asm.iter_rows(min_row=4, min_col=3, max_col=3) for c in row if isinstance(c.value, (int, float))]
    assert inputs and all(c.font.color.rgb.endswith(fm.COLOR_INPUT) for c in inputs)
    assert len(wb["利润表"]._charts) == 1 and len(wb["现金流量表"]._charts) == 1
    assert wb["检查"]["B13"].value.startswith("=IF(")
    assert wb["说明"]["B4"].value.startswith("=检查!")
    assert "#,##0" in is_["C2"].number_format and "%" in is_["C5"].number_format


def test_summary_sections(startup: Dict[str, Any]) -> None:
    text = Path(startup["files"]["summary"]).read_text(encoding="utf-8")
    for sec in ("五年预测", "盈亏平衡", "单位经济", "三情景", "敏感性排序", "估值区间", "融资三件套", "检查结果", "逻辑解释"):
        assert sec in text, sec
    assert "| 2026 |" in text and "出让" in text


def test_charts_json_renders(startup: Dict[str, Any], tmp_path: Path) -> None:
    cfg = read_json(startup["files"]["charts"])
    ids = [c["id"] for c in cfg["charts"]]
    assert ids == ["fig_fin_revenue_mix", "fig_fin_revenue_profit", "fig_fin_cashflow", "fig_fin_breakeven", "fig_fin_tornado", "fig_fin_use_of_funds"]
    assert [c["type"] for c in cfg["charts"]] == ["stacked_bar", "line", "cashflow_combo", "breakeven", "tornado", "pie_as_bar"]
    assert all(c.get("source") for c in cfg["charts"])
    cfg["outputDir"] = str(tmp_path / "04_图表")
    jp = tmp_path / "charts.json"
    write_json(jp, cfg)
    r = charts.render_charts(cfg, jp)
    assert r["ok"], r["failed"]
    assert r["count"] == 6 and all((tmp_path / "04_图表" / f"{i}.png").exists() for i in ids)


def test_scenarios_and_sensitivity(ws: Path, res: fm.StartupResult) -> None:
    a = read_json(ws / "03_财务" / "assumptions.json")
    sc = {s.key: s for s in fm.compute_scenarios(a)}
    assert set(sc) == {"pessimistic", "base", "optimistic"}
    assert sc["pessimistic"].revenue[0] < sc["base"].revenue[0] < sc["optimistic"].revenue[0]
    assert sc["base"].breakeven_month == res.breakeven.month
    se = fm.compute_sensitivity(a, res)
    assert se.target_year == 3 and len(se.cells) == len(se.drivers) * len(se.deltas) == 20
    swings = [s for _, s in se.ranking]
    assert swings == sorted(swings, reverse=True)
    assert se.cells[("price", 0.3)].net_profit_y > se.base_net_profit > se.cells[("price", -0.3)].net_profit_y
    assert se.cells[("cogsRatio", 0.3)].net_profit_y < se.base_net_profit


def test_cli_last_line_json(ws: Path, tmp_path: Path) -> None:
    out = tmp_path / "cli_out"
    proc = subprocess.run([PY, "-m", "bizkit.finmodel", str(ws / "03_财务" / "assumptions.json"), "--out", str(out)],
                          cwd=TOOLKIT, capture_output=True, text=True, encoding="utf-8", env={"PYTHONUTF8": "1", "PATH": "/usr/bin:/bin"})
    assert proc.returncode == 0, proc.stderr
    result = json.loads(proc.stdout.strip().splitlines()[-1])
    assert result["ok"] and result["mode"] == "startup"
    assert set(result) >= {"ok", "mode", "files", "checks", "warnings"}
    assert (out / "财务模型.xlsx").exists() and (out / "numbers.json").exists()


@pytest.mark.parametrize("mode", ["startup", "budget"])
def test_cli_example_templates_run(mode: str, tmp_path: Path) -> None:
    """--example 打印的模板必须是完整可运行的 assumptions（建模师照抄改数即可，不必读源码）"""
    proc = subprocess.run([PY, "-m", "bizkit.finmodel", "--example", mode], cwd=TOOLKIT, capture_output=True, text=True,
                          encoding="utf-8", env={"PYTHONUTF8": "1", "PATH": "/usr/bin:/bin"})
    assert proc.returncode == 0, proc.stderr
    tpl = json.loads(proc.stdout)
    assert tpl["mode"] == mode
    p = tmp_path / "assumptions.json"
    write_json(p, tpl)
    r = fm.run(p, str(tmp_path / "out"))
    assert r["ok"] and r["mode"] == mode and not r["warnings"], r


# --------------------------------------------------------------------------- budget
def test_budget_totals_recompute(ws: Path, budget: Dict[str, Any]) -> None:
    a = read_json(ws / "03_财务" / "assumptions_budget.json")
    subtotal = sum(it["unitPrice"] * it["quantity"] for ph in a["phases"] for it in ph["items"])
    total = subtotal * (1 + a["contingencyPct"])
    assert budget["ok"] and budget["mode"] == "budget" and budget["checks"]["withinCap"] is True
    ent = {e["key"]: e for e in read_json(budget["files"]["numbers"])["entries"]}
    assert ent["budget_total"]["value"] == pytest.approx(total / 10000, abs=0.051)
    assert ent["budget_contingency"]["value"] == pytest.approx(subtotal * a["contingencyPct"] / 10000, abs=0.051)
    assert ent["budget_per_person"]["value"] == pytest.approx(total / a["headcount"] / 10000, abs=0.01)
    assert total <= a["totalCap"]
    kpis = [k for k in ent if k.startswith("kpi_")]
    assert len(kpis) == len(a["kpis"]) and ent["kpi_1"]["label"] == a["kpis"][0]["metric"] and ent["kpi_1"]["value"] == 5000
    assert ent["kpi_2"]["value"] == pytest.approx(1.2) and ent["kpi_2"]["unit"] == "%"
    assert ent["roi"]["value"] == pytest.approx((6800000 * 0.55 - total) / total, abs=1e-3)
    assert not budget["warnings"]


def test_budget_xlsx_and_charts(budget: Dict[str, Any], tmp_path: Path) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(budget["files"]["xlsx"])
    assert wb.sheetnames == fm.BUDGET_SHEETS + ["ROI"]
    ws = wb["预算总表"]
    item_formulas = [c.value for row in ws.iter_rows(min_row=4, min_col=6, max_col=6) for c in row if isinstance(c.value, str) and re.fullmatch(r"=C\d+\*D\d+", c.value)]
    assert len(item_formulas) == 12
    labels = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    assert any(l and "预算总计" in str(l) for l in labels) and any(l and "机动金" in str(l) for l in labels) and any(l and "预算上限" in str(l) for l in labels)
    assert wb["排期"].max_row >= 6 and wb["KPI"].max_row >= 7
    text = Path(budget["files"]["summary"]).read_text(encoding="utf-8")
    assert "预算总览" in text and "KPI" in text and "ROI" in text
    cfg = read_json(budget["files"]["charts"])
    assert [c["type"] for c in cfg["charts"]] == ["pie_as_bar", "stacked_bar"]
    cfg["outputDir"] = str(tmp_path / "04_图表")
    jp = tmp_path / "charts.json"
    write_json(jp, cfg)
    r = charts.render_charts(cfg, jp)
    assert r["ok"] and r["count"] == 2, r["failed"]


def test_budget_over_cap_warning(ws: Path, tmp_path: Path) -> None:
    a = read_json(ws / "03_财务" / "assumptions_budget.json")
    a["totalCap"] = 2000000
    p = tmp_path / "assumptions_budget.json"
    write_json(p, a)
    r = fm.run(p, str(tmp_path / "out"))
    assert r["ok"] and r["checks"]["withinCap"] is False
    assert any("超过上限" in w for w in r["warnings"])
    assert "超出上限" in Path(r["files"]["summary"]).read_text(encoding="utf-8")


# --------------------------------------------------------------------------- LibreOffice 重算（可选）
@pytest.mark.skipif(shutil.which("soffice") is None, reason="本机没有 LibreOffice，跳过公式重算核对")
def test_formulas_recalculate_to_python_values(startup: Dict[str, Any], tmp_path: Path) -> None:
    from openpyxl import load_workbook

    src = Path(startup["files"]["xlsx"])
    proc = subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path), str(src)], capture_output=True, text=True, timeout=180)
    out = tmp_path / src.name
    if not out.exists():
        pytest.skip(f"LibreOffice 转换失败：{proc.stderr[:200]}")
    wb = load_workbook(out, data_only=True)
    lay_check_first, lay_year_first = 21, 15        # 数值校验列 U…、年度列 O…（12 月 + 首年合计 + 4 年）
    compared = mism = 0
    for name in ("利润表", "现金流量表", "资产负债表"):
        ws = wb[name]
        for r in range(2, ws.max_row + 1):
            for y in range(5):
                py = ws.cell(row=r, column=lay_check_first + y).value
                xl = ws.cell(row=r, column=lay_year_first + y).value
                if isinstance(py, (int, float)):
                    compared += 1
                    if not isinstance(xl, (int, float)) or abs(xl - py) > 1e-6:
                        mism += 1
    assert compared > 100 and mism == 0
    assert wb["检查"]["B13"].value == "通过" and wb["说明"]["B4"].value == "通过"
    assert wb["融资"]["F4"].value == pytest.approx(0.2)
