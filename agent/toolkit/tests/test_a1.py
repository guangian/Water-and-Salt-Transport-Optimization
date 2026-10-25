"""A1 工具链（charts / docx_render / preread / pdf / package）端到端测试。

运行：cd agent/toolkit && python -m pytest -q tests/test_a1.py

策略：把 examples/ws_demo 复制到临时目录（不带生成物），按真实流程依次跑
charts → docx_render → preread → package，再用 python-docx / pypdf / zipfile 重新打开产物做结构性断言。
PDF 相关断言只在机器上有可用转换器（Word / WPS / docx2pdf / LibreOffice）时执行。
"""
from __future__ import annotations

import json
import shutil
import sys
import zipfile
from pathlib import Path
from typing import Any, Dict, List

import pytest

TOOLKIT = Path(__file__).resolve().parents[1]
if str(TOOLKIT) not in sys.path:
    sys.path.insert(0, str(TOOLKIT))
if str(TOOLKIT / "examples") not in sys.path:
    sys.path.insert(0, str(TOOLKIT / "examples"))

from bizkit import charts, docx_render, package as package_mod, pdf as pdf_mod, preread  # noqa: E402
from bizkit._util import read_json, write_json, write_text  # noqa: E402

DEMO = TOOLKIT / "examples" / "ws_demo"
ALL_TYPES = {
    "bar", "grouped_bar", "stacked_bar", "hbar", "line", "area", "pie_as_bar", "donut", "waterfall", "funnel", "tornado",
    "breakeven", "cashflow_combo", "scatter_matrix", "positioning_map", "gantt", "radar", "bmc",
}
HAS_PDF = bool(pdf_mod.detect())


# --------------------------------------------------------------------------- fixtures
@pytest.fixture(scope="session")
def ws(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """演示工作区的干净副本：不带已生成的 PNG / docx / 提交包，让测试真正走一遍生成流程。"""
    dst = tmp_path_factory.mktemp("ws_demo")
    ignore = shutil.ignore_patterns("*.png", "index.json", "产出", "08_提交包", "_preread", "__pycache__")
    shutil.copytree(DEMO, dst, dirs_exist_ok=True, ignore=ignore)
    from make_demo_inputs import ensure_inputs  # 示例 inputs 由脚本生成，缺了就补

    ensure_inputs(dst / "inputs")
    return dst


@pytest.fixture(scope="session")
def charts_result(ws: Path) -> Dict[str, Any]:
    cfg_path = ws / "04_图表" / "charts.json"
    return charts.render_charts(read_json(cfg_path), cfg_path)


@pytest.fixture(scope="session")
def book_result(ws: Path, charts_result: Dict[str, Any]) -> Dict[str, Any]:
    assert charts_result["ok"], charts_result["failed"]
    return docx_render.render_book(ws / "05_正文" / "book.json", make_pdf=HAS_PDF)


@pytest.fixture(scope="session")
def doc(book_result: Dict[str, Any]):
    from docx import Document

    return Document(book_result["output"])


def _paras(doc, style_prefix: str) -> List[str]:
    return [p.text for p in doc.paragraphs if p.style.name.startswith(style_prefix)]


def _rfonts(style) -> Dict[str, str]:
    from docx.oxml.ns import qn

    rpr = style.element.rPr
    rf = rpr.find(qn("w:rFonts")) if rpr is not None else None
    assert rf is not None, f"样式 {style.name} 没有 rFonts"
    return {k.split('}')[-1]: v for k, v in rf.attrib.items()}


# --------------------------------------------------------------------------- charts
def test_charts_cover_all_types_and_render(ws: Path, charts_result: Dict[str, Any]) -> None:
    cfg = read_json(ws / "04_图表" / "charts.json")
    assert {c["type"] for c in cfg["charts"]} == ALL_TYPES, "示例 charts.json 必须覆盖全部 18 种图型"
    assert charts_result["ok"], charts_result["failed"]
    assert charts_result["count"] == len(cfg["charts"])
    out_dir = Path(charts_result["outputDir"])
    assert out_dir == (ws / "04_图表").resolve()
    for c in cfg["charts"]:
        assert (out_dir / f"{c['id']}.png").stat().st_size > 5_000, c["id"]


def test_charts_png_size_matches_mm_at_300dpi(charts_result: Dict[str, Any]) -> None:
    from PIL import Image

    out_dir = Path(charts_result["outputDir"])
    for info in charts_result["rendered"]:
        with Image.open(out_dir / info["file"]) as im:
            w, h = im.size
        exp_w = info["widthMm"] / 25.4 * 300
        exp_h = info["heightMm"] / 25.4 * 300
        assert abs(w - exp_w) <= 3 and abs(h - exp_h) <= 3, (info["id"], (w, h), (exp_w, exp_h))
        assert im.info.get("dpi", (300, 300))[0] >= 299, info["id"]


def test_charts_index_json(ws: Path, charts_result: Dict[str, Any]) -> None:
    index = read_json(charts_result["index"])
    cfg = read_json(ws / "04_图表" / "charts.json")
    assert set(index) == {c["id"] for c in cfg["charts"]}
    for cid, meta in index.items():
        assert meta["file"] == f"{cid}.png"
        assert meta["title"] and "px" in meta and "source" in meta
    # 超过 4 类的 donut 退化为条形图并给出警告
    assert any("fig_motivation" in w and "4" in w for w in charts_result["warnings"]), charts_result["warnings"]


def test_charts_bad_specs_are_isolated(tmp_path: Path) -> None:
    cfg = {
        "outputDir": str(tmp_path / "out"),
        "charts": [
            {"id": "bad id!", "type": "bar", "title": "x", "data": {"categories": ["a"], "values": [1]}},
            {"id": "no_values", "type": "line", "title": "x", "data": {"categories": ["a", "b"]}},
            {"id": "unknown_type", "type": "sunburst", "title": "x", "data": {}},
            {"id": "ok_bar", "type": "bar", "title": "正常图", "source": "测试", "data": {"categories": ["甲", "乙"], "values": [1200, 3400], "unit": "万元"}},
        ],
    }
    path = tmp_path / "charts.json"
    write_json(path, cfg)
    res = charts.render_charts(cfg, path)
    assert not res["ok"] and res["count"] == 1
    assert {f["id"] for f in res["failed"]} == {"bad id!", "no_values", "unknown_type"}
    for f in res["failed"]:
        assert any("\u4e00" <= ch <= "\u9fff" for ch in f["error"]), f"错误信息应为中文：{f}"
    assert (tmp_path / "out" / "ok_bar.png").exists()
    assert set(read_json(tmp_path / "out" / "index.json")) == {"ok_bar"}


def test_charts_cli_only_filter(ws: Path, charts_result: Dict[str, Any], capsys: pytest.CaptureFixture[str]) -> None:
    code = charts.main([str(ws / "04_图表" / "charts.json"), "--only", "fig_budget"])
    out = capsys.readouterr().out
    assert code == 0
    payload = json.loads(out[out.index("{"):])
    assert payload["count"] == 1 and payload["rendered"][0]["id"] == "fig_budget"
    assert len(read_json(payload["index"])) >= 19, "--only 不应清空 index.json 里其他图的记录"


# --------------------------------------------------------------------------- docx
def test_docx_result_summary(book_result: Dict[str, Any]) -> None:
    assert book_result["ok"]
    assert Path(book_result["output"]).suffix == ".docx" and Path(book_result["output"]).exists()
    assert book_result["chapters"] == 3
    assert book_result["figures"] >= 8 and book_result["tables"] >= 5
    assert book_result["words"] > 3000
    assert book_result["pagesEstimate"] > 5
    assert book_result["missingImages"] == []


def test_docx_structure(doc, book_result: Dict[str, Any]) -> None:
    assert len(doc.inline_shapes) == book_result["figures"]
    assert len(doc.tables) == book_result["tables"]
    h1 = _paras(doc, "Heading 1")
    h2 = _paras(doc, "Heading 2")
    assert "1  引言" in h1 and "2  品牌现状分析" in h1 and "3  品牌定位与传播策略" in h1
    assert "附录 A  消费者调研问卷" in h1 and "参考文献" in h1
    assert "1.1  策划背景" in h2
    captions = _paras(doc, "Caption")
    assert any(c.startswith("图 1-1  ") for c in captions)
    assert any(c.startswith("表 1-1  ") for c in captions)
    assert any(c.startswith("图 A-1  ") for c in captions), captions
    assert len(doc.sections) == 3, "封面 / 前置（罗马页码）/ 正文 三个节"
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "概要" in body and "宋韵落天青" in body and "目  录" in body
    assert "某某大学" not in body and "XX" in body, "anonymize 应把禁用词替换为 XX"
    assert "河南财经" in body, "白名单词不应被替换"
    hits = book_result["anonymizeHits"]
    assert any(h["term"] == "某某大学" and h["count"] >= 1 and h["line"] for h in hits), hits


def test_docx_fonts_and_layout(doc) -> None:
    from docx.shared import Cm

    normal = _rfonts(doc.styles["Normal"])
    assert normal["eastAsia"] == "宋体" and normal["ascii"] == "Times New Roman" and "asciiTheme" not in normal
    assert abs(doc.styles["Normal"].font.size.pt - 12) < 0.01
    h1 = _rfonts(doc.styles["Heading 1"])
    assert h1["eastAsia"] == "黑体"
    body_style = doc.styles["BizBody"]
    xml = body_style.element.xml
    assert 'w:firstLineChars="200"' in xml, "首行缩进 2 字符应使用 firstLineChars"
    assert body_style.base_style.name == "Normal" and doc.styles["Normal"].paragraph_format.line_spacing == 1.5
    sec = doc.sections[-1]
    assert abs(sec.left_margin - Cm(3.17)) < Cm(0.01) and abs(sec.top_margin - Cm(2.54)) < Cm(0.01)
    assert abs(sec.page_width - Cm(21.0)) < Cm(0.01)
    for shape in doc.inline_shapes:
        assert shape.width <= Cm(15.5) + 1


def test_docx_fields_header_footer(doc) -> None:
    body_xml = doc.element.body.xml
    assert 'TOC \\o "1-2"' in body_xml, "需要可更新的目录域"
    assert "PAGE" in body_xml or any("PAGE" in s.footer._element.xml for s in doc.sections)
    assert "w:updateFields" in doc.settings.element.xml
    assert any("弘宝汝瓷品牌策划书" in p.text for p in doc.sections[-1].header.paragraphs)
    cover_texts = "\n".join(p.text for p in doc.paragraphs[:40])
    assert "学校" not in cover_texts and "大学" not in cover_texts, "封面不得出现学校信息"
    assert "指导企业" in cover_texts and "青瓷工作室" in cover_texts


def test_docx_tables_have_header_repeat_and_banding(doc) -> None:
    t = doc.tables[0]
    xml = t._tbl.xml
    assert "w:tblHeader" in xml and "w:cantSplit" in xml
    assert "w:shd" in xml, "banded 表格应有行底纹"
    for row in t.rows:
        assert all(c.text.strip() for c in row.cells[:1]), "首列不应为空"


def test_docx_missing_image_placeholder(ws: Path, tmp_path: Path) -> None:
    md = tmp_path / "c.md"
    write_text(md, "# 测试章\n\n正文一段。\n\n![不存在的图](../不存在/none.png)\n")
    book = {
        "cover": {"title": "T", "teamName": "队"},
        "chapters": [{"file": "c.md"}],
        "style": {"toc": False},
        "output": "out/t.docx",
    }
    write_json(tmp_path / "book.json", book)
    write_json(tmp_path / "project.json", {"name": "t"})
    res = docx_render.render_book(tmp_path / "book.json", make_pdf=False)
    assert res["ok"] and res["figures"] == 0
    assert len(res["missingImages"]) == 1 and any("找不到图片" in w for w in res["warnings"])
    assert res["pagesEstimate"] < 1, "缺图占位不应按整图计入页数估算"
    from docx import Document

    text = "\n".join(p.text for p in Document(res["output"]).paragraphs)
    assert "【缺图：" in text


def test_markdown_parser_subset() -> None:
    text = (
        "# 1 标题一\n\n## 1.1 子标题\n\n正文 **粗体** 与 *斜体*，引用[1]。{资料来源：测试}\n\n"
        "- 项目甲\n- 项目乙\n\n1. 第一\n2. 第二\n\n> 引言块\n\n"
        "表：示例表\n\n| A | B |\n|:-:|--|\n| 1 | 2 |\n| 3 | 4 |\n\n"
        "![图题](x.png)\n\n---\n\n<!-- pagebreak -->\n\n[^1]: 脚注内容\n"
    )
    parsed = docx_render.parse_markdown(text)
    kinds = [type(b).__name__ for b in parsed.blocks]
    assert kinds.count("Heading") == 2 and "Table" in kinds and "Figure" in kinds and "Quote" in kinds
    assert "ListBlock" in kinds and "PageBreak" in kinds
    tbl = next(b for b in parsed.blocks if type(b).__name__ == "Table")
    assert tbl.header == ["A", "B"] and tbl.rows == [["1", "2"], ["3", "4"]] and tbl.caption == "示例表"
    heads = [b for b in parsed.blocks if type(b).__name__ == "Heading"]
    # 作者手写的编号在渲染时剥掉，由渲染器统一重新编号
    assert heads[0].level == 1 and docx_render._strip_numbering(heads[0].text) == "标题一"
    assert heads[1].level == 2 and docx_render._strip_numbering(heads[1].text) == "子标题"
    assert docx_render._strip_numbering("第一章 引言") == "引言" and docx_render._strip_numbering("附录 A：问卷") == "问卷"
    assert docx_render._strip_numbering("2026 年预算") == "2026 年预算", "年份不能被当作编号剥掉"
    assert parsed.footnotes.get("1") == "脚注内容"


@pytest.mark.parametrize(
    "ref,marker,expect",
    [
        ({"type": "report", "author": "艾媒咨询", "title": "报告", "year": 2025, "url": "https://x", "accessed": "2026-01-01"}, "[R/OL]", "(2025)[2026-01-01]. https://x."),
        ({"type": "web", "author": "某站", "title": "页面", "year": 2025, "url": "https://x", "accessed": "2026-01-01"}, "[EB/OL]", "https://x."),
        ({"type": "paper", "author": "张三, 李四", "title": "论文", "journal": "期刊", "year": 2024, "volume": "3", "issue": "2", "pages": "1-9"}, "[J]", "期刊, 2024, 3(2): 1-9."),
        ({"type": "book", "author": "王五", "title": "书名", "publisher": "出版社", "place": "北京", "year": 2020}, "[M]", "北京: 出版社, 2020."),
        ({"type": "news", "author": "记者", "title": "新闻", "newspaper": "日报", "date": "2025-01-02", "edition": "5"}, "[N]", "日报, 2025-01-02(5)."),
        ({"type": "standard", "title": "标准名", "number": "GB/T 1-2015", "publisher": "标准出版社", "year": 2015}, "[S]", "标准名: GB/T 1-2015[S]"),
        ({"type": "thesis", "author": "赵六", "title": "学位论文", "publisher": "某校", "year": 2021}, "[D]", "某校, 2021."),
        ({"type": "gov", "author": "汝州市人民政府", "title": "方案", "number": "汝政〔2024〕12 号", "year": 2024, "url": "https://g"}, "[A/OL]", "方案: 汝政〔2024〕12 号[A/OL]"),
        ({"type": "annual-report", "author": "某公司", "title": "2025 年年度报告", "year": 2026}, "[R]", "2026."),
    ],
)
def test_format_reference_markers(ref: Dict[str, Any], marker: str, expect: str) -> None:
    s = docx_render.format_reference(ref)
    assert marker in s and ref["title"] in s and expect in s, s
    assert s.startswith(str(ref.get("author") or ref.get("publisher") or "佚名")) and s.endswith(".")
    assert s.count(str(ref.get("publisher") or "\0")) <= 1, "无作者时出版者顶替作者，不应再重复出现"


def test_references_in_docx(doc) -> None:
    refs = [p.text for p in doc.paragraphs if p.style.name == "BizReference"]
    assert len(refs) >= 7
    assert refs[0].startswith("[1]") and any("[J]" in r for r in refs) and any("[M]" in r for r in refs)


@pytest.mark.skipif(not HAS_PDF, reason="本机没有 docx→pdf 转换器")
def test_pdf_conversion_and_toc_backfill(book_result: Dict[str, Any], doc) -> None:
    assert book_result["pdf"] and Path(book_result["pdf"]).exists()
    assert book_result["pages"] and book_result["pages"] > 10
    assert book_result["tocPageNumbers"] == "pdf", book_result["warnings"]
    toc_lines = [p.text for p in doc.paragraphs if p.style.name.startswith("toc ")]
    assert toc_lines, "静态目录文字应存在"
    assert any(line.rstrip().endswith("1") and "引言" in line for line in toc_lines), toc_lines
    assert pdf_mod.page_count(Path(book_result["pdf"])) == book_result["pages"]


# --------------------------------------------------------------------------- preread
def test_preread_pdf_xlsx_docx(ws: Path, book_result: Dict[str, Any], tmp_path: Path) -> None:
    out = tmp_path / "pre"
    r_pdf = preread.preread(ws / "inputs" / "示例企业资料.pdf", out)
    assert r_pdf["ok"] and r_pdf["kind"] == "pdf" and r_pdf["pages"] >= 1 and r_pdf["words"] > 50
    assert Path(r_pdf["out"]).name == "示例企业资料.pdf.txt"
    r_x = preread.preread(ws / "inputs" / "示例销售数据.xlsx", out)
    assert r_x["ok"] and r_x["kind"] == "xlsx" and r_x["tables"] >= 1 and r_x["sheets"]
    txt = Path(r_x["out"]).read_text(encoding="utf-8")
    assert "2025-01" in txt and "景区门店" in txt
    r_d = preread.preread(Path(book_result["output"]), out)
    assert r_d["ok"] and r_d["kind"] == "docx" and r_d["tables"] == book_result["tables"] and r_d["words"] > 1000
    r_md = preread.preread(ws / "05_正文" / "01_引言.md", out)
    assert r_md["ok"] and r_md["kind"] in ("text", "md") and r_md["words"] > 100
    assert not preread.preread(ws / "没有.pdf", out)["ok"]


# --------------------------------------------------------------------------- package
def test_package_zip(ws: Path, book_result: Dict[str, Any], tmp_path: Path) -> None:
    res = package_mod.package(ws, "brand", out_dir=tmp_path / "pkg")
    assert res["ok"] and res["zip"] and Path(res["zip"]).exists()
    base = res["baseName"]
    # 品牌策划赛道命名规则（tracks/brand/naming.json）：{学校名称}-{团队名称}-{作品名称}，学校名由团队提交时填写
    assert base.startswith("学校名称-青瓷工作室-") and "学校名称" in res["placeholders"]
    names = set(zipfile.ZipFile(res["zip"]).namelist())
    assert f"{base}.docx" in names and "README.md" in names and "数据真实性清单.md" in names
    assert "声明与授权/企业授权书.docx" in names
    assert any(n.startswith("调研工具包/") for n in names) and any(n.startswith("评审报告/") for n in names)
    if HAS_PDF:
        assert f"{base}.pdf" in names and res["pdfPages"] == book_result["pages"]
    assert res["syntheticCount"] >= 2
    assert "03_财务/财务模型.xlsx" in res["missing"]
    readme = (tmp_path / "pkg" / "README.md").read_text(encoding="utf-8")
    assert "学校名称" in readme and "0 KB" not in readme
    truth = (tmp_path / "pkg" / "数据真实性清单.md").read_text(encoding="utf-8")
    assert "示例数据" in truth
