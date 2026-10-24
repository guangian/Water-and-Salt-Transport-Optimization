"""A2 工具 pptx_render（slides.json → 16:9 路演 PPT）端到端测试。

运行：cd agent/toolkit && python -m pytest -q tests/test_pptx.py

策略：把 examples/ws_demo2 复制到临时目录（不带 产出/），缺图时先跑 charts 生成 PNG，再渲染 slides.json；
然后用 python-pptx 重新打开产物做结构性断言（页数 / 备注 / 字号下限 / 原生图表与表格 / 预览 JSON / 匿名替换），
最后用一份人造的「全部超限」slides.json 验证校验 → 警告 → 截断 / 缩字号的兜底逻辑。
"""
from __future__ import annotations

import json
import re
import shutil
import sys
import zipfile
from pathlib import Path
from typing import Any, Dict, Iterable, List

import pytest

TOOLKIT = Path(__file__).resolve().parents[1]
if str(TOOLKIT) not in sys.path:
    sys.path.insert(0, str(TOOLKIT))

from bizkit import charts, pptx_render  # noqa: E402
from bizkit._util import read_json, write_json  # noqa: E402

DEMO = TOOLKIT / "examples" / "ws_demo2"
MIN_FONT_PT = 14
NAVY = "14305C"


# --------------------------------------------------------------------------- fixtures
@pytest.fixture(scope="session")
def ws(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """示例工作区的干净副本；图表 PNG 不在时先生成，保证测试不依赖已提交的生成物。"""
    dst = tmp_path_factory.mktemp("ws_demo2")
    shutil.copytree(DEMO, dst, dirs_exist_ok=True, ignore=shutil.ignore_patterns("产出", "__pycache__"))
    if not any((dst / "04_图表").glob("*.png")):
        cfg_path = dst / "04_图表" / "charts.json"
        res = charts.render_charts(read_json(cfg_path), cfg_path)
        assert res["ok"], res["failed"]
    return dst


@pytest.fixture(scope="session")
def slides_cfg(ws: Path) -> Dict[str, Any]:
    return read_json(ws / "06_路演" / "slides.json")


@pytest.fixture(scope="session")
def result(ws: Path) -> Dict[str, Any]:
    return pptx_render.render_deck(ws / "06_路演" / "slides.json")


@pytest.fixture(scope="session")
def prs(result: Dict[str, Any]):
    from pptx import Presentation

    return Presentation(result["output"])


@pytest.fixture(scope="session")
def preview(result: Dict[str, Any]) -> List[Dict[str, Any]]:
    return read_json(result["preview"])


# --------------------------------------------------------------------------- helpers
def _slide_text(slide) -> str:
    parts = []
    for shp in slide.shapes:
        if shp.has_text_frame:
            parts.append(shp.text_frame.text)
        if getattr(shp, "has_table", False) and shp.has_table:
            for row in shp.table.rows:
                parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def _font_sizes_pt(slide) -> List[float]:
    """页面上所有显式写出的字号（文本框 / 表格 / 图表 XML 里的 sz 属性，单位百分之一磅）。"""
    sizes: List[float] = []
    xml_blobs = [slide._element.xml]
    for shp in slide.shapes:
        if getattr(shp, "has_chart", False) and shp.has_chart:
            xml_blobs.append(shp.chart._chartSpace.xml)
    for xml in xml_blobs:
        sizes.extend(int(m) / 100 for m in re.findall(r'\bsz="(\d+)"', xml))
    return sizes


def _charts(prs) -> Iterable:
    for slide in prs.slides:
        for shp in slide.shapes:
            if getattr(shp, "has_chart", False) and shp.has_chart:
                yield slide, shp.chart


def _tables(prs) -> Iterable:
    for slide in prs.slides:
        for shp in slide.shapes:
            if getattr(shp, "has_table", False) and shp.has_table:
                yield slide, shp.table


# --------------------------------------------------------------------------- 示例覆盖面
def test_example_covers_all_layouts_and_chart_types(slides_cfg: Dict[str, Any]) -> None:
    layouts = {s["layout"] for s in slides_cfg["slides"]}
    assert layouts == set(pptx_render.LAYOUTS), f"示例必须覆盖全部 14 种页型，缺：{set(pptx_render.LAYOUTS) - layouts}"
    chart_types = {s["chart"]["type"] for s in slides_cfg["slides"] if s["layout"] == "chart"}
    assert chart_types == {"bar", "column_clustered", "line", "pie", "stacked"}
    assert 18 <= len(slides_cfg["slides"]) <= 22
    appendix = [s["layout"] for s in slides_cfg["slides"] if s["layout"] == "appendix"]
    assert appendix and slides_cfg["slides"][-1]["layout"] == "appendix", "备用页应统一放最后"
    assert sum(s.get("duration", 0) for s in slides_cfg["slides"]) <= slides_cfg["minutes"] * 60


# --------------------------------------------------------------------------- 结果 JSON
def test_result_summary(ws: Path, result: Dict[str, Any], slides_cfg: Dict[str, Any]) -> None:
    assert result["ok"]
    out = Path(result["output"])
    assert out == (ws / "产出" / "路演.pptx").resolve() and out.stat().st_size > 50_000
    assert result["preview"] == str(out) + ".preview.json" and Path(result["preview"]).exists()
    assert result["pages"] == len(slides_cfg["slides"]) == 20
    assert result["warnings"] == [], f"示例应零警告：{result['warnings']}"
    assert result["missingImages"] == []
    assert result["durationSeconds"] == 540
    per = result["perPage"]
    assert [p["index"] for p in per] == list(range(1, 21))
    assert [p["layout"] for p in per] == [s["layout"] for s in slides_cfg["slides"]]
    for p in per:
        assert p["chars"] >= 0 and p["minFont"] is not None and p["minFont"] >= MIN_FONT_PT, p
    # 正文字数在限制内（表格页不计入 120 字限制，但要有字数）
    for p, s in zip(per, slides_cfg["slides"]):
        if s["layout"] not in ("table", "appendix"):
            assert p["chars"] <= pptx_render.LIMITS["bodyChars"] + len(s.get("title", "")), p


# --------------------------------------------------------------------------- pptx 结构
def test_presentation_is_16_9_with_expected_pages(prs, slides_cfg: Dict[str, Any]) -> None:
    from pptx.util import Inches

    assert len(prs.slides) == len(slides_cfg["slides"]) == 20
    assert abs(prs.slide_width - Inches(13.333)) < Inches(0.01) and prs.slide_height == Inches(7.5)
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.01


def test_every_slide_has_speaker_notes(prs, slides_cfg: Dict[str, Any]) -> None:
    for i, (slide, cfg) in enumerate(zip(prs.slides, slides_cfg["slides"]), 1):
        assert slide.has_notes_slide, f"第 {i} 页没有备注"
        notes = slide.notes_slide.notes_text_frame.text
        assert notes.strip(), f"第 {i} 页备注为空"
        if cfg.get("duration"):
            assert f"建议时长：{cfg['duration']} 秒" in notes
    first = prs.slides[0].notes_slide.notes_text_frame.text
    assert "开场 20 秒" in first


def test_min_font_size_is_14pt_everywhere(prs) -> None:
    for i, slide in enumerate(prs.slides, 1):
        sizes = _font_sizes_pt(slide)
        assert sizes, f"第 {i} 页没有任何显式字号"
        assert min(sizes) >= MIN_FONT_PT, f"第 {i} 页最小字号 {min(sizes)}pt < {MIN_FONT_PT}"


def test_titles_are_at_least_28pt_and_footer_has_page_numbers(prs, slides_cfg: Dict[str, Any]) -> None:
    from pptx.util import Pt

    total = len(prs.slides)
    for i, (slide, cfg) in enumerate(zip(prs.slides, slides_cfg["slides"]), 1):
        text = _slide_text(slide)
        assert f"{i} / {total}" in text, f"第 {i} 页缺页码"
        assert slides_cfg["footer"] in text, f"第 {i} 页缺页脚"
        title_runs = [
            r for shp in slide.shapes if shp.has_text_frame
            for p in shp.text_frame.paragraphs for r in p.runs
            if r.font.size is not None and r.font.size >= Pt(28) and r.font.bold
        ]
        assert title_runs, f"第 {i} 页没有 ≥28pt 的粗体标题"
        if cfg["layout"] not in ("cover", "closing"):
            assert any(r.text and r.text in cfg["title"] for r in title_runs), f"第 {i} 页标题文字不符"


def test_native_charts(prs) -> None:
    from pptx.enum.chart import XL_CHART_TYPE

    found = list(_charts(prs))
    assert len(found) == 5, "示例有 5 页原生图表"
    types = {chart.chart_type for _, chart in found}
    assert types == {
        XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.COLUMN_STACKED,
        XL_CHART_TYPE.LINE_MARKERS, XL_CHART_TYPE.PIE,
    }
    for _, chart in found:
        xml = chart._chartSpace.xml
        assert 'typeface="微软雅黑"' in xml, "图表文本要补 a:ea 中文字体"
        assert '<c:autoTitleDeleted val="1"/>' in xml, "单系列图表不应自动显示系列名作标题"
        if chart.chart_type != XL_CHART_TYPE.PIE:
            assert chart.has_legend == (len(list(chart.plots[0].series)) > 1)
            assert chart.plots[0].categories
        if chart.has_legend:
            from pptx.enum.chart import XL_LEGEND_POSITION

            assert chart.legend.position == XL_LEGEND_POSITION.BOTTOM
    # 系列颜色来自主题序列色
    column = next(c for _, c in found if c.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED)
    assert str(column.plots[0].series[0].format.fill.fore_color.rgb) == NAVY
    pie = next(c for _, c in found if c.chart_type == XL_CHART_TYPE.PIE)
    assert len(pie.plots[0].categories) <= 4
    hbar = next(c for _, c in found if c.chart_type == XL_CHART_TYPE.BAR_CLUSTERED)
    assert list(hbar.plots[0].categories)[-1] == "设计感", "横条图数据倒序喂入，使第一类显示在最上"


def test_native_tables(prs) -> None:
    found = list(_tables(prs))
    assert len(found) == 2, "示例有 2 张表（table 页 + appendix 表）"
    for _, table in found:
        assert len(table.columns) <= 7 and len(table.rows) <= 9
        header = table.cell(0, 0)
        assert str(header.fill.fore_color.rgb) == NAVY
        assert header.text_frame.paragraphs[0].runs[0].font.bold
        assert str(header.text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"
        fills = {str(table.cell(i, 0).fill.fore_color.rgb) for i in range(1, len(table.rows))}
        assert fills == {"FFFFFF", "F5F7FA"}, f"正文行应为斑马纹：{fills}"
        tbl_pr = table._tbl.tblPr
        assert tbl_pr.get("bandRow") == "0" and tbl_pr.get("firstRow") == "1"
        for row in table.rows:
            for cell in row.cells:
                for r in cell.text_frame.paragraphs[0].runs:
                    assert r.font.size.pt >= MIN_FONT_PT
    comp = found[0][1]
    assert comp.cell(0, 3).text == "弘宝文创（我们）"
    assert comp.cell(1, 0).text == "主力价格带"


def test_fonts_and_theme(prs, result: Dict[str, Any]) -> None:
    with zipfile.ZipFile(result["output"]) as z:
        theme_xml = z.read("ppt/theme/theme1.xml").decode("utf-8")
    assert 'typeface="微软雅黑"' in theme_xml and 'typeface="Arial"' in theme_xml
    assert f'<a:srgbClr val="{NAVY}"/>' in theme_xml, "主题 dk2/accent1 应为海军蓝"
    run = next(r for shp in prs.slides[1].shapes if shp.has_text_frame
               for p in shp.text_frame.paragraphs for r in p.runs if r.text.strip())
    rpr = run._r.rPr
    assert rpr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea").get("typeface") == "微软雅黑"
    assert run.font.name == "Arial"


def test_dark_pages_and_appendix_badge(prs, slides_cfg: Dict[str, Any]) -> None:
    for slide, cfg in zip(prs.slides, slides_cfg["slides"]):
        fills = [str(s.fill.fore_color.rgb) for s in slide.shapes
                 if s.shape_type is not None and hasattr(s, "fill") and s.fill.type == 1]
        text = _slide_text(slide)
        if cfg["layout"] in ("cover", "section", "closing"):
            assert fills and fills[0] == NAVY, f"{cfg['layout']} 页应为深蓝整页底"
        if cfg["layout"] == "appendix":
            assert "备用页" in text
        else:
            assert "备用页" not in text or cfg["title"].startswith("备用页")


# --------------------------------------------------------------------------- preview.json
def test_preview_json(ws: Path, preview: List[Dict[str, Any]], slides_cfg: Dict[str, Any], result: Dict[str, Any]) -> None:
    assert len(preview) == result["pages"] == len(slides_cfg["slides"])
    keys = {"index", "layout", "title", "texts", "image", "chart", "table", "kpis", "steps", "members", "quote", "note"}
    for i, (pv, cfg) in enumerate(zip(preview, slides_cfg["slides"]), 1):
        assert set(pv) == keys, pv.keys()
        assert pv["index"] == i and pv["layout"] == cfg["layout"] and pv["title"] == cfg["title"]
        assert isinstance(pv["texts"], list) and pv["note"].strip()
        if pv["image"]:
            assert not Path(pv["image"]).is_absolute() and "/" in pv["image"] and ".." not in pv["image"]
            assert (ws / pv["image"]).exists(), "image 应为相对工作区根的路径"
    by_layout = {pv["layout"]: pv for pv in preview}
    assert by_layout["image"]["image"] == "04_图表/fig_market_size.png"
    assert by_layout["chart"]["chart"]["categories"] and by_layout["chart"]["chart"]["series"][0]["values"]
    assert by_layout["chart"]["chart"]["insight"]
    assert by_layout["table"]["table"]["columns"] == ["维度", "传统窑口", "景德镇文创", "弘宝文创（我们）"]
    assert len(by_layout["table"]["table"]["rows"]) == 5
    assert len(by_layout["kpi"]["kpis"]) == 4 and by_layout["kpi"]["kpis"][0]["value"] == "33.5 亿元"
    assert len(by_layout["timeline"]["steps"]) == 5 and by_layout["timeline"]["steps"][0]["when"] == "2026 Q2"
    assert len(by_layout["team"]["members"]) == 4
    assert by_layout["quote"]["quote"] == {"text": "宋韵落天青，一釉生万象", "by": "Big Idea · 把宋代美学变成可日用的生活方式"}
    assert by_layout["bullets"]["kpis"] == [{"value": "18.6 亿元", "label": "2025E 线下市场规模（示例数据）"}]
    assert by_layout["agenda"]["texts"] == slides_cfg["slides"][2]["items"]
    appendix_pages = [pv for pv in preview if pv["layout"] == "appendix"]
    assert appendix_pages[0]["image"] == "04_图表/fig_tornado.png" and appendix_pages[1]["table"]


# --------------------------------------------------------------------------- 匿名
def test_anonymize_applied_everywhere(prs, preview: List[Dict[str, Any]], result: Dict[str, Any]) -> None:
    hits = result["anonymizeHits"]
    assert {(h["page"], h["field"], h["term"]) for h in hits} == {(17, "members[0].duty", "某某大学"), (17, "note", "某某大学")}
    team = prs.slides[16]
    text = _slide_text(team)
    assert "某某大学" not in text and "曾获XX创业大赛金奖" in text
    notes = team.notes_slide.notes_text_frame.text
    assert "某某大学" not in notes and "XX" in notes
    pv = preview[16]
    assert all("某某大学" not in t for t in pv["texts"]) and "某某大学" not in pv["note"]
    assert "某某大学" not in json.dumps(preview, ensure_ascii=False)


def test_anonymize_value_whitelist_and_paths() -> None:
    hits: List[Dict[str, Any]] = []
    data = {"title": "某某大学与河南财经某某大学", "bullets": ["XX学院 的同学", "无"], "x": {"y": ["某某大学"]}, "n": 3}
    out = pptx_render.anonymize_value(data, ["某某大学", "XX学院"], ["河南财经某某大学"], 5, "", hits)
    assert out["title"] == "XX与河南财经某某大学", "白名单词应受保护"
    assert out["bullets"] == ["XX 的同学", "无"] and out["x"]["y"] == ["XX"] and out["n"] == 3
    assert {(h["field"], h["term"], h["count"]) for h in hits} == {("title", "某某大学", 1), ("bullets[0]", "XX学院", 1), ("x.y[0]", "某某大学", 1)}
    assert all(h["page"] == 5 for h in hits)


# --------------------------------------------------------------------------- 文字度量
def test_text_metrics() -> None:
    assert pptx_render.em_width("中文") == 2.0
    assert 0.5 < pptx_render.em_width("A1") < 1.4
    assert pptx_render.count_chars("汝瓷市场 6 年增长 2.3 倍") == 12, "字数 = 去空白后的字符数"
    assert pptx_render.lines_needed("十个字十个字十个字十", 20, 200) == 1
    assert pptx_render.lines_needed("十个字十个字十个字十一", 20, 200) == 2
    size, ok = pptx_render.fit_size(["十" * 30], 300, 60, 24, 18)
    assert ok and 18 <= size <= 24
    size, ok = pptx_render.fit_size(["十" * 300], 300, 60, 24, 18)
    assert not ok and size == 18
    t = pptx_render.truncate_em("一二三四五六七八九十", 5)
    assert t.endswith("…") and pptx_render.em_width(t) <= 5
    shrunk = pptx_render.shrink_paras(["短", "很长" * 40], 200, 100, 18)
    assert shrunk[0] == "短" and shrunk[1].endswith("…") and len(shrunk[1]) < 80
    split = pptx_render.balanced_split("一釉生万象：汝瓷文创让宋代美学走进日常", 40, 540)
    assert split == "一釉生万象：\n汝瓷文创让宋代美学走进日常"
    assert pptx_render.balanced_split("没有标点的一句很长很长的标题文字", 40, 200) is None
    assert pptx_render.one_line_size("十" * 10, 400, 44, 28) == 40


# --------------------------------------------------------------------------- 超限兜底
def test_limits_warn_and_clamp(ws: Path, tmp_path: Path) -> None:
    write_json(tmp_path / "project.json", {"name": "t"})
    cfg = {
        "title": "超限测试", "footer": "页脚", "anonymize": {"forbidden": ["某某大学"]},
        "slides": [
            {"layout": "cover", "title": "这是一个非常非常非常非常非常非常非常非常非常非常非常非常长的封面标题，用来测试标题截断逻辑是否正常工作", "note": "n"},
            {"layout": "bullets", "title": "这是一个超过二十八个字的标题这是一个超过二十八个字的标题这是一个超过二十八个字的标题",
             "bullets": [f"第 {k} 条要点，" + "很长的内容" * 12 for k in range(9)], "highlight": "1234567890123456789 亿元", "highlightLabel": "标签"},
            {"layout": "table", "title": "大表", "columns": [f"列{k}" for k in range(9)], "rows": [[f"r{i}c{j}" for j in range(9)] for i in range(12)]},
            {"layout": "chart", "title": "饼图太多类", "chart": {"type": "pie", "categories": list("ABCDEF"), "series": [{"name": "s", "values": [1, 2, 3, 4, 5, 6]}]}},
            {"layout": "chart", "title": "坏图", "chart": {"type": "sunburst", "categories": ["a", "b"], "series": [{"name": "s", "values": [1]}]}},
            {"layout": "kpi", "title": "太多 KPI", "kpis": [{"value": str(k), "label": f"L{k}"} for k in range(6)]},
            {"layout": "timeline", "title": "太多步", "steps": [{"when": f"T{k}", "what": f"W{k}"} for k in range(8)]},
            {"layout": "image", "title": "缺图", "image": "../04_图表/不存在.png", "bullets": ["a", "b", "c", "d"]},
            {"layout": "hologram", "title": "未知页型", "bullets": ["x"]},
            {"layout": "appendix", "title": "空备用页"},
            {"layout": "team", "title": "团队", "members": [{"role": f"R{k}", "duty": "某某大学毕业"} for k in range(7)]},
        ],
        "output": "out/t.pptx",
    }
    path = tmp_path / "06_路演" / "slides.json"
    write_json(path, cfg)
    res = pptx_render.render_deck(path)
    assert res["ok"] and res["pages"] == 11 and Path(res["output"]) == (tmp_path / "out" / "t.pptx").resolve()
    w = "\n".join(res["warnings"])
    for frag in ("标题超过两行，已截断", "字 > 28", "要点 9 条 > 6", "条要点", "字 > 40", "正文", "字 > 120", "数字「", "过长，已截断",
                 "9 列 > 7", "12 行 > 8", "饼图有 6 类 > 4", "不支持的图表类型「sunburst」", "有 1 个值，类目有 2 个",
                 "kpi 6 张 > 4", "timeline 8 步 > 6", "找不到图片", "读图要点 4 条 > 3", "未知 layout「hologram」",
                 "appendix 页没有", "team 7 人 > 6"):
        assert frag in w, f"缺少警告「{frag}」：\n{w}"
    assert res["missingImages"] == ["../04_图表/不存在.png"]
    # 匿名在截断之前做：7 名成员都被清洗（第 7 名随后被截掉）
    assert len(res["anonymizeHits"]) == 7 and all(h["page"] == 11 for h in res["anonymizeHits"])
    assert all(p["minFont"] >= MIN_FONT_PT for p in res["perPage"]), res["perPage"]

    from pptx import Presentation
    from pptx.enum.chart import XL_CHART_TYPE

    prs2 = Presentation(res["output"])
    assert len(prs2.slides) == 11
    for i, slide in enumerate(prs2.slides, 1):
        assert min(_font_sizes_pt(slide)) >= MIN_FONT_PT, f"第 {i} 页字号过小"
        assert slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()
    bullets_boxes = [s for s in prs2.slides[1].shapes if s.has_text_frame and len(s.text_frame.paragraphs) > 3]
    assert bullets_boxes and len(bullets_boxes[0].text_frame.paragraphs) == 6, "要点截到 6 条"
    assert all(r.font.size.pt >= 18 for p in bullets_boxes[0].text_frame.paragraphs for r in p.runs), "要点字号不低于 18pt"
    table = next(t for _, t in _tables(prs2))
    assert len(table.columns) == 7 and len(table.rows) == 9
    chart_types = [c.chart_type for _, c in _charts(prs2)]
    assert chart_types == [XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.COLUMN_CLUSTERED], "饼图 >4 类与未知类型都退化为簇状柱"
    assert "【缺图：" in _slide_text(prs2.slides[7])
    assert "备用页" in _slide_text(prs2.slides[9])
    assert "某某大学" not in _slide_text(prs2.slides[10])
    preview2 = read_json(res["preview"])
    assert len(preview2) == 11 and preview2[8]["layout"] == "bullets" and preview2[7]["image"] is None
    assert len(preview2[5]["kpis"]) == 4 and len(preview2[6]["steps"]) == 6 and len(preview2[10]["members"]) == 6


# --------------------------------------------------------------------------- CLI
def test_cli_main_prints_json(ws: Path, capsys: pytest.CaptureFixture[str]) -> None:
    code = pptx_render.main([str(ws / "06_路演" / "slides.json")])
    out = capsys.readouterr().out
    assert code == 0
    payload = json.loads(out[out.index("{"):])
    assert payload["ok"] and payload["pages"] == 20 and payload["warnings"] == []
    assert pptx_render.main([str(ws / "没有.json")]) == 1
    err = json.loads(capsys.readouterr().out)
    assert not err["ok"] and "文件不存在" in err["error"]
