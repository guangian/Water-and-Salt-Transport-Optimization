"""生成示例工作区 inputs/ 下的两个测试文件（供 preread 单测使用）：

    python examples/make_demo_inputs.py            # 在 agent/toolkit 下运行

- inputs/示例销售数据.xlsx：两个工作表（月度销售、产品清单），含数值列；
- inputs/示例企业资料.pdf：含中文段落与一张表格。优先用 python-docx + 本机 PDF 转换器生成（文本与表格线可被 pdfplumber 识别），
  没有转换器时退化为 matplotlib 画一页带表格线的 PDF。
"""
from __future__ import annotations

import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

INPUTS = HERE / "ws_demo" / "inputs"

MONTHS = [f"2025-{m:02d}" for m in range(1, 13)]
CHANNELS = ["景区门店", "电商平台"]


def make_xlsx(path: Path) -> Path:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "月度销售"
    ws.append(["月份", "渠道", "销量（件）", "销售额（元）", "客单价（元）"])
    base = {"景区门店": (420, 760), "电商平台": (310, 540)}
    for i, m in enumerate(MONTHS):
        for ch in CHANNELS:
            qty = base[ch][0] + i * 9 + (23 if ch == "电商平台" else 0) * (i % 3)
            price = base[ch][1] + (i % 4) * 12
            ws.append([m, ch, qty, qty * price, price])
    ws2 = wb.create_sheet("产品清单")
    ws2.append(["产品编号", "产品名称", "系列", "零售价（元）", "单位成本（元）", "是否在售"])
    rows = [("HB-001", "天青雅集茶器套装", "茶器", 1280, 560, "是"), ("HB-002", "莲花香插", "香器", 268, 96, "是"),
            ("HB-003", "三足洗", "收藏", 3600, 1500, "是"), ("HB-004", "斗笠盏（对）", "茶器", 680, 280, "否")]
    for r in rows:
        ws2.append(list(r))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path


PARAGRAPHS = [
    "汝州市弘宝汝瓷坊有限公司成立于 1999 年，是集汝瓷研发、烧制、销售与文化传播于一体的省级非遗保护单位。公司现有员工 86 人，其中省级非遗代表性传承人 2 名、高级工艺美术师 5 名。",
    "2024 年公司实现销售收入 1,150 万元，同比增长 11.6%；其中景区门店与政企礼品渠道合计占比 68%，线上渠道占比 28%，其余为展会与拍卖。",
    "公司拥有天青釉配方相关专利 3 项，年产量约 6 万件，主力产品为茶器与香器，零售价集中在 500—1,000 元区间。",
]
TABLE = [["年份", "销售收入（万元）", "线上占比", "员工数"], ["2022", "920", "19%", "72"], ["2023", "1,030", "24%", "80"], ["2024", "1,150", "28%", "86"]]


def _pdf_via_docx(path: Path) -> bool:
    try:
        from docx import Document
        from docx.shared import Pt
        from bizkit import pdf as pdf_mod
    except Exception:
        return False
    if not pdf_mod.detect():
        return False
    tmp_docx = path.with_suffix(".tmp.docx")
    doc = Document()
    doc.add_heading("弘宝汝瓷企业基本资料（示例）", level=1)
    for text in PARAGRAPHS:
        p = doc.add_paragraph(text)
        p.paragraph_format.space_after = Pt(6)
    doc.add_paragraph("表 1 近三年经营数据")
    table = doc.add_table(rows=len(TABLE), cols=len(TABLE[0]))
    table.style = "Table Grid"
    for r, row in enumerate(TABLE):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    doc.save(str(tmp_docx))
    try:
        result = pdf_mod.convert(tmp_docx, out=path)
        return bool(result.get("ok"))
    finally:
        tmp_docx.unlink(missing_ok=True)


def _pdf_via_matplotlib(path: Path) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    from bizkit import fonts
    fonts.apply_matplotlib("sans")
    import textwrap
    with PdfPages(str(path)) as pdf:
        fig = plt.figure(figsize=(8.27, 11.69))
        ax = fig.add_axes([0, 0, 1, 1])
        ax.axis("off")
        y = 0.94
        ax.text(0.1, y, "弘宝汝瓷企业基本资料（示例）", fontsize=16, weight="bold")
        y -= 0.05
        for text in PARAGRAPHS:
            for line in textwrap.wrap(text, 38):
                ax.text(0.1, y, line, fontsize=10.5)
                y -= 0.022
            y -= 0.012
        ax.text(0.1, y, "表 1 近三年经营数据", fontsize=10.5)
        y -= 0.03
        x0, x1, rh = 0.1, 0.9, 0.028
        cols = len(TABLE[0])
        for r, row in enumerate(TABLE):
            yy = y - r * rh
            ax.plot([x0, x1], [yy, yy], color="black", lw=0.8)
            for c, val in enumerate(row):
                ax.text(x0 + (c + 0.5) * (x1 - x0) / cols, yy - rh / 2, val, fontsize=10, ha="center", va="center")
        yy = y - len(TABLE) * rh
        ax.plot([x0, x1], [yy, yy], color="black", lw=0.8)
        for c in range(cols + 1):
            xx = x0 + c * (x1 - x0) / cols
            ax.plot([xx, xx], [y, yy], color="black", lw=0.8)
        pdf.savefig(fig)
        plt.close(fig)


def make_pdf(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not _pdf_via_docx(path):
        _pdf_via_matplotlib(path)
    return path


def ensure_inputs(inputs_dir: Path = INPUTS) -> dict:
    xlsx = inputs_dir / "示例销售数据.xlsx"
    pdf = inputs_dir / "示例企业资料.pdf"
    if not xlsx.exists():
        make_xlsx(xlsx)
    if not pdf.exists():
        make_pdf(pdf)
    return {"xlsx": xlsx, "pdf": pdf}


if __name__ == "__main__":
    out = ensure_inputs()
    for k, v in out.items():
        print(k, v, v.stat().st_size, "bytes")
